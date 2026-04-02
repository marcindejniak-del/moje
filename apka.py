from __future__ import annotations
import atexit
import calendar
import csv
import json
import os
import re
import smtplib
import ssl
import statistics
import threading
import sqlite3
import tkinter as tk
from dataclasses import asdict, dataclass, field
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
try:
    from PIL import ImageGrab
    PIL_AVAILABLE = True
except Exception:
    PIL_AVAILABLE = False
EMAIL_RE = re.compile('^[^@\\s]+@[^@\\s]+\\.[^@\\s]+$')

@dataclass
class EmailAlarmConfig:
    enabled: bool = False
    smtp_host: str = 'smtp.gmail.com'
    smtp_port: int = 465
    smtp_username: str = ''
    smtp_password: str = ''
    sender_name: str = 'Analiza inercji'
    sender_email: str = ''
    recipients: list[str] = field(default_factory=list)
    trigger_mode: str = 'above'

    def validate(self):
        if not self.enabled:
            return (True, 'Powiadomienia e-mail są wyłączone')
        if not self.smtp_host.strip():
            return (False, 'Brak hosta SMTP')
        if not self.smtp_port:
            return (False, 'Brak portu SMTP')
        if not self.smtp_username.strip():
            return (False, 'Brak loginu SMTP')
        if not self.smtp_password.strip():
            return (False, 'Brak hasła SMTP')
        if not self.sender_email.strip() or not EMAIL_RE.match(self.sender_email.strip()):
            return (False, 'Niepoprawny adres nadawcy')
        if not self.recipients:
            return (False, 'Lista odbiorców jest pusta')
        bad = [x for x in self.recipients if not EMAIL_RE.match(x)]
        if bad:
            return (False, f"Niepoprawne adresy odbiorców: {', '.join(bad)}")
        if self.trigger_mode not in {'above', 'below'}:
            return (False, 'Niepoprawny tryb alarmu')
        return (True, 'OK')

class EmailAlarmService:

    def __init__(self, config_path='~/.analiza_inercji_email.json'):
        self.config_path = Path(config_path).expanduser()
        self.config = EmailAlarmConfig()
        self.last_alarm_key = None
        self.load_config()

    @staticmethod
    def parse_recipients(raw):
        unique = {}
        for item in (x.strip() for x in re.split('[;,\\n]', raw)):
            if item:
                unique.setdefault(item.lower(), item)
        return list(unique.values())

    def load_config(self):
        if not self.config_path.exists():
            return self.config
        try:
            with self.config_path.open('r', encoding='utf-8') as f:
                raw = json.load(f)
            self.config = EmailAlarmConfig(**raw)
        except Exception:
            self.config = EmailAlarmConfig()
        return self.config

    def save_config(self):
        self.config_path.parent.mkdir(parents=True, exist_ok=True)
        with self.config_path.open('w', encoding='utf-8') as f:
            json.dump(asdict(self.config), f, ensure_ascii=False, indent=2)

    def should_trigger(self, prog, wartosc):
        return wartosc > prog if self.config.trigger_mode == 'above' else wartosc < prog

    def send_alarm(self, *, prog, wartosc, data_modelu, godzina_modelu, source_file=''):
        ok, msg = self.config.validate()
        if not ok:
            return (False, msg)
        alarm_key = f'{data_modelu}|{godzina_modelu}|{prog}|{wartosc:.3f}|{self.config.trigger_mode}'
        if alarm_key == self.last_alarm_key:
            return (True, 'Alarm dla tego samego modelu już został wysłany')
        delta = wartosc - prog
        direction = 'przekroczenie' if self.config.trigger_mode == 'above' else 'spadek poniżej'
        when_txt = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        email = EmailMessage()
        email['Subject'] = f'[ALARM INERCJI] {direction} progu | {data_modelu} {godzina_modelu}'
        email['From'] = f'{self.config.sender_name} <{self.config.sender_email}>'
        email['To'] = ', '.join(self.config.recipients)
        email.set_content(f"Data modelu: {data_modelu}\nGodzina modelu: {godzina_modelu}\nWartość inercji: {wartosc:,.2f} MVA·s\nPróg: {prog:,.2f} MVA·s\nRóżnica: {delta:,.2f} MVA·s\nPlik: {source_file or '-'}\nWysłano: {when_txt}\n")
        ctx = ssl.create_default_context()
        with smtplib.SMTP_SSL(self.config.smtp_host, int(self.config.smtp_port), context=ctx, timeout=20) as smtp:
            smtp.login(self.config.smtp_username, self.config.smtp_password)
            smtp.send_message(email)
        self.last_alarm_key = alarm_key
        return (True, 'Mail alarmowy wysłany')

def center_toplevel(window: tk.Toplevel, width: int, height: int):
    try:
        window.update_idletasks()
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()
        x = max(40, int((screen_w - width) / 2))
        y = max(40, int((screen_h - height) / 2))
        window.geometry(f'{width}x{height}+{x}+{y}')
    except Exception:
        window.geometry(f'{width}x{height}')

def get_app_from_widget(widget):
    current = widget
    while current is not None:
        app = getattr(current, '_analiza_inercji_app', None)
        if app is not None:
            return app
        current = getattr(current, 'master', None)
    return None

def init_managed_toplevel(window: tk.Toplevel, layout_key: str, width: int, height: int):
    app = get_app_from_widget(window)
    if app is not None:
        app.apply_saved_window_geometry(window, layout_key, width, height)
        app.register_managed_window(window, layout_key)
    else:
        center_toplevel(window, width, height)
    try:
        window.resizable(True, True)
    except Exception:
        pass

class DatePickerWindow(tk.Toplevel):

    WEEKDAY_LABELS = ['Pn', 'Wt', 'Sr', 'Cz', 'Pt', 'Sb', 'Nd']
    MONTH_LABELS = {1: 'Styczen', 2: 'Luty', 3: 'Marzec', 4: 'Kwiecien', 5: 'Maj', 6: 'Czerwiec', 7: 'Lipiec', 8: 'Sierpien', 9: 'Wrzesien', 10: 'Pazdziernik', 11: 'Listopad', 12: 'Grudzien'}

    def __init__(self, master, *, initial_date=None, on_select=None, title='Wybierz date'):
        super().__init__(master)
        self.on_select = on_select
        current_date = initial_date or datetime.now().date()
        self.selected_date = initial_date
        self.display_year = int(current_date.year)
        self.display_month = int(current_date.month)
        self._calendar = calendar.Calendar(firstweekday=0)
        self._day_buttons = []
        self.title(title)
        init_managed_toplevel(self, 'date_picker', 360, 320)
        self.transient(master)
        try:
            self.grab_set()
        except Exception:
            pass
        self._build_ui()
        self._render_calendar()

    def _build_ui(self):
        app = get_app_from_widget(self)
        pal = app.get_palette() if app is not None else {'bg': '#f5f7fb', 'card': '#ffffff', 'text': '#111827', 'muted': '#667085', 'accent': '#3b82f6', 'neutral_btn': '#eef2f7', 'neutral_btn_active': '#dbe4f0', 'border': '#d7e0ec'}
        font_ui = app.font_ui if app is not None else ('Segoe UI', 10)
        font_ui_bold = app.font_ui_bold if app is not None else ('Segoe UI Semibold', 10)
        self.configure(bg=pal['bg'])
        outer = tk.Frame(self, bg=pal['bg'], padx=10, pady=10)
        outer.pack(fill=tk.BOTH, expand=True)
        header = tk.Frame(outer, bg=pal['card'])
        header.pack(fill=tk.X)
        self.btn_prev = tk.Button(header, text='<', command=self.prev_month, width=3)
        self.btn_prev.pack(side=tk.LEFT)
        self.lbl_month = tk.Label(header, text='', bg=pal['card'], fg=pal['text'], font=font_ui_bold)
        self.lbl_month.pack(side=tk.LEFT, expand=True)
        self.btn_next = tk.Button(header, text='>', command=self.next_month, width=3)
        self.btn_next.pack(side=tk.RIGHT)
        if app is not None:
            app.stylizuj_przycisk(self.btn_prev, 'neutral')
            app.stylizuj_przycisk(self.btn_next, 'neutral')

        action_row = tk.Frame(outer, bg=pal['bg'])
        action_row.pack(fill=tk.X, pady=(8, 8))
        self.btn_today = tk.Button(action_row, text='Dzis', command=self.select_today)
        self.btn_today.pack(side=tk.LEFT, padx=(0, 6))
        self.btn_clear = tk.Button(action_row, text='Wyczysc', command=self.clear_value)
        self.btn_clear.pack(side=tk.LEFT, padx=(0, 6))
        self.btn_close = tk.Button(action_row, text='Zamknij', command=self.destroy)
        self.btn_close.pack(side=tk.RIGHT)
        if app is not None:
            app.stylizuj_przycisk(self.btn_today, 'primary')
            app.stylizuj_przycisk(self.btn_clear, 'neutral')
            app.stylizuj_przycisk(self.btn_close, 'neutral')

        grid = tk.Frame(outer, bg=pal['card'], highlightthickness=1, highlightbackground=pal['border'])
        grid.pack(fill=tk.BOTH, expand=True)
        for col, label in enumerate(self.WEEKDAY_LABELS):
            lbl = tk.Label(grid, text=label, bg=pal['card'], fg=pal['muted'], font=font_ui_bold, padx=6, pady=6)
            lbl.grid(row=0, column=col, sticky='nsew')
            grid.grid_columnconfigure(col, weight=1)
        for row in range(1, 7):
            grid.grid_rowconfigure(row, weight=1)
            row_buttons = []
            for col in range(7):
                btn = tk.Button(grid, text='', relief='flat', bd=0, font=font_ui, command=lambda r=row, c=col: self.on_day_click(r - 1, c))
                btn.grid(row=row, column=col, sticky='nsew', padx=2, pady=2)
                row_buttons.append(btn)
            self._day_buttons.append(row_buttons)
        self._grid_frame = grid

    def _render_calendar(self):
        app = get_app_from_widget(self)
        pal = app.get_palette() if app is not None else {'card': '#ffffff', 'text': '#111827', 'muted': '#667085', 'accent': '#3b82f6', 'selection_fg': '#ffffff', 'neutral_btn': '#eef2f7'}
        month_name = self.MONTH_LABELS.get(self.display_month, f'M{self.display_month:02d}')
        self.lbl_month.config(text=f'{month_name} {self.display_year}')
        weeks = self._calendar.monthdayscalendar(self.display_year, self.display_month)
        while len(weeks) < 6:
            weeks.append([0] * 7)
        for row_idx, week in enumerate(weeks):
            for col_idx, day_num in enumerate(week):
                btn = self._day_buttons[row_idx][col_idx]
                if day_num <= 0:
                    btn.config(text='', state='disabled', bg=pal['card'], fg=pal['muted'], activebackground=pal['card'])
                    continue
                btn.config(text=str(day_num), state='normal')
                is_selected = self.selected_date is not None and self.selected_date.year == self.display_year and self.selected_date.month == self.display_month and self.selected_date.day == day_num
                if is_selected:
                    btn.config(bg=pal['accent'], fg='white', activebackground=pal['accent'], activeforeground='white')
                else:
                    btn.config(bg=pal['neutral_btn'], fg=pal['text'], activebackground=pal['neutral_btn_active'] if 'neutral_btn_active' in pal else pal['neutral_btn'], activeforeground=pal['text'])

    def prev_month(self):
        if self.display_month == 1:
            self.display_month = 12
            self.display_year -= 1
        else:
            self.display_month -= 1
        self._render_calendar()

    def next_month(self):
        if self.display_month == 12:
            self.display_month = 1
            self.display_year += 1
        else:
            self.display_month += 1
        self._render_calendar()

    def on_day_click(self, row_idx, col_idx):
        weeks = self._calendar.monthdayscalendar(self.display_year, self.display_month)
        while len(weeks) < 6:
            weeks.append([0] * 7)
        day_num = weeks[row_idx][col_idx]
        if day_num <= 0:
            return
        self.selected_date = datetime(self.display_year, self.display_month, day_num).date()
        if callable(self.on_select):
            self.on_select(self.selected_date.strftime('%d.%m.%Y'))
        self.destroy()

    def select_today(self):
        today = datetime.now().date()
        if callable(self.on_select):
            self.on_select(today.strftime('%d.%m.%Y'))
        self.destroy()

    def clear_value(self):
        if callable(self.on_select):
            self.on_select('')
        self.destroy()

class EmailAlarmConfigWindow(tk.Toplevel):

    def __init__(self, master, service):
        super().__init__(master)
        self.service = service
        self.title('Konfiguracja powiadomień e-mail')
        init_managed_toplevel(self, 'email_alarm', 820, 470)
        cfg = service.config
        self.enabled_var = tk.BooleanVar(value=cfg.enabled)
        self.smtp_host_var = tk.StringVar(value=cfg.smtp_host)
        self.smtp_port_var = tk.StringVar(value=str(cfg.smtp_port))
        self.smtp_username_var = tk.StringVar(value=cfg.smtp_username)
        self.smtp_password_var = tk.StringVar(value=cfg.smtp_password)
        self.sender_name_var = tk.StringVar(value=cfg.sender_name)
        self.sender_email_var = tk.StringVar(value=cfg.sender_email)
        self.recipients_var = tk.StringVar(value=', '.join(cfg.recipients))
        self.trigger_mode_var = tk.StringVar(value=cfg.trigger_mode)
        self._build_ui()

    def _style_entry(self, entry: tk.Entry):
        entry.configure(relief='flat', bd=0, highlightthickness=1, highlightbackground='#d0d7e2', highlightcolor='#4f7cff', bg='white', fg='#111827', insertbackground='#111827', font=('Segoe UI', 10))

    def _style_button(self, button: tk.Button, role='neutral'):
        schemes = {'primary': ('#3b82f6', 'white', '#2563eb'), 'success': ('#16a34a', 'white', '#15803d'), 'neutral': ('#eef2f7', '#1f2937', '#dbe4f0')}
        bg, fg, active = schemes.get(role, schemes['neutral'])
        button.configure(bg=bg, fg=fg, activebackground=active, activeforeground=fg, relief='flat', bd=0, padx=12, pady=8, cursor='hand2', font=('Segoe UI Semibold', 10), highlightthickness=0)

    def _build_ui(self):
        self.configure(bg='#f5f7fb')
        fr = tk.Frame(self, padx=16, pady=16, bg='#f5f7fb')
        fr.pack(fill=tk.BOTH, expand=True)
        head = tk.Frame(fr, bg='#eaf1ff', padx=14, pady=12)
        head.pack(fill=tk.X, pady=(0, 14))
        tk.Label(head, text='Powiadomienia e-mail', bg='#eaf1ff', fg='#0f172a', font=('Segoe UI Semibold', 15)).pack(anchor='w')
        tk.Label(head, text='Wprowadź dane SMTP i odbiorców alarmów dla progów inercji.', bg='#eaf1ff', fg='#475467', font=('Segoe UI', 10)).pack(anchor='w', pady=(4, 0))
        card = tk.LabelFrame(fr, text='Ustawienia wysyłki', padx=12, pady=12, bg='white', fg='#334155', font=('Segoe UI Semibold', 10), bd=1, relief='solid')
        card.pack(fill=tk.BOTH, expand=True)
        tk.Checkbutton(card, text='Włącz powiadomienia e-mail', variable=self.enabled_var, bg='white', fg='#0f172a', activebackground='white', activeforeground='#0f172a', selectcolor='white', font=('Segoe UI', 10)).grid(row=0, column=0, columnspan=2, sticky='w', pady=(0, 8))
        fields = [('SMTP host', self.smtp_host_var, None), ('SMTP port', self.smtp_port_var, None), ('SMTP login', self.smtp_username_var, None), ('SMTP hasło / App Password', self.smtp_password_var, '*'), ('Nazwa nadawcy', self.sender_name_var, None), ('Adres nadawcy', self.sender_email_var, None), ('Adresaci (przecinek/średnik)', self.recipients_var, None)]
        for i, (lbl, var, mask) in enumerate(fields, start=1):
            tk.Label(card, text=lbl, bg='white', fg='#344054', font=('Segoe UI Semibold', 9)).grid(row=i, column=0, sticky='w', pady=5)
            ent = tk.Entry(card, textvariable=var, width=72, show=mask)
            ent.grid(row=i, column=1, sticky='we', pady=5)
            self._style_entry(ent)
        tk.Label(card, text='Tryb alarmu', bg='white', fg='#344054', font=('Segoe UI Semibold', 9)).grid(row=8, column=0, sticky='w', pady=5)
        om = tk.OptionMenu(card, self.trigger_mode_var, 'above', 'below')
        om.grid(row=8, column=1, sticky='w')
        om.configure(bg='#eef2f7', fg='#1f2937', activebackground='#dbe4f0', activeforeground='#1f2937', relief='flat', bd=0, cursor='hand2', font=('Segoe UI', 10), highlightthickness=0)
        try:
            om['menu'].configure(bg='white', fg='#111827', activebackground='#eaf1ff', activeforeground='#111827')
        except Exception:
            pass
        row = tk.Frame(fr, bg='#f5f7fb')
        row.pack(anchor='w', pady=(14, 0))
        btn_save = tk.Button(row, text='Zapisz', command=self.on_save)
        btn_test = tk.Button(row, text='Wyślij test', command=self.on_test)
        btn_close = tk.Button(row, text='Zamknij', command=self.destroy)
        self._style_button(btn_save, 'success')
        self._style_button(btn_test, 'primary')
        self._style_button(btn_close, 'neutral')
        btn_save.pack(side=tk.LEFT, padx=4)
        btn_test.pack(side=tk.LEFT, padx=4)
        btn_close.pack(side=tk.LEFT, padx=4)
        card.columnconfigure(1, weight=1)

    def _apply(self):
        c = self.service.config
        c.enabled = self.enabled_var.get()
        c.smtp_host = self.smtp_host_var.get().strip()
        c.smtp_port = int(self.smtp_port_var.get().strip() or '0')
        c.smtp_username = self.smtp_username_var.get().strip()
        c.smtp_password = self.smtp_password_var.get().strip()
        c.sender_name = self.sender_name_var.get().strip() or 'Analiza inercji'
        c.sender_email = self.sender_email_var.get().strip()
        c.recipients = self.service.parse_recipients(self.recipients_var.get())
        c.trigger_mode = self.trigger_mode_var.get().strip() or 'above'

    def on_save(self):
        try:
            self._apply()
            ok, msg = self.service.config.validate()
            if not ok:
                messagebox.showerror('Błąd konfiguracji', msg)
                return
            self.service.save_config()
            messagebox.showinfo('E-mail', 'Zapisano konfigurację.')
        except Exception as e:
            messagebox.showerror('Błąd', str(e))

    def on_test(self):
        try:
            self._apply()
            self.service.last_alarm_key = None
            ok, msg = self.service.send_alarm(prog=10000.0, wartosc=12000.0 if self.service.config.trigger_mode == 'above' else 9000.0, data_modelu='TEST', godzina_modelu='00:30', source_file='test.epc')
            if ok:
                messagebox.showinfo('E-mail', msg)
            else:
                messagebox.showwarning('E-mail', msg)
        except Exception as e:
            messagebox.showerror('Błąd wysyłki', str(e))

class AlarmHistoryWindow(tk.Toplevel):

    def __init__(self, master, app):
        super().__init__(master)
        self.app = app
        self.title('Historia alarmow')
        init_managed_toplevel(self, 'alarm_history', 1540, 900)
        pal = self.app.get_palette()
        self.configure(bg=pal['bg'])
        self.query_var = tk.StringVar()
        self.trigger_var = tk.StringVar(value='Wszystkie')
        self.mail_var = tk.StringVar(value='Wszystkie')
        self.type_var = tk.StringVar(value='Wszystkie')
        self.limit_var = tk.StringVar(value='500')
        self.current_rows = []
        self._build_ui()
        self.refresh()

    def _create_filter_combo(self, parent, row, col, label, variable, values, width=16):
        pal = self.app.get_palette()
        box = tk.Frame(parent, bg=pal['card'])
        box.grid(row=row, column=col, padx=4, pady=4, sticky='ew')
        tk.Label(box, text=label, bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
        combo = ttk.Combobox(box, textvariable=variable, values=list(values), state='readonly', width=width)
        combo.pack(fill=tk.X)
        combo.bind('<<ComboboxSelected>>', lambda *_: self.refresh())
        return combo

    def _build_ui(self):
        pal = self.app.get_palette()
        outer = tk.Frame(self, bg=pal['bg'], padx=10, pady=10)
        outer.pack(fill=tk.BOTH, expand=True)

        filters = tk.LabelFrame(outer, text='Filtry i sterowanie', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        filters.pack(fill=tk.X)
        query_box = tk.Frame(filters, bg=pal['card'])
        query_box.grid(row=0, column=0, columnspan=2, padx=4, pady=4, sticky='ew')
        tk.Label(query_box, text='Szukaj', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
        entry_query = tk.Entry(query_box, textvariable=self.query_var, font=self.app.font_code_small)
        entry_query.pack(fill=tk.X)
        self.app.stylizuj_entry(entry_query)
        entry_query.bind('<KeyRelease>', lambda *_: self.refresh())
        self._create_filter_combo(filters, 0, 2, 'Tryb alarmu', self.trigger_var, ['Wszystkie', 'above', 'below'], width=12)
        self._create_filter_combo(filters, 0, 3, 'Stan maila', self.mail_var, ['Wszystkie', 'Sukces', 'Blad', 'Wylaczone'], width=12)
        self._create_filter_combo(filters, 0, 4, 'Typ modelu', self.type_var, ['Wszystkie', self.app.get_model_type_label('manual'), self.app.get_model_type_label('manual_batch'), self.app.get_model_type_label('auto')], width=12)
        limit_box = tk.Frame(filters, bg=pal['card'])
        limit_box.grid(row=0, column=5, padx=4, pady=4, sticky='ew')
        tk.Label(limit_box, text='Limit rekordow', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
        entry_limit = tk.Entry(limit_box, textvariable=self.limit_var, width=10, font=self.app.font_code_small)
        entry_limit.pack(fill=tk.X)
        self.app.stylizuj_entry(entry_limit)
        entry_limit.bind('<KeyRelease>', lambda *_: self.refresh())
        for idx in range(6):
            filters.grid_columnconfigure(idx, weight=1)
        btn_row = tk.Frame(filters, bg=pal['card'])
        btn_row.grid(row=1, column=0, columnspan=6, sticky='w', pady=(8, 0))
        self.app.utworz_przycisk(btn_row, 'Odswiez', self.refresh, role='primary').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Wyczysc filtry', self.clear_filters, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Eksport CSV', self.export_csv, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Eksport Excel', self.export_xlsx, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Kopiuj szczegoly', self.copy_selected, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Wyczysc historie', self.clear_history, role='danger').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btn_row, 'Zamknij', self.destroy, role='neutral').pack(side=tk.RIGHT, padx=3)

        frame_summary = tk.LabelFrame(outer, text='Podsumowanie', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        frame_summary.pack(fill=tk.X, pady=(10, 0))
        self.summary_label = tk.Label(frame_summary, text='Brak alarmow', justify='left', anchor='w', bg=pal['card'], fg=pal['accent'], font=self.app.font_code_small)
        self.summary_label.pack(fill=tk.X)

        body = ttk.PanedWindow(outer, orient=tk.HORIZONTAL)
        body.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        left = tk.Frame(body, bg=pal['bg'])
        right = tk.Frame(body, bg=pal['bg'])
        body.add(left, weight=5)
        body.add(right, weight=3)

        frame_tree = tk.LabelFrame(left, text='Lista alarmow', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        frame_tree.pack(fill=tk.BOTH, expand=True)
        cols = ('czas', 'typ_alarmu', 'typ_modelu', 'data', 'godzina', 'wartosc', 'prog', 'delta', 'plik', 'mail', 'status')
        self.tree = ttk.Treeview(frame_tree, columns=cols, show='headings', height=24)
        headers = {'czas': 'Czas alarmu', 'typ_alarmu': 'Tryb', 'typ_modelu': 'Model', 'data': 'Data modelu', 'godzina': 'Godzina', 'wartosc': 'Inercja [MVA*s]', 'prog': 'Prog [MVA*s]', 'delta': 'Odchylka', 'plik': 'Plik', 'mail': 'Mail', 'status': 'Status wysylki'}
        widths = {'czas': 150, 'typ_alarmu': 70, 'typ_modelu': 95, 'data': 100, 'godzina': 82, 'wartosc': 110, 'prog': 105, 'delta': 105, 'plik': 230, 'mail': 60, 'status': 220}
        anchors = {'wartosc': 'e', 'prog': 'e', 'delta': 'e'}
        for col in cols:
            self.tree.heading(col, text=headers[col])
            self.tree.column(col, width=widths[col], anchor=anchors.get(col, 'w'))
        yscroll = ttk.Scrollbar(frame_tree, orient='vertical', command=self.tree.yview)
        xscroll = ttk.Scrollbar(frame_tree, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.grid(row=0, column=0, sticky='nsew')
        yscroll.grid(row=0, column=1, sticky='ns')
        xscroll.grid(row=1, column=0, sticky='ew')
        frame_tree.rowconfigure(0, weight=1)
        frame_tree.columnconfigure(0, weight=1)
        self.tree.bind('<<TreeviewSelect>>', lambda *_: self.show_details())
        self.tree.tag_configure('mail_ok', background='#e8f5e9', foreground='#1b5e20')
        self.tree.tag_configure('mail_err', background='#ffebee', foreground='#b71c1c')
        self.tree.tag_configure('mail_off', background='#eeeeee', foreground='#424242')

        right_pane = ttk.PanedWindow(right, orient=tk.VERTICAL)
        right_pane.pack(fill=tk.BOTH, expand=True)
        frame_details = tk.LabelFrame(right_pane, text='Szczegoly alarmu', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        self.details_text = scrolledtext.ScrolledText(frame_details, wrap=tk.WORD, height=14, font=self.app.font_code_small)
        self.details_text.pack(fill=tk.BOTH, expand=True)
        self.details_text.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'])
        right_pane.add(frame_details, weight=3)
        frame_chart = tk.LabelFrame(right_pane, text='Diagnostyka alarmow', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        self.fig_alarm, (self.ax_alarm_status, self.ax_alarm_recent) = plt.subplots(2, 1, figsize=(7.0, 6.2), gridspec_kw={'height_ratios': [1.0, 1.1]})
        self.fig_alarm.subplots_adjust(left=0.10, right=0.98, top=0.94, bottom=0.12, hspace=0.48)
        self.canvas_alarm = FigureCanvasTkAgg(self.fig_alarm, master=frame_chart)
        self.canvas_alarm.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        right_pane.add(frame_chart, weight=2)

    def parse_limit(self):
        try:
            value = int(str(self.limit_var.get()).strip() or '500')
        except Exception:
            value = 500
        return max(20, min(5000, value))

    def get_mail_state(self, alarm):
        if alarm.get('mail_attempted') and alarm.get('mail_ok'):
            return 'Sukces'
        if alarm.get('mail_attempted'):
            return 'Blad'
        return 'Wylaczone'

    def get_alarm_tag(self, alarm):
        state = self.get_mail_state(alarm)
        if state == 'Sukces':
            return 'mail_ok'
        if state == 'Blad':
            return 'mail_err'
        return 'mail_off'

    def get_filtered_rows(self):
        query = self.query_var.get().strip().lower()
        trigger_filter = self.trigger_var.get().strip()
        mail_filter = self.mail_var.get().strip()
        type_filter = self.type_var.get().strip()
        rows = []
        for alarm in reversed(self.app.alarm_history):
            model_label = str(alarm.get('model_type_label') or self.app.get_model_type_label(alarm.get('model_type'))).strip() or '-'
            if trigger_filter != 'Wszystkie' and str(alarm.get('trigger_mode', '')).strip() != trigger_filter:
                continue
            if mail_filter != 'Wszystkie' and self.get_mail_state(alarm) != mail_filter:
                continue
            if type_filter != 'Wszystkie' and model_label != type_filter:
                continue
            if query:
                haystack = ' | '.join([str(alarm.get('czas_alarmu', '')), str(alarm.get('data_modelu', '')), str(alarm.get('godzina_modelu', '')), str(alarm.get('source_file', '')), str(alarm.get('mail_status', '')), str(alarm.get('opis', '')), model_label]).lower()
                if query not in haystack:
                    continue
            rows.append(alarm)
            if len(rows) >= self.parse_limit():
                break
        return rows

    def clear_filters(self):
        self.query_var.set('')
        self.trigger_var.set('Wszystkie')
        self.mail_var.set('Wszystkie')
        self.type_var.set('Wszystkie')
        self.limit_var.set('500')
        self.refresh()

    def update_summary(self):
        if not self.current_rows:
            self.summary_label.config(text='Brak alarmow dla wybranych filtrow.', fg=self.app.get_palette()['warning'])
            return
        success = sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Sukces')
        errors = sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Blad')
        disabled = sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Wylaczone')
        unique_files = len({str(row.get('source_file', '')).strip() for row in self.current_rows if str(row.get('source_file', '')).strip()})
        last_alarm = self.current_rows[0].get('czas_alarmu', '-')
        avg_margin = statistics.mean(float(row.get('wartosc', 0.0) or 0.0) - float(row.get('prog', 0.0) or 0.0) for row in self.current_rows)
        lines = [
            f"Rekordy po filtrach: {len(self.current_rows)} | unikalne pliki: {unique_files} | ostatni alarm: {last_alarm}",
            f"Mail: sukces {success} | blad {errors} | wylaczone {disabled} | srednia odchylka od progu: {avg_margin:+,.2f} MVA*s",
        ]
        self.summary_label.config(text='\n'.join(lines), fg=self.app.get_palette()['accent'])

    def update_chart(self):
        pal = self.app.get_palette()
        self.fig_alarm.patch.set_facecolor(pal['bg'])
        self.ax_alarm_status.clear()
        self.ax_alarm_recent.clear()
        self.app.stylizuj_os(self.ax_alarm_status, grid=True, grid_axis='y')
        self.app.stylizuj_os(self.ax_alarm_recent, grid=True, grid_axis='y')
        if not self.current_rows:
            self.ax_alarm_status.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=self.ax_alarm_status.transAxes, color=pal['muted'])
            self.ax_alarm_recent.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=self.ax_alarm_recent.transAxes, color=pal['muted'])
            self.canvas_alarm.draw_idle()
            return
        status_labels = ['Sukces', 'Blad', 'Wylaczone']
        status_values = [sum(1 for row in self.current_rows if self.get_mail_state(row) == label) for label in status_labels]
        bars = self.ax_alarm_status.bar(status_labels, status_values, color=[pal['success'], pal['danger'], pal['muted']], edgecolor=pal['border'], linewidth=0.8, width=0.56)
        for bar, value in zip(bars, status_values):
            self.ax_alarm_status.annotate(f'{value}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', fontsize=8, color=pal['text'])
        self.ax_alarm_status.set_title('Struktura alarmow wg statusu wysylki')
        self.ax_alarm_status.set_ylabel('Liczba alarmow')
        daily_counts = {}
        for row in self.current_rows:
            day_key = str(row.get('czas_alarmu', '')).split(' ', 1)[0]
            daily_counts[day_key] = daily_counts.get(day_key, 0) + 1
        recent_days = sorted(daily_counts.keys())[-12:]
        recent_values = [daily_counts[day] for day in recent_days]
        x_vals = list(range(len(recent_days)))
        self.ax_alarm_recent.plot(x_vals, recent_values, color=pal['line_main'], linewidth=2.1, marker='o', markersize=4.5)
        if recent_days:
            self.ax_alarm_recent.fill_between(x_vals, recent_values, color=pal['accent'], alpha=0.12)
        self.ax_alarm_recent.set_title('Alarmy w czasie')
        self.ax_alarm_recent.set_ylabel('Liczba / dzien')
        self.ax_alarm_recent.set_xlabel('Dzien alarmu')
        self.ax_alarm_recent.set_xticks(x_vals)
        self.ax_alarm_recent.set_xticklabels(recent_days, rotation=28, ha='right', fontsize=8, color=pal['text'])
        self.canvas_alarm.draw_idle()

    def refresh(self):
        self.current_rows = self.get_filtered_rows()
        self.tree.delete(*self.tree.get_children())
        self.details_text.delete('1.0', tk.END)
        for idx, alarm in enumerate(self.current_rows):
            model_label = str(alarm.get('model_type_label') or self.app.get_model_type_label(alarm.get('model_type'))).strip() or '-'
            wartosc = float(alarm.get('wartosc', 0.0) or 0.0)
            prog = float(alarm.get('prog', 0.0) or 0.0)
            delta = wartosc - prog
            iid = str(idx)
            self.tree.insert('', tk.END, iid=iid, values=(alarm.get('czas_alarmu', ''), alarm.get('trigger_mode', ''), model_label, alarm.get('data_modelu', ''), alarm.get('godzina_modelu', ''), f'{wartosc:,.2f}', f'{prog:,.2f}', f'{delta:+,.2f}', alarm.get('source_file', ''), 'TAK' if alarm.get('mail_attempted') else 'NIE', alarm.get('mail_status', '')), tags=(self.get_alarm_tag(alarm),))
        self.update_summary()
        self.update_chart()
        if self.current_rows:
            self.tree.selection_set('0')
            self.show_details()

    def show_details(self):
        self.details_text.delete('1.0', tk.END)
        selection = self.tree.selection()
        if not selection:
            self.details_text.insert(tk.END, 'Wybierz alarm z listy, aby zobaczyc szczegoly.')
            return
        try:
            alarm = self.current_rows[int(selection[0])]
        except Exception:
            self.details_text.insert(tk.END, 'Nie udalo sie odczytac wybranego alarmu.')
            return
        wartosc = float(alarm.get('wartosc', 0.0) or 0.0)
        prog = float(alarm.get('prog', 0.0) or 0.0)
        delta = wartosc - prog
        model_label = str(alarm.get('model_type_label') or self.app.get_model_type_label(alarm.get('model_type'))).strip() or '-'
        lines = [
            f"Czas alarmu: {alarm.get('czas_alarmu', '-')}",
            f"Tryb alarmu: {alarm.get('trigger_mode', '-')}",
            f"Typ modelu: {model_label}",
            f"Data modelu: {alarm.get('data_modelu', '-')}",
            f"Godzina modelu: {alarm.get('godzina_modelu', '-')}",
            f"Inercja: {wartosc:,.2f} MVA*s",
            f"Prog: {prog:,.2f} MVA*s",
            f"Odchylka: {delta:+,.2f} MVA*s",
            f"Mail: {self.get_mail_state(alarm)}",
            f"Status wysylki: {alarm.get('mail_status', '-')}",
            f"Plik: {alarm.get('source_file', '-')}",
            '',
            'Opis:',
            str(alarm.get('opis', '-') or '-'),
        ]
        self.details_text.insert(tk.END, '\n'.join(lines))

    def copy_selected(self):
        selection = self.tree.selection()
        if not selection:
            messagebox.showwarning('Historia alarmow', 'Najpierw wybierz alarm z listy.')
            return
        text = self.details_text.get('1.0', tk.END).strip()
        if not text:
            return
        self.clipboard_clear()
        self.clipboard_append(text)
        messagebox.showinfo('Historia alarmow', 'Szczegoly alarmu skopiowano do schowka.')

    def export_csv(self):
        if not self.current_rows:
            messagebox.showwarning('CSV', 'Brak alarmow do eksportu.')
            return
        path = filedialog.asksaveasfilename(title='Eksport historii alarmow do CSV', defaultextension='.csv', initialfile='historia_alarmow.csv', filetypes=[('CSV', '*.csv')])
        if not path:
            return
        try:
            with open(path, 'w', encoding='utf-8', newline='') as f:
                writer = csv.writer(f, delimiter=';')
                writer.writerow(['czas_alarmu', 'trigger_mode', 'model_type', 'data_modelu', 'godzina_modelu', 'wartosc', 'prog', 'delta', 'source_file', 'mail_state', 'mail_status', 'opis'])
                for row in self.current_rows:
                    wartosc = float(row.get('wartosc', 0.0) or 0.0)
                    prog = float(row.get('prog', 0.0) or 0.0)
                    writer.writerow([row.get('czas_alarmu', ''), row.get('trigger_mode', ''), row.get('model_type_label') or self.app.get_model_type_label(row.get('model_type')), row.get('data_modelu', ''), row.get('godzina_modelu', ''), wartosc, prog, wartosc - prog, row.get('source_file', ''), self.get_mail_state(row), row.get('mail_status', ''), row.get('opis', '')])
            messagebox.showinfo('CSV', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('CSV', str(e))

    def export_xlsx(self):
        if not self.current_rows:
            messagebox.showwarning('Excel', 'Brak alarmow do eksportu.')
            return
        path = filedialog.asksaveasfilename(title='Eksport historii alarmow do Excel', defaultextension='.xlsx', initialfile='historia_alarmow.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = 'Alarmy'
            bold = Font(bold=True)
            fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            headers = ['Czas alarmu', 'Tryb', 'Typ modelu', 'Data modelu', 'Godzina modelu', 'Inercja [MVA*s]', 'Prog [MVA*s]', 'Odchylka [MVA*s]', 'Plik', 'Mail', 'Status wysylki', 'Opis']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = bold
                cell.fill = fill
            for row_idx, row in enumerate(self.current_rows, start=2):
                wartosc = float(row.get('wartosc', 0.0) or 0.0)
                prog = float(row.get('prog', 0.0) or 0.0)
                values = [row.get('czas_alarmu', ''), row.get('trigger_mode', ''), row.get('model_type_label') or self.app.get_model_type_label(row.get('model_type')), row.get('data_modelu', ''), row.get('godzina_modelu', ''), wartosc, prog, wartosc - prog, row.get('source_file', ''), self.get_mail_state(row), row.get('mail_status', ''), row.get('opis', '')]
                for col_idx, value in enumerate(values, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            for col, width in {'A': 21, 'B': 10, 'C': 14, 'D': 14, 'E': 14, 'F': 16, 'G': 15, 'H': 17, 'I': 32, 'J': 12, 'K': 26, 'L': 40}.items():
                ws.column_dimensions[col].width = width
            ws_summary = wb.create_sheet('Podsumowanie')
            ws_summary['A1'] = 'Stan maila'
            ws_summary['B1'] = 'Liczba'
            ws_summary['A1'].font = bold
            ws_summary['B1'].font = bold
            summary_counts = [('Sukces', sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Sukces')), ('Blad', sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Blad')), ('Wylaczone', sum(1 for row in self.current_rows if self.get_mail_state(row) == 'Wylaczone'))]
            for idx, (label, value) in enumerate(summary_counts, start=2):
                ws_summary.cell(row=idx, column=1, value=label)
                ws_summary.cell(row=idx, column=2, value=value)
            chart = BarChart()
            chart.title = 'Alarmy wg statusu wysylki'
            chart.y_axis.title = 'Liczba alarmow'
            chart.height = 7
            chart.width = 11
            chart.add_data(Reference(ws_summary, min_col=2, min_row=1, max_row=4), titles_from_data=True)
            chart.set_categories(Reference(ws_summary, min_col=1, min_row=2, max_row=4))
            ws_summary.add_chart(chart, 'D2')
            wb.save(path)
            messagebox.showinfo('Excel', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('Excel', str(e))

    def clear_history(self):
        if not self.app.alarm_history:
            return
        if not messagebox.askyesno('Historia alarmow', 'Czy na pewno wyczyscic cala historie alarmow?'):
            return
        self.app.alarm_history = []
        self.refresh()



@dataclass
class ArchiveSearchResult:
    id: int
    archived_at: str
    model_type: str
    model_date: str
    model_hour: str
    source_file: str
    total_inertia: float
    original_inertia: float
    demand: float
    connected_count: int
    data_json: str

class InertiaArchiveRepository:

    def __init__(self, db_path='~/.analiza_inercji_archive.db'):
        self.db_path = Path(db_path).expanduser()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        self.ensure_schema()

    def ensure_schema(self):
        self.conn.execute("""
        CREATE TABLE IF NOT EXISTS inertia_archive (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            archived_at TEXT NOT NULL,
            model_type TEXT NOT NULL,
            model_date TEXT,
            model_hour TEXT,
            source_file TEXT,
            total_inertia REAL NOT NULL,
            original_inertia REAL NOT NULL,
            demand REAL NOT NULL,
            connected_count INTEGER NOT NULL,
            data_json TEXT NOT NULL
        )
        """)
        for stmt in [
            'CREATE INDEX IF NOT EXISTS idx_inertia_archive_date_hour ON inertia_archive(model_date, model_hour)',
            'CREATE INDEX IF NOT EXISTS idx_inertia_archive_inertia ON inertia_archive(total_inertia)',
            'CREATE INDEX IF NOT EXISTS idx_inertia_archive_type_archived ON inertia_archive(model_type, archived_at)',
            'CREATE INDEX IF NOT EXISTS idx_inertia_archive_source_file ON inertia_archive(source_file)',
        ]:
            try:
                self.conn.execute(stmt)
            except sqlite3.OperationalError:
                pass
        self.conn.commit()

    def close(self):
        try:
            self.conn.close()
        except Exception:
            pass

    def insert_snapshot(self, payload):
        cur = self.conn.execute(
            """
            INSERT INTO inertia_archive (
                archived_at, model_type, model_date, model_hour, source_file,
                total_inertia, original_inertia, demand, connected_count, data_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload.get('archived_at', ''),
                payload.get('model_type', ''),
                payload.get('model_date', ''),
                payload.get('model_hour', ''),
                payload.get('source_file', ''),
                float(payload.get('total_inertia', 0.0)),
                float(payload.get('original_inertia', 0.0)),
                float(payload.get('demand', 0.0)),
                int(payload.get('connected_count', 0)),
                json.dumps(payload, ensure_ascii=False),
            )
        )
        self.conn.commit()
        return int(cur.lastrowid)

    def get_record(self, record_id):
        row = self.conn.execute('SELECT * FROM inertia_archive WHERE id=?', (int(record_id),)).fetchone()
        return dict(row) if row else None

    def fetch_records(self, filters=None, sort_by='archived_at', sort_desc=True):
        filters = filters or {}
        valid_sort = {
            'id': 'id',
            'archived_at': 'archived_at',
            'model_date': 'model_date',
            'model_hour': 'model_hour',
            'source_file': 'source_file',
            'total_inertia': 'total_inertia',
            'original_inertia': 'original_inertia',
            'demand': 'demand',
            'connected_count': 'connected_count',
        }
        sort_col = valid_sort.get(sort_by, 'archived_at')
        order = 'DESC' if sort_desc else 'ASC'
        date_text = (filters.get('model_date') or '').strip().lower()
        hour_text = (filters.get('model_hour') or '').strip().lower()
        file_text = (filters.get('source_file') or '').strip().lower()
        model_text = (filters.get('model_type') or '').strip().lower()
        expr = (filters.get('query') or '').strip()
        clauses = []
        params = []
        if date_text:
            clauses.append("LOWER(COALESCE(model_date, '')) LIKE ?")
            params.append(f'%{date_text}%')
        if hour_text:
            clauses.append("LOWER(COALESCE(model_hour, '')) LIKE ?")
            params.append(f'%{hour_text}%')
        if file_text:
            clauses.append("LOWER(COALESCE(source_file, '')) LIKE ?")
            params.append(f'%{file_text}%')
        if model_text:
            clauses.append("LOWER(COALESCE(model_type, '')) LIKE ?")
            params.append(f'%{model_text}%')
        sql = 'SELECT * FROM inertia_archive'
        if clauses:
            sql += ' WHERE ' + ' AND '.join(clauses)
        sql += f' ORDER BY {sort_col} {order}, id DESC'
        rows = self.conn.execute(sql, params).fetchall()
        results = []
        for row in rows:
            rec = dict(row)
            if expr and (not self._matches_expression(rec, expr)):
                continue
            results.append(rec)
        return results

    def fetch_history_rows(self, model_type=''):
        clauses = []
        params = []
        model_text = str(model_type or '').strip().lower()
        if model_text:
            clauses.append("LOWER(COALESCE(model_type, '')) = ?")
            params.append(model_text)
        sql = 'SELECT id, archived_at, model_type, model_date, model_hour, source_file, total_inertia, original_inertia, demand, connected_count FROM inertia_archive'
        if clauses:
            sql += ' WHERE ' + ' AND '.join(clauses)
        sql += ' ORDER BY model_date ASC, model_hour ASC, archived_at ASC, id ASC'
        rows = self.conn.execute(sql, params).fetchall()
        return [dict(row) for row in rows]

    def _matches_expression(self, rec, expr):
        expr = expr.strip()
        if not expr:
            return True
        aliases = {
            'inercja': rec.get('total_inertia', 0.0),
            'energia': rec.get('total_inertia', 0.0),
            'e_sys': rec.get('total_inertia', 0.0),
            'oryginalna': rec.get('original_inertia', 0.0),
            'zapotrzebowanie': rec.get('demand', 0.0),
            'load': rec.get('demand', 0.0),
            'liczba': rec.get('connected_count', 0),
            'connected_count': rec.get('connected_count', 0),
        }
        text_expr = expr
        for alias, value in aliases.items():
            text_expr = re.sub(rf'\b{re.escape(alias)}\b', str(value), text_expr, flags=re.IGNORECASE)
        if not re.fullmatch(r'[0-9\s<>!=().,+\-/*]+', text_expr):
            joined = ' | '.join(str(x or '') for x in (rec.get('model_date'), rec.get('model_hour'), rec.get('source_file'))).lower()
            return expr.lower() in joined
        try:
            return bool(eval(text_expr, {'__builtins__': {}}, {}))
        except Exception:
            return False

class HCoefficientRepository:

    DEFAULT_GROUP_TYPES = {101: 'JWCDc', 102: 'JWCDp', 201: 'JWCKc', 202: 'JWCKw', 203: 'JWCKp', 204: 'JWCKf', 205: 'JWCKpv', 301: 'MC', 302: 'PR', 303: 'MW', 304: 'FW', 305: 'PV'}
    DEFAULT_GROUP_H = 3.0

    def __init__(self, db_path='~/.analiza_inercji_h.db'):
        self.db_path = Path(db_path).expanduser()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        self.ensure_schema()
        self.seed_default_group_types()

    def ensure_schema(self):
        self.conn.execute("""
        CREATE TABLE IF NOT EXISTS h_coefficients (
            generator_name TEXT PRIMARY KEY,
            dyd_h REAL,
            effective_h REAL NOT NULL,
            is_manual INTEGER NOT NULL DEFAULT 0,
            source_file TEXT,
            imported_at TEXT,
            updated_at TEXT NOT NULL
        )
        """)
        self.conn.execute('CREATE INDEX IF NOT EXISTS idx_h_coefficients_manual ON h_coefficients(is_manual)')
        self.conn.execute("""
        CREATE TABLE IF NOT EXISTS h_group_defaults (
            group_code INTEGER PRIMARY KEY,
            group_label TEXT,
            effective_h REAL NOT NULL,
            updated_at TEXT NOT NULL
        )
        """)
        self.conn.commit()

    def close(self):
        try:
            self.conn.close()
        except Exception:
            pass

    def seed_default_group_types(self):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        changed = False
        for group_code, group_label in self.DEFAULT_GROUP_TYPES.items():
            row = self.conn.execute('SELECT group_code, group_label, effective_h FROM h_group_defaults WHERE group_code=?', (int(group_code),)).fetchone()
            if row is None:
                self.conn.execute(
                    'INSERT INTO h_group_defaults(group_code, group_label, effective_h, updated_at) VALUES (?, ?, ?, ?)',
                    (int(group_code), group_label, float(self.DEFAULT_GROUP_H), now)
                )
                changed = True
            else:
                current_label = str(row['group_label'] or '').strip()
                current_h = float(row['effective_h']) if row['effective_h'] is not None else None
                if current_label != group_label or current_h is None:
                    self.conn.execute(
                        'UPDATE h_group_defaults SET group_label=?, effective_h=COALESCE(effective_h, ?), updated_at=? WHERE group_code=?',
                        (group_label, float(self.DEFAULT_GROUP_H), now, int(group_code))
                    )
                    changed = True
        if changed:
            self.conn.commit()
    def upsert_from_dyd_map(self, h_map, source_file=''):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        for gen_name, h_value in (h_map or {}).items():
            try:
                h_value = float(h_value)
            except Exception:
                continue
            row = self.conn.execute('SELECT generator_name, is_manual, effective_h FROM h_coefficients WHERE generator_name=?', (gen_name,)).fetchone()
            if row is None:
                self.conn.execute(
                    'INSERT INTO h_coefficients(generator_name, dyd_h, effective_h, is_manual, source_file, imported_at, updated_at) VALUES (?, ?, ?, 0, ?, ?, ?)',
                    (gen_name, h_value, h_value, source_file, now, now)
                )
            else:
                is_manual = int(row['is_manual'])
                effective_h = float(row['effective_h']) if is_manual else h_value
                self.conn.execute(
                    'UPDATE h_coefficients SET dyd_h=?, effective_h=?, source_file=?, imported_at=?, updated_at=? WHERE generator_name=?',
                    (h_value, effective_h, source_file, now, now, gen_name)
                )
        self.conn.commit()

    def set_manual_value(self, generator_name, h_value):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        row = self.conn.execute('SELECT dyd_h, source_file, imported_at FROM h_coefficients WHERE generator_name=?', (generator_name,)).fetchone()
        if row is None:
            self.conn.execute(
                'INSERT INTO h_coefficients(generator_name, dyd_h, effective_h, is_manual, source_file, imported_at, updated_at) VALUES (?, NULL, ?, 1, ?, ?, ?)',
                (generator_name, float(h_value), '', '', now)
            )
        else:
            self.conn.execute(
                'UPDATE h_coefficients SET effective_h=?, is_manual=1, updated_at=? WHERE generator_name=?',
                (float(h_value), now, generator_name)
            )
        self.conn.commit()

    def restore_dyd_value(self, generator_name):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        row = self.conn.execute('SELECT dyd_h FROM h_coefficients WHERE generator_name=?', (generator_name,)).fetchone()
        if row is None or row['dyd_h'] is None:
            return False
        self.conn.execute(
            'UPDATE h_coefficients SET effective_h=dyd_h, is_manual=0, updated_at=? WHERE generator_name=?',
            (now, generator_name)
        )
        self.conn.commit()
        return True

    def fetch_all_rows(self, filter_text=''):
        rows = self.conn.execute('SELECT * FROM h_coefficients ORDER BY generator_name COLLATE NOCASE ASC').fetchall()
        filter_text = (filter_text or '').strip().lower()
        result = []
        for row in rows:
            rec = dict(row)
            if filter_text and filter_text not in (rec.get('generator_name') or '').lower():
                continue
            result.append(rec)
        return result

    def get_effective_map(self, only_names=None):
        rows = self.conn.execute('SELECT generator_name, effective_h FROM h_coefficients').fetchall()
        names = set(only_names) if only_names else None
        out = {}
        for row in rows:
            name = row['generator_name']
            if names is not None and name not in names:
                continue
            out[name] = float(row['effective_h'])
        return out

    def get_dyd_map(self, only_names=None):
        rows = self.conn.execute('SELECT generator_name, dyd_h FROM h_coefficients WHERE dyd_h IS NOT NULL').fetchall()
        names = set(only_names) if only_names else None
        out = {}
        for row in rows:
            name = row['generator_name']
            if names is not None and name not in names:
                continue
            out[name] = float(row['dyd_h'])
        return out

    def export_to_csv(self, path):
        rows = self.fetch_all_rows()
        with open(path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f, delimiter=';')
            writer.writerow(['generator_name', 'dyd_h', 'effective_h', 'is_manual', 'source_file', 'imported_at', 'updated_at'])
            for row in rows:
                writer.writerow([
                    row.get('generator_name', ''),
                    '' if row.get('dyd_h') is None else row.get('dyd_h'),
                    row.get('effective_h', ''),
                    int(row.get('is_manual') or 0),
                    row.get('source_file', ''),
                    row.get('imported_at', ''),
                    row.get('updated_at', ''),
                ])

    def export_to_xlsx(self, path):
        rows = self.fetch_all_rows()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Baza H'
        headers = ['generator_name', 'dyd_h', 'effective_h', 'is_manual', 'source_file', 'imported_at', 'updated_at']
        ws.append(headers)
        for row in rows:
            ws.append([
                row.get('generator_name', ''),
                None if row.get('dyd_h') is None else float(row.get('dyd_h')),
                None if row.get('effective_h') is None else float(row.get('effective_h')),
                int(row.get('is_manual') or 0),
                row.get('source_file', ''),
                row.get('imported_at', ''),
                row.get('updated_at', ''),
            ])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type='solid', fgColor='DCE6F1')
            cell.alignment = Alignment(horizontal='center')
        widths = {'A': 32, 'B': 12, 'C': 14, 'D': 10, 'E': 34, 'F': 20, 'G': 20}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        wb.save(path)

    def import_from_csv(self, path):
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f, delimiter=';')
            rows = list(reader)
        return self._import_rows(rows, source_hint=Path(path).name)

    def import_from_xlsx(self, path):
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows_raw = list(ws.iter_rows(values_only=True))
        if not rows_raw:
            return 0
        headers = [str(x).strip() if x is not None else '' for x in rows_raw[0]]
        rows = []
        for row in rows_raw[1:]:
            item = {}
            for idx, header in enumerate(headers):
                if header:
                    item[header] = row[idx] if idx < len(row) else None
            rows.append(item)
        return self._import_rows(rows, source_hint=Path(path).name)

    def _import_rows(self, rows, source_hint=''):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        imported = 0
        for rec in rows:
            name = str(rec.get('generator_name') or rec.get('Generator') or '').strip()
            if not name:
                continue
            def parse_float(value):
                if value in (None, ''):
                    return None
                try:
                    return float(str(value).replace(',', '.').strip())
                except Exception:
                    return None
            dyd_h = parse_float(rec.get('dyd_h'))
            eff_h = parse_float(rec.get('effective_h'))
            if eff_h is None:
                eff_h = dyd_h
            if eff_h is None or eff_h <= 0:
                continue
            manual_raw = rec.get('is_manual', 0)
            try:
                is_manual = int(float(str(manual_raw).strip() or '0'))
            except Exception:
                is_manual = 1 if str(manual_raw).strip().lower() in {'tak', 'true', 'manual', 'ręczny', 'reczny'} else 0
            source_file = str(rec.get('source_file') or source_hint or '').strip()
            imported_at = str(rec.get('imported_at') or now).strip()
            self.conn.execute(
                'INSERT INTO h_coefficients(generator_name, dyd_h, effective_h, is_manual, source_file, imported_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?) '
                'ON CONFLICT(generator_name) DO UPDATE SET dyd_h=excluded.dyd_h, effective_h=excluded.effective_h, is_manual=excluded.is_manual, source_file=excluded.source_file, imported_at=excluded.imported_at, updated_at=excluded.updated_at',
                (name, dyd_h, eff_h, is_manual, source_file, imported_at, now)
            )
            imported += 1
        self.conn.commit()
        return imported


    def set_group_default(self, group_code, group_label, h_value):
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        self.conn.execute(
            'INSERT INTO h_group_defaults(group_code, group_label, effective_h, updated_at) VALUES (?, ?, ?, ?) '
            'ON CONFLICT(group_code) DO UPDATE SET group_label=excluded.group_label, effective_h=excluded.effective_h, updated_at=excluded.updated_at',
            (int(group_code), str(group_label or '').strip(), float(h_value), now)
        )
        self.conn.commit()

    def delete_group_default(self, group_code):
        self.conn.execute('DELETE FROM h_group_defaults WHERE group_code=?', (int(group_code),))
        self.conn.commit()

    def fetch_group_default_rows(self):
        rows = self.conn.execute('SELECT group_code, group_label, effective_h, updated_at FROM h_group_defaults ORDER BY group_code ASC').fetchall()
        return [dict(r) for r in rows]

    def get_group_default_map(self):
        rows = self.conn.execute('SELECT group_code, effective_h FROM h_group_defaults').fetchall()
        return {int(r['group_code']): float(r['effective_h']) for r in rows}


class HDatabaseWindow(tk.Toplevel):

    def __init__(self, master, app):
        super().__init__(master)
        self.app = app
        self.repo = app.h_repo
        self.title('Baza H generatorów')
        init_managed_toplevel(self, 'h_database', 1180, 760)
        self.configure(bg=app.get_palette()['bg'])
        self.filter_var = tk.StringVar()
        self.selected_name_var = tk.StringVar()
        self.edit_value_var = tk.StringVar()
        self.group_code_var = tk.StringVar()
        self.group_label_var = tk.StringVar()
        self.group_h_var = tk.StringVar()
        self._wheel_widgets = []
        self._group_wheel_widgets = []
        self._build_ui()
        self._bind_local_mousewheel()
        self.refresh()

    def _build_ui(self):
        pal = self.app.get_palette()
        top = tk.Frame(self, bg=pal['bg'], padx=10, pady=10)
        top.pack(fill=tk.BOTH, expand=True)
        controls = tk.LabelFrame(top, text='Filtrowanie i edycja', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        controls.pack(fill=tk.X)
        tk.Label(controls, text='Filtr generatora:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=0, column=0, sticky='w')
        ent_filter = tk.Entry(controls, textvariable=self.filter_var, width=30, font=self.app.font_code_small)
        ent_filter.grid(row=0, column=1, sticky='ew', padx=(6, 8))
        self.app.stylizuj_entry(ent_filter)
        ent_filter.bind('<KeyRelease>', lambda e: self.refresh())
        self.app.utworz_przycisk(controls, 'Odśwież', self.refresh, role='primary').grid(row=0, column=2, padx=4)
        self.app.utworz_przycisk(controls, 'Wyczyść filtr', lambda: (self.filter_var.set(''), self.refresh()), role='neutral').grid(row=0, column=3, padx=4)
        tk.Label(controls, text='Generator:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=1, column=0, sticky='w', pady=(10, 0))
        ent_name = tk.Entry(controls, textvariable=self.selected_name_var, width=32, font=self.app.font_code_small)
        ent_name.grid(row=1, column=1, sticky='ew', padx=(6, 8), pady=(10, 0))
        self.app.stylizuj_entry(ent_name)
        tk.Label(controls, text='H efektywne [s]:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=1, column=2, sticky='w', pady=(10, 0))
        ent_val = tk.Entry(controls, textvariable=self.edit_value_var, width=12, font=self.app.font_code_small)
        ent_val.grid(row=1, column=3, sticky='w', padx=(6, 8), pady=(10, 0))
        self.app.stylizuj_entry(ent_val)
        ent_val.bind('<Return>', lambda e: self.save_manual())
        self.app.utworz_przycisk(controls, 'Zapisz do bazy', self.save_manual, role='success').grid(row=1, column=4, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Przywróć DYD', self.restore_dyd, role='neutral').grid(row=1, column=5, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Import CSV', self.import_csv, role='primary').grid(row=1, column=6, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Eksport CSV', self.export_csv, role='neutral').grid(row=1, column=7, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Import XLSX', self.import_xlsx, role='primary').grid(row=1, column=8, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Eksport XLSX', self.export_xlsx, role='neutral').grid(row=1, column=9, padx=4, pady=(10, 0))
        self.app.utworz_przycisk(controls, 'Zamknij', self.destroy, role='neutral').grid(row=1, column=10, padx=4, pady=(10, 0))
        controls.grid_columnconfigure(1, weight=1)

        group_fr = tk.LabelFrame(top, text='Domyślne H wg typu / grupy generatora', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        group_fr.pack(fill=tk.X, pady=(10, 0))
        self.group_frame = group_fr
        tk.Label(group_fr, text='Kod typu:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=0, column=0, sticky='w')
        ent_group_code = tk.Entry(group_fr, textvariable=self.group_code_var, width=10, font=self.app.font_code_small)
        ent_group_code.grid(row=0, column=1, sticky='w', padx=(6, 8))
        self.app.stylizuj_entry(ent_group_code)
        tk.Label(group_fr, text='Typ / opis:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=0, column=2, sticky='w')
        ent_group_label = tk.Entry(group_fr, textvariable=self.group_label_var, width=16, font=self.app.font_code_small)
        ent_group_label.grid(row=0, column=3, sticky='w', padx=(6, 8))
        self.app.stylizuj_entry(ent_group_label)
        tk.Label(group_fr, text='H domyślne [s]:', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).grid(row=0, column=4, sticky='w')
        ent_group_h = tk.Entry(group_fr, textvariable=self.group_h_var, width=10, font=self.app.font_code_small)
        ent_group_h.grid(row=0, column=5, sticky='w', padx=(6, 8))
        self.app.stylizuj_entry(ent_group_h)
        ent_group_h.bind('<Return>', lambda e: self.save_group_default())
        self.app.utworz_przycisk(group_fr, 'Zapisz typ', self.save_group_default, role='success').grid(row=0, column=6, padx=4)
        self.app.utworz_przycisk(group_fr, 'Usuń typ', self.delete_group_default, role='warning').grid(row=0, column=7, padx=4)
        tk.Label(group_fr, text='Typy 101..305 są domyślnie wpisane do bazy z H=3. Są używane jako fallback, gdy generator nie ma własnego H z DYD.', bg=pal['card'], fg=pal['muted'], font=self.app.font_ui_small).grid(row=1, column=0, columnspan=8, sticky='w', pady=(6, 6))
        self.group_tree = ttk.Treeview(group_fr, columns=('group_code', 'group_label', 'effective_h', 'updated_at'), show='headings', height=5, selectmode='browse')
        for col, header, width in [('group_code', 'Kod', 70), ('group_label', 'Typ', 140), ('effective_h', 'H [s]', 90), ('updated_at', 'Ostatnia zmiana', 150)]:
            self.group_tree.heading(col, text=header)
            self.group_tree.column(col, width=width, anchor='center' if col != 'group_label' else 'w')
        group_scroll = ttk.Scrollbar(group_fr, orient='vertical', command=self.group_tree.yview)
        self.group_tree.configure(yscrollcommand=group_scroll.set)
        self.group_tree.grid(row=2, column=0, columnspan=7, sticky='nsew', pady=(0, 4))
        group_scroll.grid(row=2, column=7, sticky='ns', pady=(0, 4))
        self.group_tree.bind('<<TreeviewSelect>>', self.on_group_select)
        group_fr.grid_columnconfigure(6, weight=1)

        table_fr = tk.LabelFrame(top, text='Trwała baza SQLite współczynników H', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        table_fr.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        cols = ('generator_name', 'dyd_h', 'effective_h', 'is_manual', 'source_file', 'imported_at', 'updated_at')
        self.tree = ttk.Treeview(table_fr, columns=cols, show='headings', height=22)
        headers = {'generator_name': 'Generator', 'dyd_h': 'H z DYD', 'effective_h': 'H do obliczeń', 'is_manual': 'Tryb', 'source_file': 'Plik źródłowy', 'imported_at': 'Import DYD', 'updated_at': 'Ostatnia zmiana'}
        widths = {'generator_name': 250, 'dyd_h': 100, 'effective_h': 110, 'is_manual': 90, 'source_file': 220, 'imported_at': 150, 'updated_at': 150}
        for col in cols:
            self.tree.heading(col, text=headers[col])
            self.tree.column(col, width=widths[col], anchor='w' if col in {'generator_name', 'source_file'} else 'center')
        ys = ttk.Scrollbar(table_fr, orient='vertical', command=self.tree.yview)
        xs = ttk.Scrollbar(table_fr, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ys.pack(side=tk.RIGHT, fill=tk.Y)
        xs.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree.bind('<<TreeviewSelect>>', self.on_select)
        self.tree.tag_configure('manual', background='#fff3cd' if not self.app.dark_mode_var.get() else '#3a3212', foreground='#e65100' if not self.app.dark_mode_var.get() else '#fcd34d')
        self._wheel_widgets = [self, self.tree, self.group_tree, table_fr, top, controls, self.group_frame]

    def refresh(self):
        rows = self.repo.fetch_all_rows(self.filter_var.get())
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in rows:
            tag = ('manual',) if int(row.get('is_manual') or 0) else ()
            self.tree.insert('', tk.END, iid=row['generator_name'], values=(
                row['generator_name'],
                '' if row.get('dyd_h') is None else f"{float(row['dyd_h']):.4f}",
                f"{float(row['effective_h']):.4f}",
                'Ręczny' if int(row.get('is_manual') or 0) else 'DYD',
                row.get('source_file', ''),
                row.get('imported_at', ''),
                row.get('updated_at', ''),
            ), tags=tag)
        if hasattr(self, 'group_tree'):
            for item in self.group_tree.get_children():
                self.group_tree.delete(item)
            for row in self.repo.fetch_group_default_rows():
                self.group_tree.insert('', tk.END, iid=str(row['group_code']), values=(row['group_code'], row.get('group_label', ''), f"{float(row['effective_h']):.4f}", row.get('updated_at', '')))

    def on_select(self, event=None):
        sel = self.tree.selection()
        if not sel:
            return
        name = sel[0]
        self.selected_name_var.set(name)
        vals = self.tree.item(name, 'values')
        if len(vals) >= 3:
            self.edit_value_var.set(str(vals[2]))

    def save_manual(self):
        name = self.selected_name_var.get().strip()
        if not name:
            messagebox.showwarning('Baza H', 'Wybierz generator albo wpisz jego nazwę.')
            return
        try:
            val = float(self.edit_value_var.get().strip().replace(',', '.'))
        except Exception:
            messagebox.showerror('Baza H', 'Podaj poprawną wartość liczbową H.')
            return
        if val < 0 or val >= 100:
            messagebox.showerror('Baza H', 'Wartość H musi należeć do zakresu 0 ≤ H < 100 s.')
            return
        self.repo.set_manual_value(name, val)
        self.app.reload_h_from_database()
        self.app.przelicz_wszystkie_modele_po_zmianie_h()
        self.refresh()
        try:
            self.tree.selection_set(name)
            self.tree.see(name)
        except Exception:
            pass

    def restore_dyd(self):
        name = self.selected_name_var.get().strip()
        if not name:
            messagebox.showwarning('Baza H', 'Wybierz generator.')
            return
        if not self.repo.restore_dyd_value(name):
            messagebox.showwarning('Baza H', 'Dla tego generatora brak wartości źródłowej DYD.')
            return
        self.app.reload_h_from_database()
        self.app.przelicz_wszystkie_modele_po_zmianie_h()
        self.refresh()
        try:
            self.tree.selection_set(name)
            self.tree.see(name)
        except Exception:
            pass

    def on_group_select(self, event=None):
        sel = self.group_tree.selection() if hasattr(self, 'group_tree') else ()
        if not sel:
            return
        iid = sel[0]
        vals = self.group_tree.item(iid, 'values')
        if len(vals) >= 3:
            self.group_code_var.set(str(vals[0]))
            self.group_label_var.set(str(vals[1]))
            self.group_h_var.set(str(vals[2]))

    def save_group_default(self):
        try:
            code = int(str(self.group_code_var.get()).strip())
        except Exception:
            messagebox.showerror('Baza H', 'Podaj poprawny kod typu generatora, np. 301 albo 305.')
            return
        label = str(self.group_label_var.get()).strip() or self.app.GENERATOR_GROUP_LABELS.get(code, '')
        try:
            h_val = float(str(self.group_h_var.get()).strip().replace(',', '.'))
        except Exception:
            messagebox.showerror('Baza H', 'Podaj poprawną wartość H dla typu generatora.')
            return
        if h_val < 0 or h_val >= 100:
            messagebox.showerror('Baza H', 'Wartość H musi należeć do zakresu 0 ≤ H < 100 s.')
            return
        self.repo.set_group_default(code, label, h_val)
        self.app.reload_h_from_database()
        self.app.przelicz_wszystkie_modele_po_zmianie_h()
        self.refresh()
        try:
            self.group_tree.selection_set(str(code))
            self.group_tree.see(str(code))
        except Exception:
            pass

    def delete_group_default(self):
        try:
            code = int(str(self.group_code_var.get()).strip())
        except Exception:
            messagebox.showwarning('Baza H', 'Wybierz typ generatora do usunięcia.')
            return
        self.repo.delete_group_default(code)
        self.app.reload_h_from_database()
        self.app.przelicz_wszystkie_modele_po_zmianie_h()
        self.refresh()

    def export_csv(self):
        path = filedialog.asksaveasfilename(title='Eksport bazy H do CSV', defaultextension='.csv', initialfile='baza_H_generatorow.csv', filetypes=[('CSV', '*.csv')])
        if not path:
            return
        try:
            self.repo.export_to_csv(path)
            messagebox.showinfo('Baza H', f'Wyeksportowano bazę H do pliku\n{path}')
        except Exception as e:
            messagebox.showerror('Baza H', str(e))

    def export_xlsx(self):
        path = filedialog.asksaveasfilename(title='Eksport bazy H do XLSX', defaultextension='.xlsx', initialfile='baza_H_generatorow.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            self.repo.export_to_xlsx(path)
            messagebox.showinfo('Baza H', f'Wyeksportowano bazę H do pliku\n{path}')
        except Exception as e:
            messagebox.showerror('Baza H', str(e))

    def import_csv(self):
        path = filedialog.askopenfilename(title='Import bazy H z CSV', filetypes=[('CSV', '*.csv')])
        if not path:
            return
        try:
            imported = self.repo.import_from_csv(path)
            self.app.reload_h_from_database()
            self.app.przelicz_wszystkie_modele_po_zmianie_h()
            self.refresh()
            messagebox.showinfo('Baza H', f'Zaimportowano rekordy: {imported}\nŹródło: {path}')
        except Exception as e:
            messagebox.showerror('Baza H', str(e))

    def import_xlsx(self):
        path = filedialog.askopenfilename(title='Import bazy H z XLSX', filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            imported = self.repo.import_from_xlsx(path)
            self.app.reload_h_from_database()
            self.app.przelicz_wszystkie_modele_po_zmianie_h()
            self.refresh()
            messagebox.showinfo('Baza H', f'Zaimportowano rekordy: {imported}\nŹródło: {path}')
        except Exception as e:
            messagebox.showerror('Baza H', str(e))

    def _bind_local_mousewheel(self):
        for widget in self._wheel_widgets:
            try:
                widget.bind('<MouseWheel>', self._on_mousewheel, add='+')
                widget.bind('<Shift-MouseWheel>', self._on_shift_mousewheel, add='+')
                widget.bind('<Button-4>', self._on_mousewheel_linux, add='+')
                widget.bind('<Button-5>', self._on_mousewheel_linux, add='+')
            except Exception:
                pass

    def _wheel_target(self, x_root, y_root):
        widget = self.winfo_containing(x_root, y_root)
        if widget is None:
            return self.tree
        current = widget
        while current is not None:
            if current is self.group_tree:
                return self.group_tree
            if current is self.tree:
                return self.tree
            if hasattr(self, 'group_frame') and current is self.group_frame:
                return self.group_tree
            current = getattr(current, 'master', None)
        return self.tree

    def _on_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if getattr(event, 'delta', 0) else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        self._wheel_target(event.x_root, event.y_root).yview_scroll(units, 'units')
        return 'break'

    def _on_shift_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if getattr(event, 'delta', 0) else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        self.tree.xview_scroll(units, 'units')
        return 'break'

    def _on_mousewheel_linux(self, event):
        step = -1 if event.num == 4 else 1
        self._wheel_target(event.x_root, event.y_root).yview_scroll(step, 'units')
        return 'break'


class ArchiveBrowserWindow(tk.Toplevel):

    def __init__(self, master, app):
        super().__init__(master)
        self.app = app
        self.repo = app.archive_repo
        self.sort_by = 'archived_at'
        self.sort_desc = True
        self.title('Archiwum inercji')
        init_managed_toplevel(self, 'archive_browser', 1680, 960)
        self.configure(bg=app.get_palette()['bg'])
        self.date_var = tk.StringVar()
        self.model_date_from_var = tk.StringVar()
        self.model_date_to_var = tk.StringVar()
        self.archived_from_var = tk.StringVar()
        self.archived_to_var = tk.StringVar()
        self.hour_var = tk.StringVar()
        self.file_var = tk.StringVar()
        self.query_var = tk.StringVar()
        self.inertia_min_var = tk.StringVar()
        self.inertia_max_var = tk.StringVar()
        self.demand_min_var = tk.StringVar()
        self.demand_max_var = tk.StringVar()
        self.connected_min_var = tk.StringVar()
        self.connected_max_var = tk.StringVar()
        self.type_options = {'Wszystkie': '', '24h': 'manual', 'Reczny': 'manual_batch', 'Biezace': 'auto'}
        self.type_var = tk.StringVar(value='Wszystkie')
        self.current_rows = []
        self._plot_annotation = None
        self._plot_points = []
        self._build_ui()
        self._bind_local_mousewheel()
        self.refresh()

    def _build_ui(self):
        pal = self.app.get_palette()
        top = tk.Frame(self, bg=pal['bg'], padx=10, pady=10)
        top.pack(fill=tk.X)
        filters = tk.LabelFrame(top, text='Filtry archiwum', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        filters.pack(fill=tk.X)
        def add_entry(row, col, label, variable, width=18):
            cell = tk.Frame(filters, bg=pal['card'])
            cell.grid(row=row, column=col, padx=5, pady=2, sticky='ew')
            tk.Label(cell, text=label, bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
            ent = tk.Entry(cell, textvariable=variable, width=width, font=self.app.font_code_small)
            ent.pack(fill=tk.X)
            self.app.stylizuj_entry(ent)
            return ent

        def add_combo(row, col, label, variable, values):
            cell = tk.Frame(filters, bg=pal['card'])
            cell.grid(row=row, column=col, padx=5, pady=2, sticky='ew')
            tk.Label(cell, text=label, bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
            combo = ttk.Combobox(cell, textvariable=variable, values=list(values), state='readonly')
            combo.pack(fill=tk.X)
            return combo

        def add_date_entry(row, col, label, variable):
            cell = tk.Frame(filters, bg=pal['card'])
            cell.grid(row=row, column=col, padx=5, pady=2, sticky='ew')
            tk.Label(cell, text=label, bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
            row_box = tk.Frame(cell, bg=pal['card'])
            row_box.pack(fill=tk.X)
            ent = tk.Entry(row_box, textvariable=variable, font=self.app.font_code_small)
            ent.pack(side=tk.LEFT, fill=tk.X, expand=True)
            self.app.stylizuj_entry(ent)
            btn = self.app.utworz_przycisk(row_box, 'Kalendarz', lambda v=variable: self.open_date_picker(v), role='neutral')
            btn.pack(side=tk.LEFT, padx=(6, 0))
            return ent

        add_entry(0, 0, 'Szukaj / zapytanie', self.query_var, width=24)
        add_entry(0, 1, 'Nazwa modelu / pliku', self.file_var, width=24)
        add_entry(0, 2, 'Godzina modelu', self.hour_var, width=12)
        self.combo_type = add_combo(0, 3, 'Typ modelu', self.type_var, self.type_options.keys())

        add_date_entry(1, 0, 'Data modelu od', self.model_date_from_var)
        add_date_entry(1, 1, 'Data modelu do', self.model_date_to_var)
        add_date_entry(1, 2, 'Archiwizowano od', self.archived_from_var)
        add_date_entry(1, 3, 'Archiwizowano do', self.archived_to_var)

        add_entry(2, 0, 'Inercja min [MVA*s]', self.inertia_min_var, width=14)
        add_entry(2, 1, 'Inercja max [MVA*s]', self.inertia_max_var, width=14)
        add_entry(2, 2, 'Load min [MW]', self.demand_min_var, width=14)
        add_entry(2, 3, 'Load max [MW]', self.demand_max_var, width=14)

        add_entry(3, 0, 'Liczba gen min', self.connected_min_var, width=12)
        add_entry(3, 1, 'Liczba gen max', self.connected_max_var, width=12)
        self.label_filter_hint = tk.Label(filters, text='Zmien filtry i kliknij Odswiez. Daty mozesz tez wpisac recznie, np. 02.04.2026 albo 2026-04-02.', bg=pal['card'], fg=pal['muted'], anchor='w', justify='left', font=self.app.font_ui_bold_small)
        self.label_filter_hint.grid(row=3, column=2, columnspan=2, sticky='ew', padx=5, pady=(8, 2))
        for i in range(4):
            filters.grid_columnconfigure(i, weight=1)
        btns = tk.Frame(filters, bg=pal['card'])
        btns.grid(row=4, column=0, columnspan=4, sticky='w', pady=(8, 0))
        self.app.utworz_przycisk(btns, 'Odswiez', self.refresh, role='primary').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Wyczysc filtry', self.clear_filters, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Wczytaj zaznaczony rekord', self.restore_selected, role='success').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport CSV', self.export_csv, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport Excel', self.export_xlsx, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport wykresu', self.export_chart, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Zamknij', self.destroy, role='neutral').pack(side=tk.LEFT, padx=3)

        frame_summary = tk.LabelFrame(top, text='Podsumowanie filtrow', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        frame_summary.pack(fill=tk.X, pady=(10, 0))
        self.summary_label = tk.Label(frame_summary, text='Brak danych archiwalnych', justify='left', anchor='w', bg=pal['card'], fg=pal['accent'], font=self.app.font_code_small)
        self.summary_label.pack(fill=tk.X)

        main = tk.Frame(self, bg=pal['bg'], padx=10, pady=0)
        main.pack(fill=tk.BOTH, expand=True)
        left = tk.Frame(main, bg=pal['bg'])
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        right = tk.Frame(main, bg=pal['bg'])
        right.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(10, 0))

        frame_tree = tk.LabelFrame(left, text='Rekordy archiwum', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=6, pady=6)
        frame_tree.pack(fill=tk.BOTH, expand=True)
        cols = ('id', 'archived_at', 'model_date', 'model_hour', 'model_type', 'total_inertia', 'demand', 'connected_count', 'source_file')
        self.tree = ttk.Treeview(frame_tree, columns=cols, show='headings', height=18)
        headers = {
            'id': 'ID', 'archived_at': 'Archiwizowano', 'model_date': 'Data', 'model_hour': 'Godzina', 'model_type': 'Typ',
            'total_inertia': 'Inercja [MVA·s]', 'demand': 'Load [MW]', 'connected_count': 'L. gen', 'source_file': 'Plik'
        }
        widths = {'id': 60, 'archived_at': 150, 'model_date': 110, 'model_hour': 90, 'model_type': 70, 'total_inertia': 125, 'demand': 105, 'connected_count': 70, 'source_file': 300}
        anchors = {'id': 'center', 'archived_at': 'center', 'model_date': 'center', 'model_hour': 'center', 'model_type': 'center', 'total_inertia': 'e', 'demand': 'e', 'connected_count': 'e', 'source_file': 'w'}
        for col in cols:
            self.tree.heading(col, text=headers[col], command=lambda c=col: self.change_sort(c))
            self.tree.column(col, width=widths[col], anchor=anchors.get(col, 'w'))
        ys = ttk.Scrollbar(frame_tree, orient='vertical', command=self.tree.yview)
        xs = ttk.Scrollbar(frame_tree, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ys.pack(side=tk.RIGHT, fill=tk.Y)
        xs.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree.bind('<<TreeviewSelect>>', lambda e: self.show_details())
        self.tree.bind('<Double-1>', lambda e: self.restore_selected())

        frame_plot = tk.LabelFrame(right, text='Historia inercji systemu dla przefiltrowanych rekordów', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=6, pady=6)
        frame_plot.pack(fill=tk.BOTH, expand=True)
        self.fig_hist, self.ax_hist = plt.subplots(figsize=(11.8, 4.8))
        self.fig_hist.subplots_adjust(left=0.12, right=0.985, top=0.90, bottom=0.17)
        self.canvas_hist = FigureCanvasTkAgg(self.fig_hist, master=frame_plot)
        self.canvas_hist.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.canvas_hist.mpl_connect('motion_notify_event', self.on_plot_hover)

        frame_rank = tk.LabelFrame(right, text='Ranking generatorów w archiwum', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=6, pady=6)
        frame_rank.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        self.rank_text = scrolledtext.ScrolledText(frame_rank, wrap=tk.WORD, height=12, font=self.app.font_code_small)
        self.rank_text.pack(fill=tk.BOTH, expand=True)
        self.rank_text.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'])

        frame_details = tk.LabelFrame(right, text='Szczegóły rekordu', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=6, pady=6)
        frame_details.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        self.details_text = scrolledtext.ScrolledText(frame_details, wrap=tk.WORD, height=20, font=self.app.font_code)
        self.details_text.pack(fill=tk.BOTH, expand=True)
        self.details_text.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'])

    def _bind_local_mousewheel(self):
        widgets = [self, self.tree, self.rank_text, self.details_text, self.canvas_hist.get_tk_widget()]
        for widget in widgets:
            widget.bind('<Enter>', self._activate_mousewheel_scope, add='+')
            widget.bind('<Leave>', self._deactivate_mousewheel_scope, add='+')
        self.bind('<MouseWheel>', self._on_mousewheel, add='+')
        self.bind('<Shift-MouseWheel>', self._on_shift_mousewheel, add='+')
        self.bind('<Button-4>', self._on_mousewheel_linux, add='+')
        self.bind('<Button-5>', self._on_mousewheel_linux, add='+')

    def _activate_mousewheel_scope(self, event=None):
        try:
            self.grab_set()
        except Exception:
            pass

    def _deactivate_mousewheel_scope(self, event=None):
        try:
            if self.grab_current() == self:
                self.grab_release()
        except Exception:
            pass

    def _wheel_target(self, x_root, y_root):
        widget = self.winfo_containing(x_root, y_root)
        if widget is None:
            return self.tree
        current = widget
        while current is not None:
            if current in (self.tree, self.rank_text, self.details_text):
                return current
            current = getattr(current, 'master', None)
        return self.tree

    def _scroll_widget(self, widget, units, horizontal=False):
        try:
            if horizontal:
                widget.xview_scroll(units, 'units')
            else:
                widget.yview_scroll(units, 'units')
            return 'break'
        except Exception:
            return 'break'

    def _on_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if getattr(event, 'delta', 0) else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        return self._scroll_widget(self._wheel_target(event.x_root, event.y_root), units)

    def _on_shift_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if getattr(event, 'delta', 0) else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        return self._scroll_widget(self.tree, units, horizontal=True)

    def _on_mousewheel_linux(self, event):
        step = -1 if event.num == 4 else 1
        return self._scroll_widget(self._wheel_target(event.x_root, event.y_root), step)

    def open_date_picker(self, variable):
        current_value = str(variable.get() or '').strip()
        initial_date = self.app.parse_model_date_text(current_value)
        if initial_date is None:
            parsed_dt = self.parse_archive_datetime_filter(current_value)
            initial_date = parsed_dt.date() if parsed_dt is not None else None
        picker = DatePickerWindow(self, initial_date=initial_date, on_select=lambda value, var=variable: var.set(value))
        try:
            picker.lift()
            picker.focus_force()
        except Exception:
            pass

    def get_selected_type_key(self):
        return self.type_options.get(self.type_var.get().strip(), '')

    def parse_float_filter(self, raw):
        txt = str(raw or '').strip().replace(',', '.')
        if not txt:
            return None
        try:
            return float(txt)
        except Exception:
            return None

    def parse_int_filter(self, raw):
        txt = str(raw or '').strip()
        if not txt:
            return None
        try:
            return int(float(txt.replace(',', '.')))
        except Exception:
            return None

    def parse_archive_datetime_filter(self, raw):
        txt = str(raw or '').strip()
        if not txt:
            return None
        for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%Y/%m/%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(txt, fmt)
            except Exception:
                pass
        try:
            return datetime.strptime(txt, '%Y-%m-%d %H:%M:%S')
        except Exception:
            return None

    def row_matches_advanced_filters(self, row):
        model_date = self.app.parse_model_date_text(row.get('model_date', ''))
        model_date_from = self.app.parse_model_date_text(self.model_date_from_var.get())
        model_date_to = self.app.parse_model_date_text(self.model_date_to_var.get())
        if model_date_from is not None and (model_date is None or model_date < model_date_from):
            return False
        if model_date_to is not None and (model_date is None or model_date > model_date_to):
            return False
        archived_dt = None
        try:
            archived_dt = datetime.strptime(str(row.get('archived_at', '')).strip(), '%Y-%m-%d %H:%M:%S')
        except Exception:
            archived_dt = None
        archived_from = self.parse_archive_datetime_filter(self.archived_from_var.get())
        archived_to = self.parse_archive_datetime_filter(self.archived_to_var.get())
        if archived_from is not None and (archived_dt is None or archived_dt < archived_from):
            return False
        if archived_to is not None and (archived_dt is None or archived_dt.date() > archived_to.date()):
            return False
        inertia_min = self.parse_float_filter(self.inertia_min_var.get())
        inertia_max = self.parse_float_filter(self.inertia_max_var.get())
        demand_min = self.parse_float_filter(self.demand_min_var.get())
        demand_max = self.parse_float_filter(self.demand_max_var.get())
        connected_min = self.parse_int_filter(self.connected_min_var.get())
        connected_max = self.parse_int_filter(self.connected_max_var.get())
        total_inertia = float(row.get('total_inertia', 0.0) or 0.0)
        demand = float(row.get('demand', 0.0) or 0.0)
        connected = int(row.get('connected_count', 0) or 0)
        if inertia_min is not None and total_inertia < inertia_min:
            return False
        if inertia_max is not None and total_inertia > inertia_max:
            return False
        if demand_min is not None and demand < demand_min:
            return False
        if demand_max is not None and demand > demand_max:
            return False
        if connected_min is not None and connected < connected_min:
            return False
        if connected_max is not None and connected > connected_max:
            return False
        return True

    def update_summary(self):
        pal = self.app.get_palette()
        if not self.current_rows:
            self.summary_label.config(text='Brak rekordow dla zadanych filtrow.', fg=pal['warning'])
            return
        inertias = [float(row.get('total_inertia', 0.0) or 0.0) for row in self.current_rows]
        demands = [float(row.get('demand', 0.0) or 0.0) for row in self.current_rows]
        dates = [self.app.parse_model_date_text(row.get('model_date', '')) for row in self.current_rows]
        dates = [date for date in dates if date is not None]
        lines = [
            f"Rekordy po filtrach: {len(self.current_rows)} | inercja min/sr/max: {min(inertias):,.2f} / {statistics.mean(inertias):,.2f} / {max(inertias):,.2f} MVA*s",
            f"Load min/sr/max: {min(demands):,.2f} / {statistics.mean(demands):,.2f} / {max(demands):,.2f} MW | srednia liczba generatorow: {statistics.mean(int(row.get('connected_count', 0) or 0) for row in self.current_rows):.1f}",
        ]
        if dates:
            lines.append(f"Zakres dat modeli: {min(dates).strftime('%d.%m.%Y')} - {max(dates).strftime('%d.%m.%Y')}")
        self.summary_label.config(text='\n'.join(lines), fg=pal['accent'])

    def clear_filters(self):
        self.date_var.set('')
        self.model_date_from_var.set('')
        self.model_date_to_var.set('')
        self.archived_from_var.set('')
        self.archived_to_var.set('')
        self.hour_var.set('')
        self.file_var.set('')
        self.query_var.set('')
        self.inertia_min_var.set('')
        self.inertia_max_var.set('')
        self.demand_min_var.set('')
        self.demand_max_var.set('')
        self.connected_min_var.set('')
        self.connected_max_var.set('')
        self.type_var.set('Wszystkie')
        self.refresh()

    def change_sort(self, col):
        if self.sort_by == col:
            self.sort_desc = not self.sort_desc
        else:
            self.sort_by = col
            self.sort_desc = True
        self.refresh()

    def get_filters(self):
        return {
            'model_date': self.date_var.get(),
            'model_hour': self.hour_var.get(),
            'source_file': self.file_var.get(),
            'query': self.query_var.get(),
            'model_type': self.get_selected_type_key(),
        }

    def refresh(self):
        fetched_rows = self.repo.fetch_records(self.get_filters(), sort_by=self.sort_by, sort_desc=self.sort_desc)
        self.current_rows = [row for row in fetched_rows if self.row_matches_advanced_filters(row)]
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in self.current_rows:
            model_type_label = self.app.get_model_type_label(row['model_type']) if row.get('model_type') in {'manual', 'manual_batch', 'auto'} else row.get('model_type', '')
            self.tree.insert('', tk.END, iid=str(row['id']), values=(
                row['id'], row['archived_at'], row['model_date'], row['model_hour'], model_type_label,
                f"{row['total_inertia']:,.2f}", f"{row['demand']:,.2f}", row['connected_count'], row['source_file']
            ))
        self.update_summary()
        self.update_history_plot()
        self.update_archive_ranking()
        self.show_details()

    def update_history_plot(self):
        pal = self.app.get_palette()
        self.fig_hist.patch.set_facecolor(pal['bg'])
        self.ax_hist.clear()
        self.ax_hist.set_title('Historia inercji systemu')
        self.ax_hist.set_xlabel('Kolejność rekordów')
        self.ax_hist.set_ylabel('Inercja [MVA·s]')
        self.app.stylizuj_os(self.ax_hist, grid=True, grid_axis='both')
        self._plot_points = []
        if self._plot_annotation is not None:
            try:
                self._plot_annotation.remove()
            except Exception:
                pass
            self._plot_annotation = None
        if not self.current_rows:
            self.ax_hist.text(0.5, 0.5, 'Brak rekordów po filtrach', ha='center', va='center', transform=self.ax_hist.transAxes, color=pal['muted'])
        else:
            rows = list(reversed(self.current_rows[-250:]))
            xs = list(range(1, len(rows) + 1))
            ys = [row['total_inertia'] for row in rows]
            line, = self.ax_hist.plot(xs, ys, 'o-', linewidth=2.2, markersize=5, color=pal['line_main'])
            self._plot_points = list(zip(xs, ys, rows))
            if ys:
                max_idx = ys.index(max(ys))
                min_idx = ys.index(min(ys))
                self.ax_hist.plot(xs[max_idx], ys[max_idx], 'o', markersize=10, color=pal['success'])
                self.ax_hist.plot(xs[min_idx], ys[min_idx], 'o', markersize=10, color=pal['danger'])
                ymin = min(ys)
                ymax = max(ys)
                pad = max((ymax - ymin) * 0.12, max(ymax, 1.0) * 0.08, 100.0)
                self.ax_hist.set_ylim(max(0, ymin - pad), ymax + pad)
            self.ax_hist.tick_params(axis='y', pad=6)
            self._plot_annotation = self.ax_hist.annotate('', xy=(0, 0), xytext=(14, 14), textcoords='offset points', bbox=dict(boxstyle='round,pad=0.25', fc=pal['card'], ec=pal['border'], alpha=0.95), color=pal['text'])
            self._plot_annotation.set_visible(False)
        self.canvas_hist.draw_idle()

    def on_plot_hover(self, event):
        if self._plot_annotation is None or event.inaxes != self.ax_hist or not self._plot_points:
            if self._plot_annotation is not None and self._plot_annotation.get_visible():
                self._plot_annotation.set_visible(False)
                self.canvas_hist.draw_idle()
            return
        nearest = None
        min_dx = float('inf')
        for x, y, row in self._plot_points:
            dx = abs(event.xdata - x) if event.xdata is not None else float('inf')
            if dx < min_dx:
                min_dx = dx
                nearest = (x, y, row)
        if nearest is None or min_dx > 0.6:
            if self._plot_annotation.get_visible():
                self._plot_annotation.set_visible(False)
                self.canvas_hist.draw_idle()
            return
        x, y, row = nearest
        self._plot_annotation.xy = (x, y)
        self._plot_annotation.set_text(f"{row['model_date']} {row['model_hour']}\nInercja: {y:,.2f} MVA·s")
        self._plot_annotation.set_visible(True)
        self.canvas_hist.draw_idle()

    def update_archive_ranking(self):
        self.rank_text.delete('1.0', tk.END)
        aggregate = {}
        counts = {}
        for row in self.current_rows[:500]:
            payload = json.loads(row['data_json'])
            for gen in payload.get('generators', []):
                if gen.get('status') != 1:
                    continue
                name = gen.get('reg_name', '')
                ei = float(gen.get('Ei', 0.0))
                aggregate[name] = aggregate.get(name, 0.0) + ei
                counts[name] = counts.get(name, 0) + 1
        items = sorted(aggregate.items(), key=lambda x: x[1], reverse=True)[:20]
        if not items:
            self.rank_text.insert(tk.END, 'Brak danych rankingowych dla aktualnych filtrów.')
            return
        lines = ['TOP generatorów wg skumulowanego wkładu w archiwum\n', '-' * 90 + '\n']
        for idx, (name, total_ei) in enumerate(items, 1):
            lines.append(f'{idx:>2}. {name:<28} suma Ei={total_ei:>14,.2f} MVA·s | wystąpienia={counts.get(name, 0):>4}')
        self.rank_text.insert(tk.END, '\n'.join(lines))

    def show_details(self):
        self.details_text.delete('1.0', tk.END)
        sel = self.tree.selection()
        if not sel:
            self.details_text.insert(tk.END, 'Wybierz rekord z tabeli archiwum.')
            return
        row = self.repo.get_record(int(sel[0]))
        if not row:
            self.details_text.insert(tk.END, 'Nie znaleziono rekordu.')
            return
        payload = json.loads(row['data_json'])
        rocof = float(payload.get('rocof', 0.0) or 0.0)
        largest_loss = float(payload.get('largest_loss', 0.0) or 0.0)
        lines = [
            f"ID: {row['id']}",
            f"Archiwizowano: {row['archived_at']}",
            f"Typ modelu: {self.app.get_model_type_label(row['model_type']) if row.get('model_type') in {'manual', 'manual_batch', 'auto'} else row['model_type']}",
            f"Data modelu: {row['model_date']}",
            f"Godzina modelu: {row['model_hour']}",
            f"Plik źródłowy: {row['source_file']}",
            f"Inercja: {row['total_inertia']:,.2f} MVA·s",
            f"Inercja oryginalna: {row['original_inertia']:,.2f} MVA·s",
            f"Zapotrzebowanie: {row['demand']:,.2f} MW",
            f"Liczba aktywnych generatorów: {row['connected_count']}",
            f"ROCOF: {rocof:,.4f} Hz/s | krytyczne dP: {largest_loss:,.2f} MW",
            '',
            'TOP 15 generatorow w rekordzie:',
        ]
        gens = payload.get('generators', [])
        ranked = sorted(((g.get('reg_name', ''), float(g.get('Ei', 0.0)), g.get('status', 0), g.get('group_label', '-')) for g in gens), key=lambda x: x[1], reverse=True)[:15]
        for idx, (name, ei, status, group_label) in enumerate(ranked, 1):
            lines.append(f'{idx:>2}. {name:<32} Ei={ei:>14,.2f} | status={status} | grupa={group_label}')
        self.details_text.insert(tk.END, '\n'.join(lines))

    def restore_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning('Archiwum', 'Najpierw wybierz rekord archiwum.')
            return
        row = self.repo.get_record(int(sel[0]))
        if not row:
            messagebox.showerror('Archiwum', 'Nie znaleziono rekordu w archiwum.')
            return
        payload = json.loads(row['data_json'])
        payload['id'] = row['id']
        self.app.restore_from_archive_payload(payload)
        messagebox.showinfo('Archiwum', 'Rekord został odtworzony i wczytany do aplikacji.')

    def export_csv(self):
        if not self.current_rows:
            messagebox.showwarning('CSV', 'Brak rekordów do eksportu.')
            return
        path = filedialog.asksaveasfilename(title='Eksport archiwum do CSV', defaultextension='.csv', initialfile='archiwum_inercji.csv', filetypes=[('CSV', '*.csv')])
        if not path:
            return
        try:
            with open(path, 'w', encoding='utf-8') as f:
                f.write('id;archived_at;model_type;model_date;model_hour;source_file;total_inertia;original_inertia;demand;connected_count\n')
                for row in self.current_rows:
                    model_type_label = self.app.get_model_type_label(row['model_type']) if row.get('model_type') in {'manual', 'manual_batch', 'auto'} else row.get('model_type', '')
                    vals = [row['id'], row['archived_at'], model_type_label, row['model_date'], row['model_hour'], row['source_file'], row['total_inertia'], row['original_inertia'], row['demand'], row['connected_count']]
                    f.write(';'.join(str(v) for v in vals) + '\n')
            messagebox.showinfo('CSV', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('CSV', str(e))

    def export_xlsx(self):
        if not self.current_rows:
            messagebox.showwarning('Excel', 'Brak rekordow do eksportu.')
            return
        path = filedialog.asksaveasfilename(title='Eksport archiwum do Excel', defaultextension='.xlsx', initialfile='archiwum_inercji.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = 'Archiwum'
            bold = Font(bold=True)
            fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            headers = ['ID', 'Archiwizowano', 'Typ modelu', 'Data modelu', 'Godzina modelu', 'Plik', 'Inercja [MVA*s]', 'Inercja org. [MVA*s]', 'Load [MW]', 'Liczba gen']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = bold
                cell.fill = fill
            for row_idx, row in enumerate(self.current_rows, start=2):
                model_type_label = self.app.get_model_type_label(row['model_type']) if row.get('model_type') in {'manual', 'manual_batch', 'auto'} else row.get('model_type', '')
                values = [row.get('id'), row.get('archived_at', ''), model_type_label, row.get('model_date', ''), row.get('model_hour', ''), row.get('source_file', ''), float(row.get('total_inertia', 0.0) or 0.0), float(row.get('original_inertia', 0.0) or 0.0), float(row.get('demand', 0.0) or 0.0), int(row.get('connected_count', 0) or 0)]
                for col_idx, value in enumerate(values, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            for col, width in {'A': 8, 'B': 22, 'C': 12, 'D': 14, 'E': 12, 'F': 34, 'G': 17, 'H': 17, 'I': 14, 'J': 10}.items():
                ws.column_dimensions[col].width = width
            ws_summary = wb.create_sheet('Podsumowanie')
            ws_summary['A1'] = 'Rekordy po filtrach'
            ws_summary['B1'] = len(self.current_rows)
            ws_summary['A2'] = 'Srednia inercja [MVA*s]'
            ws_summary['B2'] = statistics.mean(float(row.get('total_inertia', 0.0) or 0.0) for row in self.current_rows)
            ws_summary['A3'] = 'Sredni load [MW]'
            ws_summary['B3'] = statistics.mean(float(row.get('demand', 0.0) or 0.0) for row in self.current_rows)
            ws_summary['A1'].font = bold
            ws_summary['A2'].font = bold
            ws_summary['A3'].font = bold
            chart = LineChart()
            chart.title = 'Historia inercji'
            chart.y_axis.title = 'Inercja [MVA*s]'
            chart.x_axis.title = 'Rekord'
            chart.height = 8
            chart.width = 16
            max_row = max(2, len(self.current_rows) + 1)
            chart.add_data(Reference(ws, min_col=7, min_row=1, max_row=max_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=max_row))
            ws_summary.add_chart(chart, 'D2')
            wb.save(path)
            messagebox.showinfo('Excel', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('Excel', str(e))

    def export_chart(self):
        if not self.current_rows:
            messagebox.showwarning('Wykres', 'Brak rekordow do eksportu.')
            return
        self.app.export_chart(self.fig_hist, 'archiwum_inercji.png')


class HistoricalInertiaWindow(tk.Toplevel):

    def __init__(self, master, app):
        super().__init__(master)
        self.app = app
        self.title('Analiza historyczna inercji')
        init_managed_toplevel(self, 'history_window', 1600, 980)
        self.configure(bg=app.get_palette()['bg'])
        self.mode_var = tk.StringVar(value='Miesieczny')
        self.year_var = tk.StringVar(value='')
        self.month_var = tk.StringVar(value='01 - Styczen')
        self.season_var = tk.StringVar(value='Zima')
        self.granularity_var = tk.StringVar(value='Dobowa')
        self.metric_var = tk.StringVar(value='Srednia')
        self.type_options = {'Wszystkie': '', '24h manual': 'manual', 'Reczne': 'manual_batch', 'Biezace': 'auto'}
        self.type_var = tk.StringVar(value='Wszystkie')
        self.current_analysis = None
        self._closing = False
        self.protocol('WM_DELETE_WINDOW', self.on_close)
        self.bind('<Destroy>', self._on_destroy, add='+')
        self._build_ui()
        self.refresh_years()
        self.refresh()

    def _build_ui(self):
        pal = self.app.get_palette()
        outer = tk.Frame(self, bg=pal['bg'], padx=10, pady=10)
        outer.pack(fill=tk.BOTH, expand=True)

        filters = tk.LabelFrame(outer, text='Zakres i agregacja', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        filters.pack(fill=tk.X)

        def add_combo(parent, row, col, label, variable, values, width=18):
            cell = tk.Frame(parent, bg=pal['card'])
            cell.grid(row=row, column=col, padx=5, pady=2, sticky='ew')
            tk.Label(cell, text=label, bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold_small).pack(anchor='w')
            combo = ttk.Combobox(cell, textvariable=variable, values=list(values), state='readonly', width=width)
            combo.pack(fill=tk.X)
            combo.bind('<<ComboboxSelected>>', lambda *_: self.refresh())
            return combo

        self.combo_mode = add_combo(filters, 0, 0, 'Widok', self.mode_var, ['Miesieczny', 'Sezonowy', 'Roczny', 'Wieloletni'])
        self.combo_year = add_combo(filters, 0, 1, 'Rok', self.year_var, [], width=12)
        month_values = [f'{idx:02d} - {self.app.MONTH_NAMES_PL[idx]}' for idx in range(1, 13)]
        self.combo_month = add_combo(filters, 0, 2, 'Miesiac', self.month_var, month_values, width=18)
        self.combo_season = add_combo(filters, 0, 3, 'Sezon', self.season_var, list(self.app.SEASON_MONTHS.keys()), width=14)
        self.combo_granularity = add_combo(filters, 0, 4, 'Agregacja', self.granularity_var, ['Dobowa', 'Godzinowa'], width=12)
        self.combo_metric = add_combo(filters, 0, 5, 'Metryka', self.metric_var, ['Srednia', 'Mediana', 'Minimum', 'Maksimum'], width=14)
        self.combo_type = add_combo(filters, 0, 6, 'Typ modelu', self.type_var, list(self.type_options.keys()), width=14)
        self.combo_granularity.bind('<<ComboboxSelected>>', lambda *_: (self.refresh_years(), self.refresh()))
        self.combo_type.bind('<<ComboboxSelected>>', lambda *_: (self.refresh_years(), self.refresh()))
        for idx in range(7):
            filters.grid_columnconfigure(idx, weight=1)
        btns = tk.Frame(filters, bg=pal['card'])
        btns.grid(row=1, column=0, columnspan=7, sticky='w', pady=(8, 0))
        self.app.utworz_przycisk(btns, 'Odswiez', self.refresh, role='primary').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Odswiez cache', self.refresh_with_cache_clear, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport CSV analizy', self.export_csv, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport Excel analizy', self.export_xlsx, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport wykresu PNG', self.export_chart_png, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Eksport wykresu PDF', self.export_chart_pdf, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Kopiuj statystyki', self.copy_stats, role='neutral').pack(side=tk.LEFT, padx=3)
        self.app.utworz_przycisk(btns, 'Zamknij', self.on_close, role='neutral').pack(side=tk.LEFT, padx=3)

        frame_summary = tk.LabelFrame(outer, text='Podsumowanie', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        frame_summary.pack(fill=tk.X, pady=(10, 0))
        self.summary_label = tk.Label(frame_summary, text='Brak danych archiwalnych', justify='left', anchor='w', bg=pal['card'], fg=pal['accent'], font=self.app.font_code_small)
        self.summary_label.pack(fill=tk.X)

        body = ttk.PanedWindow(outer, orient=tk.HORIZONTAL)
        body.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        left = tk.Frame(body, bg=pal['bg'])
        right = tk.Frame(body, bg=pal['bg'])
        body.add(left, weight=4)
        body.add(right, weight=2)

        frame_plot = tk.LabelFrame(left, text='Wykresy historyczne', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        frame_plot.pack(fill=tk.BOTH, expand=True)
        self.fig_hist_analysis, (self.ax_hist_main, self.ax_hist_secondary, self.ax_hist_distribution) = plt.subplots(3, 1, figsize=(13.8, 10.1), gridspec_kw={'height_ratios': [2.35, 1.15, 1.0]})
        self.fig_hist_analysis.subplots_adjust(left=0.07, right=0.99, top=0.96, bottom=0.08, hspace=0.48)
        self.canvas_hist_analysis = FigureCanvasTkAgg(self.fig_hist_analysis, master=frame_plot)
        self.canvas_hist_analysis.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        notebook = ttk.Notebook(right)
        notebook.pack(fill=tk.BOTH, expand=True)
        self.analysis_notebook = notebook

        frame_stats = tk.Frame(notebook, bg=pal['bg'])
        frame_weak = tk.Frame(notebook, bg=pal['bg'])
        frame_strong = tk.Frame(notebook, bg=pal['bg'])
        frame_threshold = tk.Frame(notebook, bg=pal['bg'])
        notebook.add(frame_stats, text='Statystyki')
        notebook.add(frame_weak, text='Najslabsze')
        notebook.add(frame_strong, text='Najmocniejsze')
        notebook.add(frame_threshold, text='Ponizej progu')

        stats_box = tk.LabelFrame(frame_stats, text='Statystyki i diagnostyka', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        stats_box.pack(fill=tk.BOTH, expand=True)
        self.stats_text = scrolledtext.ScrolledText(stats_box, wrap=tk.WORD, height=12, font=self.app.font_code_small)
        self.stats_text.pack(fill=tk.BOTH, expand=True)
        self.stats_text.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'])

        rank_columns = [('point', 'Punkt', 142), ('metric', 'Metryka', 96), ('min', 'Min', 86), ('max', 'Max', 86), ('range', 'Rozstep', 88), ('count', 'Probki', 72)]
        threshold_columns = [('point', 'Punkt', 142), ('metric', 'Metryka', 96), ('threshold', 'Prog', 86), ('margin', 'Odchylka', 96), ('count', 'Probki', 72)]
        self.tree_weak = self._create_rank_tree(frame_weak, rank_columns)
        self.tree_strong = self._create_rank_tree(frame_strong, rank_columns)
        self.tree_threshold = self._create_rank_tree(frame_threshold, threshold_columns)

    def get_selected_model_type(self):
        return self.type_options.get(self.type_var.get().strip(), '')

    def on_close(self):
        if self._closing:
            return
        self._closing = True
        try:
            if getattr(self.app, 'history_window', None) is self:
                self.app.history_window = None
        except Exception:
            pass
        try:
            self.destroy()
        except Exception:
            pass

    def _on_destroy(self, event=None):
        if event is not None and event.widget is not self:
            return
        try:
            if getattr(self.app, 'history_window', None) is self:
                self.app.history_window = None
        except Exception:
            pass

    def _create_rank_tree(self, parent, columns):
        pal = self.app.get_palette()
        box = tk.LabelFrame(parent, text='Tabela wynikow', bg=pal['card'], fg=pal['text'], font=self.app.font_ui_bold, padx=8, pady=8)
        box.pack(fill=tk.BOTH, expand=True, padx=2, pady=2)
        tree = ttk.Treeview(box, columns=[col[0] for col in columns], show='headings', height=14)
        for col, header, width in columns:
            tree.heading(col, text=header)
            tree.column(col, width=width, anchor='center')
        yscroll = ttk.Scrollbar(box, orient='vertical', command=tree.yview)
        xscroll = ttk.Scrollbar(box, orient='horizontal', command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky='nsew')
        yscroll.grid(row=0, column=1, sticky='ns')
        xscroll.grid(row=1, column=0, sticky='ew')
        box.rowconfigure(0, weight=1)
        box.columnconfigure(0, weight=1)
        return tree

    def refresh_years(self):
        years = self.app.get_archive_available_years(model_type=self.get_selected_model_type(), granularity=self.granularity_var.get())
        values = [str(y) for y in years] if years else ['']
        self.combo_year.configure(values=values)
        if self.year_var.get() not in values:
            self.year_var.set(values[-1] if values else '')

    def refresh_with_cache_clear(self):
        self.app.clear_history_analysis_cache()
        self.refresh_years()
        self.refresh()

    def _set_control_state(self):
        mode = self.mode_var.get()
        self.combo_year.configure(state='readonly' if mode in {'Miesieczny', 'Sezonowy', 'Roczny'} else 'disabled')
        self.combo_month.configure(state='readonly' if mode == 'Miesieczny' else 'disabled')
        self.combo_season.configure(state='readonly' if mode == 'Sezonowy' else 'disabled')

    def clear_rank_tables(self):
        for tree in (self.tree_weak, self.tree_strong, self.tree_threshold):
            for item in tree.get_children():
                tree.delete(item)

    def refresh(self):
        self._set_control_state()
        month_raw = self.month_var.get().strip()
        month_value = int(month_raw.split('-', 1)[0].strip()) if month_raw else ''
        analysis = self.app.build_historical_analysis(mode=self.mode_var.get(), year=self.year_var.get(), month=month_value, season=self.season_var.get(), metric_label=self.metric_var.get(), model_type=self.get_selected_model_type(), granularity=self.granularity_var.get())
        self.current_analysis = analysis
        pal = self.app.get_palette()
        self.fig_hist_analysis.patch.set_facecolor(pal['bg'])
        self.ax_hist_main.clear()
        self.ax_hist_secondary.clear()
        self.ax_hist_distribution.clear()
        self.app.stylizuj_os(self.ax_hist_main, grid=True, grid_axis='y')
        self.app.stylizuj_os(self.ax_hist_secondary, grid=True, grid_axis='y')
        self.app.stylizuj_os(self.ax_hist_distribution, grid=True, grid_axis='y')
        self.stats_text.delete('1.0', tk.END)
        self.clear_rank_tables()
        if analysis is None:
            self.summary_label.config(text='Brak danych archiwalnych dla wybranego zakresu.', fg=pal['warning'])
            self.ax_hist_main.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=self.ax_hist_main.transAxes, color=pal['muted'])
            self.ax_hist_secondary.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=self.ax_hist_secondary.transAxes, color=pal['muted'])
            self.ax_hist_distribution.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=self.ax_hist_distribution.transAxes, color=pal['muted'])
            self.canvas_hist_analysis.draw_idle()
            return
        self.summary_label.config(text='\n'.join(analysis['summary_lines']), fg=pal['accent'])
        self.stats_text.insert(tk.END, '\n'.join(analysis['stats_lines']))

        main_points = analysis['main_points']
        main_labels = [pt['label'] for pt in main_points]
        main_values = [pt['value'] for pt in main_points]
        xs = list(range(len(main_points)))
        main_markevery = 1 if len(main_points) <= 24 else 2 if len(main_points) <= 72 else 4 if len(main_points) <= 168 else 8
        if len(main_points) > 20:
            self.ax_hist_main.plot(xs, main_values, color=pal['line_main'], linewidth=2.2, marker='o', markersize=4.2 if len(main_points) <= 96 else 3.4, markevery=main_markevery)
            mins = [pt.get('min', pt['value']) for pt in main_points]
            maxs = [pt.get('max', pt['value']) for pt in main_points]
            self.ax_hist_main.fill_between(xs, mins, maxs, color=pal['accent'], alpha=0.10)
        else:
            bars = self.ax_hist_main.bar(xs, main_values, color=pal['line_main'], edgecolor=pal['border'], linewidth=0.8)
            for bar, value in zip(bars, main_values):
                self.ax_hist_main.annotate(f'{value:,.0f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=7.3, color=pal['text'])
        tick_step = max(1, len(main_labels) // 18) if len(main_labels) > 18 else 1
        tick_pos = xs[::tick_step] if xs else []
        tick_labels = main_labels[::tick_step] if main_labels else []
        if xs and xs[-1] not in tick_pos:
            tick_pos = tick_pos + [xs[-1]]
            tick_labels = tick_labels + [main_labels[-1]]
        self.ax_hist_main.set_title(analysis['main_title'])
        self.ax_hist_main.set_ylabel('Inercja [MVA·s]')
        self.ax_hist_main.set_xticks(tick_pos)
        self.ax_hist_main.set_xlabel(analysis.get('main_xlabel', 'Punkt czasu'))
        self.ax_hist_main.set_xticklabels(tick_labels, rotation=50 if len(main_points) > 72 else 34 if len(main_points) > 20 else 18, ha='right', fontsize=8, color=pal['text'])
        self.ax_hist_main.margins(x=0.02)

        secondary_points = analysis['secondary_points']
        sec_labels = [pt['label'] for pt in secondary_points]
        sec_values = [pt['value'] for pt in secondary_points]
        sec_alt_values = [pt.get('value_alt') for pt in secondary_points]
        sec_x = list(range(len(secondary_points)))
        if any((val is not None for val in sec_alt_values)):
            width = 0.36
            self.ax_hist_secondary.bar([x - width / 2 for x in sec_x], sec_values, width=width, color=pal['accent'], edgecolor=pal['border'], linewidth=0.8, label='Srednia')
            alt_clean = [float(v or 0.0) for v in sec_alt_values]
            self.ax_hist_secondary.bar([x + width / 2 for x in sec_x], alt_clean, width=width, color=pal['warning'], edgecolor=pal['border'], linewidth=0.8, label='Minimum')
            legend = self.ax_hist_secondary.legend(loc='upper right', fontsize=7.5, frameon=True)
            if legend is not None:
                legend.get_frame().set_facecolor(pal['card'])
                legend.get_frame().set_edgecolor(pal['border'])
        else:
            self.ax_hist_secondary.bar(sec_x, sec_values, color=pal['line_secondary'], edgecolor=pal['border'], linewidth=0.8)
        if len(sec_x) <= 16:
            for x_val, value in zip(sec_x, sec_values):
                self.ax_hist_secondary.annotate(f'{value:,.0f}', xy=(x_val, value), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=7.0, color=pal['text'])
        self.ax_hist_secondary.set_title(analysis['secondary_title'])
        self.ax_hist_secondary.set_ylabel('Inercja [MVA·s]')
        self.ax_hist_secondary.set_xlabel(analysis.get('secondary_xlabel', 'Przekroj'))
        self.ax_hist_secondary.set_xticks(sec_x)
        self.ax_hist_secondary.set_xticklabels(sec_labels, rotation=26 if len(sec_labels) > 12 else 12, ha='right', fontsize=8, color=pal['text'])
        self.ax_hist_secondary.margins(x=0.02)
        if len(sec_labels) > 18:
            sec_tick_step = max(1, len(sec_labels) // 18)
            sec_tick_pos = sec_x[::sec_tick_step]
            sec_tick_labels = sec_labels[::sec_tick_step]
            if sec_x and sec_x[-1] not in sec_tick_pos:
                sec_tick_pos = sec_tick_pos + [sec_x[-1]]
                sec_tick_labels = sec_tick_labels + [sec_labels[-1]]
            self.ax_hist_secondary.set_xticks(sec_tick_pos)
            self.ax_hist_secondary.set_xticklabels(sec_tick_labels, rotation=32, ha='right', fontsize=8, color=pal['text'])

        distribution_values = analysis.get('distribution_values', [])
        if distribution_values:
            bins = max(8, min(22, int(len(distribution_values) ** 0.5) + 2))
            self.ax_hist_distribution.hist(distribution_values, bins=bins, color=pal['line_secondary'], alpha=0.72, edgecolor=pal['border'])
            self.ax_hist_distribution.axvline(analysis.get('metric_mean', 0.0), color=pal['accent'], linewidth=2.0, linestyle='-', label='Srednia')
            self.ax_hist_distribution.axvline(analysis.get('metric_median', 0.0), color=pal['warning'], linewidth=1.8, linestyle='--', label='Mediana')
            threshold = analysis.get('threshold')
            if threshold is not None:
                self.ax_hist_distribution.axvline(threshold, color=pal['danger'], linewidth=1.8, linestyle=':', label='Prog')
            legend = self.ax_hist_distribution.legend(loc='upper right', fontsize=7.6, frameon=True)
            if legend is not None:
                legend.get_frame().set_facecolor(pal['card'])
                legend.get_frame().set_edgecolor(pal['border'])
        else:
            self.ax_hist_distribution.text(0.5, 0.5, 'Brak rozkladu do pokazania', ha='center', va='center', transform=self.ax_hist_distribution.transAxes, color=pal['muted'])
        self.ax_hist_distribution.set_title('Rozklad metryki w wybranym okresie')
        self.ax_hist_distribution.set_xlabel(f"{analysis.get('metric_label', self.metric_var.get())} [MVA*s]")
        self.ax_hist_distribution.set_ylabel('Liczba punktow')

        metric_key = analysis['metric_key']
        row_label_key = analysis.get('row_label_key', 'date_label')
        row_label_header = analysis.get('row_label_header', 'Punkt')
        point_width = 152 if row_label_key == 'time_label' else 118
        self.tree_weak.heading('point', text=row_label_header)
        self.tree_strong.heading('point', text=row_label_header)
        self.tree_threshold.heading('point', text=row_label_header)
        self.tree_weak.column('point', width=point_width, anchor='center')
        self.tree_strong.column('point', width=point_width, anchor='center')
        self.tree_threshold.column('point', width=point_width, anchor='center')
        for row in analysis['weakest_rows']:
            self.tree_weak.insert('', tk.END, values=(row.get(row_label_key, row.get('date_label', '')), f"{row[metric_key]:.2f}", f"{row['min']:.2f}", f"{row['max']:.2f}", f"{row['range']:.2f}", row['count_records']))
        for row in analysis.get('strongest_rows', []):
            self.tree_strong.insert('', tk.END, values=(row.get(row_label_key, row.get('date_label', '')), f"{row[metric_key]:.2f}", f"{row['min']:.2f}", f"{row['max']:.2f}", f"{row['range']:.2f}", row['count_records']))
        threshold_value = analysis.get('threshold')
        for row in analysis.get('threshold_rows', []):
            margin = float(row.get(metric_key, 0.0) or 0.0) - float(threshold_value or 0.0)
            self.tree_threshold.insert('', tk.END, values=(row.get(row_label_key, row.get('date_label', '')), f"{row[metric_key]:.2f}", f"{float(threshold_value or 0.0):.2f}", f"{margin:+.2f}", row['count_records']))
        self.canvas_hist_analysis.draw_idle()

    def export_csv(self):
        if not self.current_analysis:
            messagebox.showwarning('CSV', 'Najpierw policz analize historyczna.')
            return
        path = filedialog.asksaveasfilename(title='Eksport analizy historycznej do CSV', defaultextension='.csv', initialfile='analiza_historyczna_inercji.csv', filetypes=[('CSV', '*.csv')])
        if not path:
            return
        metric_key = self.current_analysis['metric_key']
        try:
            with open(path, 'w', encoding='utf-8', newline='') as f:
                metric_key = self.current_analysis['metric_key']
                writer = csv.writer(f, delimiter=';')
                writer.writerow(['punkt', 'data', 'godzina', 'wartosc', 'min', 'max', 'mediana', 'odchylenie', 'rozstep', 'p05', 'p25', 'p75', 'p95', 'sredni_load', 'liczba_probek'])
                for row in self.current_analysis.get('selected_rows', []):
                    point_label = row.get('time_label', row.get('date_label', ''))
                    writer.writerow([point_label, row.get('date_label', ''), row.get('hour_label', ''), f"{row[metric_key]:.4f}", f"{row['min']:.4f}", f"{row['max']:.4f}", f"{row['median']:.4f}", f"{row['std']:.4f}", f"{row['range']:.4f}", f"{row['p05']:.4f}", f"{row.get('p25', 0.0):.4f}", f"{row.get('p75', 0.0):.4f}", f"{row['p95']:.4f}", f"{row['demand_mean']:.4f}", row['count_records']])
            messagebox.showinfo('CSV', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('CSV', str(e))

    def _write_points_sheet(self, ws, rows, metric_key):
        headers = ['Punkt', 'Data', 'Godzina', 'Wartosc', 'Min', 'Max', 'Mediana', 'Std', 'Rozstep', 'P05', 'P25', 'P75', 'P95', 'Sredni load', 'Probki']
        bold = Font(bold=True)
        fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold
            cell.fill = fill
        for row_idx, row in enumerate(rows, start=2):
            point_label = row.get('time_label', row.get('date_label', ''))
            values = [point_label, row.get('date_label', ''), row.get('hour_label', ''), float(row.get(metric_key, 0.0) or 0.0), float(row.get('min', 0.0) or 0.0), float(row.get('max', 0.0) or 0.0), float(row.get('median', 0.0) or 0.0), float(row.get('std', 0.0) or 0.0), float(row.get('range', 0.0) or 0.0), float(row.get('p05', 0.0) or 0.0), float(row.get('p25', 0.0) or 0.0), float(row.get('p75', 0.0) or 0.0), float(row.get('p95', 0.0) or 0.0), float(row.get('demand_mean', 0.0) or 0.0), int(row.get('count_records', 0) or 0)]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        for col, width in {'A': 18, 'B': 14, 'C': 11, 'D': 14, 'E': 12, 'F': 12, 'G': 12, 'H': 11, 'I': 12, 'J': 11, 'K': 11, 'L': 11, 'M': 11, 'N': 14, 'O': 9}.items():
            ws.column_dimensions[col].width = width

    def _write_rank_sheet(self, ws, rows, metric_key, threshold=None):
        headers = ['Punkt', 'Wartosc', 'Min', 'Max', 'Rozstep', 'Sredni load', 'Probki']
        if threshold is not None:
            headers.extend(['Prog', 'Odchylka'])
        bold = Font(bold=True)
        fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold
            cell.fill = fill
        for row_idx, row in enumerate(rows, start=2):
            point_label = row.get('time_label', row.get('date_label', ''))
            values = [point_label, float(row.get(metric_key, 0.0) or 0.0), float(row.get('min', 0.0) or 0.0), float(row.get('max', 0.0) or 0.0), float(row.get('range', 0.0) or 0.0), float(row.get('demand_mean', 0.0) or 0.0), int(row.get('count_records', 0) or 0)]
            if threshold is not None:
                values.extend([float(threshold), float(row.get(metric_key, 0.0) or 0.0) - float(threshold)])
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        for col, width in {'A': 18, 'B': 14, 'C': 11, 'D': 11, 'E': 12, 'F': 14, 'G': 9, 'H': 12, 'I': 12}.items():
            ws.column_dimensions[col].width = width

    def export_xlsx(self):
        if not self.current_analysis:
            messagebox.showwarning('Excel', 'Najpierw policz analize historyczna.')
            return
        path = filedialog.asksaveasfilename(title='Eksport analizy historycznej do Excel', defaultextension='.xlsx', initialfile='analiza_historyczna_inercji.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            analysis = self.current_analysis
            wb = Workbook()
            ws_summary = wb.active
            ws_summary.title = 'Podsumowanie'
            bold = Font(bold=True)
            fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            ws_summary['A1'] = 'Analiza'
            ws_summary['B1'] = analysis.get('title', '')
            ws_summary['A1'].font = bold
            ws_summary['B1'].font = bold
            row_idx = 3
            for line in analysis.get('summary_lines', []):
                ws_summary.cell(row=row_idx, column=1, value=line)
                row_idx += 1
            row_idx += 1
            for line in analysis.get('stats_lines', []):
                ws_summary.cell(row=row_idx, column=1, value=line)
                row_idx += 1
            ws_summary.column_dimensions['A'].width = 64
            ws_summary.column_dimensions['B'].width = 28

            ws_points = wb.create_sheet('Punkty')
            self._write_points_sheet(ws_points, analysis.get('selected_rows', []), analysis['metric_key'])
            max_row = max(2, len(analysis.get('selected_rows', [])) + 1)
            chart = LineChart()
            chart.title = 'Przebieg inercji'
            chart.y_axis.title = 'Inercja [MVA*s]'
            chart.x_axis.title = 'Punkt czasu'
            chart.height = 8
            chart.width = 18
            chart.add_data(Reference(ws_points, min_col=4, min_row=1, max_row=max_row), titles_from_data=True)
            chart.set_categories(Reference(ws_points, min_col=1, min_row=2, max_row=max_row))
            ws_summary.add_chart(chart, 'C3')

            ws_profile = wb.create_sheet('Profil')
            for col, header in enumerate(['Etykieta', 'Wartosc', 'Alternatywna'], start=1):
                cell = ws_profile.cell(row=1, column=col, value=header)
                cell.font = bold
                cell.fill = fill
            for idx, point in enumerate(analysis.get('secondary_points', []), start=2):
                ws_profile.cell(row=idx, column=1, value=point.get('label', ''))
                ws_profile.cell(row=idx, column=2, value=float(point.get('value', 0.0) or 0.0))
                if point.get('value_alt') is not None:
                    ws_profile.cell(row=idx, column=3, value=float(point.get('value_alt', 0.0) or 0.0))
            for col, width in {'A': 16, 'B': 14, 'C': 14}.items():
                ws_profile.column_dimensions[col].width = width
            profile_rows = max(2, len(analysis.get('secondary_points', [])) + 1)
            chart_profile = BarChart()
            chart_profile.title = analysis.get('secondary_title', 'Profil')
            chart_profile.y_axis.title = 'Inercja [MVA*s]'
            chart_profile.height = 7
            chart_profile.width = 15
            max_col = 3 if any(point.get('value_alt') is not None for point in analysis.get('secondary_points', [])) else 2
            chart_profile.add_data(Reference(ws_profile, min_col=2, max_col=max_col, min_row=1, max_row=profile_rows), titles_from_data=True)
            chart_profile.set_categories(Reference(ws_profile, min_col=1, min_row=2, max_row=profile_rows))
            ws_summary.add_chart(chart_profile, 'C20')

            ws_weak = wb.create_sheet('Najslabsze')
            self._write_rank_sheet(ws_weak, analysis.get('weakest_rows', []), analysis['metric_key'])
            ws_strong = wb.create_sheet('Najmocniejsze')
            self._write_rank_sheet(ws_strong, analysis.get('strongest_rows', []), analysis['metric_key'])
            threshold_rows = analysis.get('threshold_rows', [])
            if threshold_rows:
                ws_thr = wb.create_sheet('Ponizej progu')
                self._write_rank_sheet(ws_thr, threshold_rows, analysis['metric_key'], threshold=analysis.get('threshold'))
            wb.save(path)
            messagebox.showinfo('Excel', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('Excel', str(e))

    def _export_figure(self, suffix):
        if self.current_analysis is None:
            messagebox.showwarning('Wykres', 'Najpierw policz analize historyczna.')
            return
        path = filedialog.asksaveasfilename(title='Eksport wykresu analizy historycznej', defaultextension=suffix, initialfile=f'analiza_historyczna_inercji{suffix}', filetypes=[('PNG', '*.png'), ('PDF', '*.pdf')])
        if not path:
            return
        try:
            if path.lower().endswith('.pdf'):
                with PdfPages(path) as pdf:
                    pdf.savefig(self.fig_hist_analysis, bbox_inches='tight')
            else:
                self.fig_hist_analysis.savefig(path, dpi=180, bbox_inches='tight')
            messagebox.showinfo('Wykres', f'Zapisano plik:\n{path}')
        except Exception as e:
            messagebox.showerror('Wykres', str(e))

    def export_chart_png(self):
        self._export_figure('.png')

    def export_chart_pdf(self):
        self._export_figure('.pdf')

    def copy_stats(self):
        if not self.current_analysis:
            messagebox.showwarning('Statystyki', 'Najpierw policz analize historyczna.')
            return
        text = self.stats_text.get('1.0', tk.END).strip()
        if not text:
            return
        self.clipboard_clear()
        self.clipboard_append(text)
        messagebox.showinfo('Statystyki', 'Statystyki skopiowano do schowka.')

class TreeviewTooltip:

    def __init__(self, tree: ttk.Treeview):
        self.tree = tree
        self.tipwindow = None
        self.current_item = None
        self.text = ''
        self.tree.bind('<Motion>', self.on_motion)
        self.tree.bind('<Leave>', self.hide)

    def show(self, x, y, text):
        if self.tipwindow and self.text == text:
            try:
                self.tipwindow.geometry(f'+{x + 18}+{y + 12}')
                return
            except Exception:
                pass
        self.hide()
        self.text = text
        self.tipwindow = tw = tk.Toplevel(self.tree)
        tw.wm_overrideredirect(True)
        tw.wm_geometry(f'+{x + 18}+{y + 12}')
        lbl = tk.Label(tw, text=text, justify='left', background='#fff8e1', relief='solid', borderwidth=1, font=('Cascadia Mono', 9), padx=8, pady=6)
        lbl.pack()

    def hide(self, event=None):
        if self.tipwindow:
            try:
                self.tipwindow.destroy()
            except Exception:
                pass
            self.tipwindow = None
            self.current_item = None
            self.text = ''

    def on_motion(self, event):
        item = self.tree.identify_row(event.y)
        if not item:
            self.hide()
            return
        if hasattr(self.tree, '_tooltip_provider'):
            txt = self.tree._tooltip_provider(item)
            if txt:
                self.current_item = item
                self.show(event.x_root, event.y_root, txt)
            else:
                self.hide()

class AnalizaInercjiExcelApp:
    GENERATOR_GROUP_LABELS = {101: 'JWCDc', 102: 'JWCDp', 201: 'JWCKc', 202: 'JWCKw', 203: 'JWCKp', 204: 'JWCKf', 205: 'JWCKpv', 301: 'MC', 302: 'PR', 303: 'MW', 304: 'FW', 305: 'PV'}
    ODM_CODE_TO_NAME = {'1': 'Warszawa', '2': 'Radom', '3': 'Katowice', '4': 'Poznan', '5': 'Bydgoszcz'}
    ODM_ORDER = ['1', '2', '3', '4', '5']
    GROUP_COLORS = ['#2563eb', '#f97316', '#22c55e', '#a855f7', '#ec4899', '#14b8a6', '#8b5cf6', '#0ea5e9', '#eab308', '#ef4444', '#6366f1', '#84cc16']
    NOMINAL_FREQUENCY_HZ = 50.0
    MANUAL_IMPORT_CHART_WINDOW = 36
    AUTO_IMPORT_CHART_WINDOW = 36
    MONTH_NAMES_PL = {1: 'Styczen', 2: 'Luty', 3: 'Marzec', 4: 'Kwiecien', 5: 'Maj', 6: 'Czerwiec', 7: 'Lipiec', 8: 'Sierpien', 9: 'Wrzesien', 10: 'Pazdziernik', 11: 'Listopad', 12: 'Grudzien'}
    SEASON_MONTHS = {'Zima': (12, 1, 2), 'Wiosna': (3, 4, 5), 'Lato': (6, 7, 8), 'Jesien': (9, 10, 11)}
    LIGHT_PALETTE = {'bg': '#f5f7fb', 'card': '#ffffff', 'card_alt': '#eef2f7', 'header': '#eaf1ff', 'border': '#d7e0ec', 'text': '#0f172a', 'muted': '#667085', 'input_bg': '#ffffff', 'input_fg': '#111827', 'accent': '#3b82f6', 'accent_active': '#2563eb', 'success': '#16a34a', 'success_active': '#15803d', 'warning': '#ea580c', 'warning_active': '#c2410c', 'danger': '#dc2626', 'danger_active': '#b91c1c', 'neutral_btn': '#eef2f7', 'neutral_btn_active': '#dbe4f0', 'chart_grid': '#d9e3f0', 'line_main': '#2563eb', 'line_secondary': '#64748b', 'input_border': '#cfd8e3', 'selection': '#dbeafe', 'selection_fg': '#0f172a'}
    DARK_PALETTE = {'bg': '#0f172a', 'card': '#172033', 'card_alt': '#1e293b', 'header': '#12223f', 'border': '#334155', 'text': '#e5e7eb', 'muted': '#9aa8bd', 'input_bg': '#111827', 'input_fg': '#f8fafc', 'accent': '#60a5fa', 'accent_active': '#3b82f6', 'success': '#4ade80', 'success_active': '#22c55e', 'warning': '#fb923c', 'warning_active': '#f97316', 'danger': '#f87171', 'danger_active': '#ef4444', 'neutral_btn': '#243145', 'neutral_btn_active': '#334155', 'chart_grid': '#334155', 'line_main': '#60a5fa', 'line_secondary': '#94a3b8', 'input_border': '#334155', 'selection': '#1d4ed8', 'selection_fg': '#f8fafc'}

    def __init__(self, root):
        self.root = root
        self.root._analiza_inercji_app = self
        self.root.title('Analiza inercji – nowoczesny interfejs')
        self.root.configure(bg=self.LIGHT_PALETTE['bg'])
        self.root.bind('<Destroy>', self._on_root_destroy, add='+')
        self.font_ui = ('Segoe UI', 10)
        self.font_ui_small = ('Segoe UI', 9)
        self.font_ui_bold = ('Segoe UI Semibold', 10)
        self.font_ui_bold_small = ('Segoe UI Semibold', 9)
        self.font_title = ('Segoe UI Semibold', 20)
        self.font_subtitle = ('Segoe UI', 10)
        self.font_code = ('Cascadia Mono', 10)
        self.font_code_small = ('Cascadia Mono', 9)
        self.wszystkie_generatory = {}
        self.statusy_oryginalne = {}
        self.opis_modeli = {}
        self.inercja_na_godzine = {h: 0.0 for h in range(24)}
        self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
        self.odm_metrics_na_godzine = {h: self.build_empty_odm_metrics() for h in range(24)}
        self.rocof_na_godzine = {h: 0.0 for h in range(24)}
        self.largest_loss_na_godzine = {h: 0.0 for h in range(24)}
        self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
        self.zapotrzebowanie_na_godzine = {h: 0.0 for h in range(24)}
        self.modele_reczne = {}
        self.reczne_order = []
        self.reczne_counter = 0
        self.modele_auto = {}
        self.auto_order = []
        self.auto_seen_files = {}
        self.auto_counter = 0
        self.obliczenia_wykonane = False
        self.tekst_filtrowania = ''
        self.H_generatory = {}
        self.H_generatory_original = {}
        self.H_group_defaults = {}
        self.h_filter_var = tk.StringVar(value='')
        self.h_edit_value_var = tk.StringVar(value='')
        self.h_selected_generator_var = tk.StringVar(value='')
        self.threshold_update_job = None
        self.sent_alarm_tokens = set()
        self.alarm_history = []
        self.aktywny_typ_modelu = None
        self.aktualna_godzina = None
        self.aktualny_reczny_id = None
        self.aktualny_auto_id = None
        self.tryb_filtrowania_var = tk.StringVar(value='Wszystkie')
        self.sortowanie_var = tk.StringVar(value='Kod generatora A-Z')
        self.udzialy_scope_var = tk.StringVar(value='Cały model')
        self.rocof_mode_var = tk.StringVar(value='Największa jednostka online (Pgen)')
        self.rocof_custom_dp_var = tk.StringVar(value='')
        self.pokaz_etykiety_inercji_var = tk.BooleanVar(value=True)
        self.pokaz_etykiety_zapotrzebowania_var = tk.BooleanVar(value=True)
        self.pokaz_linie_progu_var = tk.BooleanVar(value=True)
        self.secondary_chart_scope_var = tk.StringVar(value='Modele ręczne')
        self.progress_var = tk.DoubleVar(value=0)
        self.monitor_folder_var = tk.BooleanVar(value=False)
        self.monitor_folder_path = ''
        self.monitor_interval_ms = 5000
        self.email_alarm = EmailAlarmService()
        self.archive_repo = InertiaArchiveRepository()
        self.h_repo = HCoefficientRepository()
        self.h_db_window = None
        self.archive_restore_counter = 0
        self.ui_config_path = Path('~/.analiza_inercji_ui.json').expanduser()
        self.ui_config = {}
        self.managed_windows = {}
        self.dark_mode_var = tk.BooleanVar(value=False)
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        self.column_filters: dict[str, tk.StringVar] = {}
        self.filter_entries: dict[str, tk.Entry] = {}
        self.bliski_prog_proc = 0.05
        self.ui_refresh_job = None
        self.pane_layout_apply_job = None
        self.rocof_window = None
        self.rocof_widgets = {}
        self.history_window = None
        self.alarm_history_window = None
        self.history_analysis_cache = {}
        self.main_content_panes = None
        self.panel_srodkowy = None
        self.left_model_panes = None
        self.right_content_panes = None
        self.action_groups = {}
        self.ttk_style = ttk.Style(self.root)
        self.button_roles: dict[tk.Button, str] = {}
        self.label_neutral = set()
        self.skonfiguruj_przewijane_okno()
        self.utworz_interfejs()
        self.load_ui_config()
        self.apply_ui_config()
        self.reload_h_from_database()
        self.tooltip = TreeviewTooltip(self.tree_generatory)
        self.tree_generatory._tooltip_provider = self.get_generator_tooltip
        atexit.register(self.safe_save_ui_config)

    def get_palette(self):
        return self.DARK_PALETTE if self.dark_mode_var.get() else self.LIGHT_PALETTE

    def clear_history_analysis_cache(self):
        self.history_analysis_cache.clear()

    def _on_root_destroy(self, event=None):
        try:
            if event is not None and event.widget is not self.root:
                return
        except Exception:
            return
        self.cancel_pending_ui_jobs()

    def cancel_pending_ui_jobs(self):
        for attr_name in ('threshold_update_job', 'ui_refresh_job', 'pane_layout_apply_job'):
            job_id = getattr(self, attr_name, None)
            if not job_id:
                continue
            try:
                self.root.after_cancel(job_id)
            except Exception:
                pass
            setattr(self, attr_name, None)

    def shutdown(self):
        self.cancel_pending_ui_jobs()
        try:
            self.save_ui_config()
        except Exception:
            pass
        try:
            self.archive_repo.close()
        except Exception:
            pass
        try:
            self.h_repo.close()
        except Exception:
            pass

    def parse_model_date_text(self, date_txt):
        txt = str(date_txt or '').strip()
        if not txt:
            return None
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%Y%m%d', '%d-%m-%Y', '%d/%m/%Y'):
            try:
                return datetime.strptime(txt, fmt).date()
            except Exception:
                pass
        return None

    def parse_model_hour_text(self, hour_txt):
        txt = str(hour_txt or '').strip()
        if not txt:
            return None
        match = re.search(r'(?P<h>\d{1,2})[:.](?P<m>\d{2})', txt)
        if match:
            try:
                hh = int(match.group('h'))
                mm = int(match.group('m'))
                if 0 <= hh <= 23 and 0 <= mm <= 59:
                    return (hh, mm)
            except Exception:
                return None
        digits = re.sub(r'\D+', '', txt)
        if len(digits) >= 4:
            try:
                hh = int(digits[:2])
                mm = int(digits[2:4])
                if 0 <= hh <= 23 and 0 <= mm <= 59:
                    return (hh, mm)
            except Exception:
                return None
        return None

    def get_season_for_date(self, dt_value):
        month = int(dt_value.month)
        year = int(dt_value.year)
        if month == 12:
            return (year + 1, 'Zima')
        if month in (1, 2):
            return (year, 'Zima')
        if month in (3, 4, 5):
            return (year, 'Wiosna')
        if month in (6, 7, 8):
            return (year, 'Lato')
        return (year, 'Jesien')

    def calculate_percentile(self, values, percentile):
        seq = sorted(float(v) for v in values if v is not None)
        if not seq:
            return 0.0
        if len(seq) == 1:
            return seq[0]
        pos = (len(seq) - 1) * (float(percentile) / 100.0)
        low = int(pos)
        high = min(low + 1, len(seq) - 1)
        frac = pos - low
        return seq[low] * (1.0 - frac) + seq[high] * frac

    def calculate_linear_trend(self, rows, metric_key):
        if len(rows) < 2:
            return 0.0
        xs = []
        for row in rows:
            dt_value = row.get('datetime')
            if dt_value is not None:
                xs.append(float(dt_value.timestamp()))
            else:
                xs.append(float(row['date'].toordinal() * 86400.0))
        ys = [float(row.get(metric_key, 0.0) or 0.0) for row in rows]
        mean_x = statistics.mean(xs)
        mean_y = statistics.mean(ys)
        denom = sum((x - mean_x) ** 2 for x in xs)
        if denom <= 0:
            return 0.0
        num = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
        return num / denom

    def calculate_daily_change_stats(self, rows, metric_key):
        if len(rows) < 2:
            return (0.0, 0.0)
        deltas = []
        for prev_row, next_row in zip(rows, rows[1:]):
            delta = float(next_row.get(metric_key, 0.0) or 0.0) - float(prev_row.get(metric_key, 0.0) or 0.0)
            deltas.append(delta)
        return (statistics.mean(deltas), statistics.mean(abs(delta) for delta in deltas))

    def calculate_correlation(self, xs, ys):
        if len(xs) < 2 or len(xs) != len(ys):
            return 0.0
        mean_x = statistics.mean(xs)
        mean_y = statistics.mean(ys)
        centered_x = [x - mean_x for x in xs]
        centered_y = [y - mean_y for y in ys]
        denom_x = sum(value ** 2 for value in centered_x)
        denom_y = sum(value ** 2 for value in centered_y)
        if denom_x <= 0 or denom_y <= 0:
            return 0.0
        num = sum(x * y for x, y in zip(centered_x, centered_y))
        return num / ((denom_x * denom_y) ** 0.5)

    def calculate_rolling_window_extremes(self, rows, metric_key, window=7):
        if len(rows) < max(2, int(window)):
            return None
        window = int(window)
        label_key = 'time_label' if rows and rows[0].get('time_label') else 'date_label'
        results = []
        for idx in range(window - 1, len(rows)):
            chunk = rows[idx - window + 1:idx + 1]
            values = [float(row.get(metric_key, 0.0) or 0.0) for row in chunk]
            results.append({
                'start': chunk[0].get(label_key, chunk[0].get('date_label', '')),
                'end': chunk[-1].get(label_key, chunk[-1].get('date_label', '')),
                'mean': statistics.mean(values),
            })
        weakest = min(results, key=lambda row: row['mean'])
        strongest = max(results, key=lambda row: row['mean'])
        return {'window': window, 'weakest': weakest, 'strongest': strongest}

    def calculate_tail_mean(self, values, fraction=0.1, from_top=False):
        seq = sorted(float(v) for v in values if v is not None)
        if not seq:
            return 0.0
        count = max(1, int(round(len(seq) * float(fraction))))
        subset = seq[-count:] if from_top else seq[:count]
        return statistics.mean(subset)

    def calculate_threshold_streaks(self, rows, metric_key, threshold):
        if threshold is None or not rows:
            return None
        label_key = 'time_label' if rows and rows[0].get('time_label') else 'date_label'
        streaks = []
        current = []
        for row in rows:
            if float(row.get(metric_key, 0.0) or 0.0) < float(threshold):
                current.append(row)
            elif current:
                streaks.append(list(current))
                current = []
        if current:
            streaks.append(list(current))
        if not streaks:
            return None
        longest = max(streaks, key=len)
        worst = min(streaks, key=lambda seq: statistics.mean(float(item.get(metric_key, 0.0) or 0.0) for item in seq))
        return {
            'count': len(streaks),
            'longest': {
                'length': len(longest),
                'start': longest[0].get(label_key, longest[0].get('date_label', '')),
                'end': longest[-1].get(label_key, longest[-1].get('date_label', '')),
                'mean': statistics.mean(float(item.get(metric_key, 0.0) or 0.0) for item in longest),
            },
            'worst': {
                'length': len(worst),
                'start': worst[0].get(label_key, worst[0].get('date_label', '')),
                'end': worst[-1].get(label_key, worst[-1].get('date_label', '')),
                'mean': statistics.mean(float(item.get(metric_key, 0.0) or 0.0) for item in worst),
            },
        }

    def get_archive_daily_rows(self, model_type=''):
        model_key = str(model_type or '').strip().lower()
        cache_key = ('archive_daily_rows', model_key)
        if cache_key in self.history_analysis_cache:
            return self.history_analysis_cache[cache_key]
        rows = self.archive_repo.fetch_history_rows(model_type=model_key)
        grouped = {}
        for row in rows:
            dt_value = self.parse_model_date_text(row.get('model_date', ''))
            if dt_value is None:
                continue
            inertia = float(row.get('total_inertia', 0.0) or 0.0)
            if inertia <= 0:
                continue
            bucket = grouped.setdefault(dt_value, {'date': dt_value, 'values': [], 'demand_values': [], 'count_records': 0})
            bucket['values'].append(inertia)
            bucket['demand_values'].append(float(row.get('demand', 0.0) or 0.0))
            bucket['count_records'] += 1
        daily_rows = []
        for dt_value in sorted(grouped.keys()):
            payload = grouped[dt_value]
            values = payload['values']
            demand_values = payload['demand_values']
            season_year, season_name = self.get_season_for_date(dt_value)
            row = {
                'date': dt_value,
                'date_label': dt_value.strftime('%d.%m.%Y'),
                'year': dt_value.year,
                'month': dt_value.month,
                'month_label': self.MONTH_NAMES_PL.get(dt_value.month, f'M{dt_value.month:02d}'),
                'day': dt_value.day,
                'weekday': dt_value.weekday(),
                'weekday_label': ['Pon', 'Wt', 'Sro', 'Czw', 'Pt', 'Sob', 'Ndz'][dt_value.weekday()],
                'season_year': season_year,
                'season_name': season_name,
                'count_records': payload['count_records'],
                'mean': statistics.mean(values),
                'median': statistics.median(values),
                'min': min(values),
                'max': max(values),
                'std': statistics.pstdev(values) if len(values) > 1 else 0.0,
                'p05': self.calculate_percentile(values, 5),
                'p25': self.calculate_percentile(values, 25),
                'p75': self.calculate_percentile(values, 75),
                'p95': self.calculate_percentile(values, 95),
                'range': max(values) - min(values),
                'demand_mean': statistics.mean(demand_values) if demand_values else 0.0,
            }
            daily_rows.append(row)
        self.history_analysis_cache[cache_key] = daily_rows
        return daily_rows

    def get_archive_hourly_rows(self, model_type=''):
        model_key = str(model_type or '').strip().lower()
        cache_key = ('archive_hourly_rows', model_key)
        if cache_key in self.history_analysis_cache:
            return self.history_analysis_cache[cache_key]
        rows = self.archive_repo.fetch_history_rows(model_type=model_key)
        grouped = {}
        for row in rows:
            dt_value = self.parse_model_date_text(row.get('model_date', ''))
            hm = self.parse_model_hour_text(row.get('model_hour', ''))
            if dt_value is None or hm is None:
                continue
            inertia = float(row.get('total_inertia', 0.0) or 0.0)
            if inertia <= 0:
                continue
            hour_key = datetime(dt_value.year, dt_value.month, dt_value.day, hm[0], hm[1])
            bucket = grouped.setdefault(hour_key, {'dt': hour_key, 'values': [], 'demand_values': [], 'count_records': 0})
            bucket['values'].append(inertia)
            bucket['demand_values'].append(float(row.get('demand', 0.0) or 0.0))
            bucket['count_records'] += 1
        hourly_rows = []
        for hour_key in sorted(grouped.keys()):
            payload = grouped[hour_key]
            values = payload['values']
            demand_values = payload['demand_values']
            season_year, season_name = self.get_season_for_date(hour_key.date())
            hourly_rows.append({
                'date': hour_key.date(),
                'datetime': hour_key,
                'date_label': hour_key.strftime('%d.%m.%Y'),
                'time_label': hour_key.strftime('%d.%m.%Y %H:%M'),
                'hour_label': hour_key.strftime('%H:%M'),
                'year': hour_key.year,
                'month': hour_key.month,
                'month_label': self.MONTH_NAMES_PL.get(hour_key.month, f'M{hour_key.month:02d}'),
                'day': hour_key.day,
                'hour': hour_key.hour,
                'minute': hour_key.minute,
                'weekday': hour_key.weekday(),
                'weekday_label': ['Pon', 'Wt', 'Sro', 'Czw', 'Pt', 'Sob', 'Ndz'][hour_key.weekday()],
                'season_year': season_year,
                'season_name': season_name,
                'count_records': payload['count_records'],
                'mean': statistics.mean(values),
                'median': statistics.median(values),
                'min': min(values),
                'max': max(values),
                'std': statistics.pstdev(values) if len(values) > 1 else 0.0,
                'p05': self.calculate_percentile(values, 5),
                'p25': self.calculate_percentile(values, 25),
                'p75': self.calculate_percentile(values, 75),
                'p95': self.calculate_percentile(values, 95),
                'range': max(values) - min(values),
                'demand_mean': statistics.mean(demand_values) if demand_values else 0.0,
            })
        self.history_analysis_cache[cache_key] = hourly_rows
        return hourly_rows

    def get_archive_available_years(self, model_type='', granularity='Dobowa'):
        gran_key = str(granularity or 'Dobowa').strip()
        rows = self.get_archive_hourly_rows(model_type=model_type) if gran_key == 'Godzinowa' else self.get_archive_daily_rows(model_type=model_type)
        return sorted({int(row['year']) for row in rows})

    def get_daily_metric_key(self, metric_label):
        mapping = {'Srednia': 'mean', 'Mediana': 'median', 'Minimum': 'min', 'Maksimum': 'max', 'Srednia dobowa': 'mean', 'Mediana dobowa': 'median', 'Minimum dobowe': 'min', 'Maksimum dobowe': 'max'}
        return mapping.get(str(metric_label or '').strip(), 'mean')

    def build_historical_analysis(self, *, mode, year=None, month=None, season=None, metric_label='Srednia', model_type='', granularity='Dobowa'):
        mode_key = str(mode or 'Miesieczny').strip()
        metric_key = self.get_daily_metric_key(metric_label)
        model_type_key = str(model_type or '').strip().lower()
        granularity_key = str(granularity or 'Dobowa').strip()
        cache_key = ('history_analysis', model_type_key, granularity_key, mode_key, str(year or ''), str(month or ''), str(season or ''), metric_key)
        if cache_key in self.history_analysis_cache:
            return self.history_analysis_cache[cache_key]
        analysis_rows = self.get_archive_hourly_rows(model_type=model_type) if granularity_key == 'Godzinowa' else self.get_archive_daily_rows(model_type=model_type)
        if not analysis_rows:
            return None
        threshold = self.pobierz_wartosc_graniczna()
        available_years = self.get_archive_available_years(model_type=model_type, granularity=granularity_key)
        selected_year = int(year) if str(year or '').strip().isdigit() else available_years[-1]
        selected_month = int(month) if str(month or '').strip().isdigit() else analysis_rows[-1]['month']
        selected_season = str(season or 'Zima').strip()
        if selected_season not in self.SEASON_MONTHS:
            selected_season = 'Zima'

        if mode_key == 'Miesieczny':
            selected_rows = [row for row in analysis_rows if row['year'] == selected_year and row['month'] == selected_month]
            title = f"{metric_label} ({granularity_key}) - {self.MONTH_NAMES_PL.get(selected_month, selected_month)} {selected_year}"
            if granularity_key == 'Godzinowa':
                main_points = [{'label': row['time_label'], 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                hour_buckets = {}
                for row in selected_rows:
                    hour_buckets.setdefault(row['hour'], []).append(row[metric_key])
                secondary_points = [{'label': f'{hour:02d}:00', 'value': statistics.mean(vals), 'value_alt': min(vals)} for hour, vals in sorted(hour_buckets.items())]
                secondary_title = 'Srednia i minimum wg godziny doby'
                secondary_xlabel = 'Godzina doby'
            else:
                main_points = [{'label': f"{row['day']:02d}", 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                weekday_buckets = {}
                for row in selected_rows:
                    weekday_buckets.setdefault(row['weekday'], []).append(row[metric_key])
                secondary_points = [{'label': ['Pon', 'Wt', 'Sro', 'Czw', 'Pt', 'Sob', 'Ndz'][wd], 'value': statistics.mean(vals)} for wd, vals in sorted(weekday_buckets.items())]
                secondary_title = 'Srednia wg dnia tygodnia'
                secondary_xlabel = 'Dzien tygodnia'
            main_xlabel = 'Punkt czasu'
        elif mode_key == 'Sezonowy':
            selected_rows = [row for row in analysis_rows if row['season_year'] == selected_year and row['season_name'] == selected_season]
            title = f"{metric_label} ({granularity_key}) - sezon {selected_season} {selected_year}"
            if granularity_key == 'Godzinowa':
                main_points = [{'label': row['time_label'], 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                hour_buckets = {}
                for row in selected_rows:
                    hour_buckets.setdefault(row['hour'], []).append(row[metric_key])
                secondary_points = [{'label': f'{hour:02d}:00', 'value': statistics.mean(vals), 'value_alt': min(vals)} for hour, vals in sorted(hour_buckets.items())]
                secondary_title = 'Profil godzinowy dla sezonu'
                secondary_xlabel = 'Godzina doby'
            else:
                main_points = [{'label': row['date'].strftime('%d.%m'), 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                month_buckets = {}
                for row in selected_rows:
                    month_buckets.setdefault(row['month'], []).append(row[metric_key])
                secondary_points = [{'label': self.MONTH_NAMES_PL.get(m, f'M{m:02d}')[:3], 'value': statistics.mean(vals), 'value_alt': min(vals)} for m, vals in sorted(month_buckets.items())]
                secondary_title = 'Srednia i minimum miesieczne w sezonie'
                secondary_xlabel = 'Miesiac'
            main_xlabel = 'Punkt czasu'
        elif mode_key == 'Roczny':
            selected_rows = [row for row in analysis_rows if row['year'] == selected_year]
            title = f"{metric_label} ({granularity_key}) - rok {selected_year}"
            if granularity_key == 'Godzinowa':
                main_points = [{'label': row['time_label'], 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                hour_buckets = {}
                for row in selected_rows:
                    hour_buckets.setdefault(row['hour'], []).append(row[metric_key])
                secondary_points = [{'label': f'{hour:02d}:00', 'value': statistics.mean(vals), 'value_alt': min(vals)} for hour, vals in sorted(hour_buckets.items())]
                secondary_title = 'Profil godzinowy dla roku'
                secondary_xlabel = 'Godzina doby'
            else:
                main_points = [{'label': row['date'].strftime('%d.%m'), 'value': row[metric_key], 'min': row['min'], 'max': row['max']} for row in selected_rows]
                month_buckets = {}
                for row in selected_rows:
                    month_buckets.setdefault(row['month'], []).append(row[metric_key])
                secondary_points = [{'label': self.MONTH_NAMES_PL.get(m, f'M{m:02d}')[:3], 'value': statistics.mean(vals), 'value_alt': min(vals)} for m, vals in sorted(month_buckets.items())]
                secondary_title = 'Srednia i minimum miesieczne'
                secondary_xlabel = 'Miesiac'
            main_xlabel = 'Punkt czasu'
        else:
            selected_rows = list(analysis_rows)
            title = f"{metric_label} ({granularity_key}) - analiza wieloletnia"
            year_buckets = {}
            for row in selected_rows:
                year_buckets.setdefault(row['year'], []).append(row[metric_key])
            main_points = [{'label': str(y), 'value': statistics.mean(vals), 'value_alt': min(vals)} for y, vals in sorted(year_buckets.items())]
            if granularity_key == 'Godzinowa':
                hour_buckets = {}
                for row in selected_rows:
                    hour_buckets.setdefault(row['hour'], []).append(row[metric_key])
                secondary_points = [{'label': f'{hour:02d}:00', 'value': statistics.mean(vals), 'value_alt': min(vals)} for hour, vals in sorted(hour_buckets.items())]
                secondary_title = 'Wieloletnia klimatologia godzinowa'
                secondary_xlabel = 'Godzina doby'
            else:
                month_buckets = {}
                for row in selected_rows:
                    month_buckets.setdefault(row['month'], []).append(row[metric_key])
                secondary_points = [{'label': self.MONTH_NAMES_PL.get(m, f'M{m:02d}')[:3], 'value': statistics.mean(vals)} for m, vals in sorted(month_buckets.items())]
                secondary_title = 'Klimatologia miesieczna z calego archiwum'
                secondary_xlabel = 'Miesiac'
            main_xlabel = 'Rok'

        if not selected_rows:
            return None
        metric_values = [row[metric_key] for row in selected_rows]
        metric_mean = statistics.mean(metric_values)
        metric_median = statistics.median(metric_values)
        metric_min = min(metric_values)
        metric_max = max(metric_values)
        metric_std = statistics.pstdev(metric_values) if len(metric_values) > 1 else 0.0
        p05 = self.calculate_percentile(metric_values, 5)
        p25 = self.calculate_percentile(metric_values, 25)
        p75 = self.calculate_percentile(metric_values, 75)
        p95 = self.calculate_percentile(metric_values, 95)
        iqr = p75 - p25
        low_tail_mean = self.calculate_tail_mean(metric_values, fraction=0.1, from_top=False)
        high_tail_mean = self.calculate_tail_mean(metric_values, fraction=0.1, from_top=True)
        weakest_rows = sorted(selected_rows, key=lambda row: row[metric_key])[:20]
        strongest_rows = sorted(selected_rows, key=lambda row: row[metric_key], reverse=True)[:20]
        strongest_row = strongest_rows[0]
        weakest_row = weakest_rows[0]
        threshold_rows = [row for row in selected_rows if (threshold is not None and row[metric_key] < threshold)]
        threshold_share = len(threshold_rows) / len(selected_rows) * 100.0 if selected_rows else 0.0
        cv_pct = metric_std / metric_mean * 100.0 if abs(metric_mean) > 1e-9 else 0.0
        trend_slope = self.calculate_linear_trend(selected_rows, metric_key)
        avg_delta, avg_abs_delta = self.calculate_daily_change_stats(selected_rows, metric_key)
        demand_values = [float(row.get('demand_mean', 0.0) or 0.0) for row in selected_rows]
        demand_corr = self.calculate_correlation(metric_values, demand_values) if len({round(val, 6) for val in demand_values}) > 1 else 0.0
        rolling_window = 24 if granularity_key == 'Godzinowa' else 7
        rolling_week = self.calculate_rolling_window_extremes(selected_rows, metric_key, window=rolling_window)
        threshold_streaks = self.calculate_threshold_streaks(selected_rows, metric_key, threshold)
        model_type_label = {'': 'Wszystkie', 'manual': '24h manual', 'manual_batch': 'Reczne', 'auto': 'Biezace'}.get(model_type_key, model_type_key or 'Wszystkie')
        sample_label = 'Liczba punktow godzinowych' if granularity_key == 'Godzinowa' else 'Liczba dni z danymi'
        row_label_key = 'time_label' if granularity_key == 'Godzinowa' else 'date_label'
        trend_scale = 3600.0 if granularity_key == 'Godzinowa' else 86400.0
        row_label_header = 'Godzina' if granularity_key == 'Godzinowa' else 'Data'
        transition_label = 'godzina-do-godziny' if granularity_key == 'Godzinowa' else 'dzien-do-dnia'
        trend_unit = 'MVA*s/h' if granularity_key == 'Godzinowa' else 'MVA*s/dzien'
        trend_preview = trend_slope * (24.0 if granularity_key == 'Godzinowa' else 30.0)
        trend_preview_label = '24 h' if granularity_key == 'Godzinowa' else '30 dni'
        period_start = selected_rows[0].get(row_label_key, selected_rows[0].get('date_label', ''))
        period_end = selected_rows[-1].get(row_label_key, selected_rows[-1].get('date_label', ''))
        stats_lines = [
            f"Zakres: {title}",
            f"{sample_label}: {len(selected_rows)} | liczba rekordow z archiwum: {sum(row['count_records'] for row in selected_rows)}",
            f"{metric_label}: srednia {metric_mean:,.2f} | mediana {metric_median:,.2f} | min {metric_min:,.2f} | max {metric_max:,.2f} MVA*s",
            f"Rozklad: std {metric_std:,.2f} | P05 {p05:,.2f} | P25 {p25:,.2f} | P75 {p75:,.2f} | P95 {p95:,.2f} | IQR {iqr:,.2f}",
            f"Stabilnosc: CV = {cv_pct:.2f}% | srednia zmiana {transition_label} = {avg_delta:+,.2f} MVA*s | srednia zmiana bezwzgledna = {avg_abs_delta:,.2f} MVA*s",
            f"Trend liniowy: {trend_slope * trend_scale:+,.2f} {trend_unit} | orientacyjnie {trend_preview:+,.2f} MVA*s / {trend_preview_label}",
            f"Ogony rozkladu: srednia dolnych 10% = {low_tail_mean:,.2f} | srednia gornych 10% = {high_tail_mean:,.2f} MVA*s",
            f"Korelacja ze srednim loadem: r = {demand_corr:+.3f}",
            f"Najslabszy punkt: {weakest_row.get(row_label_key, weakest_row.get('date_label', ''))} = {weakest_row[metric_key]:,.2f} MVA*s",
            f"Najmocniejszy punkt: {strongest_row.get(row_label_key, strongest_row.get('date_label', ''))} = {strongest_row[metric_key]:,.2f} MVA*s",
        ]
        if threshold is not None:
            stats_lines.append(f"Punkty ponizej progu {threshold:,.2f} MVA*s: {len(threshold_rows)} z {len(selected_rows)} ({threshold_share:.1f}%)")
        if rolling_week is not None:
            stats_lines.append(f"Najslabsze okno {rolling_week['window']}-punktowe: {rolling_week['weakest']['start']} - {rolling_week['weakest']['end']} = {rolling_week['weakest']['mean']:,.2f} MVA*s")
            stats_lines.append(f"Najmocniejsze okno {rolling_week['window']}-punktowe: {rolling_week['strongest']['start']} - {rolling_week['strongest']['end']} = {rolling_week['strongest']['mean']:,.2f} MVA*s")
        if threshold_streaks is not None:
            stats_lines.append(f"Serie ponizej progu: {threshold_streaks['count']} | najdluzsza {threshold_streaks['longest']['length']} pkt ({threshold_streaks['longest']['start']} - {threshold_streaks['longest']['end']})")
            stats_lines.append(f"Najgorsza seria: {threshold_streaks['worst']['start']} - {threshold_streaks['worst']['end']} | srednia {threshold_streaks['worst']['mean']:,.2f} MVA*s")
        summary_lines = [
            f"Analiza historyczna: {title}",
            f"Typ modelu: {model_type_label} | Agregacja: {granularity_key} | Metryka: {metric_label}",
            f"Okres danych: {period_start} - {period_end}",
        ]
        if threshold is not None:
            summary_lines.append(f"Punkty ponizej progu: {len(threshold_rows)} z {len(selected_rows)} ({threshold_share:.1f}%)")
        result = {
            'title': title,
            'metric_label': metric_label,
            'metric_key': metric_key,
            'granularity_label': granularity_key,
            'summary_lines': summary_lines,
            'stats_lines': stats_lines,
            'main_title': title,
            'main_points': main_points,
            'main_xlabel': main_xlabel,
            'secondary_title': secondary_title,
            'secondary_xlabel': secondary_xlabel,
            'secondary_points': secondary_points,
            'selected_rows': selected_rows,
            'weakest_rows': weakest_rows,
            'strongest_rows': strongest_rows,
            'threshold_rows': sorted(threshold_rows, key=lambda row: row[metric_key])[:50],
            'distribution_values': metric_values,
            'threshold': threshold,
            'metric_mean': metric_mean,
            'metric_median': metric_median,
            'row_label_key': row_label_key,
            'row_label_header': row_label_header,
            'mode': mode_key,
        }
        self.history_analysis_cache[cache_key] = result
        return result

    def get_button_scheme(self, role='neutral'):
        pal = self.get_palette()
        if role == 'primary':
            return (pal['accent'], 'white', pal['accent_active'])
        if role == 'success':
            return (pal['success'], 'white', pal['success_active'])
        if role == 'warning':
            return (pal['warning'], 'white', pal['warning_active'])
        if role == 'danger':
            return (pal['danger'], 'white', pal['danger_active'])
        return (pal['neutral_btn'], pal['text'], pal['neutral_btn_active'])

    def stylizuj_przycisk(self, button: tk.Button, role='neutral'):
        bg, fg, active = self.get_button_scheme(role)
        button.configure(bg=bg, fg=fg, activebackground=active, activeforeground=fg, relief='flat', bd=0, padx=12, pady=8, cursor='hand2', font=self.font_ui_bold, highlightthickness=0)

    def utworz_przycisk(self, parent, text, command, role='neutral', width=None):
        btn = tk.Button(parent, text=text, command=command, width=width)
        self.button_roles[btn] = role
        self.stylizuj_przycisk(btn, role)
        return btn

    def create_action_group(self, parent, key, title, items, *, columns=3, default_expanded=True):
        pal = self.get_palette()
        state_cfg = (self.ui_config or {}).get('action_group_state', {})
        expanded = bool(state_cfg.get(key, default_expanded))
        outer = tk.LabelFrame(parent, text=title, padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        outer.pack(side=tk.TOP, fill=tk.X, pady=(0, 8))
        header = tk.Frame(outer, bg=pal['card'])
        header.pack(fill=tk.X)
        desc = tk.Label(header, text='Grupa polecen operacyjnych', bg=pal['card'], fg=pal['muted'], font=self.font_ui_small, anchor='w')
        desc.pack(side=tk.LEFT, fill=tk.X, expand=True)
        toggle_btn = self.utworz_przycisk(header, '', lambda gkey=key: self.toggle_action_group(gkey), role='neutral')
        toggle_btn.pack(side=tk.RIGHT)
        content = tk.Frame(outer, bg=pal['card'])
        for col in range(columns):
            content.grid_columnconfigure(col, weight=1, uniform=f'action_{key}')
        for idx, (label, command, role) in enumerate(items):
            row = idx // columns
            col = idx % columns
            btn = self.utworz_przycisk(content, label, command, role=role)
            btn.grid(row=row, column=col, padx=4, pady=4, sticky='ew')
        self.action_groups[key] = {'outer': outer, 'content': content, 'button': toggle_btn, 'expanded': expanded}
        self._apply_action_group_state(key)
        return outer

    def _apply_action_group_state(self, key):
        meta = self.action_groups.get(key)
        if not meta:
            return
        expanded = bool(meta.get('expanded'))
        if expanded:
            meta['content'].pack(fill=tk.X, pady=(8, 0))
            meta['button'].configure(text='Zwin')
        else:
            meta['content'].pack_forget()
            meta['button'].configure(text='Rozwin')

    def toggle_action_group(self, key):
        meta = self.action_groups.get(key)
        if not meta:
            return
        meta['expanded'] = not bool(meta.get('expanded'))
        cfg = self.ui_config if isinstance(self.ui_config, dict) else {}
        group_cfg = cfg.setdefault('action_group_state', {})
        group_cfg[key] = meta['expanded']
        self.ui_config = cfg
        self._apply_action_group_state(key)
        self.save_ui_config()

    def stylizuj_etykiete_neutralna(self, label: tk.Label):
        self.label_neutral.add(label)
        parent_bg = label.master.cget('bg') if label.master else self.get_palette()['bg']
        label.configure(bg=parent_bg, fg=self.get_palette()['text'], font=self.font_ui)

    def stylizuj_entry(self, entry: tk.Entry):
        pal = self.get_palette()
        entry.configure(relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'], font=self.font_code)

    def stylizuj_option_menu(self, om: tk.Widget):
        pal = self.get_palette()
        om.configure(bg=pal['neutral_btn'], fg=pal['text'], activebackground=pal['neutral_btn_active'], activeforeground=pal['text'], relief='flat', bd=0, cursor='hand2', font=self.font_ui, highlightthickness=0)
        try:
            om['menu'].configure(bg=pal['card'], fg=pal['text'], activebackground=pal['selection'], activeforeground=pal['selection_fg'], font=self.font_ui)
        except Exception:
            pass

    def configure_ttk_styles(self):
        pal = self.get_palette()
        try:
            self.ttk_style.theme_use('clam')
        except Exception:
            pass
        self.ttk_style.configure('Treeview', background=pal['input_bg'], fieldbackground=pal['input_bg'], foreground=pal['text'], bordercolor=pal['border'], lightcolor=pal['border'], darkcolor=pal['border'], rowheight=28, relief='flat', font=self.font_ui_small)
        self.ttk_style.map('Treeview', background=[('selected', pal['accent'])], foreground=[('selected', 'white')])
        self.ttk_style.configure('Treeview.Heading', background=pal['card_alt'], foreground=pal['text'], bordercolor=pal['border'], relief='flat', font=self.font_ui_bold, padding=(8, 6))
        self.ttk_style.map('Treeview.Heading', background=[('active', pal['header'])], foreground=[('active', pal['text'])])
        self.ttk_style.configure('Modern.Horizontal.TProgressbar', troughcolor=pal['card_alt'], background=pal['accent'], bordercolor=pal['border'], lightcolor=pal['accent'], darkcolor=pal['accent'], thickness=10)
        self.ttk_style.configure('Vertical.TScrollbar', background=pal['card_alt'], troughcolor=pal['bg'], bordercolor=pal['border'], arrowcolor=pal['muted'], relief='flat')
        self.ttk_style.configure('Horizontal.TScrollbar', background=pal['card_alt'], troughcolor=pal['bg'], bordercolor=pal['border'], arrowcolor=pal['muted'], relief='flat')

    def configure_widget_tree(self, widget):
        pal = self.get_palette()
        cls = widget.winfo_class()
        if isinstance(widget, tk.Button):
            self.stylizuj_przycisk(widget, self.button_roles.get(widget, 'neutral'))
            return
        if widget is self.root:
            widget.configure(bg=pal['bg'])
        elif isinstance(widget, tk.LabelFrame):
            widget.configure(bg=pal['card'], fg=pal['text'], bd=1, relief='solid', padx=8, pady=8, font=self.font_ui_bold, highlightbackground=pal['border'], highlightcolor=pal['border'])
        elif isinstance(widget, tk.Frame):
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            widget.configure(bg=parent_bg)
        elif isinstance(widget, tk.Label):
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            fg = widget.cget('fg')
            if widget in self.label_neutral:
                fg = pal['text']
            widget.configure(bg=parent_bg, fg=fg)
            try:
                if widget.cget('font') in ('TkDefaultFont',):
                    widget.configure(font=self.font_ui)
            except Exception:
                pass
        elif isinstance(widget, tk.Entry):
            self.stylizuj_entry(widget)
        elif isinstance(widget, tk.Listbox):
            widget.configure(bg=pal['input_bg'], fg=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], selectbackground=pal['selection'], selectforeground=pal['selection_fg'], font=self.font_code)
        elif isinstance(widget, (tk.Text, scrolledtext.ScrolledText)):
            widget.configure(bg=pal['input_bg'], fg=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], insertbackground=pal['input_fg'], selectbackground=pal['selection'], selectforeground=pal['selection_fg'], font=self.font_code_small)
        elif isinstance(widget, tk.Canvas):
            if widget is self.main_canvas:
                widget.configure(bg=pal['bg'], highlightthickness=0)
            else:
                parent_bg = pal['bg']
                try:
                    parent_bg = widget.master.cget('bg')
                except Exception:
                    pass
                widget.configure(bg=parent_bg, highlightthickness=0)
        elif cls == 'Checkbutton':
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            widget.configure(bg=parent_bg, fg=pal['text'], activebackground=parent_bg, activeforeground=pal['text'], selectcolor=parent_bg, font=self.font_ui)
        elif cls == 'Menubutton':
            self.stylizuj_option_menu(widget)
        for child in widget.winfo_children():
            self.configure_widget_tree(child)

    def stylizuj_text_tags(self):
        pal = self.get_palette()
        self.pole_szczegoly.tag_config('header', foreground=pal['accent'], font=(self.font_code_small[0], self.font_code_small[1], 'bold'))
        self.pole_szczegoly.tag_config('green_line', foreground='#15803d' if not self.dark_mode_var.get() else '#86efac')
        self.pole_szczegoly.tag_config('red_line', foreground='#b91c1c' if not self.dark_mode_var.get() else '#fca5a5')
        self.pole_szczegoly.tag_config('manual_line', background='#fff3cd' if not self.dark_mode_var.get() else '#3a2f12', foreground='#c77700' if not self.dark_mode_var.get() else '#fbbf24')
        self.pole_szczegoly.tag_config('summary', foreground=pal['text'], font=(self.font_code_small[0], self.font_code_small[1], 'bold'))

    def stylizuj_tree_tags(self):
        if self.dark_mode_var.get():
            self.tree_generatory.tag_configure('zal', background='#113126', foreground='#86efac')
            self.tree_generatory.tag_configure('wyl', background='#3a1d22', foreground='#fca5a5')
            self.tree_generatory.tag_configure('manual', background='#3a3212', foreground='#fcd34d')
            if hasattr(self, 'tree_h_values'):
                self.tree_h_values.tag_configure('h_manual', background='#3a3212', foreground='#fcd34d')
        else:
            self.tree_generatory.tag_configure('zal', background='#e8f5e9', foreground='#1b5e20')
            self.tree_generatory.tag_configure('wyl', background='#ffebee', foreground='#b71c1c')
            self.tree_generatory.tag_configure('manual', background='#fff3cd', foreground='#e65100')
            if hasattr(self, 'tree_h_values'):
                self.tree_h_values.tag_configure('h_manual', background='#fff3cd', foreground='#e65100')

    def stylizuj_os(self, ax, grid=True, grid_axis='both'):
        pal = self.get_palette()
        ax.set_facecolor(pal['card'])
        for spine in ax.spines.values():
            spine.set_color(pal['border'])
        ax.tick_params(colors=pal['muted'])
        ax.title.set_color(pal['text'])
        ax.xaxis.label.set_color(pal['text'])
        ax.yaxis.label.set_color(pal['text'])
        if grid:
            ax.grid(True, axis=grid_axis, linestyle='--', alpha=0.6, color=pal['chart_grid'])

    def stylizuj_os_prawa(self, ax, color):
        pal = self.get_palette()
        ax.set_facecolor('none')
        ax.grid(False)
        for spine in ax.spines.values():
            spine.set_color(pal['border'])
        try:
            ax.spines['right'].set_color(color)
        except Exception:
            pass
        ax.tick_params(axis='y', colors=color)
        ax.yaxis.label.set_color(color)
        ax.tick_params(axis='x', which='both', bottom=False, labelbottom=False)
        ax.yaxis.set_label_position('right')
        ax.yaxis.tick_right()
        ax.yaxis.label.set_rotation(270)
        ax.yaxis.label.set_verticalalignment('bottom')
        ax.yaxis.label.set_horizontalalignment('center')
        ax.yaxis.set_label_coords(1.12, 0.5)

    def apply_theme(self):
        pal = self.get_palette()
        self.root.configure(bg=pal['bg'])
        self.main_canvas.configure(bg=pal['bg'])
        self.main_frame.configure(bg=pal['bg'])
        self.configure_ttk_styles()
        self.configure_widget_tree(self.root)
        try:
            if self.rocof_window is not None and self.rocof_window.winfo_exists():
                self.configure_widget_tree(self.rocof_window)
        except Exception:
            pass
        try:
            if self.history_window is not None and self.history_window.winfo_exists():
                self.refresh_history_window()
        except Exception:
            pass
        try:
            if self.alarm_history_window is not None and self.alarm_history_window.winfo_exists():
                self.refresh_alarm_history_window()
        except Exception:
            pass
        self.stylizuj_tree_tags()
        self.stylizuj_text_tags()
        self.progress_bar.configure(style='Modern.Horizontal.TProgressbar')
        try:
            self.figura.patch.set_facecolor(pal['bg'])
            self.figura_grupy.patch.set_facecolor(pal['bg'])
        except Exception:
            pass
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            self.przygotuj_pusty_wykres_grup()
        else:
            self.aktualizuj_wykres_grup_generatorow(model_type, key)
        self.refresh_rocof_window()

    def toggle_dark_mode(self):
        self.apply_theme()
        self.save_ui_config()

    def bind_scroll_area(self, widget):
        return

    def is_widget_or_descendant(self, widget, ancestor):
        current = widget
        while current is not None:
            if current == ancestor:
                return True
            current = getattr(current, 'master', None)
        return False

    def get_scroll_target_under_pointer(self, x_root, y_root):
        try:
            widget = self.root.winfo_containing(x_root, y_root)
        except Exception:
            widget = None
        if widget is None:
            return self.main_canvas
        try:
            top_widget = widget.winfo_toplevel()
        except Exception:
            top_widget = self.root
        if top_widget is not self.root:
            return None
        if self.is_widget_or_descendant(widget, self.tree_generatory):
            return self.tree_generatory
        if hasattr(self, 'tree_h_values') and self.is_widget_or_descendant(widget, self.tree_h_values):
            return self.tree_h_values
        if self.is_widget_or_descendant(widget, self.listbox_modele):
            return self.listbox_modele
        if hasattr(self, 'listbox_modele_reczne') and self.is_widget_or_descendant(widget, self.listbox_modele_reczne):
            return self.listbox_modele_reczne
        if self.is_widget_or_descendant(widget, self.listbox_modele_auto):
            return self.listbox_modele_auto
        if self.is_widget_or_descendant(widget, self.pole_szczegoly):
            return self.pole_szczegoly
        return self.main_canvas

    def scroll_widget(self, widget, units, horizontal=False):
        try:
            if horizontal:
                widget.xview_scroll(units, 'units')
            else:
                widget.yview_scroll(units, 'units')
            return True
        except Exception:
            return False

    def on_global_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if event.delta else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if target is None:
            return 'break'
        if not self.scroll_widget(target, units):
            self.main_canvas.yview_scroll(units, 'units')
        return 'break'

    def on_global_mousewheel_linux(self, event):
        step = -1 if event.num == 4 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if target is None:
            return 'break'
        if not self.scroll_widget(target, step):
            self.main_canvas.yview_scroll(step, 'units')
        return 'break'

    def on_shift_mousewheel_main(self, event):
        units = int(-1 * (event.delta / 120)) if event.delta else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if target is None:
            return 'break'
        if target is self.tree_generatory:
            if not self.scroll_widget(self.tree_generatory, units, horizontal=True):
                self.main_canvas.xview_scroll(units, 'units')
        else:
            self.main_canvas.xview_scroll(units, 'units')
        return 'break'

    def reset_alarm_cache(self):
        self.sent_alarm_tokens.clear()

    def konwertuj_float(self, value):
        return float(str(value).replace(',', '.'))

    def konwertuj_float_soft(self, value, default=0.0):
        try:
            return self.konwertuj_float(value)
        except Exception:
            return default

    def aktualizuj_status(self, tekst):
        try:
            self.label_status.config(text=f'Status: {tekst}')
            self.root.update_idletasks()
        except Exception:
            pass

    def pobierz_wartosc_graniczna(self):
        txt = self.entry_wartosc_graniczna.get().strip()
        if not txt:
            return None
        try:
            return self.konwertuj_float(txt)
        except Exception:
            return None

    def schedule_prog_changed(self, event=None):
        if self.threshold_update_job is not None:
            try:
                self.root.after_cancel(self.threshold_update_job)
            except Exception:
                pass
        self.threshold_update_job = self.root.after(350, self.on_prog_changed)

    def on_prog_changed(self, event=None):
        self.threshold_update_job = None
        self.odswiez_widok_dla_progu(check_alarm=False)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.save_ui_config()

    def sprawdz_prog_teraz(self):
        if self.threshold_update_job is not None:
            try:
                self.root.after_cancel(self.threshold_update_job)
            except Exception:
                pass
            self.threshold_update_job = None
        self.odswiez_widok_dla_progu(check_alarm=False)

    def odswiez_widok_dla_progu(self, check_alarm=False):
        prog = self.pobierz_wartosc_graniczna()
        trigger_mode = self.email_alarm.config.trigger_mode
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if prog is None:
            txt = self.entry_wartosc_graniczna.get().strip()
            if txt:
                self.label_limit_ocena.config(text='Ocena: niepoprawna wartość graniczna', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text='Ocena: brak wartości granicznej', fg=self.get_palette()['muted'])
        elif dodatnie:
            if trigger_mode == 'above':
                _, val_ref = max(dodatnie, key=lambda x: x[1])
                if val_ref > prog:
                    self.label_limit_ocena.config(text=f'Ocena: PRZEKROCZONO wartość graniczną o {val_ref - prog:,.2f} MVA·s', fg=self.get_palette()['danger'])
                else:
                    self.label_limit_ocena.config(text=f'Ocena: NIE przekroczono wartości granicznej (zapas {prog - val_ref:,.2f} MVA·s)', fg=self.get_palette()['success'])
            else:
                _, val_ref = min(dodatnie, key=lambda x: x[1])
                if val_ref < prog:
                    self.label_limit_ocena.config(text=f'Ocena: SPADŁO poniżej wartości granicznej o {prog - val_ref:,.2f} MVA·s', fg=self.get_palette()['danger'])
                else:
                    self.label_limit_ocena.config(text=f'Ocena: NIE spadło poniżej wartości granicznej (margines {val_ref - prog:,.2f} MVA·s)', fg=self.get_palette()['success'])
        else:
            self.label_limit_ocena.config(text='Ocena: brak wartości do porównania', fg=self.get_palette()['muted'])
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wykres()
        self.aktualizuj_wskazniki_jakosci()
        if self.aktywny_typ_modelu is not None:
            self.aktualizuj_analize_inercyjna(*self.get_current_model_ref())

    def przygotuj_pusty_wykres(self):
        pal = self.get_palette()
        self.os1.clear()
        self.os_manual.clear()
        self.os_auto.clear()
        if not hasattr(self, 'os1_load') or self.os1_load is None:
            self.os1_load = self.os1.twinx()
        if not hasattr(self, 'os_manual_load') or self.os_manual_load is None:
            self.os_manual_load = self.os_manual.twinx()
        if not hasattr(self, 'os_auto_load') or self.os_auto_load is None:
            self.os_auto_load = self.os_auto.twinx()
        self.os1_load.clear()
        self.os_manual_load.clear()
        self.os_auto_load.clear()
        self.figura.patch.set_facecolor(pal['bg'])
        self._draw_inertia_load_chart(
            self.os1,
            self.os1_load,
            list(range(24)),
            [self.inercja_na_godzine.get(h, 0.0) for h in range(24)],
            [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h in range(24)],
            [f'{h:02d}:30' for h in range(24)],
            title='Inercja systemowa E_sys(t) = Σ S_i(t) × H_i',
            xlabel='Czas [HH:30]',
            left_label_coords=(-0.09, 0.78),
            window_size=24,
        )
        manual_payload = self.get_manual_chart_payload()
        self._draw_inertia_load_chart(
            self.os_manual,
            self.os_manual_load,
            [],
            [],
            [],
            [],
            title=manual_payload['title'],
            xlabel=manual_payload['xlabel'],
            empty_message=manual_payload['empty_message'],
            left_label_coords=(-0.09, 0.78),
            window_size=manual_payload['window_size'],
        )
        auto_payload = self.get_auto_chart_payload()
        self._draw_inertia_load_chart(
            self.os_auto,
            self.os_auto_load,
            [],
            [],
            [],
            [],
            title=auto_payload['title'],
            xlabel=auto_payload['xlabel'],
            empty_message=auto_payload['empty_message'],
            left_label_coords=(-0.09, 0.78),
            window_size=auto_payload['window_size'],
        )
        self.canvas.draw_idle()

    def przygotuj_pusty_wykres_grup(self, komunikat='Brak wybranego modelu'):
        pal = self.get_palette()
        self.figura_grupy.patch.set_facecolor(pal['bg'])
        self.os_grupy.clear()
        self.os_pmax_grupy.clear()
        self.os_odm_inercja.clear()
        self.os_grupy.set_title('Udziały typów jednostek')
        self.os_grupy.set_xlabel('Typ jednostki')
        self.os_grupy.set_ylabel('Udział [%]')
        self.stylizuj_os(self.os_grupy, grid=True, grid_axis='y')
        self.os_pmax_grupy.set_title('Pmax wg typów jednostek')
        self.os_pmax_grupy.set_xlabel('Typ jednostki')
        self.os_pmax_grupy.set_ylabel('Pmax [MW]')
        self.stylizuj_os(self.os_pmax_grupy, grid=True, grid_axis='y')
        self.os_odm_inercja.set_title('Sumaryczna inercja dla ODM')
        self.os_odm_inercja.set_xlabel('ODM')
        self.os_odm_inercja.set_ylabel('Inercja [MVA·s]')
        self.stylizuj_os(self.os_odm_inercja, grid=True, grid_axis='y')
        self.os_grupy.set_ylim(0, 100)
        self.os_grupy.text(0.5, 0.5, komunikat, ha='center', va='center', transform=self.os_grupy.transAxes, fontsize=10, color=pal['muted'])
        self.os_pmax_grupy.set_ylim(0, 1)
        self.os_pmax_grupy.text(0.5, 0.5, komunikat, ha='center', va='center', transform=self.os_pmax_grupy.transAxes, fontsize=10, color=pal['muted'])
        self.os_odm_inercja.set_ylim(0, 1)
        self.os_odm_inercja.text(0.5, 0.5, komunikat, ha='center', va='center', transform=self.os_odm_inercja.transAxes, fontsize=10, color=pal['muted'])
        self.canvas_grupy.draw_idle()

    def skonfiguruj_przewijane_okno(self):
        self.main_canvas = tk.Canvas(self.root, highlightthickness=0, bg=self.get_palette()['bg'])
        self.v_scroll_main = ttk.Scrollbar(self.root, orient='vertical', command=self.main_canvas.yview)
        self.h_scroll_main = ttk.Scrollbar(self.root, orient='horizontal', command=self.main_canvas.xview)
        self.main_canvas.configure(yscrollcommand=self.v_scroll_main.set, xscrollcommand=self.h_scroll_main.set)
        self.v_scroll_main.pack(side=tk.RIGHT, fill=tk.Y)
        self.h_scroll_main.pack(side=tk.BOTTOM, fill=tk.X)
        self.main_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.main_frame = tk.Frame(self.main_canvas, bg=self.get_palette()['bg'])
        self.main_canvas.create_window((0, 0), window=self.main_frame, anchor='nw')
        self.main_frame.bind('<Configure>', lambda e: self.main_canvas.configure(scrollregion=self.main_canvas.bbox('all')))
        self.main_canvas.bind('<Configure>', lambda e: self.main_canvas.configure(scrollregion=self.main_canvas.bbox('all')))
        self.main_canvas.bind_all('<MouseWheel>', self.on_global_mousewheel)
        self.main_canvas.bind_all('<Button-4>', self.on_global_mousewheel_linux)
        self.main_canvas.bind_all('<Button-5>', self.on_global_mousewheel_linux)
        self.main_canvas.bind_all('<Shift-MouseWheel>', self.on_shift_mousewheel_main)

    def utworz_interfejs(self):
        pal = self.get_palette()
        header = tk.Frame(self.main_frame, bg=pal['header'], padx=18, pady=16)
        header.pack(side=tk.TOP, fill=tk.X, padx=10, pady=(10, 12))
        self.header_frame = header
        self.label_title = tk.Label(header, text='Analiza inercji', bg=pal['header'], fg=pal['text'], font=self.font_title)
        self.label_title.pack(anchor='w')
        self.label_subtitle = tk.Label(header, bg=pal['header'], fg=pal['muted'], font=self.font_subtitle, text='Import EPC + DYD, analiza E_sys oraz zapotrzebowania mocy z sekcji LOAD DATA')
        self.label_subtitle.pack(anchor='w', pady=(4, 0))
        panel_gorny = tk.Frame(self.main_frame, padx=10, pady=0, bg=pal['bg'])
        panel_gorny.pack(side=tk.TOP, fill=tk.X)
        frame_przyciski = tk.LabelFrame(panel_gorny, text='Akcje główne', padx=10, pady=10, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_przyciski.pack(side=tk.TOP, fill=tk.X)
        self.action_groups = {}
        self.create_action_group(frame_przyciski, 'import_calc', 'Import i obliczenia', [
            ('Import EPC 24h', self.wczytaj_wiele_plikow_epc, 'primary'),
            ('Import EPC ręczne', self.wczytaj_modele_reczne_epc, 'primary'),
            ('Import DYD / H', self.wczytaj_plik_genrou_H, 'primary'),
            ('Przelicz wszystko', self.oblicz_inercje_wszystkie, 'success'),
            ('Przelicz model', self.przelicz_aktywny_model, 'neutral'),
        ], columns=3, default_expanded=True)
        self.create_action_group(frame_przyciski, 'analysis', 'Analizy i diagnostyka', [
            ('Analiza ROCOF', self.open_rocof_window, 'warning'),
            ('Historia inercji', self.open_historical_window, 'warning'),
            ('Archiwum', self.open_archive_browser, 'primary'),
            ('Baza H', self.open_h_database_window, 'primary'),
            ('Walidacja EPC ↔ DYD', self.pokaz_walidacje_epc_dyd, 'warning'),
            ('Konfiguruj alarm', self.konfiguruj_email_alarm, 'warning'),
            ('Historia alarmów', self.pokaz_historie_alarmow, 'neutral'),
        ], columns=3, default_expanded=False)
        self.create_action_group(frame_przyciski, 'restore_export', 'Eksport i odtwarzanie', [
            ('Przywróć model', self.przywroc_oryginalne_statusy_modelu, 'neutral'),
            ('Przywróć 24h', self.przywroc_wszystkie_modele_do_oryginalu, 'neutral'),
            ('Eksport Excel', self.eksportuj_do_excel, 'neutral'),
            ('Eksport EPC po zmianach', self.eksportuj_aktualny_model_epc_po_zmianach, 'neutral'),
            ('Raport zmian', self.eksportuj_raport_zmian, 'neutral'),
            ('Raport PDF', self.eksportuj_pelny_raport_pdf, 'success'),
            ('Zapisz sesję', self.zapisz_sesje, 'neutral'),
            ('Wczytaj sesję', self.wczytaj_sesje, 'neutral'),
        ], columns=3, default_expanded=False)
        frame_postep = tk.LabelFrame(panel_gorny, text='Postęp importu', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_postep.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        tk.Label(frame_postep, text='Postęp importu EPC:', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT, padx=(0, 8))
        self.progress_bar = ttk.Progressbar(frame_postep, orient='horizontal', mode='determinate', variable=self.progress_var, length=340, style='Modern.Horizontal.TProgressbar')
        self.progress_bar.pack(side=tk.LEFT, padx=4)
        self.label_progress = tk.Label(frame_postep, text='0%', fg=pal['accent'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_progress.pack(side=tk.LEFT, padx=8)
        frame_monitor = tk.LabelFrame(panel_gorny, text='Monitoring folderu z modelami bieżącymi', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_monitor.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.utworz_przycisk(frame_monitor, 'Wybierz folder', self.wybierz_folder_monitoringu, role='neutral').pack(side=tk.LEFT, padx=4)
        chk_monitor = tk.Checkbutton(frame_monitor, text='Monitoruj folder', variable=self.monitor_folder_var, command=self.toggle_monitor_folder)
        chk_monitor.pack(side=tk.LEFT, padx=8)
        self.label_monitor_folder = tk.Label(frame_monitor, text='Folder: nie wybrano', fg=pal['accent'], bg=pal['card'], font=self.font_code_small)
        self.label_monitor_folder.pack(side=tk.LEFT, padx=10)
        frame_export_plot = tk.LabelFrame(panel_gorny, text='Szybki eksport wykresów', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_export_plot.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.utworz_przycisk(frame_export_plot, 'Eksport wykresu 24h PNG', lambda: self.export_chart(self.figura, 'wykres_24h.png'), role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_export_plot, 'Eksport wykresu grup PNG', lambda: self.export_chart(self.figura_grupy, 'wykres_grup.png'), role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_export_plot, 'Eksport dashboardu PNG', self.export_dashboard_png, role='neutral').pack(side=tk.LEFT, padx=3)
        frame_layout = tk.LabelFrame(panel_gorny, text='Układ interfejsu', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_layout.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.utworz_przycisk(frame_layout, 'Zapisz układ do pliku', self.save_layout_to_file, role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_layout, 'Wczytaj układ z pliku', self.load_layout_from_file, role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_layout, 'Zapisz jako domyślny', self.save_current_layout_as_default, role='success').pack(side=tk.LEFT, padx=3)
        self.main_content_panes = ttk.PanedWindow(self.main_frame, orient=tk.VERTICAL)
        self.main_content_panes.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=10, pady=(12, 10))
        self.main_content_panes.bind('<ButtonRelease-1>', lambda *_: self.save_ui_config(), add='+')
        frame_workspace = tk.Frame(self.main_content_panes, bg=pal['bg'])
        self.main_content_panes.add(frame_workspace, weight=4)
        self.panel_srodkowy = ttk.PanedWindow(frame_workspace, orient=tk.HORIZONTAL)
        self.panel_srodkowy.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        lewa_kolumna = tk.Frame(self.panel_srodkowy, bg=pal['bg'])
        prawa_kolumna = tk.Frame(self.panel_srodkowy, bg=pal['bg'])
        self.panel_srodkowy.add(lewa_kolumna, weight=1)
        self.panel_srodkowy.add(prawa_kolumna, weight=2)
        self.left_model_panes = ttk.PanedWindow(lewa_kolumna, orient=tk.VERTICAL)
        self.left_model_panes.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        frame_modele = tk.LabelFrame(self.left_model_panes, text='Modele godzinowe 24h', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.left_model_panes.add(frame_modele, weight=1)
        listbox_frame = tk.Frame(frame_modele, bg=pal['card'])
        listbox_frame.pack(fill=tk.BOTH, expand=True)
        self.listbox_modele = tk.Listbox(listbox_frame, font=self.font_code_small, width=40, height=5)
        scrollbar_modele = ttk.Scrollbar(listbox_frame, orient='vertical', command=self.listbox_modele.yview)
        self.listbox_modele.configure(yscrollcommand=scrollbar_modele.set)
        self.listbox_modele.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_modele.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox_modele.bind('<<ListboxSelect>>', self.wybierz_model_manual)
        self.bind_scroll_area(self.listbox_modele)
        frame_modele_reczne = tk.LabelFrame(self.left_model_panes, text='Modele ręczne importowane', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.left_model_panes.add(frame_modele_reczne, weight=2)
        listbox_reczne_frame = tk.Frame(frame_modele_reczne, bg=pal['card'])
        listbox_reczne_frame.pack(fill=tk.BOTH, expand=True)
        self.listbox_modele_reczne = tk.Listbox(listbox_reczne_frame, font=self.font_code_small, width=52, height=8)
        scrollbar_modele_reczne = ttk.Scrollbar(listbox_reczne_frame, orient='vertical', command=self.listbox_modele_reczne.yview)
        self.listbox_modele_reczne.configure(yscrollcommand=scrollbar_modele_reczne.set)
        self.listbox_modele_reczne.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_modele_reczne.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox_modele_reczne.bind('<<ListboxSelect>>', self.wybierz_model_reczny)
        self.bind_scroll_area(self.listbox_modele_reczne)
        frame_modele_auto = tk.LabelFrame(self.left_model_panes, text='Modele bieżące', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.left_model_panes.add(frame_modele_auto, weight=2)
        listbox_auto_frame = tk.Frame(frame_modele_auto, bg=pal['card'])
        listbox_auto_frame.pack(fill=tk.BOTH, expand=True)
        self.listbox_modele_auto = tk.Listbox(listbox_auto_frame, font=self.font_code_small, width=52, height=7)
        scrollbar_modele_auto = ttk.Scrollbar(listbox_auto_frame, orient='vertical', command=self.listbox_modele_auto.yview)
        self.listbox_modele_auto.configure(yscrollcommand=scrollbar_modele_auto.set)
        self.listbox_modele_auto.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_modele_auto.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox_modele_auto.bind('<<ListboxSelect>>', self.wybierz_model_auto)
        self.bind_scroll_area(self.listbox_modele_auto)
        self.panel_srodkowy.bind('<ButtonRelease-1>', lambda *_: self.save_ui_config(), add='+')
        self.left_model_panes.bind('<ButtonRelease-1>', lambda *_: self.save_ui_config(), add='+')
        frame_filtr = tk.LabelFrame(lewa_kolumna, text='Filtr generatorów', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_filtr.pack(side=tk.TOP, anchor='w', fill=tk.X, pady=(10, 0))
        tk.Label(frame_filtr, text='Kod generatora zawiera (można kilka po przecinku):', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(anchor='w')
        self.entry_filtr = tk.Entry(frame_filtr, font=self.font_code, width=18)
        self.entry_filtr.pack(fill=tk.X, padx=5, pady=4)
        self.entry_filtr.bind('<KeyRelease>', self.zastosuj_filtr)
        row1 = tk.Frame(frame_filtr, bg=pal['card'])
        row1.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(row1, text='Widok:', bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT)
        opcje_trybu = ['Wszystkie', 'Tylko zmienione ręcznie', 'Tylko załączone', 'Tylko wyłączone']
        self.option_tryb = tk.OptionMenu(row1, self.tryb_filtrowania_var, *opcje_trybu, command=lambda _: self.schedule_active_model_refresh(delay_ms=40))
        self.option_tryb.pack(side=tk.LEFT, padx=(8, 0))
        row2 = tk.Frame(frame_filtr, bg=pal['card'])
        row2.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(row2, text='Sortuj:', bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT)
        opcje_sort = ['Kod generatora A-Z', 'Kod generatora Z-A', 'ODM', 'S malejąco', 'S rosnąco', 'H malejąco', 'H rosnąco', 'Ei malejąco', 'Ei rosnąco', 'Status']
        self.option_sort = tk.OptionMenu(row2, self.sortowanie_var, *opcje_sort, command=lambda _: self.on_dropdown_sort_changed())
        self.option_sort.pack(side=tk.LEFT, padx=(8, 0))
        row3 = tk.Frame(frame_filtr, bg=pal['card'])
        row3.pack(fill=tk.X, padx=5, pady=4)
        self.utworz_przycisk(row3, 'Wyczyść', self.wyczysc_filtr, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Wszystkie', self.pokaz_wszystkie, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Załącz widoczne', self.zalacz_widoczne_z_filtra, role='success').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Wyłącz widoczne', self.wylacz_widoczne_z_filtra, role='warning').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Przywróć widoczne', self.przywroc_widoczne_z_filtra, role='neutral').pack(side=tk.LEFT, padx=2)
        self.label_wynikow = tk.Label(frame_filtr, text='0/0 generatorów po filtrze', fg=pal['success'], bg=pal['card'], font=self.font_code)
        self.label_wynikow.pack(pady=2)
        frame_limit = tk.LabelFrame(prawa_kolumna, text='Wartość graniczna i maksimum z 24 godzin', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_limit.pack(side=tk.TOP, fill=tk.X, pady=(0, 10))
        tk.Label(frame_limit, text='Wartość graniczna [MVA·s]:', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT, padx=(5, 5))
        self.entry_wartosc_graniczna = tk.Entry(frame_limit, width=14, font=self.font_code)
        self.entry_wartosc_graniczna.pack(side=tk.LEFT, padx=(0, 8))
        self.entry_wartosc_graniczna.bind('<KeyRelease>', self.schedule_prog_changed)
        self.utworz_przycisk(frame_limit, 'Sprawdź', self.sprawdz_prog_teraz, role='primary').pack(side=tk.LEFT, padx=4)
        self.label_maksimum_info = tk.Label(frame_limit, text='Maksymalna inercja z 24 godzin: brak danych', fg=pal['accent'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_maksimum_info.pack(side=tk.LEFT, padx=12)
        self.label_limit_ocena = tk.Label(frame_limit, text='Ocena: brak wartości granicznej', fg=pal['muted'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_limit_ocena.pack(side=tk.LEFT, padx=12)
        frame_summary_row = tk.Frame(prawa_kolumna, bg=pal['bg'])
        frame_summary_row.pack(side=tk.TOP, fill=tk.X, pady=(0, 10))
        frame_compare = tk.LabelFrame(frame_summary_row, text='Porównanie przed / po zmianach', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_compare.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.label_porownanie = tk.Label(frame_compare, text='Brak wybranego modelu', justify='left', anchor='w', font=('Cascadia Mono', 10, 'bold'), fg=pal['accent'], bg=pal['card'])
        self.label_porownanie.pack(fill=tk.X, padx=5, pady=3)
        frame_metrics = tk.LabelFrame(frame_summary_row, text='Wskaźniki jakości (24h)', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_metrics.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 0))
        self.label_wskazniki = tk.Label(frame_metrics, text='Brak danych 24h', justify='left', anchor='w', font=self.font_code_small, fg=pal['accent'], bg=pal['card'])
        self.label_wskazniki.pack(fill=tk.X, padx=5, pady=3)
        self.right_content_panes = ttk.PanedWindow(prawa_kolumna, orient=tk.VERTICAL)
        self.right_content_panes.pack(side=tk.TOP, fill=tk.BOTH, expand=True, pady=(0, 0))
        self.right_content_panes.bind('<ButtonRelease-1>', lambda *_: self.save_ui_config(), add='+')

        frame_h_values = tk.LabelFrame(self.right_content_panes, text='Panel H z pliku DYD – podgląd i ręczna korekta', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.label_h_summary = tk.Label(frame_h_values, text='Brak wczytanego pliku DYD', justify='left', anchor='w', font=self.font_code_small, fg=pal['accent'], bg=pal['card'])
        self.label_h_summary.pack(fill=tk.X, padx=5, pady=(0, 6))
        frame_h_filter = tk.Frame(frame_h_values, bg=pal['card'])
        frame_h_filter.pack(fill=tk.X, padx=5, pady=(0, 6))
        tk.Label(frame_h_filter, text='Filtr generatora:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).pack(side=tk.LEFT)
        self.entry_h_filter = tk.Entry(frame_h_filter, textvariable=self.h_filter_var, width=28, font=self.font_code_small)
        self.entry_h_filter.pack(side=tk.LEFT, padx=(6, 4))
        self.entry_h_filter.bind('<KeyRelease>', lambda event: self.odswiez_panel_h_dyd())
        self.utworz_przycisk(frame_h_filter, 'Wyczyść filtr', self.wyczysc_filtr_h, role='neutral').pack(side=tk.LEFT, padx=4)
        self.utworz_przycisk(frame_h_filter, 'Otwórz bazę H', self.open_h_database_window, role='primary').pack(side=tk.LEFT, padx=4)
        h_tree_frame = tk.Frame(frame_h_values, bg=pal['card'])
        h_tree_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=(0, 6))
        self.tree_h_values = ttk.Treeview(h_tree_frame, columns=('generator', 'h', 'status'), show='headings', height=8, selectmode='browse')
        self.tree_h_values.heading('generator', text='Generator')
        self.tree_h_values.heading('h', text='H [s]')
        self.tree_h_values.heading('status', text='Źródło')
        self.tree_h_values.column('generator', width=220, anchor='w')
        self.tree_h_values.column('h', width=95, anchor='e')
        self.tree_h_values.column('status', width=120, anchor='center')
        scrollbar_h_y = ttk.Scrollbar(h_tree_frame, orient='vertical', command=self.tree_h_values.yview)
        self.tree_h_values.configure(yscrollcommand=scrollbar_h_y.set)
        self.tree_h_values.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_h_y.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree_h_values.bind('<<TreeviewSelect>>', self.on_h_tree_select)
        self.bind_scroll_area(self.tree_h_values)
        frame_h_edit = tk.Frame(frame_h_values, bg=pal['card'])
        frame_h_edit.pack(fill=tk.X, padx=5, pady=(0, 4))
        tk.Label(frame_h_edit, text='Wybrany generator:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).grid(row=0, column=0, sticky='w')
        self.entry_h_selected = tk.Entry(frame_h_edit, textvariable=self.h_selected_generator_var, state='readonly', width=28, font=self.font_code_small)
        self.entry_h_selected.grid(row=0, column=1, sticky='ew', padx=(6, 10))
        tk.Label(frame_h_edit, text='Nowe H [s]:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).grid(row=0, column=2, sticky='w')
        self.entry_h_edit = tk.Entry(frame_h_edit, textvariable=self.h_edit_value_var, width=12, font=self.font_code_small)
        self.entry_h_edit.grid(row=0, column=3, sticky='w', padx=(6, 10))
        self.entry_h_edit.bind('<Return>', lambda event: self.zastosuj_reczna_zmiane_h())
        self.utworz_przycisk(frame_h_edit, 'Zapisz H', self.zastosuj_reczna_zmiane_h, role='success').grid(row=0, column=4, padx=4)
        self.utworz_przycisk(frame_h_edit, 'Przywróć H z DYD', self.przywroc_h_dla_zaznaczonego, role='neutral').grid(row=0, column=5, padx=4)
        frame_h_edit.grid_columnconfigure(1, weight=1)
        self.right_content_panes.add(frame_h_values, weight=2)

        frame_tree = tk.LabelFrame(self.right_content_panes, text='Panel generatorów – dwuklik zmienia status', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        tk.Label(frame_tree, text='Dane z generatorów mają H z DYD, a zapotrzebowanie pochodzi z aktywnych rekordów LOAD DATA (suma pola mw).', fg=pal['success'], bg=pal['card'], font=self.font_ui_bold_small).pack(anchor='w', padx=5, pady=(0, 6))
        frame_ops = tk.Frame(frame_tree, bg=pal['card'])
        frame_ops.pack(fill=tk.X, padx=5, pady=(0, 6))
        self.utworz_przycisk(frame_ops, 'Przełącz zaznaczone', self.przelacz_zaznaczone_w_tree, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Załącz zaznaczone', lambda: self.ustaw_status_zaznaczonych_w_tree(1), role='success').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Wyłącz zaznaczone', lambda: self.ustaw_status_zaznaczonych_w_tree(0), role='warning').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Przywróć zaznaczone', self.przywroc_zaznaczone_w_tree, role='neutral').pack(side=tk.LEFT, padx=2)
        frame_col_filters = tk.LabelFrame(frame_tree, text='Filtry kolumnowe', padx=4, pady=6, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small)
        frame_col_filters.pack(fill=tk.X, padx=5, pady=(0, 6))
        filter_cols = [('kod', 'Kod'), ('grupa', 'Grupa / kod'), ('odm', 'ODM'), ('stan', 'Status'), ('h', 'H min-max'), ('pmax', 'S min-max'), ('ei', 'Ei min-max')]
        for i, (col_key, label) in enumerate(filter_cols):
            cell = tk.Frame(frame_col_filters, bg=pal['card'])
            cell.grid(row=0, column=i, padx=3, pady=2, sticky='ew')
            tk.Label(cell, text=label, font=self.font_ui_bold_small, bg=pal['card'], fg=pal['text']).pack(anchor='w')
            var = tk.StringVar()
            ent = tk.Entry(cell, textvariable=var, width=16, font=self.font_code_small)
            ent.pack(fill=tk.X)
            ent.bind('<KeyRelease>', self.on_column_filter_changed)
            self.column_filters[col_key] = var
            self.filter_entries[col_key] = ent
        for i in range(len(filter_cols)):
            frame_col_filters.grid_columnconfigure(i, weight=1)
        self.label_reczne = tk.Label(frame_tree, text='Brak wybranego modelu', fg=pal['accent'], bg=pal['card'], font=self.font_code)
        self.label_reczne.pack(anchor='w', padx=5, pady=5)
        tree_frame = tk.Frame(frame_tree, bg=pal['card'])
        tree_frame.pack(fill=tk.BOTH, expand=True)
        cols = ('stan', 'kod', 'grupa', 'odm', 'pmax', 'h', 'ei', 'udzial', 'zmiana')
        self.tree_generatory = ttk.Treeview(tree_frame, columns=cols, show='headings', height=9, selectmode='extended')
        self.tree_generatory.heading('stan', text='Stan', command=lambda: self.sort_tree_by_column('stan'))
        self.tree_generatory.heading('kod', text='Kod generatora', command=lambda: self.sort_tree_by_column('kod'))
        self.tree_generatory.heading('grupa', text='Grupa', command=lambda: self.sort_tree_by_column('grupa'))
        self.tree_generatory.heading('odm', text='ODM', command=lambda: self.sort_tree_by_column('odm'))
        self.tree_generatory.heading('pmax', text='S [MVA]', command=lambda: self.sort_tree_by_column('pmax'))
        self.tree_generatory.heading('h', text='H [s]', command=lambda: self.sort_tree_by_column('h'))
        self.tree_generatory.heading('ei', text='Ei [MVA·s]', command=lambda: self.sort_tree_by_column('ei'))
        self.tree_generatory.heading('udzial', text='Udział [%]', command=lambda: self.sort_tree_by_column('udzial'))
        self.tree_generatory.heading('zmiana', text='Ręczna zmiana', command=lambda: self.sort_tree_by_column('zmiana'))
        self.tree_generatory.column('stan', width=70, anchor='center')
        self.tree_generatory.column('kod', width=220, anchor='w')
        self.tree_generatory.column('grupa', width=95, anchor='center')
        self.tree_generatory.column('odm', width=105, anchor='center')
        self.tree_generatory.column('pmax', width=90, anchor='e')
        self.tree_generatory.column('h', width=80, anchor='e')
        self.tree_generatory.column('ei', width=110, anchor='e')
        self.tree_generatory.column('udzial', width=85, anchor='e')
        self.tree_generatory.column('zmiana', width=115, anchor='center')
        scrollbar_tree_y = ttk.Scrollbar(tree_frame, orient='vertical', command=self.tree_generatory.yview)
        scrollbar_tree_x = ttk.Scrollbar(tree_frame, orient='horizontal', command=self.tree_generatory.xview)
        self.tree_generatory.configure(yscrollcommand=scrollbar_tree_y.set, xscrollcommand=scrollbar_tree_x.set)
        self.tree_generatory.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_tree_y.pack(side=tk.RIGHT, fill=tk.Y)
        scrollbar_tree_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_generatory.bind('<Double-1>', self.on_tree_double_click)
        self.bind_scroll_area(self.tree_generatory)
        self.right_content_panes.add(frame_tree, weight=4)

        frame_rank_and_group = tk.Frame(self.right_content_panes, bg=pal['bg'])
        frame_ranking = tk.LabelFrame(frame_rank_and_group, text='Ranking TOP 10 wkładu generatorów (S×H)', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_ranking.pack(side=tk.TOP, fill=tk.X, expand=False)
        self.label_ranking = tk.Label(frame_ranking, text='Brak wybranego modelu', justify='left', anchor='w', font=self.font_code_small, bg=pal['card'], fg=pal['text'])
        self.label_ranking.pack(fill=tk.BOTH, expand=True, padx=5, pady=3)
        frame_group_chart = tk.LabelFrame(frame_rank_and_group, text='Wykresy udziałów, Pmax i inercji ODM', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_group_chart.pack(side=tk.TOP, fill=tk.BOTH, expand=True, pady=(10, 0))
        frame_group_controls = tk.Frame(frame_group_chart, bg=pal['card'])
        frame_group_controls.pack(fill=tk.X, pady=(0, 6))
        tk.Label(frame_group_controls, text='Obszar wykresu udziałów:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).pack(side=tk.LEFT)
        self.option_udzialy_scope = tk.OptionMenu(frame_group_controls, self.udzialy_scope_var, 'Cały model', *[self.ODM_CODE_TO_NAME[x] for x in self.ODM_ORDER], command=lambda *_: self.on_udzialy_scope_changed())
        self.option_udzialy_scope.pack(side=tk.LEFT, padx=(8, 0))
        self.stylizuj_option_menu(self.option_udzialy_scope)
        self.figura_grupy, (self.os_grupy, self.os_pmax_grupy, self.os_odm_inercja) = plt.subplots(3, 1, figsize=(10.8, 8.6), gridspec_kw={'height_ratios': [3.2, 3.2, 2.3]})
        self.figura_grupy.subplots_adjust(bottom=0.10, left=0.08, right=0.99, top=0.95, hspace=0.82)
        self.canvas_grupy = FigureCanvasTkAgg(self.figura_grupy, master=frame_group_chart)
        self.canvas_grupy.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.przygotuj_pusty_wykres_grup()
        frame_analiza = tk.LabelFrame(frame_rank_and_group, text='Rozszerzona analiza inercyjna', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_analiza.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.label_analiza_inercyjna = tk.Label(frame_analiza, text='Brak wybranego modelu', justify='left', anchor='w', font=self.font_code_small, fg=pal['accent'], bg=pal['card'])
        self.label_analiza_inercyjna.pack(fill=tk.X, padx=5, pady=3)
        self.right_content_panes.add(frame_rank_and_group, weight=4)

        frame_details_model = tk.LabelFrame(self.right_content_panes, text='Szczegóły modelu i komentarz', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.pole_szczegoly = scrolledtext.ScrolledText(frame_details_model, width=120, height=22, wrap=tk.WORD, font=self.font_code_small)
        self.pole_szczegoly.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        self.bind_scroll_area(self.pole_szczegoly)
        self.right_content_panes.add(frame_details_model, weight=3)
        panel_wykres = tk.LabelFrame(self.main_content_panes, text='Wykresy dobowe i monitoring importów', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        self.main_content_panes.add(panel_wykres, weight=2)
        frame_opcje = tk.LabelFrame(panel_wykres, text='Opcje widoku', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_opcje.pack(side=tk.TOP, fill=tk.X, pady=(0, 8))
        tk.Checkbutton(frame_opcje, text='Etykiety inercji', variable=self.pokaz_etykiety_inercji_var, command=self.on_chart_label_toggle).pack(side=tk.LEFT, padx=10)
        tk.Checkbutton(frame_opcje, text='Etykiety zapotrzebowania', variable=self.pokaz_etykiety_zapotrzebowania_var, command=self.on_chart_label_toggle).pack(side=tk.LEFT, padx=5)
        tk.Checkbutton(frame_opcje, text='Pokaż linię progu', variable=self.pokaz_linie_progu_var, command=self.aktualizuj_wykres).pack(side=tk.LEFT, padx=5)
        tk.Label(frame_opcje, text=f'Importy ręczne i bieżące są pokazywane jako osobne okna kroczące po {self.MANUAL_IMPORT_CHART_WINDOW}/{self.AUTO_IMPORT_CHART_WINDOW} modeli.', bg=pal['card'], fg=pal['muted'], font=self.font_ui_bold_small).pack(side=tk.LEFT, padx=(12, 6))
        tk.Checkbutton(frame_opcje, text='Tryb ciemny', variable=self.dark_mode_var, command=self.toggle_dark_mode).pack(side=tk.LEFT, padx=8)
        self.label_status = tk.Label(frame_opcje, text='Status: Krok 1 → Wczytaj EPC, Krok 2 → DYD, Krok 3 → Oblicz', fg=pal['accent'], bg=pal['card'], font=self.font_ui_bold)
        self.label_status.pack(side=tk.RIGHT, padx=10)
        self.figura, (self.os1, self.os_manual, self.os_auto) = plt.subplots(3, 1, figsize=(13.8, 12.4), gridspec_kw={'height_ratios': [3.1, 2.7, 2.7]})
        self.figura.subplots_adjust(bottom=0.08, hspace=0.50, left=0.10, right=0.91, top=0.97)
        self.os1_load = self.os1.twinx()
        self.os_manual_load = self.os_manual.twinx()
        self.os_auto_load = self.os_auto.twinx()
        self.canvas = FigureCanvasTkAgg(self.figura, master=panel_wykres)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.przygotuj_pusty_wykres()
        for entry in [self.entry_filtr, self.entry_wartosc_graniczna, self.entry_h_filter, self.entry_h_selected, self.entry_h_edit, *self.filter_entries.values()]:
            self.stylizuj_entry(entry)
        self.stylizuj_option_menu(self.option_tryb)
        self.stylizuj_option_menu(self.option_sort)
        for label in [self.label_status, self.label_maksimum_info, self.label_limit_ocena, self.label_porownanie, self.label_wskazniki, self.label_h_summary, self.label_reczne, self.label_ranking, self.label_analiza_inercyjna, self.label_wynikow, self.label_title, self.label_subtitle, self.label_monitor_folder]:
            self.label_neutral.add(label)
        self.apply_theme()

    def _safe_window_state(self, window):
        try:
            return window.state()
        except Exception:
            return 'normal'

    def register_managed_window(self, window, layout_key):
        self.managed_windows[layout_key] = window

        def _remember(*_args):
            self.remember_window_geometry(layout_key, window)

        window.bind('<Configure>', _remember, add='+')
        window.bind('<Map>', _remember, add='+')
        window.bind('<Destroy>', lambda *_: self.remember_window_geometry(layout_key, window), add='+')
        self.remember_window_geometry(layout_key, window)

    def remember_window_geometry(self, layout_key, window):
        try:
            cfg = self.ui_config if isinstance(self.ui_config, dict) else {}
            windows_cfg = cfg.setdefault('window_layouts', {})
            windows_cfg[layout_key] = {
                'geometry': window.geometry(),
                'state': self._safe_window_state(window),
            }
            self.ui_config = cfg
        except Exception:
            pass

    def apply_saved_window_geometry(self, window, layout_key, default_width, default_height):
        cfg = self.ui_config if isinstance(self.ui_config, dict) else {}
        win_cfg = cfg.get('window_layouts', {}).get(layout_key, {})
        geom = str(win_cfg.get('geometry', '') or '').strip()
        if geom:
            try:
                window.geometry(geom)
            except Exception:
                center_toplevel(window, default_width, default_height)
        else:
            center_toplevel(window, default_width, default_height)
        try:
            state = str(win_cfg.get('state', 'normal') or 'normal').strip()
            if state in {'normal', 'zoomed'}:
                window.state(state)
        except Exception:
            pass
        try:
            window.minsize(max(520, int(default_width * 0.55)), max(320, int(default_height * 0.55)))
        except Exception:
            pass

    def capture_pane_layout(self):
        pane_layout = {}
        try:
            if self.main_content_panes is not None:
                pane_layout['main_vertical_sash'] = int(self.main_content_panes.sashpos(0))
        except Exception:
            pass
        try:
            if self.panel_srodkowy is not None:
                pane_layout['main_horizontal_sash'] = int(self.panel_srodkowy.sashpos(0))
        except Exception:
            pass
        try:
            if self.left_model_panes is not None:
                pane_layout['left_model_sash_0'] = int(self.left_model_panes.sashpos(0))
                pane_layout['left_model_sash_1'] = int(self.left_model_panes.sashpos(1))
        except Exception:
            pass
        try:
            if self.right_content_panes is not None:
                pane_layout['right_content_sash_0'] = int(self.right_content_panes.sashpos(0))
                pane_layout['right_content_sash_1'] = int(self.right_content_panes.sashpos(1))
                pane_layout['right_content_sash_2'] = int(self.right_content_panes.sashpos(2))
        except Exception:
            pass
        return pane_layout

    def apply_saved_pane_layouts(self):
        cfg = self.ui_config if isinstance(self.ui_config, dict) else {}
        pane_layout = cfg.get('pane_layout', {})
        if not pane_layout:
            self.apply_default_pane_layouts()
            return
        try:
            if self.main_content_panes is not None and 'main_vertical_sash' in pane_layout:
                self.main_content_panes.sashpos(0, int(pane_layout['main_vertical_sash']))
        except Exception:
            pass
        try:
            if self.panel_srodkowy is not None and 'main_horizontal_sash' in pane_layout:
                self.panel_srodkowy.sashpos(0, int(pane_layout['main_horizontal_sash']))
        except Exception:
            pass
        try:
            if self.left_model_panes is not None:
                if 'left_model_sash_0' in pane_layout:
                    self.left_model_panes.sashpos(0, int(pane_layout['left_model_sash_0']))
                if 'left_model_sash_1' in pane_layout:
                    self.left_model_panes.sashpos(1, int(pane_layout['left_model_sash_1']))
        except Exception:
            pass
        try:
            if self.right_content_panes is not None:
                if 'right_content_sash_0' in pane_layout:
                    self.right_content_panes.sashpos(0, int(pane_layout['right_content_sash_0']))
                if 'right_content_sash_1' in pane_layout:
                    self.right_content_panes.sashpos(1, int(pane_layout['right_content_sash_1']))
                if 'right_content_sash_2' in pane_layout:
                    self.right_content_panes.sashpos(2, int(pane_layout['right_content_sash_2']))
        except Exception:
            pass

    def apply_default_pane_layouts(self):
        try:
            if self.main_content_panes is not None:
                total_h = max(self.main_content_panes.winfo_height(), 720)
                self.main_content_panes.sashpos(0, max(460, int(total_h * 0.58)))
        except Exception:
            pass
        try:
            if self.panel_srodkowy is not None:
                total_w = max(self.panel_srodkowy.winfo_width(), 1100)
                self.panel_srodkowy.sashpos(0, max(360, int(total_w * 0.28)))
        except Exception:
            pass
        try:
            if self.left_model_panes is not None:
                total_h = max(self.left_model_panes.winfo_height(), 360)
                self.left_model_panes.sashpos(0, max(120, int(total_h * 0.20)))
                self.left_model_panes.sashpos(1, max(270, int(total_h * 0.62)))
        except Exception:
            pass
        try:
            if self.right_content_panes is not None:
                total_h = max(self.right_content_panes.winfo_height(), 640)
                self.right_content_panes.sashpos(0, max(170, int(total_h * 0.23)))
                self.right_content_panes.sashpos(1, max(340, int(total_h * 0.53)))
                self.right_content_panes.sashpos(2, max(470, int(total_h * 0.80)))
        except Exception:
            pass

    def build_ui_config_snapshot(self):
        cfg = dict(self.ui_config) if isinstance(self.ui_config, dict) else {}
        cfg.update({
            'window_geometry': self.root.geometry(),
            'window_state': self._safe_window_state(self.root),
            'monitor_folder_path': self.monitor_folder_path,
            'threshold': self.entry_wartosc_graniczna.get().strip() if hasattr(self, 'entry_wartosc_graniczna') else '',
            'main_filter': self.entry_filtr.get().strip() if hasattr(self, 'entry_filtr') else '',
            'tryb_filtrowania': self.tryb_filtrowania_var.get() if hasattr(self, 'tryb_filtrowania_var') else 'Wszystkie',
            'sortowanie': self.sortowanie_var.get() if hasattr(self, 'sortowanie_var') else 'Kod generatora A-Z',
            'udzialy_scope': self.udzialy_scope_var.get() if hasattr(self, 'udzialy_scope_var') else 'Cały model',
            'pokaz_etykiety_inercji': self.pokaz_etykiety_inercji_var.get() if hasattr(self, 'pokaz_etykiety_inercji_var') else True,
            'pokaz_etykiety_zapotrzebowania': self.pokaz_etykiety_zapotrzebowania_var.get() if hasattr(self, 'pokaz_etykiety_zapotrzebowania_var') else True,
            'dark_mode': self.dark_mode_var.get(),
            'tree_sort_column': self.tree_sort_column,
            'tree_sort_reverse': self.tree_sort_reverse,
            'tree_columns': {col: self.tree_generatory.column(col, option='width') for col in self.tree_generatory['columns']} if hasattr(self, 'tree_generatory') else {},
            'column_filters': {col: var.get() for col, var in self.column_filters.items()},
            'secondary_chart_scope': self.secondary_chart_scope_var.get() if hasattr(self, 'secondary_chart_scope_var') else 'Modele ręczne',
            'pane_layout': self.capture_pane_layout(),
        })
        for key, window in list(self.managed_windows.items()):
            try:
                if window is not None and window.winfo_exists():
                    self.remember_window_geometry(key, window)
            except Exception:
                pass
        return cfg

    def apply_open_managed_window_layouts(self):
        for key, window in list(self.managed_windows.items()):
            try:
                if window is None or not window.winfo_exists():
                    continue
                self.apply_saved_window_geometry(window, key, max(window.winfo_width(), 680), max(window.winfo_height(), 420))
            except Exception:
                pass

    def schedule_apply_saved_pane_layouts(self, delay_ms=140):
        try:
            if self.pane_layout_apply_job is not None:
                self.root.after_cancel(self.pane_layout_apply_job)
        except Exception:
            pass
        try:
            self.pane_layout_apply_job = self.root.after(delay_ms, self._run_apply_saved_pane_layouts)
        except Exception:
            self.pane_layout_apply_job = None

    def _run_apply_saved_pane_layouts(self):
        self.pane_layout_apply_job = None
        try:
            if not self.root.winfo_exists():
                return
        except Exception:
            return
        self.apply_saved_pane_layouts()

    def safe_save_ui_config(self):
        try:
            self.save_ui_config()
        except Exception:
            pass

    def load_ui_config(self):
        if not self.ui_config_path.exists():
            self.ui_config = {}
            return
        try:
            with self.ui_config_path.open('r', encoding='utf-8') as f:
                self.ui_config = json.load(f)
        except Exception:
            self.ui_config = {}

    def save_ui_config(self):
        try:
            data = self.build_ui_config_snapshot()
            self.ui_config = data
            self.ui_config_path.parent.mkdir(parents=True, exist_ok=True)
            with self.ui_config_path.open('w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def apply_ui_config(self):
        cfg = self.ui_config or {}
        try:
            geom = cfg.get('window_geometry')
            if geom:
                self.root.geometry(geom)
        except Exception:
            pass
        try:
            state = str(cfg.get('window_state', 'normal') or 'normal').strip()
            if state in {'normal', 'zoomed'}:
                self.root.state(state)
        except Exception:
            pass
        self.monitor_folder_path = cfg.get('monitor_folder_path', self.monitor_folder_path)
        if self.monitor_folder_path:
            self.label_monitor_folder.config(text=f'Folder: {self.monitor_folder_path}')
        threshold = cfg.get('threshold', '')
        if threshold:
            self.entry_wartosc_graniczna.delete(0, tk.END)
            self.entry_wartosc_graniczna.insert(0, threshold)
        main_filter = cfg.get('main_filter', '')
        if main_filter:
            self.entry_filtr.delete(0, tk.END)
            self.entry_filtr.insert(0, main_filter)
        self.tryb_filtrowania_var.set(cfg.get('tryb_filtrowania', 'Wszystkie'))
        self.sortowanie_var.set(cfg.get('sortowanie', 'Kod generatora A-Z'))
        self.udzialy_scope_var.set(cfg.get('udzialy_scope', 'Cały model'))
        self.secondary_chart_scope_var.set(cfg.get('secondary_chart_scope', self.secondary_chart_scope_var.get()))
        self.pokaz_etykiety_inercji_var.set(cfg.get('pokaz_etykiety_inercji', True))
        self.pokaz_etykiety_zapotrzebowania_var.set(cfg.get('pokaz_etykiety_zapotrzebowania', True))
        self.dark_mode_var.set(cfg.get('dark_mode', False))
        self.tree_sort_column = cfg.get('tree_sort_column')
        self.tree_sort_reverse = cfg.get('tree_sort_reverse', False)
        widths = cfg.get('tree_columns', {})
        for col, width in widths.items():
            try:
                self.tree_generatory.column(col, width=width)
            except Exception:
                pass
        saved_filters = cfg.get('column_filters', {})
        for col, value in saved_filters.items():
            if col in self.column_filters:
                self.column_filters[col].set(value)
        self.apply_theme()
        self.apply_open_managed_window_layouts()
        self.schedule_apply_saved_pane_layouts()

    def save_current_layout_as_default(self):
        self.save_ui_config()
        messagebox.showinfo('Układ interfejsu', f'Zapisano bieżący układ jako domyślny:\n{self.ui_config_path}')

    def save_layout_to_file(self):
        path = filedialog.asksaveasfilename(title='Zapisz układ interfejsu', defaultextension='.json', initialfile='uklad_interfejsu_analiza_inercji.json', filetypes=[('JSON', '*.json')])
        if not path:
            return
        try:
            data = self.build_ui_config_snapshot()
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            messagebox.showinfo('Układ interfejsu', f'Zapisano układ do pliku:\n{path}')
        except Exception as e:
            messagebox.showerror('Układ interfejsu', str(e))

    def load_layout_from_file(self):
        path = filedialog.askopenfilename(title='Wczytaj układ interfejsu', filetypes=[('JSON', '*.json')])
        if not path:
            return
        try:
            with open(path, 'r', encoding='utf-8') as f:
                self.ui_config = json.load(f)
            self.apply_ui_config()
            self.apply_open_managed_window_layouts()
            self.aktualizuj_wykres()
            self.aktualizuj_status(f'Wczytano układ interfejsu z pliku: {os.path.basename(path)}')
        except Exception as e:
            messagebox.showerror('Układ interfejsu', str(e))

    def on_chart_label_toggle(self):
        self.aktualizuj_wykres()
        self.save_ui_config()

    def przeskaluj_wartosc_miedzy_osiami(self, wartosc, src_min, src_max, dst_min, dst_max):
        if abs(src_max - src_min) < 1e-09:
            return (dst_min + dst_max) / 2
        return dst_min + (wartosc - src_min) / (src_max - src_min) * (dst_max - dst_min)

    def wyznacz_offset_etykiety(self, indeks, wartosc, seria, *, prefer='up', other_scaled=None, is_edge=False):
        cycle = [(-22, 24), (22, 34), (-28, 44), (28, 54)] if prefer == 'up' else [(-22, -24), (22, -34), (-28, -44), (28, -54)]
        base_x, base_y = cycle[indeks % len(cycle)]
        prev_val = seria[indeks - 1] if indeks > 0 else None
        next_val = seria[indeks + 1] if indeks < len(seria) - 1 else None
        near_span = max(abs(wartosc) * 0.035, 700.0)
        if prev_val is not None and abs(wartosc - prev_val) < near_span:
            base_y += 12 if prefer == 'up' else -12
            base_x += -10
        if next_val is not None and abs(wartosc - next_val) < near_span:
            base_y += 12 if prefer == 'up' else -12
            base_x += 10
        if other_scaled is not None and abs(wartosc - other_scaled) < near_span:
            base_y += 18 if prefer == 'up' else -18
        if is_edge:
            base_x = 26 if indeks == 0 else -26
            base_y += 10 if prefer == 'up' else -10
        if prev_val is not None and next_val is not None:
            if wartosc >= prev_val and wartosc >= next_val:
                base_y += 10 if prefer == 'up' else -6
            elif wartosc <= prev_val and wartosc <= next_val:
                base_y += 6 if prefer == 'up' else -10
        return (base_x, base_y)

    def dodaj_etykiete_punktu(self, ax, x, y, text, x_off, y_off, *, fg, facecolor, edgecolor, valign, fontsize=8):
        x_min, x_max = ax.get_xlim()
        y_min, y_max = ax.get_ylim()
        x_range = (x_max - x_min) or 1.0
        y_range = (y_max - y_min) or 1.0
        x_frac = (x - x_min) / x_range
        y_frac = (y - y_min) / y_range

        x_off = max(-28, min(28, x_off))
        y_off = max(-24, min(24, y_off))

        if x_frac <= 0.10:
            x_off = abs(x_off) + 8
        elif x_frac >= 0.90:
            x_off = -(abs(x_off) + 8)

        if y_frac >= 0.88:
            y_off = -(abs(y_off) + 10)
            valign = 'top'
        elif y_frac <= 0.12:
            y_off = abs(y_off) + 10
            valign = 'bottom'

        ha = 'left' if x_off > 10 else 'right' if x_off < -10 else 'center'
        ann = ax.annotate(
            text,
            xy=(x, y),
            xytext=(x_off, y_off),
            textcoords='offset points',
            ha=ha,
            va=valign,
            fontsize=fontsize,
            fontweight='bold',
            color=fg,
            bbox=dict(boxstyle='round,pad=0.25', facecolor=facecolor, edgecolor=edgecolor, alpha=0.92),
            arrowprops=dict(arrowstyle='-', color=edgecolor, lw=0.9, shrinkA=4, shrinkB=4, alpha=0.85),
            annotation_clip=True,
            zorder=6,
        )
        ann.set_clip_on(True)
        try:
            ann.set_clip_path(ax.patch)
        except Exception:
            pass

    def get_auto_monitor_chart_data(self, limit=None):
        rows = []
        for auto_id in self.auto_order:
            model = self.modele_auto.get(auto_id, {})
            opis = model.get('opis', {})
            label = opis.get('godzina', '--:--')
            data_txt = opis.get('data', '')
            if data_txt:
                label = f'{label}\n{data_txt}'
            rows.append({
                'auto_id': auto_id,
                'label': label,
                'godzina': opis.get('godzina', '--:--'),
                'data': data_txt,
                'inercja': float(model.get('inercja', 0.0)),
                'zapotrzebowanie': float(model.get('zapotrzebowanie', 0.0)),
            })
        max_rows = int(limit or self.AUTO_IMPORT_CHART_WINDOW)
        return rows[-max_rows:]

    def get_manual_import_chart_data(self, limit=None):
        rows = []
        for model_id in self.reczne_order:
            model = self.modele_reczne.get(model_id, {})
            opis = model.get('opis', {})
            label = opis.get('godzina', '--:--')
            data_txt = opis.get('data', '')
            if data_txt:
                label = f'{label}\n{data_txt}'
            rows.append({
                'model_id': model_id,
                'label': label,
                'godzina': opis.get('godzina', '--:--'),
                'data': data_txt,
                'inercja': float(model.get('inercja', 0.0)),
                'zapotrzebowanie': float(model.get('zapotrzebowanie', 0.0)),
            })
        max_rows = int(limit or self.MANUAL_IMPORT_CHART_WINDOW)
        return rows[-max_rows:]

    def get_manual_chart_payload(self):
        return {
            'title': f'Ręcznie importowane modele – okno kroczące {self.MANUAL_IMPORT_CHART_WINDOW} modeli',
            'xlabel': 'Kolejność ręcznych importów',
            'empty_message': 'Brak ręcznie importowanych modeli',
            'rows': self.get_manual_import_chart_data(self.MANUAL_IMPORT_CHART_WINDOW),
            'window_size': self.MANUAL_IMPORT_CHART_WINDOW,
        }

    def get_auto_chart_payload(self):
        return {
            'title': f'Monitoring modeli bieżących – okno kroczące {self.AUTO_IMPORT_CHART_WINDOW} modeli',
            'xlabel': 'Kolejność modeli bieżących',
            'empty_message': 'Brak modeli bieżących z monitorowanego folderu',
            'rows': self.get_auto_monitor_chart_data(self.AUTO_IMPORT_CHART_WINDOW),
            'window_size': self.AUTO_IMPORT_CHART_WINDOW,
        }

    def _draw_inertia_load_chart(self, ax, ax_load, xs, wartosci, zapotrzebowania, xlabels, *, title, xlabel, empty_message='Brak danych do wykresu', left_label_coords=None, window_size=None):
        pal = self.get_palette()
        ax.clear()
        ax_load.clear()
        ax.set_title(title, pad=18)
        ax.set_xlabel(xlabel)
        ax.set_ylabel('Inercja [MVA·s]')
        self.stylizuj_os(ax, grid=True, grid_axis='both')
        if left_label_coords is not None:
            try:
                ax.yaxis.set_label_coords(*left_label_coords)
            except Exception:
                pass
        ax_load.set_ylabel('Zapotrzebowanie [MW]', labelpad=18)
        self.stylizuj_os_prawa(ax_load, pal['warning'])
        ax.set_zorder(2)
        ax.patch.set_alpha(1.0)
        ax_load.set_zorder(3)
        ax_load.patch.set_alpha(0.0)
        display_count = max(int(window_size or 0), len(xs), 1)
        if not xs:
            ax.set_xlim(-0.5, display_count - 0.5)
            ax.set_ylim(0, 1)
            ax_load.set_xlim(-0.5, display_count - 0.5)
            ax_load.set_ylim(0, 1)
            ax.text(0.5, 0.5, empty_message, ha='center', va='center', transform=ax.transAxes, color=pal['muted'])
            return
        point_count = len(xs)
        if window_size is not None and window_size > point_count:
            offset = window_size - point_count
            plot_xs = [idx + offset for idx in range(point_count)]
        else:
            plot_xs = list(xs)
        display_labels = ['' for _ in range(display_count)]
        for pos, label in zip(plot_xs, xlabels):
            if 0 <= pos < display_count:
                display_labels[pos] = label
        tick_step = max(1, point_count // 14) if point_count > 14 else 1
        label_step = 1 if point_count <= 18 else 2 if point_count <= 36 else 4 if point_count <= 72 else 6
        ax.plot(plot_xs, wartosci, 'o-', linewidth=3, markersize=8, color=pal['line_main'], markerfacecolor=pal['selection'], markeredgecolor=pal['accent_active'], label='Inercja [MVA·s]', zorder=3)
        ax.fill_between(plot_xs, [0.0 for _ in plot_xs], wartosci, color=pal['line_main'], alpha=0.08, zorder=2)
        ax_load.plot(plot_xs, zapotrzebowania, 's--', linewidth=2.1, markersize=5.3 if point_count <= 36 else 4.5, color=pal['warning'], markerfacecolor=pal['card'], markeredgecolor=pal['warning_active'], label='Zapotrzebowanie [MW]', zorder=5, alpha=0.98)
        prog = self.pobierz_wartosc_graniczna()
        if self.pokaz_linie_progu_var.get() and prog is not None:
            ax.axhline(prog, linestyle='-', linewidth=1.2, color=pal['danger'], label=f'Próg alarmowy: {prog:.0f}', zorder=1)
        tick_candidates = [pos for pos, label in enumerate(display_labels) if label]
        tick_positions = tick_candidates[::tick_step] if tick_candidates else []
        tick_labels = [display_labels[pos] for pos in tick_positions]
        if tick_candidates and tick_candidates[-1] not in tick_positions:
            tick_positions = tick_positions + [tick_candidates[-1]]
            tick_labels = tick_labels + [display_labels[tick_candidates[-1]]]
        ax.set_xticks(tick_positions)
        ax.set_xticklabels(tick_labels, rotation=55 if point_count > 72 else 42 if point_count > 32 else 18, ha='right', fontsize=8.2)
        ax.set_xlim(-0.5, display_count - 0.5)
        ax.margins(x=0.02)
        max_source = list(wartosci)
        if prog is not None:
            max_source.append(prog)
        max_val = max(max_source) if max_source else 0.0
        min_val = min(wartosci) if wartosci else 0.0
        y_pad = max((max_val - min_val) * 0.18, max(max_val, 1.0) * 0.12, 50.0)
        ax.set_ylim(max(0.0, min_val - y_pad * 0.2), max_val + y_pad if max_val > 0 else 1.0)
        zap_dost = [z for z in zapotrzebowania if z > 0]
        if zap_dost:
            min_load = min(zap_dost)
            max_load = max(zap_dost)
            if abs(max_load - min_load) < 1e-09:
                pad = max(1.0, abs(max_load) * 0.2)
            else:
                pad = max(1.0, (max_load - min_load) * 0.15)
            dol = min(0.0, min_load - pad)
            gora = max(1.0, max_load + pad)
        else:
            dol, gora = (0.0, 1.0)
        ax_load.set_xlim(-0.5, display_count - 0.5)
        ax_load.set_ylim(dol, gora)
        ax_load.margins(x=0.02)
        wartosci_dodatnie = [(plot_xs[idx], val, idx) for idx, val in enumerate(wartosci) if val > 0]
        if wartosci_dodatnie:
            godz_min, val_min, idx_min = min(wartosci_dodatnie, key=lambda x: x[1])
            godz_max, val_max, idx_max = max(wartosci_dodatnie, key=lambda x: x[1])
            ax.plot(godz_min, val_min, 'o', markersize=14, color=pal['danger'], label=f'Minimum: {xlabels[idx_min].splitlines()[0]}', zorder=4)
            ax.plot(godz_max, val_max, 'o', markersize=14, color=pal['success'], label=f'Maksimum: {xlabels[idx_max].splitlines()[0]}', zorder=4)
            self.dodaj_etykiete_punktu(ax, godz_min, val_min, f'MIN {xlabels[idx_min].splitlines()[0]}\n{val_min:.0f}', -42 if int(godz_min) % 2 == 0 else -30, -34, fg=pal['danger'], facecolor=pal['card'], edgecolor=pal['danger'], valign='top', fontsize=9)
            self.dodaj_etykiete_punktu(ax, godz_max, val_max, f'MAX {xlabels[idx_max].splitlines()[0]}\n{val_max:.0f}', 42 if int(godz_max) % 2 == 0 else 30, 30, fg=pal['success'], facecolor=pal['card'], edgecolor=pal['success'], valign='bottom', fontsize=9)
        y1_min, y1_max = ax.get_ylim()
        y2_min, y2_max = ax_load.get_ylim()
        load_scaled_values = [self.przeskaluj_wartosc_miedzy_osiami(load_val, y2_min, y2_max, y1_min, y1_max) for load_val in zapotrzebowania]
        if self.pokaz_etykiety_inercji_var.get():
            for idx, (x, val) in enumerate(zip(plot_xs, wartosci)):
                if val <= 0:
                    continue
                if idx % label_step != 0 and idx not in {0, len(plot_xs) - 1}:
                    continue
                x_off, y_off = self.wyznacz_offset_etykiety(idx, val, wartosci, prefer='up', other_scaled=load_scaled_values[idx], is_edge=(idx == 0 or idx == len(plot_xs) - 1))
                self.dodaj_etykiete_punktu(ax, x, val, f'{val:.0f}', x_off, y_off, fg=pal['text'], facecolor=pal['header'], edgecolor=pal['border'], valign='bottom')
        if self.pokaz_etykiety_zapotrzebowania_var.get():
            for idx, (x, load_val) in enumerate(zip(plot_xs, zapotrzebowania)):
                if load_val <= 0:
                    continue
                if idx % label_step != 0 and idx not in {0, len(plot_xs) - 1}:
                    continue
                x_off, y_off = self.wyznacz_offset_etykiety(idx, load_scaled_values[idx], load_scaled_values, prefer='down', other_scaled=wartosci[idx], is_edge=(idx == 0 or idx == len(plot_xs) - 1))
                self.dodaj_etykiete_punktu(ax_load, x, load_val, f'{load_val:.0f}', x_off, y_off, fg=pal['warning_active'], facecolor=pal['card_alt'], edgecolor=pal['warning'], valign='top')
        h1, l1 = ax.get_legend_handles_labels()
        h2, l2 = ax_load.get_legend_handles_labels()
        legend = ax.legend(h1 + h2, l1 + l2, loc='upper left', ncol=1 if point_count <= 24 else 2, fontsize=8.5, borderaxespad=0.8)
        if legend:
            legend.get_frame().set_facecolor(pal['card'])
            legend.get_frame().set_edgecolor(pal['border'])
            for txt in legend.get_texts():
                txt.set_color(pal['text'])

    def konfiguruj_email_alarm(self):
        EmailAlarmConfigWindow(self.root, self.email_alarm)

    def pokaz_historie_alarmow(self):
        try:
            if self.alarm_history_window is not None and self.alarm_history_window.winfo_exists():
                self.alarm_history_window.lift()
                self.alarm_history_window.focus_force()
                self.refresh_alarm_history_window()
                return
        except Exception:
            pass
        self.alarm_history_window = AlarmHistoryWindow(self.root, self)
        self.alarm_history_window.protocol('WM_DELETE_WINDOW', self.close_alarm_history_window)

    def dodaj_do_historii_alarmow(self, **kwargs):
        self.alarm_history.append(kwargs)
        if len(self.alarm_history) > 5000:
            self.alarm_history = self.alarm_history[-5000:]
        self.refresh_alarm_history_window()

    def sprawdz_i_wyslij_alarm_progu(self, model_type, key, wartosc):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            return
        if not self.email_alarm.should_trigger(prog, wartosc):
            return
        opis = self.get_model_meta(model_type, key)
        trigger_mode = self.email_alarm.config.trigger_mode
        token = (model_type, str(key), opis.get('plik', ''), opis.get('data', ''), opis.get('godzina', ''), trigger_mode, f'{prog:.6f}', f'{wartosc:.6f}')
        if token in self.sent_alarm_tokens:
            return
        self.sent_alarm_tokens.add(token)
        opis_txt = 'Przekroczenie progu' if trigger_mode == 'above' else 'Spadek poniżej progu'
        base_entry = {'czas_alarmu': datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 'trigger_mode': trigger_mode, 'model_type': model_type, 'model_type_label': self.get_model_type_label(model_type), 'data_modelu': opis.get('data', 'brak daty'), 'godzina_modelu': opis.get('godzina', '--:--'), 'wartosc': wartosc, 'prog': prog, 'source_file': opis.get('plik', ''), 'mail_attempted': False, 'mail_ok': False, 'mail_status': 'Powiadomienia wyłączone', 'opis': opis_txt}
        if not self.email_alarm.config.enabled:
            self.dodaj_do_historii_alarmow(**base_entry)
            self.aktualizuj_status(f"⚠️ Alarm spełniony dla modelu {opis.get('data', '')} {opis.get('godzina', '')}, ale powiadomienia e-mail są wyłączone")
            return
        base_entry['mail_attempted'] = True

        def worker():
            try:
                ok, msg = self.email_alarm.send_alarm(prog=prog, wartosc=wartosc, data_modelu=opis.get('data', 'brak daty'), godzina_modelu=opis.get('godzina', '--:--'), source_file=opis.get('plik', ''))
                entry = dict(base_entry)
                entry['mail_ok'] = ok
                entry['mail_status'] = msg
                self.root.after(0, lambda e=entry, m=msg, o=ok: self._finalize_alarm_send(e, m, o))
            except Exception as e:
                entry = dict(base_entry)
                entry['mail_ok'] = False
                entry['mail_status'] = str(e)
                self.root.after(0, lambda en=entry, err=e: self._finalize_alarm_send(en, f'Błąd wysyłki e-mail: {err}', False))
        threading.Thread(target=worker, daemon=True).start()

    def sprawdz_alarm_dla_modelu(self, model_type, key):
        wartosc = self.get_model_energy_now(model_type, key)
        if wartosc <= 0:
            return
        self.sprawdz_i_wyslij_alarm_progu(model_type, key, wartosc)

    def sprawdz_alarmy_dla_wszystkich_modeli(self):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            return
        for godzina in sorted(self.wszystkie_generatory.keys()):
            self.sprawdz_alarm_dla_modelu('manual', godzina)
        for model_id in self.reczne_order:
            self.sprawdz_alarm_dla_modelu('manual_batch', model_id)
        for auto_id in self.auto_order:
            self.sprawdz_alarm_dla_modelu('auto', auto_id)

    def _finalize_alarm_send(self, entry, status_msg, ok):
        self.dodaj_do_historii_alarmow(**entry)
        self.aktualizuj_status(f'📧 {status_msg}' if ok else f'⚠️ {status_msg}')

    def close_alarm_history_window(self):
        try:
            if self.alarm_history_window is not None and self.alarm_history_window.winfo_exists():
                self.alarm_history_window.destroy()
        except Exception:
            pass
        self.alarm_history_window = None

    def refresh_alarm_history_window(self):
        try:
            if self.alarm_history_window is not None and self.alarm_history_window.winfo_exists():
                self.configure_widget_tree(self.alarm_history_window)
                self.alarm_history_window.refresh()
                return
        except Exception:
            pass
        self.alarm_history_window = None

    def get_current_model_ref(self):
        if self.aktywny_typ_modelu == 'manual' and self.aktualna_godzina is not None:
            return ('manual', self.aktualna_godzina)
        if self.aktywny_typ_modelu == 'manual_batch' and self.aktualny_reczny_id is not None:
            return ('manual_batch', self.aktualny_reczny_id)
        if self.aktywny_typ_modelu == 'auto' and self.aktualny_auto_id is not None:
            return ('auto', self.aktualny_auto_id)
        return (None, None)

    def get_model_type_label(self, model_type):
        if model_type == 'manual_batch':
            return 'Reczny'
        return {'manual': '24h', 'auto': 'Bieżące'}.get(model_type, str(model_type or '-'))

    def clear_other_model_selections(self, keep=None):
        widgets = {
            'manual': getattr(self, 'listbox_modele', None),
            'manual_batch': getattr(self, 'listbox_modele_reczne', None),
            'auto': getattr(self, 'listbox_modele_auto', None),
        }
        for key, widget in widgets.items():
            if key == keep or widget is None:
                continue
            try:
                widget.selection_clear(0, tk.END)
            except Exception:
                pass

    def build_session_model_entry(self, gen_data, *, data_txt='', godz_txt='', sciezka='', source_tag='', godzina_idx=None):
        statusy_org = {g['reg_name']: g['status'] for g in gen_data}
        return {
            'generatory': gen_data,
            'statusy_oryginalne': statusy_org,
            'opis': {
                'data': data_txt if data_txt else 'brak daty',
                'godzina': godz_txt if godz_txt else '--:--',
                'plik': os.path.basename(sciezka) if sciezka else '',
                'sciezka': sciezka,
                'zrodlo_modelu': source_tag,
                'godzina_idx': godzina_idx,
            },
            'inercja': 0.0,
            'inercja_oryginalna': 0.0,
            'odm_metrics': self.build_empty_odm_metrics(),
            'rocof': 0.0,
            'largest_loss': 0.0,
            'liczba_zalaczonych': 0,
            'zapotrzebowanie': 0.0,
        }

    def _model_mapping(self, model_type, field):
        manual_maps = {'generatory': self.wszystkie_generatory, 'statusy_oryginalne': self.statusy_oryginalne, 'opis': self.opis_modeli, 'inercja': self.inercja_na_godzine, 'inercja_oryginalna': self.inercja_oryginalna_na_godzine, 'odm_metrics': self.odm_metrics_na_godzine, 'rocof': self.rocof_na_godzine, 'largest_loss': self.largest_loss_na_godzine, 'liczba_zalaczonych': self.liczba_zalaczonych_na_godzine, 'zapotrzebowanie': self.zapotrzebowanie_na_godzine}
        defaults = {'generatory': [], 'statusy_oryginalne': {}, 'opis': {}, 'inercja': 0.0, 'inercja_oryginalna': 0.0, 'odm_metrics': {}, 'rocof': 0.0, 'largest_loss': 0.0, 'liczba_zalaczonych': 0, 'zapotrzebowanie': 0.0}
        return (manual_maps.get(field) if model_type == 'manual' else None, defaults[field])

    def get_model_value(self, model_type, key, field):
        mapping, default = self._model_mapping(model_type, field)
        if mapping is not None:
            return mapping.get(key, default)
        if model_type == 'manual_batch':
            return self.modele_reczne.get(key, {}).get(field, default)
        if model_type == 'auto':
            return self.modele_auto.get(key, {}).get(field, default)
        return default

    def set_model_value(self, model_type, key, field, value):
        mapping, _ = self._model_mapping(model_type, field)
        if mapping is not None:
            mapping[key] = value
        elif model_type == 'manual_batch' and key in self.modele_reczne:
            self.modele_reczne[key][field] = value
        elif model_type == 'auto' and key in self.modele_auto:
            self.modele_auto[key][field] = value

    def get_model_generators(self, model_type, key):
        return self.get_model_value(model_type, key, 'generatory')

    def get_model_original_statuses(self, model_type, key):
        return self.get_model_value(model_type, key, 'statusy_oryginalne')

    def get_model_meta(self, model_type, key):
        return self.get_model_value(model_type, key, 'opis')

    def get_model_energy_now(self, model_type, key):
        return self.get_model_value(model_type, key, 'inercja')

    def get_model_energy_original(self, model_type, key):
        return self.get_model_value(model_type, key, 'inercja_oryginalna')

    def get_model_demand_now(self, model_type, key):
        return self.get_model_value(model_type, key, 'zapotrzebowanie')

    def get_model_odm_metrics(self, model_type, key):
        value = self.get_model_value(model_type, key, 'odm_metrics')
        return value if value else self.build_empty_odm_metrics()

    def get_model_rocof_now(self, model_type, key):
        return self.get_model_value(model_type, key, 'rocof')

    def get_model_largest_loss(self, model_type, key):
        return self.get_model_value(model_type, key, 'largest_loss')

    def set_model_energy_now(self, model_type, key, value):
        self.set_model_value(model_type, key, 'inercja', value)

    def set_model_energy_original(self, model_type, key, value):
        self.set_model_value(model_type, key, 'inercja_oryginalna', value)

    def set_model_odm_metrics(self, model_type, key, value):
        self.set_model_value(model_type, key, 'odm_metrics', value)

    def set_model_rocof_now(self, model_type, key, value):
        self.set_model_value(model_type, key, 'rocof', value)

    def set_model_largest_loss(self, model_type, key, value):
        self.set_model_value(model_type, key, 'largest_loss', value)

    def set_model_connected_count(self, model_type, key, value):
        self.set_model_value(model_type, key, 'liczba_zalaczonych', value)

    def is_model_epc_exportable(self, model_type, key):
        meta = self.get_model_meta(model_type, key)
        sciezka = meta.get('sciezka', '')
        return bool(sciezka and os.path.exists(sciezka))

    def wybierz_model_manual(self, event=None):
        selection = self.listbox_modele.curselection()
        if not selection:
            return
        try:
            linia = self.listbox_modele.get(selection[0])
            match = re.match('(?:\\[MIN\\]\\s*)?(?:\\[MAX\\]\\s*)?(\\d{2}):30', linia)
            if not match:
                return
            godzina = int(match.group(1))
            self.aktywny_typ_modelu = 'manual'
            self.aktualna_godzina = godzina
            self.aktualny_reczny_id = None
            self.aktualny_auto_id = None
            self.clear_other_model_selections(keep='manual')
            if not self.obliczenia_wykonane:
                self.pole_szczegoly.delete('1.0', tk.END)
                self.pole_szczegoly.insert(tk.END, '⚠️ Najpierw: 1. EPC → 2. DYD → 3. OBLICZ')
                self.przygotuj_pusty_wykres_grup('Brak obliczeń dla modelu')
                return
            self.odswiez_aktywny_model()
        except Exception:
            pass

    def wybierz_model_auto(self, event=None):
        selection = self.listbox_modele_auto.curselection()
        if not selection:
            return
        try:
            idx = selection[0]
            if idx < 0 or idx >= len(self.auto_order):
                return
            auto_id = self.auto_order[idx]
            self.aktywny_typ_modelu = 'auto'
            self.aktualny_auto_id = auto_id
            self.aktualna_godzina = None
            self.aktualny_reczny_id = None
            self.clear_other_model_selections(keep='auto')
            self.odswiez_aktywny_model()
        except Exception:
            pass

    def wybierz_model_reczny(self, event=None):
        selection = self.listbox_modele_reczne.curselection()
        if not selection:
            return
        try:
            idx = selection[0]
            if idx < 0 or idx >= len(self.reczne_order):
                return
            model_id = self.reczne_order[idx]
            self.aktywny_typ_modelu = 'manual_batch'
            self.aktualny_reczny_id = model_id
            self.aktualna_godzina = None
            self.aktualny_auto_id = None
            self.clear_other_model_selections(keep='manual_batch')
            self.odswiez_aktywny_model()
        except Exception:
            pass

    def wyciagnij_date_i_godzine_z_nazwy(self, sciezka):
        nazwa = os.path.basename(sciezka)
        stem = Path(nazwa).stem
        nazwa_lower = stem.lower()
        data_txt = ''
        matched_span = None
        date_patterns = ['(?<!\\d)(?P<y>(?:20\\d{2}|19\\d{2}))[\\s,._-](?P<m>0?[1-9]|1[0-2])[\\s,._-](?P<d>0?[1-9]|[12]\\d|3[01])(?!\\d)', '(?<!\\d)(?P<d>0?[1-9]|[12]\\d|3[01])[\\s,._-](?P<m>0?[1-9]|1[0-2])[\\s,._-](?P<y>(?:20\\d{2}|19\\d{2}))(?!\\d)', '(?<!\\d)(?P<y>(?:20\\d{2}|19\\d{2}))(?P<m>0[1-9]|1[0-2])(?P<d>0[1-9]|[12]\\d|3[01])(?!\\d)', '(?<!\\d)(?P<d>0[1-9]|[12]\\d|3[01])(?P<m>0[1-9]|1[0-2])(?P<y>(?:20\\d{2}|19\\d{2}))(?!\\d)']
        candidates = []
        for wz in date_patterns:
            for m in re.finditer(wz, nazwa_lower):
                try:
                    d = int(m.group('d'))
                    mm = int(m.group('m'))
                    yyyy = int(m.group('y'))
                    candidates.append((m.start(), m.end(), f'{d:02d}.{mm:02d}.{yyyy:04d}'))
                except Exception:
                    continue
        if candidates:
            candidates.sort(key=lambda x: (x[0], -(x[1] - x[0])))
            start_idx, end_idx, data_txt = candidates[0]
            matched_span = (start_idx, end_idx)
        scope_for_hour = nazwa_lower
        if matched_span:
            a, b = matched_span
            scope_for_hour = nazwa_lower[:a] + ' ' + nazwa_lower[b:]
        godzina = None
        time_patterns = ['(?<!\\d)([01]?\\d|2[0-3])[\\s,._:-]?30(?!\\d)', '(?<!\\d)([01]?\\d|2[0-3])30(?!\\d)']
        for wz in time_patterns:
            matches = list(re.finditer(wz, scope_for_hour))
            if not matches:
                continue
            for match in reversed(matches):
                try:
                    hh = int(match.group(1))
                except Exception:
                    continue
                if 0 <= hh <= 23:
                    godzina = hh
                    break
            if godzina is not None:
                break
        godz_txt = f'{godzina:02d}:30' if godzina is not None else ''
        return (data_txt, godz_txt, godzina)

    def parsuj_linie_generatora(self, line):
        gen = self.parsuj_rekord_generatora(line)
        return [gen] if gen else []

    def build_empty_odm_metrics(self):
        return {
            code: {
                'code': code,
                'name': name,
                'energy': 0.0,
                'rocof_energy': 0.0,
                'share_pct': 0.0,
                'connected_count': 0,
                'largest_loss': 0.0,
                'rocof': 0.0,
            }
            for code, name in self.ODM_CODE_TO_NAME.items()
        }

    def formatuj_odm_z_kodu(self, code):
        return self.ODM_CODE_TO_NAME.get(str(code or '').strip(), 'Nieprzypisany')

    def formatuj_odm_generatora(self, gen):
        code = str(gen.get('odm_code', '') or '').strip()
        if not code:
            return str(gen.get('odm_name') or 'Nieprzypisany')
        return str(gen.get('odm_name') or self.ODM_CODE_TO_NAME.get(code, f'ODM {code}'))

    def tokenizuj_segment_epc(self, text):
        tokens = []
        for match in re.finditer('"([^"]*)"|(\\S+)', text or ''):
            token = match.group(1) if match.group(1) is not None else match.group(2)
            token = (token or '').strip()
            if token:
                tokens.append(token)
        return tokens

    def normalizuj_kod_odm(self, value):
        raw = str(value or '').strip()
        if not raw:
            return ''
        try:
            return str(int(float(raw.replace(',', '.'))))
        except Exception:
            return raw

    def wyciagnij_info_odm_generatora(self, full_record, zone_code=''):
        try:
            zone_code = self.normalizuj_kod_odm(zone_code)
            if zone_code in self.ODM_CODE_TO_NAME:
                return {'odm_code': zone_code, 'odm_name': self.ODM_CODE_TO_NAME[zone_code], 'odm_source': f'ZONE={zone_code}'}
            if ':' not in full_record:
                return {'odm_code': '', 'odm_name': 'Nieprzypisany', 'odm_source': ''}
            _, prawa = full_record.split(':', 1)
            kandydaci = [x.strip() for x in re.findall('"([^"]*)"', prawa) if x.strip()]
            if not kandydaci:
                kandydaci = self.tokenizuj_segment_epc(prawa)
            for kandydat in kandydaci:
                normalized = re.sub('\\s+', '', kandydat.upper())
                if len(normalized) < 5 or (not any((ch.isalpha() for ch in normalized))):
                    continue
                odm_code = normalized[4]
                if odm_code in self.ODM_CODE_TO_NAME:
                    return {'odm_code': odm_code, 'odm_name': self.ODM_CODE_TO_NAME[odm_code], 'odm_source': normalized}
        except Exception:
            pass
        return {'odm_code': '', 'odm_name': 'Nieprzypisany', 'odm_source': ''}

    def oblicz_rocof(self, inertia_mws, power_loss_mw):
        inertia = float(inertia_mws or 0.0)
        power_loss = float(power_loss_mw or 0.0)
        if inertia <= 0 or power_loss <= 0:
            return 0.0
        return self.NOMINAL_FREQUENCY_HZ * power_loss / (2.0 * inertia)

    def get_generator_apparent_power(self, gen):
        pmax = float(gen.get('pmax', 0.0) or 0.0)
        qmax = float(gen.get('qmax', 0.0) or 0.0)
        return (pmax ** 2 + qmax ** 2) ** 0.5

    def get_generator_active_power_for_rocof(self, gen):
        pgen = float(gen.get('pgen', 0.0) or 0.0)
        if pgen > 0:
            return pgen
        return max(float(gen.get('pmax', 0.0) or 0.0), 0.0)

    def is_generator_rocof_eligible(self, gen):
        area = self.normalizuj_kod_odm(gen.get('area', '1'))
        return area == '1'

    def get_generator_online_status(self, gen, model_type, key, use_original=False):
        statusy = self.get_model_original_statuses(model_type, key) if use_original else None
        if statusy is None:
            return int(gen.get('status', 0))
        return int(statusy.get(gen.get('reg_name', ''), gen.get('status', 0)))

    def build_rocof_analysis(self, model_type, key, custom_delta_p=None):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return None
        total_energy = 0.0
        online_rows = []
        eligible_count = 0
        inertia_count = 0
        for gen in gen_list:
            status = self.get_generator_online_status(gen, model_type, key, use_original=False)
            if status != 1:
                continue
            if not self.is_generator_rocof_eligible(gen):
                continue
            eligible_count += 1
            reg_name = gen.get('reg_name', '')
            delta_p = self.get_generator_active_power_for_rocof(gen)
            h_value = float(self.get_h_value_for_generator(gen, prefer_original=False))
            s_apparent = self.get_generator_apparent_power(gen)
            energy = s_apparent * h_value if h_value > 0 else 0.0
            if energy > 0:
                total_energy += energy
                inertia_count += 1
            online_rows.append({
                'reg_name': reg_name,
                'group_label': self.formatuj_grupe_generatora(gen),
                'delta_p': delta_p,
                'energy': energy,
                'h': h_value,
                's_apparent': s_apparent,
            })
        if not online_rows:
            return None
        online_rows.sort(key=lambda row: row['delta_p'], reverse=True)
        largest_unit = dict(online_rows[0])
        default_dp = float(custom_delta_p if custom_delta_p is not None else largest_unit['delta_p'])
        system_rocof = self.oblicz_rocof(total_energy, default_dp)
        time_to_49hz = (1.0 / system_rocof) if system_rocof > 0 else 0.0
        contingencies = []
        for row in online_rows:
            rocof = self.oblicz_rocof(total_energy, row['delta_p'])
            contingencies.append({
                'reg_name': row['reg_name'],
                'group_label': row['group_label'],
                'delta_p': row['delta_p'],
                'rocof': rocof,
                'time_to_49hz': (1.0 / rocof) if rocof > 0 else 0.0,
                'energy': row['energy'],
                'h': row['h'],
            })
        contingencies.sort(key=lambda row: row['rocof'], reverse=True)
        return {
            'energy': total_energy,
            'default_dp': default_dp,
            'rocof': system_rocof,
            'time_to_49hz': time_to_49hz,
            'contingencies': contingencies,
            'eligible_count': eligible_count,
            'inertia_count': inertia_count,
            'largest_unit': largest_unit,
            'source_note': 'Panel ROCOF pracuje w ujęciu całego systemu (COI / obszar synchroniczny), bez rozbicia na ODM; do analizy są brane tylko generatory z area = 1.',
        }

    def znajdz_kod_grupy_generatora(self, full_record):
        try:
            if '/' not in full_record:
                return None
            after_slash = full_record.split('/', 1)[1]
            tokens = re.findall('[-+]?\\d+(?:[.,]\\d+)?', after_slash)
            int_tokens = []
            for t in tokens:
                if '.' in t or ',' in t:
                    continue
                try:
                    int_tokens.append(int(t))
                except Exception:
                    pass
            for value in reversed(int_tokens):
                if value in self.GENERATOR_GROUP_LABELS:
                    return value
        except Exception:
            pass
        return None

    def parsuj_rekord_generatora(self, full_record):
        try:
            if ':' not in full_record:
                return None
            lewa, prawa = full_record.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            if not quoted:
                return None
            reg_name = quoted[-1].strip()
            if not reg_name or reg_name.isdigit():
                return None
            czesci_prawe = self.tokenizuj_segment_epc(prawa)
            status = int(float(czesci_prawe[0].replace(',', '.'))) if czesci_prawe else 0
            area_code = self.normalizuj_kod_odm(czesci_prawe[6] if len(czesci_prawe) > 6 else '1')
            zone_code = self.normalizuj_kod_odm(czesci_prawe[7] if len(czesci_prawe) > 7 else '')
            pgen = self.konwertuj_float_soft(czesci_prawe[8], 0.0) if len(czesci_prawe) > 8 else 0.0
            pmax = self.konwertuj_float_soft(czesci_prawe[9], 0.0) if len(czesci_prawe) > 9 else 0.0
            pmin = self.konwertuj_float_soft(czesci_prawe[10], 0.0) if len(czesci_prawe) > 10 else 0.0
            qgen = self.konwertuj_float_soft(czesci_prawe[11], 0.0) if len(czesci_prawe) > 11 else 0.0
            qmax = self.konwertuj_float_soft(czesci_prawe[12], 0.0) if len(czesci_prawe) > 12 else 0.0
            qmin = self.konwertuj_float_soft(czesci_prawe[13], 0.0) if len(czesci_prawe) > 13 else 0.0
            if pmax <= 0:
                for w in reversed(czesci_prawe):
                    test = self.konwertuj_float_soft(w, 0.0)
                    if test > 0:
                        pmax = test
                        break
            group_code = self.znajdz_kod_grupy_generatora(full_record)
            odm_info = self.wyciagnij_info_odm_generatora(full_record, zone_code=zone_code)
            if pmax > 0:
                return {'reg_name': reg_name, 'status': status, 'area': area_code, 'zone': zone_code, 'pgen': pgen, 'pmax': pmax, 'pmin': pmin, 'qgen': qgen, 'qmax': qmax, 'qmin': qmin, 'group_code': group_code, 'odm_code': odm_info['odm_code'], 'odm_name': odm_info['odm_name'], 'odm_source': odm_info['odm_source']}
        except Exception:
            pass
        return None

    def parsuj_rekord_loada(self, full_record):
        try:
            if ':' not in full_record:
                return None
            lewa, prawa = full_record.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            load_name = quoted[-1].strip() if quoted else ''
            czesci = [p.strip().strip('"') for p in prawa.split() if p.strip()]
            if len(czesci) < 8:
                return None
            status = int(float(czesci[0].replace(',', '.')))
            mw = float(czesci[1].replace(',', '.'))
            ar = int(float(czesci[7]))
            if status == 1 and ar == 1:
                mw_val = mw
            else:
                mw_val = 0.0
            return {'load_name': load_name, 'status': status, 'mw': mw_val, 'mw_raw': mw, 'area': ar}
        except Exception:
            return None

    def wczytaj_model_epc(self, sciezka):
        gen_data = []
        zapotrzebowanie_mw = 0.0
        current_section = None
        pending_record = None

        def flush_pending(section_name, record_text):
            nonlocal zapotrzebowanie_mw
            if not record_text:
                return
            if section_name == 'generator':
                gen = self.parsuj_rekord_generatora(record_text)
                if gen and gen['reg_name'].strip() and (gen['pmax'] > 0):
                    gen_data.append(gen)
            elif section_name == 'load':
                load = self.parsuj_rekord_loada(record_text)
                if load:
                    zapotrzebowanie_mw += load['mw']
        try:
            with open(sciezka, 'r', encoding='utf-8', errors='ignore') as f:
                for raw_line in f:
                    line = raw_line.strip()
                    if not line:
                        continue
                    line_lower = line.lower()
                    if 'generator data' in line_lower:
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = 'generator'
                        continue
                    if 'load data' in line_lower:
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = 'load'
                        continue
                    if current_section in {'generator', 'load'} and ':' not in line and ('data' in line_lower):
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = None
                        continue
                    if current_section not in {'generator', 'load'}:
                        continue
                    if ':' in line:
                        flush_pending(current_section, pending_record)
                        pending_record = line
                    elif pending_record:
                        pending_record += ' ' + line
                flush_pending(current_section, pending_record)
        except Exception as e:
            print(f'Błąd EPC {sciezka}: {e}')
        return {'generatory': gen_data, 'zapotrzebowanie_mw': zapotrzebowanie_mw}

    def wczytaj_generatory_z_pliku(self, sciezka):
        return self.wczytaj_model_epc(sciezka)['generatory']

    def wczytaj_wiele_plikow_epc(self):
        sciezki = filedialog.askopenfilenames(title='Wybierz pliki EPC', filetypes=[('*.epc', '*.epc'), ('*.*', '*.*')])
        if not sciezki:
            return
        self.wszystkie_generatory = {}
        self.statusy_oryginalne = {}
        self.opis_modeli = {}
        self.inercja_na_godzine = {h: 0.0 for h in range(24)}
        self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
        self.odm_metrics_na_godzine = {h: self.build_empty_odm_metrics() for h in range(24)}
        self.rocof_na_godzine = {h: 0.0 for h in range(24)}
        self.largest_loss_na_godzine = {h: 0.0 for h in range(24)}
        self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
        self.zapotrzebowanie_na_godzine = {h: 0.0 for h in range(24)}
        self.obliczenia_wykonane = False
        self.aktywny_typ_modelu = None
        self.aktualna_godzina = None
        self.aktualny_reczny_id = None
        self.aktualny_auto_id = None
        self.reset_alarm_cache()
        self.listbox_modele.delete(0, tk.END)
        self.pole_szczegoly.delete('1.0', tk.END)
        self.tree_generatory.delete(*self.tree_generatory.get_children())
        self.progress_var.set(0)
        self.label_progress.config(text='0%')
        self.aktualizuj_status('Wczytuję EPC...')
        przetworzone = 0
        suma = len(sciezki)
        for i, sciezka in enumerate(sciezki, start=1):
            data_txt, godz_txt, godzina = self.wyciagnij_date_i_godzine_z_nazwy(sciezka)
            if godzina is not None:
                model_epc = self.wczytaj_model_epc(sciezka)
                gen_data = model_epc['generatory']
                zapotrzebowanie_mw = model_epc['zapotrzebowanie_mw']
                if gen_data:
                    self.wszystkie_generatory[godzina] = gen_data
                    self.statusy_oryginalne[godzina] = {g['reg_name']: g['status'] for g in gen_data}
                    self.opis_modeli[godzina] = {'data': data_txt if data_txt else 'brak daty', 'godzina': godz_txt if godz_txt else f'{godzina:02d}:30', 'plik': os.path.basename(sciezka), 'sciezka': sciezka}
                    self.zapotrzebowanie_na_godzine[godzina] = zapotrzebowanie_mw
                    przetworzone += 1
            proc = i / suma * 100
            self.progress_var.set(proc)
            self.label_progress.config(text=f'{proc:.0f}% ({i}/{suma})')
            self.aktualizuj_status(f'Wczytywanie EPC... {i}/{suma}')
            self.root.update_idletasks()
        self.odswiez_liste_modeli()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.przygotuj_pusty_wykres_grup()
        self.aktualizuj_status(f'✅ EPC: {przetworzone} modeli. Teraz wczytaj DYD PSLF!')

    def wczytaj_modele_reczne_epc(self):
        sciezki = filedialog.askopenfilenames(title='Wybierz ręcznie importowane pliki EPC', filetypes=[('*.epc', '*.epc'), ('*.*', '*.*')])
        if not sciezki:
            return
        inserted = 0
        duplicate_paths = {str(self.modele_reczne[key].get('opis', {}).get('sciezka', '')).lower(): key for key in self.reczne_order if key in self.modele_reczne}
        self.progress_var.set(0)
        self.label_progress.config(text='0%')
        self.aktualizuj_status('WczytujÄ™ rÄ™cznie importowane modele EPC...')
        total = len(sciezki)
        for idx, sciezka in enumerate(sciezki, start=1):
            path_key = str(sciezka).lower()
            model_epc = self.wczytaj_model_epc(sciezka)
            gen_data = model_epc['generatory']
            zapotrzebowanie_mw = model_epc['zapotrzebowanie_mw']
            if not gen_data:
                continue
            data_txt, godz_txt, godzina_idx = self.wyciagnij_date_i_godzine_z_nazwy(sciezka)
            model_id = duplicate_paths.get(path_key)
            if not model_id:
                self.reczne_counter += 1
                model_id = f'manual_batch_{self.reczne_counter:05d}'
                self.modele_reczne[model_id] = self.build_session_model_entry(gen_data, data_txt=data_txt, godz_txt=godz_txt, sciezka=sciezka, source_tag='manual_import', godzina_idx=godzina_idx)
                self.reczne_order.append(model_id)
                duplicate_paths[path_key] = model_id
                inserted += 1
            entry = self.modele_reczne[model_id]
            entry['generatory'] = gen_data
            entry['statusy_oryginalne'] = {g['reg_name']: g['status'] for g in gen_data}
            entry['opis'].update({'data': data_txt if data_txt else 'brak daty', 'godzina': godz_txt if godz_txt else '--:--', 'plik': os.path.basename(sciezka), 'sciezka': sciezka, 'zrodlo_modelu': 'manual_import', 'godzina_idx': godzina_idx, 'czas_importu': datetime.now().strftime('%Y-%m-%d %H:%M:%S')})
            entry['zapotrzebowanie'] = zapotrzebowanie_mw
            entry.setdefault('odm_metrics', self.build_empty_odm_metrics())
            entry.setdefault('rocof', 0.0)
            entry.setdefault('largest_loss', 0.0)
            entry.setdefault('liczba_zalaczonych', 0)
            if self.H_generatory:
                self.przelicz_inercje_dla_modelu('manual_batch', model_id)
                self.sprawdz_alarm_dla_modelu('manual_batch', model_id)
            proc = idx / max(1, total) * 100.0
            self.progress_var.set(proc)
            self.label_progress.config(text=f'{proc:.0f}% ({idx}/{total})')
            self.root.update_idletasks()
        self.odswiez_liste_modeli_reczne()
        self.aktualizuj_wykres()
        if inserted > 0:
            self.aktualizuj_status(f'✅ Dodano {inserted} modeli do listy ręcznej')
        else:
            self.aktualizuj_status('✅ Zaktualizowano istniejące modele ręczne')

    def wczytaj_plik_genrou_H(self):
        sciezka = filedialog.askopenfilename(title='Wybierz plik DYD PSLF (.dyd)', filetypes=[('DYD PSLF', '*.dyd'), ('Tekst', '*.txt *.raw *.dyr *.dat'), ('Wszystkie', '*.*')])
        if not sciezka:
            return
        self.aktualizuj_status('Wczytuję współczynniki H...')
        H_map = {}
        current_block = []
        current_name = None
        try:
            with open(sciezka, 'r', encoding='utf-8', errors='ignore') as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    if re.match('GENROU|GENCLS|GENSAL', line.upper()):
                        if current_block and current_name:
                            H = self.parsuj_H_z_bloku_genrou(current_block)
                            if H:
                                H_map[current_name] = H
                        current_name = self.wyciagnij_nazwe_z_genrou(line)
                        current_block = [line]
                        continue
                    if current_name:
                        current_block.append(line)
            if current_block and current_name:
                H = self.parsuj_H_z_bloku_genrou(current_block)
                if H:
                    H_map[current_name] = H
        except Exception as e:
            messagebox.showerror('Błąd DYD', f'Błąd: {e}')
            return
        self.h_repo.upsert_from_dyd_map(H_map, os.path.basename(sciezka))
        self.reload_h_from_database(only_names=H_map.keys())
        self.odswiez_panel_h_dyd()
        self.aktualizuj_status(f'✅ Wczytano współczynniki H i zapisano do bazy SQLite: {len(H_map)} rekordów')
        for godzina in self.wszystkie_generatory.keys():
            self.przelicz_inercje_dla_modelu('manual', godzina)
        for model_id in self.reczne_order:
            self.przelicz_inercje_dla_modelu('manual_batch', model_id)
        for auto_id in self.auto_order:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_reczne()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()

    def parsuj_H_z_bloku_genrou(self, block_lines):
        try:
            full_text = ' '.join(block_lines)
            match = re.search('/\\s*([0-9.-]+)', full_text)
            if match:
                H = self.konwertuj_float_soft(match.group(1), 0.0)
                if 0 < H < 20:
                    return H
            for line in block_lines:
                if ':' in line:
                    tokens = re.findall('\\d+\\.?\\d*', line)
                    for t in tokens:
                        H = self.konwertuj_float_soft(t, 0.0)
                        if 0 < H < 20:
                            return H
            return None
        except Exception:
            return None

    def wyciagnij_nazwe_z_genrou(self, line):
        m = re.search('"([^"]+)"', line)
        return m.group(1).strip() if m else None

    def get_h_value_for_generator(self, gen, prefer_original=False, return_source=False):
        reg_name = gen.get('reg_name', '')
        if prefer_original:
            if reg_name in self.H_generatory_original:
                value = float(self.H_generatory_original.get(reg_name, 0.0))
                return (value, 'DYD') if return_source else value
        else:
            if reg_name in self.H_generatory:
                source = 'Ręczny' if (reg_name in self.H_generatory_original and abs(float(self.H_generatory.get(reg_name, 0.0)) - float(self.H_generatory_original.get(reg_name, 0.0))) > 1e-9) else 'DYD'
                value = float(self.H_generatory.get(reg_name, 0.0))
                return (value, source) if return_source else value
        group_code = gen.get('group_code')
        if group_code in self.H_group_defaults:
            value = float(self.H_group_defaults.get(group_code, 0.0))
            label = self.GENERATOR_GROUP_LABELS.get(group_code, f'KOD {group_code}')
            return (value, f'Typ {group_code} {label}') if return_source else value
        return (0.0, 'Brak H') if return_source else 0.0

    def policz_metryki_modelu(self, model_type, key, gen_list, use_original=False):
        energia = 0.0
        rocof_energy = 0.0
        liczba = 0
        largest_loss = 0.0
        odm_metrics = self.build_empty_odm_metrics()
        statusy_org = self.get_model_original_statuses(model_type, key)
        for gen in gen_list:
            status = statusy_org.get(gen['reg_name'], gen['status']) if use_original else gen['status']
            if status == 1:
                odm_code = str(gen.get('odm_code', '') or '').strip()
                if self.is_generator_rocof_eligible(gen):
                    delta_p = self.get_generator_active_power_for_rocof(gen)
                    largest_loss = max(largest_loss, delta_p)
                    if odm_code in odm_metrics:
                        bucket = odm_metrics[odm_code]
                        bucket['largest_loss'] = max(bucket['largest_loss'], delta_p)
            h_value = self.get_h_value_for_generator(gen, prefer_original=use_original)
            if h_value <= 0 or status != 1:
                continue
            s_apparent = self.get_generator_apparent_power(gen)
            ei = s_apparent * h_value
            energia += ei
            liczba += 1
            if odm_code in odm_metrics:
                bucket = odm_metrics[odm_code]
                bucket['energy'] += ei
                bucket['connected_count'] += 1
            if self.is_generator_rocof_eligible(gen):
                rocof_energy += ei
                if odm_code in odm_metrics:
                    bucket = odm_metrics[odm_code]
                    bucket['rocof_energy'] += ei
        for code, bucket in odm_metrics.items():
            bucket['share_pct'] = bucket['energy'] / energia * 100 if energia > 0 else 0.0
            bucket['rocof'] = self.oblicz_rocof(bucket.get('rocof_energy', 0.0), bucket['largest_loss'])
            bucket['name'] = self.ODM_CODE_TO_NAME.get(code, bucket.get('name', f'ODM {code}'))
        return {'energia': energia, 'liczba': liczba, 'largest_loss': largest_loss, 'rocof': self.oblicz_rocof(rocof_energy, largest_loss), 'odm_metrics': odm_metrics}

    def policz_inercje_dla_listy(self, model_type, key, gen_list, use_original=False):
        metryki = self.policz_metryki_modelu(model_type, key, gen_list, use_original=use_original)
        return (metryki['energia'], metryki['liczba'])

    def przelicz_inercje_dla_modelu(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return
        metryki_now = self.policz_metryki_modelu(model_type, key, gen_list, use_original=False)
        metryki_org = self.policz_metryki_modelu(model_type, key, gen_list, use_original=True)
        self.set_model_energy_now(model_type, key, metryki_now['energia'])
        self.set_model_energy_original(model_type, key, metryki_org['energia'])
        self.set_model_connected_count(model_type, key, metryki_now['liczba'])
        self.set_model_odm_metrics(model_type, key, metryki_now['odm_metrics'])
        self.set_model_rocof_now(model_type, key, metryki_now['rocof'])
        self.set_model_largest_loss(model_type, key, metryki_now['largest_loss'])

    def oblicz_inercje_wszystkie(self):
        if not self.wszystkie_generatory and not self.modele_reczne and not self.modele_auto:
            messagebox.showwarning('!', 'Wczytaj pliki EPC')
            return
        if not self.H_generatory:
            messagebox.showwarning('!', 'Wczytaj DYD PSLF albo ustaw H wg typu generatora')
            return
        self.aktualizuj_status('Obliczanie E_sys = Σ S_i × H_i ...')
        for godzina in self.wszystkie_generatory.keys():
            self.przelicz_inercje_dla_modelu('manual', godzina)
        for model_id in self.reczne_order:
            self.przelicz_inercje_dla_modelu('manual_batch', model_id)
        for auto_id in self.auto_order:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
        self.obliczenia_wykonane = True
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_reczne()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        if self.aktywny_typ_modelu is not None:
            self.odswiez_aktywny_model()
        else:
            self.przygotuj_pusty_wykres_grup()
        inserted = self.save_all_models_to_archive()
        self.sprawdz_alarmy_dla_wszystkich_modeli()
        self.aktualizuj_status(f'✅ GOTOWE! Kliknij modele po szczegóły. Zapisano do archiwum: {inserted}')
        messagebox.showinfo('✅', f'Obliczono dla {len(self.wszystkie_generatory)} modeli 24h, {len(self.reczne_order)} modeli recznych oraz {len(self.auto_order)} modeli biezacych.\nZapisano rekordow do archiwum: {inserted}')

    def przelicz_aktywny_model(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        if not self.H_generatory:
            messagebox.showwarning('Brak H', 'Najpierw wczytaj plik DYD z H albo ustaw H wg typu generatora.')
            return
        self.przelicz_inercje_dla_modelu(model_type, key)
        if model_type == 'manual':
            self.obliczenia_wykonane = True
        self.odswiez_aktywny_model()
        self.save_model_to_archive(model_type, key)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_reczne()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.aktualizuj_status('✅ Przeliczono wybrany model')

    def przelicz_i_odswiez_po_zmianie(self, model_type, key):
        self.przelicz_inercje_dla_modelu(model_type, key)
        if model_type == 'manual':
            self.obliczenia_wykonane = True
        self.odswiez_aktywny_model()
        self.save_model_to_archive(model_type, key)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_reczne()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.aktualizuj_status('✅ Przeliczono model po zmianie statusów')

    def wybierz_folder_monitoringu(self):
        folder = filedialog.askdirectory(title='Wybierz folder z modelami EPC')
        if not folder:
            return
        self.monitor_folder_path = folder
        self.label_monitor_folder.config(text=f'Folder: {folder}')
        self.aktualizuj_status('Wybrano folder monitoringu EPC')
        self.save_ui_config()

    def toggle_monitor_folder(self):
        if self.monitor_folder_var.get():
            if not self.monitor_folder_path:
                messagebox.showwarning('Brak folderu', 'Najpierw wybierz folder do monitorowania.')
                self.monitor_folder_var.set(False)
                return
            self.aktualizuj_status('Monitoring folderu włączony')
            self.sprawdz_folder_monitoringu()
        else:
            self.aktualizuj_status('Monitoring folderu wyłączony')

    def pobierz_wszystkie_pliki_epc_w_folderze(self):
        if not self.monitor_folder_path or not os.path.isdir(self.monitor_folder_path):
            return []
        pliki = []
        for nazwa in os.listdir(self.monitor_folder_path):
            if nazwa.lower().endswith('.epc'):
                pelna = os.path.join(self.monitor_folder_path, nazwa)
                try:
                    pliki.append((os.path.getmtime(pelna), pelna))
                except Exception:
                    pass
        pliki.sort()
        return pliki

    def sprawdz_folder_monitoringu(self):
        if not self.monitor_folder_var.get():
            return
        pliki = self.pobierz_wszystkie_pliki_epc_w_folderze()
        for mtime, sciezka in pliki:
            last_seen_mtime = self.auto_seen_files.get(sciezka)
            if last_seen_mtime is None or mtime > last_seen_mtime:
                self.auto_seen_files[sciezka] = mtime
                self.wczytaj_model_auto_z_folderu(sciezka)
        self.root.after(self.monitor_interval_ms, self.sprawdz_folder_monitoringu)

    def wczytaj_model_auto_z_folderu(self, sciezka):
        model_epc = self.wczytaj_model_epc(sciezka)
        gen_data = model_epc['generatory']
        zapotrzebowanie_mw = model_epc['zapotrzebowanie_mw']
        if not gen_data:
            return
        data_txt, godz_txt, godzina_idx = self.wyciagnij_date_i_godzine_z_nazwy(sciezka)
        self.auto_counter += 1
        auto_id = f'auto_{self.auto_counter:05d}'
        statusy_org = {g['reg_name']: g['status'] for g in gen_data}
        self.modele_auto[auto_id] = {'generatory': gen_data, 'statusy_oryginalne': statusy_org, 'opis': {'data': data_txt if data_txt else 'brak daty', 'godzina': godz_txt if godz_txt else '--:--', 'plik': os.path.basename(sciezka), 'sciezka': sciezka, 'czas_wykrycia': datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 'godzina_idx': godzina_idx}, 'inercja': 0.0, 'inercja_oryginalna': 0.0, 'odm_metrics': self.build_empty_odm_metrics(), 'rocof': 0.0, 'largest_loss': 0.0, 'liczba_zalaczonych': 0, 'zapotrzebowanie': zapotrzebowanie_mw}
        self.auto_order.append(auto_id)
        if self.H_generatory:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
            self.sprawdz_alarm_dla_modelu('auto', auto_id)
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.save_model_to_archive('auto', auto_id)
        self.aktualizuj_status(f'✅ Wykryto nowy lub zaktualizowany model w folderze i dodano do listy modeli bieżących: {os.path.basename(sciezka)}')

    def czy_status_zmieniony_recznie(self, model_type, key, reg_name, status_biezacy):
        statusy_org = self.get_model_original_statuses(model_type, key)
        status_org = statusy_org.get(reg_name)
        return status_org is not None and status_org != status_biezacy

    def pobierz_widoczne_generatory(self, model_type, key):
        wszystkie_gen = self.get_model_generators(model_type, key)
        gen_z_H = [g for g in wszystkie_gen if self.get_h_value_for_generator(g) > 0]
        return self.filtrowanie_generatory(model_type, key, gen_z_H)

    def parse_range_filter(self, text):
        txt = text.strip().replace(',', '.')
        if not txt:
            return None
        m = re.match('^\\s*(-?\\d+(?:\\.\\d+)?)\\s*-\\s*(-?\\d+(?:\\.\\d+)?)\\s*$', txt)
        if m:
            a, b = (float(m.group(1)), float(m.group(2)))
            return ('range', min(a, b), max(a, b))
        m = re.match('^\\s*(>=|<=|>|<|=)\\s*(-?\\d+(?:\\.\\d+)?)\\s*$', txt)
        if m:
            return ('cmp', m.group(1), float(m.group(2)))
        try:
            val = float(txt)
            return ('cmp', '=', val)
        except Exception:
            return ('text', txt.lower())

    def match_numeric_filter(self, value, flt):
        if not flt:
            return True
        if flt[0] == 'range':
            return flt[1] <= value <= flt[2]
        if flt[0] == 'cmp':
            op, ref = (flt[1], flt[2])
            if op == '>':
                return value > ref
            if op == '<':
                return value < ref
            if op == '>=':
                return value >= ref
            if op == '<=':
                return value <= ref
            return abs(value - ref) < 1e-09
        return True

    def passes_column_filters(self, gen, model_type, key):
        H = self.get_h_value_for_generator(gen)
        s_apparent = self.get_generator_apparent_power(gen)
        Ei = s_apparent * H
        grupa_txt = self.formatuj_grupe_generatora(gen).lower()
        odm_txt = self.formatuj_odm_generatora(gen).lower()
        stan_txt = 'zał' if gen['status'] == 1 else 'wył'
        f_kod = self.column_filters['kod'].get().strip().lower()
        if f_kod and f_kod not in gen['reg_name'].lower():
            return False
        f_grupa = self.column_filters['grupa'].get().strip().lower()
        if f_grupa and f_grupa not in grupa_txt:
            return False
        f_odm = self.column_filters['odm'].get().strip().lower()
        if f_odm and f_odm not in odm_txt:
            return False
        f_stan = self.column_filters['stan'].get().strip().lower()
        if f_stan:
            mapy = {'1': 'zał', 'zal': 'zał', 'zał': 'zał', '0': 'wył', 'wyl': 'wył', 'wył': 'wył'}
            wanted = mapy.get(f_stan, f_stan)
            if wanted not in stan_txt:
                return False
        for field_key, value in [('h', H), ('pmax', s_apparent), ('ei', Ei)]:
            raw = self.column_filters[field_key].get().strip()
            if raw:
                flt = self.parse_range_filter(raw)
                if flt and flt[0] in {'range', 'cmp'}:
                    if not self.match_numeric_filter(value, flt):
                        return False
                elif flt and flt[0] == 'text':
                    if flt[1] not in f'{value}'.lower():
                        return False
        return True

    def filtrowanie_generatory(self, model_type, key, gen_data):
        tekst = self.entry_filtr.get().strip().lower()
        self.tekst_filtrowania = tekst
        frazy = [x.strip() for x in tekst.split(',') if x.strip()]
        wynik = list(gen_data)
        if frazy:
            wynik = [g for g in wynik if any((f in g['reg_name'].lower() for f in frazy))]
        tryb = self.tryb_filtrowania_var.get()
        if tryb == 'Tylko zmienione ręcznie':
            wynik = [g for g in wynik if self.czy_status_zmieniony_recznie(model_type, key, g['reg_name'], g['status'])]
        elif tryb == 'Tylko załączone':
            wynik = [g for g in wynik if g['status'] == 1]
        elif tryb == 'Tylko wyłączone':
            wynik = [g for g in wynik if g['status'] == 0]
        wynik = [g for g in wynik if self.passes_column_filters(g, model_type, key)]

        def energia(g):
            return self.get_generator_apparent_power(g) * self.get_h_value_for_generator(g)

        def apparent_power(g):
            return self.get_generator_apparent_power(g)
        if self.tree_sort_column:
            col = self.tree_sort_column
            rev = self.tree_sort_reverse
            if col == 'stan':
                wynik.sort(key=lambda g: g['status'], reverse=rev)
            elif col == 'kod':
                wynik.sort(key=lambda g: g['reg_name'], reverse=rev)
            elif col == 'grupa':
                wynik.sort(key=lambda g: self.formatuj_grupe_generatora(g), reverse=rev)
            elif col == 'odm':
                wynik.sort(key=lambda g: self.formatuj_odm_generatora(g), reverse=rev)
            elif col == 'pmax':
                wynik.sort(key=apparent_power, reverse=rev)
            elif col == 'h':
                wynik.sort(key=lambda g: self.get_h_value_for_generator(g), reverse=rev)
            elif col == 'ei':
                wynik.sort(key=energia, reverse=rev)
            elif col == 'udzial':
                e_total = max(self.get_model_energy_now(model_type, key), 0.0)
                wynik.sort(key=lambda g: energia(g) / e_total * 100 if e_total > 0 else 0.0, reverse=rev)
            elif col == 'zmiana':
                wynik.sort(key=lambda g: self.czy_status_zmieniony_recznie(model_type, key, g['reg_name'], g['status']), reverse=rev)
            return wynik
        sort_mode = self.sortowanie_var.get()
        if sort_mode == 'Kod generatora A-Z':
            wynik.sort(key=lambda g: g['reg_name'])
        elif sort_mode == 'Kod generatora Z-A':
            wynik.sort(key=lambda g: g['reg_name'], reverse=True)
        elif sort_mode == 'ODM':
            wynik.sort(key=lambda g: (self.formatuj_odm_generatora(g), g['reg_name']))
        elif sort_mode == 'S malejąco':
            wynik.sort(key=apparent_power, reverse=True)
        elif sort_mode == 'S rosnąco':
            wynik.sort(key=apparent_power)
        elif sort_mode == 'H malejąco':
            wynik.sort(key=lambda g: self.get_h_value_for_generator(g), reverse=True)
        elif sort_mode == 'H rosnąco':
            wynik.sort(key=lambda g: self.get_h_value_for_generator(g))
        elif sort_mode == 'Ei malejąco':
            wynik.sort(key=energia, reverse=True)
        elif sort_mode == 'Ei rosnąco':
            wynik.sort(key=energia)
        elif sort_mode == 'Status':
            wynik.sort(key=lambda g: (g['status'], g['reg_name']), reverse=True)
        return wynik

    def zastosuj_filtr(self, event=None):
        self.save_ui_config()
        self.schedule_active_model_refresh()

    def on_column_filter_changed(self, event=None):
        self.save_ui_config()
        self.schedule_active_model_refresh()

    def sort_tree_by_column(self, col):
        if self.tree_sort_column == col:
            self.tree_sort_reverse = not self.tree_sort_reverse
        else:
            self.tree_sort_column = col
            self.tree_sort_reverse = False
        self.save_ui_config()
        self.schedule_active_model_refresh(delay_ms=40)

    def on_dropdown_sort_changed(self):
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        self.save_ui_config()
        self.schedule_active_model_refresh(delay_ms=40)

    def on_udzialy_scope_changed(self):
        self.save_ui_config()
        self.schedule_active_model_refresh(delay_ms=40)

    def schedule_active_model_refresh(self, delay_ms=120):
        if self.ui_refresh_job is not None:
            try:
                self.root.after_cancel(self.ui_refresh_job)
            except Exception:
                pass
        self.ui_refresh_job = self.root.after(delay_ms, self.flush_active_model_refresh)

    def flush_active_model_refresh(self):
        self.ui_refresh_job = None
        self.odswiez_aktywny_model()

    def wyczysc_filtr(self):
        self.entry_filtr.delete(0, tk.END)
        self.tryb_filtrowania_var.set('Wszystkie')
        self.sortowanie_var.set('Kod generatora A-Z')
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        for var in self.column_filters.values():
            var.set('')
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def pokaz_wszystkie(self):
        self.wyczysc_filtr()

    def odswiez_aktywny_model(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            self.refresh_rocof_window()
            return
        self.odswiez_panel_generatorow(model_type, key)
        self.pokaz_szczegoly_modelu(model_type, key)
        self.aktualizuj_porownanie_modelu(model_type, key)
        self.aktualizuj_ranking(model_type, key)
        self.aktualizuj_wykres_grup_generatorow(model_type, key)
        self.aktualizuj_analize_inercyjna(model_type, key)
        self.refresh_rocof_window()

    def formatuj_grupe_generatora(self, gen):
        code = gen.get('group_code')
        if code is None:
            return '-'
        label = self.GENERATOR_GROUP_LABELS.get(code, f'KOD {code}')
        return f'{code} ({label})'

    def formatuj_etykiete_grupy_wykres(self, code, label, max_len=16):
        txt = str(label or '').strip()
        if len(txt) > max_len:
            txt = txt[:max_len - 1] + '…'
        return f'{code} - {txt}'

    def odswiez_panel_generatorow(self, model_type, key):
        self.tree_generatory.delete(*self.tree_generatory.get_children())
        gen_data = self.pobierz_widoczne_generatory(model_type, key)
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        typ_txt = self.get_model_type_label(model_type)
        demand_txt = self.get_model_demand_now(model_type, key)
        self.label_reczne.config(text=f'Model [{typ_txt}] {data_txt} {godz_txt} | Widoczne: {len(gen_data)} | Load: {demand_txt:,.2f} MW')
        e_total = max(self.get_model_energy_now(model_type, key), 0.0)
        for gen in gen_data:
            H = self.get_h_value_for_generator(gen)
            s_apparent = self.get_generator_apparent_power(gen)
            Ei = s_apparent * H
            udzial = Ei / e_total * 100 if e_total > 0 else 0.0
            zmieniony = self.czy_status_zmieniony_recznie(model_type, key, gen['reg_name'], gen['status'])
            stan = 'ZAŁ' if gen['status'] == 1 else 'WYŁ'
            tag = 'manual' if zmieniony else 'zal' if gen['status'] == 1 else 'wyl'
            self.tree_generatory.insert('', tk.END, iid=gen['reg_name'], values=(stan, gen['reg_name'], self.formatuj_grupe_generatora(gen), self.formatuj_odm_generatora(gen), f"{s_apparent:.3f}", f'{H:.3f}', f'{Ei:.2f}', f'{udzial:.2f}', 'TAK' if zmieniony else '-'), tags=(tag,))

    def get_generator_tooltip(self, item_id):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            return ''
        gen_list = self.get_model_generators(model_type, key)
        e_total = max(self.get_model_energy_now(model_type, key), 0.0)
        for gen in gen_list:
            if gen['reg_name'] == item_id:
                H = self.get_h_value_for_generator(gen)
                s_apparent = self.get_generator_apparent_power(gen)
                Ei = s_apparent * H
                udzial = Ei / e_total * 100 if e_total > 0 else 0.0
                grupa = self.formatuj_grupe_generatora(gen)
                odm = self.formatuj_odm_generatora(gen)
                org = self.get_model_original_statuses(model_type, key).get(gen['reg_name'], gen['status'])
                org_txt = 'ZAŁ' if org == 1 else 'WYŁ'
                now_txt = 'ZAŁ' if gen['status'] == 1 else 'WYŁ'
                return f"Generator: {gen['reg_name']}\nPełna nazwa: {gen['reg_name']}\nStatus oryginalny: {org_txt}\nStatus bieżący: {now_txt}\nTyp grupy: {grupa}\nODM: {odm}\nH: {H:.3f} s\nPmax: {gen.get('pmax', 0.0):.3f} MW\nQmax: {gen.get('qmax', 0.0):.3f} MVAr\nS: {s_apparent:.3f} MVA\nEi = S×H: {Ei:.2f} MVA·s\nUdział: {udzial:.2f}%"
        return ''

    def on_tree_double_click(self, event):
        item = self.tree_generatory.identify_row(event.y)
        model_type, key = self.get_current_model_ref()
        if not item or model_type is None:
            return
        gen_list = self.get_model_generators(model_type, key)
        for gen in gen_list:
            if gen['reg_name'] == item:
                gen['status'] = 0 if gen['status'] == 1 else 1
                break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przelacz_zaznaczone_w_tree(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa:
                    gen['status'] = 0 if gen['status'] == 1 else 1
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def ustaw_status_zaznaczonych_w_tree(self, nowy_status):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa:
                    gen['status'] = nowy_status
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przywroc_zaznaczone_w_tree(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        oryg = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa and nazwa in oryg:
                    gen['status'] = oryg[nazwa]
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przywroc_oryginalne_statusy_modelu(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        oryginalne = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        liczba_zmian = 0
        for gen in gen_list:
            nazwa = gen['reg_name']
            if nazwa in oryginalne and gen['status'] != oryginalne[nazwa]:
                gen['status'] = oryginalne[nazwa]
                liczba_zmian += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        self.aktualizuj_status(f'✅ Przywrócono oryginalne statusy modelu ({liczba_zmian} zmian)')

    def przywroc_wszystkie_modele_do_oryginalu(self):
        if not self.wszystkie_generatory or not self.statusy_oryginalne:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj modele EPC.')
            return
        lacznie_zmian = 0
        for godzina, gen_list in self.wszystkie_generatory.items():
            oryginalne = self.statusy_oryginalne.get(godzina, {})
            for gen in gen_list:
                nazwa = gen['reg_name']
                if nazwa in oryginalne and gen['status'] != oryginalne[nazwa]:
                    gen['status'] = oryginalne[nazwa]
                    lacznie_zmian += 1
            self.przelicz_inercje_dla_modelu('manual', godzina)
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.odswiez_liste_modeli()
        if self.aktywny_typ_modelu is not None:
            self.odswiez_aktywny_model()
        else:
            self.przygotuj_pusty_wykres_grup()
        self.aktualizuj_status(f'✅ Przywrócono wszystkie modele 24h do statusów domyślnych ({lacznie_zmian} zmian)')
        messagebox.showinfo('Przywracanie zakończone', f'Przywrócono statusy domyślne we wszystkich modelach 24h.\nLiczba zmian: {lacznie_zmian}')

    def zalacz_widoczne_z_filtra(self):
        self.ustaw_status_widocznych(1)

    def wylacz_widoczne_z_filtra(self):
        self.ustaw_status_widocznych(0)

    def przywroc_widoczne_z_filtra(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        widoczne = {g['reg_name'] for g in self.pobierz_widoczne_generatory(model_type, key)}
        oryginalne = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        liczba = 0
        for gen in gen_list:
            if gen['reg_name'] in widoczne and gen['reg_name'] in oryginalne:
                if gen['status'] != oryginalne[gen['reg_name']]:
                    gen['status'] = oryginalne[gen['reg_name']]
                    liczba += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        self.aktualizuj_status(f'✅ Przywrócono widoczne generatory do oryginału ({liczba} zmian)')

    def ustaw_status_widocznych(self, nowy_status):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        widoczne = {g['reg_name'] for g in self.pobierz_widoczne_generatory(model_type, key)}
        gen_list = self.get_model_generators(model_type, key)
        liczba = 0
        for gen in gen_list:
            if gen['reg_name'] in widoczne and gen['status'] != nowy_status:
                gen['status'] = nowy_status
                liczba += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        txt = 'załączono' if nowy_status == 1 else 'wyłączono'
        self.aktualizuj_status(f'✅ {txt.capitalize()} widoczne generatory ({liczba} zmian)')

    def pobierz_dane_szczegolowe(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return None
        gen_z_H = [g for g in gen_list if self.get_h_value_for_generator(g) > 0]
        gen_data = self.filtrowanie_generatory(model_type, key, gen_z_H)
        energia_calkowita = self.get_model_energy_now(model_type, key)
        zapotrzebowanie = self.get_model_demand_now(model_type, key)
        odm_metrics = self.get_model_odm_metrics(model_type, key)
        rocof = self.get_model_rocof_now(model_type, key)
        largest_loss = self.get_model_largest_loss(model_type, key)
        wiersze = []
        wlaczonych = 0
        energia_wlaczonych = 0.0
        energia_filtrowanych = 0.0
        energia_wylaczonych = 0.0
        for i, gen in enumerate(gen_data, 1):
            pmax_raw = float(gen.get('pmax', 0.0))
            qmax_raw = float(gen.get('qmax', 0.0))
            s_apparent = self.get_generator_apparent_power(gen)
            H = self.get_h_value_for_generator(gen)
            Ei = s_apparent * H
            energia_filtrowanych += Ei
            uwzgledniony = 'TAK' if gen['status'] == 1 else 'NIE'
            status_tekst = 'ZAŁ.' if gen['status'] == 1 else 'WYŁ.'
            status_org = self.get_model_original_statuses(model_type, key).get(gen['reg_name'], gen['status'])
            reczna_zmiana = status_org != gen['status']
            if gen['status'] == 1:
                wlaczonych += 1
                energia_wlaczonych += Ei
            else:
                energia_wylaczonych += Ei
            udzial = Ei / energia_calkowita * 100 if energia_calkowita > 0 else 0.0
            wiersze.append({'lp': i, 'kod_generatora': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'odm': self.formatuj_odm_generatora(gen), 'status': status_tekst, 'pmax_raw': pmax_raw, 'qmax_raw': qmax_raw, 's_obliczeniowa': s_apparent, 'H': H, 'Ei': Ei, 'udzial': udzial, 'uwzgledniony': uwzgledniony, 'reczna_zmiana': reczna_zmiana, 'status_oryginalny': status_org})
        podsumowanie = {'liczba_gen_odfiltrowanych': len(gen_data), 'liczba_wlaczonych': wlaczonych, 'energia_odfiltrowanych': energia_filtrowanych, 'energia_wlaczonych': energia_wlaczonych, 'energia_wylaczonych': energia_wylaczonych, 'energia_calkowita': energia_calkowita, 'udzial_filtra_proc': energia_wlaczonych / energia_calkowita * 100 if energia_calkowita > 0 else 0.0, 'liczba_gen_przed_filtrem': len(gen_z_H), 'zapotrzebowanie': zapotrzebowanie, 'odm_metrics': odm_metrics, 'rocof': rocof, 'largest_loss': largest_loss}
        return {'wiersze': wiersze, 'podsumowanie': podsumowanie}

    def pokaz_szczegoly_modelu(self, model_type, key):
        dane = self.pobierz_dane_szczegolowe(model_type, key)
        if dane is None:
            self.pole_szczegoly.delete('1.0', tk.END)
            self.pole_szczegoly.insert(tk.END, '❌ Brak danych dla wybranego modelu')
            return
        wiersze = dane['wiersze']
        p = dane['podsumowanie']
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        plik_txt = opis.get('plik', '')
        typ_txt = self.get_model_type_label(model_type)
        self.pole_szczegoly.delete('1.0', tk.END)
        filtr_info = f"(filtr: '{self.tekst_filtrowania}' – {p['liczba_gen_odfiltrowanych']}/{p['liczba_gen_przed_filtrem']})" if self.tekst_filtrowania else ''
        self.label_wynikow.config(text=f"{p['liczba_gen_odfiltrowanych']}/{p['liczba_gen_przed_filtrem']} generatorów po filtrze")
        self.pole_szczegoly.insert(tk.END, f'🔍 Model [{typ_txt}]: {data_txt} {godz_txt} {filtr_info}\n', 'header')
        if plik_txt:
            self.pole_szczegoly.insert(tk.END, f'📄 Plik: {plik_txt}\n', 'header')
        self.pole_szczegoly.insert(tk.END, f"{'=' * 190}\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"📈 Całkowita inercja E_sys = {p['energia_calkowita']:,.2f} MVA·s\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"⚡ Zapotrzebowanie (LOAD DATA, suma mw aktywnych loadów) = {p['zapotrzebowanie']:,.2f} MW\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"ROCOF N-1 (tylko area=1, utrata największej aktywnej jednostki ΔP = {p['largest_loss']:,.2f} MW) = {p['rocof']:.4f} Hz/s\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"{'=' * 190}\n\n", 'header')
        naglowek = f"{'LP':<4} {'Kod generatora':<28} {'Grupa':<14} {'ODM':<12} {'Status':<10} {'S [MVA]':<12} {'H [s]':<10} {'E_i=S×H [MVA·s]':<18} {'Uwzględniony':<18} {'Ręczna zmiana':<16}\n"
        self.pole_szczegoly.insert(tk.END, naglowek, 'summary')
        self.pole_szczegoly.insert(tk.END, f"{'-' * 190}\n", 'summary')
        for w in wiersze:
            znacznik = 'TAK' if w['reczna_zmiana'] else '-'
            line = f"{w['lp']:<4} {w['kod_generatora'][:26]:<28} {w['grupa'][:12]:<14} {w['odm'][:10]:<12} {w['status']:<10} {w['s_obliczeniowa']:>10.3f}   {w['H']:>8.3f}   {w['Ei']:>16.3f}   {w['uwzgledniony']:<18} {znacznik:<16}\n"
            if w['reczna_zmiana']:
                tag = 'manual_line'
            elif w['status'] == 'ZAŁ.':
                tag = 'green_line'
            else:
                tag = 'red_line'
            self.pole_szczegoly.insert(tk.END, line, tag)
        self.pole_szczegoly.insert(tk.END, f"{'-' * 190}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, '📊 PODSUMOWANIE:\n', 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Odfiltrowane generatory:                           {p['liczba_gen_odfiltrowanych']:>6}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Załączone (po filtrze):                            {p['liczba_wlaczonych']:>6}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji odfiltrowanych generatorów:          {p['energia_odfiltrowanych']:>12.2f} MVA·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji uwzględnionych generatorów w obliczeniach: {p['energia_wlaczonych']:>12.2f} MVA·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji nieuwzględnionych generatorów:       {p['energia_wylaczonych']:>12.2f} MVA·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Całkowita inercja systemu:                        {p['energia_calkowita']:>12.2f} MVA·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Zapotrzebowanie systemu (LOAD DATA):              {p['zapotrzebowanie']:>12.2f} MW\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Udział odfiltrowanych w E_sys:                    {p['udzial_filtra_proc']:>11.2f} %\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• ROCOF N-1 dla całego systemu:                    {p['rocof']:>11.4f} Hz/s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, '\nODM - porownanie inercji i ROCOF:\n', 'summary')
        for odm_code in self.ODM_ORDER:
            meta = p['odm_metrics'].get(odm_code, {})
            self.pole_szczegoly.insert(tk.END, f"• {self.ODM_CODE_TO_NAME[odm_code]:<10} | E = {meta.get('energy', 0.0):>11.2f} MVA·s | Udział = {meta.get('share_pct', 0.0):>6.2f} % | Gen = {int(meta.get('connected_count', 0)):>3} | ROCOF = {meta.get('rocof', 0.0):>8.4f} Hz/s\n", 'summary')
        self.pole_szczegoly.see('1.0')

    def aktualizuj_porownanie_modelu(self, model_type, key):
        E_org = self.get_model_energy_original(model_type, key)
        E_now = self.get_model_energy_now(model_type, key)
        P_load = self.get_model_demand_now(model_type, key)
        rocof = self.get_model_rocof_now(model_type, key)
        largest_loss = self.get_model_largest_loss(model_type, key)
        diff = E_now - E_org
        diff_proc = diff / E_org * 100 if E_org > 0 else 0.0
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        typ_txt = self.get_model_type_label(model_type)
        txt = f'Model [{typ_txt}]: {data_txt} {godz_txt}\nInercja oryginalna: {E_org:,.2f} MVA·s\nInercja po zmianach: {E_now:,.2f} MVA·s\nRóżnica: {diff:,.2f} MVA·s\nZmiana procentowa: {diff_proc:,.2f} %\nZapotrzebowanie mocy: {P_load:,.2f} MW\nROCOF N-1 (area=1): {rocof:.4f} Hz/s (ΔP krytyczne {largest_loss:,.2f} MW)'
        self.label_porownanie.config(text=txt, fg=self.get_palette()['success'] if diff >= 0 else self.get_palette()['danger'])

    def aktualizuj_wskazniki_jakosci(self):
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if not dodatnie:
            self.label_wskazniki.config(text='Brak danych 24h', fg=self.get_palette()['accent'])
            return
        wartosci = [v for _, v in dodatnie]
        srednia = statistics.mean(wartosci)
        mediana = statistics.median(wartosci)
        minimum = min(wartosci)
        maksimum = max(wartosci)
        odchylenie = statistics.pstdev(wartosci) if len(wartosci) > 1 else 0.0
        zap = [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h, _ in dodatnie]
        rocof_vals = [self.rocof_na_godzine.get(h, 0.0) for h, _ in dodatnie]
        srednie_zap = statistics.mean(zap) if zap else 0.0
        min_zap = min(zap) if zap else 0.0
        max_zap = max(zap) if zap else 0.0
        sredni_rocof = statistics.mean(rocof_vals) if rocof_vals else 0.0
        max_rocof = max(rocof_vals) if rocof_vals else 0.0
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            godziny_naruszenia = 'brak progu'
            liczba_ponizej = 'brak progu'
        else:
            godziny = [h for h, v in dodatnie if v < prog]
            liczba_ponizej = len(godziny)
            godziny_naruszenia = ', '.join((f'{h:02d}:30' for h in godziny)) if godziny else 'brak'
        udzialy = self.policz_udzialy_top_globalnie()
        txt = f"Średnia inercja dobowa: {srednia:,.2f} MVA·s\nMinimum / maksimum: {minimum:,.2f} / {maksimum:,.2f} MVA·s\nMediana: {mediana:,.2f} MVA·s\nOdchylenie standardowe: {odchylenie:,.2f} MVA·s\nŚrednie zapotrzebowanie: {srednie_zap:,.2f} MW\nZapotrzebowanie min / max: {min_zap:,.2f} / {max_zap:,.2f} MW\nŚredni / maks. ROCOF N-1 (area=1): {sredni_rocof:.4f} / {max_rocof:.4f} Hz/s\nLiczba godzin poniżej progu: {liczba_ponizej}\nGodziny poniżej progu: {godziny_naruszenia}\nUdział TOP 1 / TOP 3 / TOP 5 generatorów: {udzialy['top1']:.2f}% / {udzialy['top3']:.2f}% / {udzialy['top5']:.2f}%"
        self.label_wskazniki.config(text=txt, fg=self.get_palette()['accent'])

    def policz_udzialy_top_globalnie(self):
        wkłady = []
        suma = 0.0
        for godzina, gen_list in self.wszystkie_generatory.items():
            for gen in gen_list:
                h_value = self.get_h_value_for_generator(gen)
                if gen['status'] == 1 and h_value > 0:
                    ei = self.get_generator_apparent_power(gen) * h_value
                    wkłady.append(ei)
                    suma += ei
        if suma <= 0 or not wkłady:
            return {'top1': 0.0, 'top3': 0.0, 'top5': 0.0}
        wkłady.sort(reverse=True)
        return {'top1': sum(wkłady[:1]) / suma * 100, 'top3': sum(wkłady[:3]) / suma * 100, 'top5': sum(wkłady[:5]) / suma * 100}

    def aktualizuj_ranking(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            self.label_ranking.config(text='Brak wybranego modelu')
            return
        ranking = []
        for gen in gen_list:
            h_value = self.get_h_value_for_generator(gen)
            if h_value > 0:
                s_apparent = self.get_generator_apparent_power(gen)
                Ei = s_apparent * h_value
                ranking.append((gen['reg_name'], gen['status'], s_apparent, h_value, Ei, self.formatuj_grupe_generatora(gen), self.formatuj_odm_generatora(gen)))
        ranking.sort(key=lambda x: x[4], reverse=True)
        ranking = ranking[:10]
        if not ranking:
            self.label_ranking.config(text='Brak generatorów z H dla wybranego modelu')
            return
        linie = [f"{'LP':<3} {'Kod':<20} {'Gr':<10} {'ODM':<10} {'St':<4} {'S':>8} {'H':>8} {'Ei':>12}"]
        for i, (nazwa, status, s_apparent, H, Ei, grupa_txt, odm_txt) in enumerate(ranking, 1):
            st = 'ZAŁ' if status == 1 else 'WYŁ'
            linie.append(f'{i:<3} {nazwa[:20]:<20} {grupa_txt[:10]:<10} {odm_txt[:10]:<10} {st:<4} {s_apparent:>8.2f} {H:>8.2f} {Ei:>12.2f}')
        self.label_ranking.config(text='\n'.join(linie))

    def oblicz_statystyki_grup_generatorow(self, model_type, key, odm_code=None):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return []
        grupy = {}
        suma_e = 0.0
        for gen in gen_list:
            if odm_code is not None and str(gen.get('odm_code', '') or '').strip() != str(odm_code):
                continue
            code = gen.get('group_code')
            if code is None:
                continue
            pmax = float(gen.get('pmax', 0.0) or 0.0)
            bucket = grupy.setdefault(code, {'ei_sum': 0.0, 'count': 0, 'default_count': 0, 'count_total': 0, 'count_active': 0, 'pmax_total': 0.0, 'pmax_active': 0.0, 'pmax_inertia': 0.0})
            bucket['count_total'] += 1
            if pmax > 0:
                bucket['pmax_total'] += pmax
            if int(gen.get('status', 0)) != 1:
                continue
            bucket['count_active'] += 1
            if pmax > 0:
                bucket['pmax_active'] += pmax
            h_value = float(self.get_h_value_for_generator(gen, prefer_original=False))
            if h_value <= 0:
                continue
            ei = self.get_generator_apparent_power(gen) * h_value
            if ei <= 0:
                continue
            suma_e += ei
            bucket['ei_sum'] += ei
            bucket['count'] += 1
            if pmax > 0:
                bucket['pmax_inertia'] += pmax
            if gen.get('reg_name') not in self.H_generatory:
                bucket['default_count'] += 1
        if not grupy:
            return []
        wynik = []
        for code, meta in sorted(grupy.items(), key=lambda x: (x[1]['ei_sum'], x[1]['pmax_total']), reverse=True):
            label = self.GENERATOR_GROUP_LABELS.get(code, f'KOD {code}')
            ei_sum = float(meta['ei_sum'])
            udzial = ei_sum / suma_e * 100 if suma_e > 0 else 0.0
            wynik.append({'code': code, 'label': label, 'ei_sum': ei_sum, 'share_pct': udzial, 'count': int(meta['count']), 'default_count': int(meta['default_count']), 'count_total': int(meta['count_total']), 'count_active': int(meta['count_active']), 'pmax_total': float(meta['pmax_total']), 'pmax_active': float(meta['pmax_active']), 'pmax_inertia': float(meta['pmax_inertia'])})
        return wynik

    def oblicz_udzialy_grup_generatorow(self, model_type, key, odm_code=None):
        return self.oblicz_statystyki_grup_generatorow(model_type, key, odm_code=odm_code)

    def aktualizuj_wykres_grup_generatorow(self, model_type, key):
        selected_scope = self.udzialy_scope_var.get().strip() or 'Cały model'
        odm_code_selected = None
        for odm_code in self.ODM_ORDER:
            if selected_scope == self.ODM_CODE_TO_NAME[odm_code]:
                odm_code_selected = odm_code
                break
        dane = self.oblicz_statystyki_grup_generatorow(model_type, key, odm_code=odm_code_selected)
        pal = self.get_palette()
        self.figura_grupy.patch.set_facecolor(pal['bg'])
        self.os_grupy.clear()
        self.os_pmax_grupy.clear()
        self.os_odm_inercja.clear()
        scope_label = selected_scope if odm_code_selected else 'Cały model'
        self.os_grupy.set_title(f'Udziały typów jednostek - {scope_label}')
        self.os_grupy.set_xlabel('Typ jednostki')
        self.os_grupy.set_ylabel('Udział [%]')
        self.stylizuj_os(self.os_grupy, grid=True, grid_axis='y')
        self.os_pmax_grupy.set_title(f'Pmax wg typów jednostek - {scope_label}')
        self.os_pmax_grupy.set_xlabel('Typ jednostki')
        self.os_pmax_grupy.set_ylabel('Pmax [MW]')
        self.stylizuj_os(self.os_pmax_grupy, grid=True, grid_axis='y')
        self.stylizuj_os(self.os_odm_inercja, grid=True, grid_axis='y')
        self.os_odm_inercja.set_title('Sumaryczna inercja dla ODM')
        self.os_odm_inercja.set_xlabel('ODM')
        self.os_odm_inercja.set_ylabel('Inercja [MVA·s]')
        self.figura_grupy.subplots_adjust(bottom=0.10, left=0.08, right=0.99, top=0.95, hspace=0.82)
        if not dane:
            self.os_grupy.set_ylim(0, 100)
            self.os_grupy.text(0.5, 0.5, 'Brak danych grup generatorów\nlub brak aktywnych generatorów z H', ha='center', va='center', transform=self.os_grupy.transAxes, fontsize=10, color=pal['muted'])
            self.os_pmax_grupy.set_ylim(0, 1)
            self.os_pmax_grupy.text(0.5, 0.5, 'Brak danych Pmax dla wybranego zakresu', ha='center', va='center', transform=self.os_pmax_grupy.transAxes, fontsize=10, color=pal['muted'])
        else:
            x_positions = list(range(len(dane)))
            x_labels = [self.formatuj_etykiete_grupy_wykres(row['code'], row['label']) for row in dane]
            y_values = [row['share_pct'] for row in dane]
            colors = [self.GROUP_COLORS[idx % len(self.GROUP_COLORS)] for idx in range(len(dane))]
            bars = self.os_grupy.bar(x_positions, y_values, color=colors, edgecolor=pal['border'], linewidth=0.8)
            self.os_grupy.set_xticks(x_positions)
            self.os_grupy.set_xticklabels(x_labels, rotation=18, ha='right', fontsize=8.0, color=pal['text'])
            self.os_grupy.set_ylim(0, max(100, (max(y_values) * 1.18) if y_values else 100))
            for bar, row in zip(bars, dane):
                self.os_grupy.annotate(f"{row['share_pct']:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=8, color=pal['text'])
            total_vals = [row['pmax_total'] for row in dane]
            inertia_vals = [row['pmax_inertia'] for row in dane]
            width = 0.38
            bars_total = self.os_pmax_grupy.bar([x - width / 2 for x in x_positions], total_vals, width=width, color=pal['line_secondary'], edgecolor=pal['border'], linewidth=0.8, label='Pmax całkowity')
            bars_inertia = self.os_pmax_grupy.bar([x + width / 2 for x in x_positions], inertia_vals, width=width, color=pal['accent'], edgecolor=pal['border'], linewidth=0.8, label='Pmax jednostek w inercji')
            self.os_pmax_grupy.set_xticks(x_positions)
            self.os_pmax_grupy.set_xticklabels(x_labels, rotation=18, ha='right', fontsize=8.0, color=pal['text'])
            ymax_pmax = max(total_vals + inertia_vals) if (total_vals or inertia_vals) else 0.0
            self.os_pmax_grupy.set_ylim(0, ymax_pmax * 1.22 if ymax_pmax > 0 else 1.0)
            for bar, value in zip(bars_total, total_vals):
                if value > 0:
                    self.os_pmax_grupy.annotate(f'{value:,.0f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=7.3, color=pal['text'])
            for bar, value, row in zip(bars_inertia, inertia_vals, dane):
                if value > 0:
                    podpis = f"{value:,.0f}\n{row['count']} gen"
                    self.os_pmax_grupy.annotate(podpis, xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=7.1, color=pal['text'])
            legend = self.os_pmax_grupy.legend(loc='upper right', fontsize=7.5, frameon=True)
            if legend is not None:
                legend.get_frame().set_facecolor(pal['card'])
                legend.get_frame().set_edgecolor(pal['border'])
        odm_metrics = self.get_model_odm_metrics(model_type, key)
        odm_labels = [self.ODM_CODE_TO_NAME[code] for code in self.ODM_ORDER]
        odm_values = [float(odm_metrics.get(code, {}).get('energy', 0.0)) for code in self.ODM_ORDER]
        odm_colors = []
        for code in self.ODM_ORDER:
            if code == odm_code_selected:
                odm_colors.append(pal['accent'])
            else:
                odm_colors.append(pal['line_secondary'])
        bars_odm = self.os_odm_inercja.bar(range(len(self.ODM_ORDER)), odm_values, color=odm_colors, edgecolor=pal['border'], linewidth=0.8)
        self.os_odm_inercja.set_xticks(range(len(self.ODM_ORDER)))
        self.os_odm_inercja.set_xticklabels(odm_labels, rotation=12, ha='right', fontsize=8.2, color=pal['text'])
        ymax_odm = max(odm_values) if odm_values else 0.0
        self.os_odm_inercja.set_ylim(0, ymax_odm * 1.18 if ymax_odm > 0 else 1.0)
        for bar, value in zip(bars_odm, odm_values):
            self.os_odm_inercja.annotate(f'{value:,.0f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=7.5, color=pal['text'])
        self.canvas_grupy.draw_idle()

    def aktualizuj_analize_inercyjna(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            self.label_analiza_inercyjna.config(text='Brak wybranego modelu')
            return
        aktywne = []
        nieaktywne = []
        for gen in gen_list:
            H = self.get_h_value_for_generator(gen)
            if not H:
                continue
            Ei = self.get_generator_apparent_power(gen) * H
            rekord = (gen['reg_name'], Ei)
            if gen['status'] == 1:
                aktywne.append(rekord)
            else:
                nieaktywne.append(rekord)
        aktywne.sort(key=lambda x: x[1], reverse=True)
        nieaktywne.sort(key=lambda x: x[1], reverse=True)
        E = self.get_model_energy_now(model_type, key)
        P = self.get_model_demand_now(model_type, key)
        rocof = self.get_model_rocof_now(model_type, key)
        largest_loss = self.get_model_largest_loss(model_type, key)
        odm_metrics = self.get_model_odm_metrics(model_type, key)
        prog = self.pobierz_wartosc_graniczna()
        trigger_mode = self.email_alarm.config.trigger_mode
        top3 = sum((v for _, v in aktywne[:3]))
        koncentracja = top3 / E * 100 if E > 0 else 0.0
        linie = [f'Bieżąca inercja: {E:,.2f} MVA·s', f'Zapotrzebowanie systemu: {P:,.2f} MW', f'ROCOF N-1 (area=1): {rocof:.4f} Hz/s (ΔP krytyczne {largest_loss:,.2f} MW)']
        if prog is not None:
            if trigger_mode == 'above':
                if E >= prog:
                    linie.append(f'Margines względem progu: +{E - prog:,.2f} MVA·s')
                else:
                    linie.append(f'Niedobór względem progu: {E - prog:,.2f} MVA·s')
            elif E < prog:
                linie.append(f'Model jest poniżej progu o {prog - E:,.2f} MVA·s')
            else:
                linie.append(f'Model jest powyżej progu o {E - prog:,.2f} MVA·s')
        linie.append(f'Koncentracja TOP3 aktywnych: {koncentracja:.2f} % E_sys')
        if aktywne:
            linie.append(f'Najbardziej krytyczny aktywny generator: {aktywne[0][0]} ({aktywne[0][1]:.2f} MVA·s)')
        if nieaktywne:
            linie.append(f'Najlepszy kandydat do załączenia: {nieaktywne[0][0]} (+{nieaktywne[0][1]:.2f} MVA·s)')
        grupy = self.oblicz_udzialy_grup_generatorow(model_type, key)
        if grupy:
            top_grupa = grupy[0]
            linie.append(f"Dominująca grupa generatorów: {top_grupa['code']} ({top_grupa['label']}) = {top_grupa['share_pct']:.2f}% E_sys")
        if odm_metrics:
            linie.append('ODM [E / udział / ROCOF]:')
            for odm_code in self.ODM_ORDER:
                meta = odm_metrics.get(odm_code, {})
                linie.append(f"{self.ODM_CODE_TO_NAME[odm_code]}: {meta.get('energy', 0.0):,.2f} MVA·s | {meta.get('share_pct', 0.0):.2f}% E_sys | ROCOF {meta.get('rocof', 0.0):.4f} Hz/s")
        if prog is not None and trigger_mode == 'above' and (E < prog):
            brak = prog - E
            suma = 0.0
            rekom = []
            for nazwa, ei in nieaktywne:
                rekom.append(f'{nazwa} (+{ei:.1f})')
                suma += ei
                if suma >= brak:
                    break
            if rekom:
                linie.append('Rekomendowane załączenia do przekroczenia progu: ' + ', '.join(rekom))
        ok_color = True
        if prog is not None:
            ok_color = E >= prog
        self.label_analiza_inercyjna.config(text='\n'.join(linie), fg=self.get_palette()['success'] if ok_color else self.get_palette()['danger'])

    def formatuj_opis_modelu(self, godzina):
        opis = self.opis_modeli.get(godzina, {})
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', f'{godzina:02d}:30')
        prefix = ''
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if dodatnie:
            godz_min, _ = min(dodatnie, key=lambda x: x[1])
            godz_max, _ = max(dodatnie, key=lambda x: x[1])
            if godzina == godz_min:
                prefix += '[MIN] '
            if godzina == godz_max:
                prefix += '[MAX] '
        return f'{prefix}{godz_txt} | {data_txt}'

    def formatuj_opis_modelu_recznego(self, model_id):
        model = self.modele_reczne.get(model_id, {})
        opis = model.get('opis', {})
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        plik = opis.get('plik', '')
        inertia = float(model.get('inercja', 0.0))
        demand = float(model.get('zapotrzebowanie', 0.0))
        return f"{godz_txt} | {data_txt} | {inertia:,.2f} MVAÂ·s | {demand:,.2f} MW | {plik}"

    def klasyfikuj_model_wg_progu(self, wartosc):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None or prog <= 0:
            return None
        if self.email_alarm.config.trigger_mode == 'above':
            if wartosc < prog:
                return 'danger'
            if wartosc <= prog * (1 + self.bliski_prog_proc):
                return 'warning'
            return 'safe'
        else:
            if wartosc > prog:
                return 'safe'
            if wartosc >= prog * (1 - self.bliski_prog_proc):
                return 'warning'
            return 'danger'

    def odswiez_liste_modeli(self):
        self.listbox_modele.delete(0, tk.END)
        indeksy = []
        for godzina in sorted(self.wszystkie_generatory.keys()):
            txt = self.formatuj_opis_modelu(godzina)
            self.listbox_modele.insert(tk.END, txt)
            indeksy.append((self.listbox_modele.size() - 1, godzina))
        self.listbox_modele.insert(tk.END, '')
        self.listbox_modele.insert(tk.END, f'📋 RAZEM: {len(self.wszystkie_generatory)} modeli EPC')
        for idx, godzina in indeksy:
            status = self.klasyfikuj_model_wg_progu(self.inercja_na_godzine.get(godzina, 0.0))
            try:
                if status == 'danger':
                    self.listbox_modele.itemconfig(idx, fg='#dc2626' if not self.dark_mode_var.get() else '#fca5a5')
                elif status == 'warning':
                    self.listbox_modele.itemconfig(idx, fg='#ea580c' if not self.dark_mode_var.get() else '#fdba74')
                elif status == 'safe':
                    self.listbox_modele.itemconfig(idx, fg='#15803d' if not self.dark_mode_var.get() else '#86efac')
            except Exception:
                pass

    def odswiez_liste_modeli_reczne(self):
        self.listbox_modele_reczne.delete(0, tk.END)
        indeksy = []
        for model_id in self.reczne_order:
            if model_id not in self.modele_reczne:
                continue
            self.listbox_modele_reczne.insert(tk.END, self.formatuj_opis_modelu_recznego(model_id))
            indeksy.append((self.listbox_modele_reczne.size() - 1, model_id))
        self.listbox_modele_reczne.insert(tk.END, '')
        self.listbox_modele_reczne.insert(tk.END, f'🗂 RAZEM: {len(self.reczne_order)} modeli ręcznych')
        for idx, model_id in indeksy:
            status = self.klasyfikuj_model_wg_progu(self.modele_reczne[model_id].get('inercja', 0.0))
            try:
                if status == 'danger':
                    self.listbox_modele_reczne.itemconfig(idx, fg='#dc2626' if not self.dark_mode_var.get() else '#fca5a5')
                elif status == 'warning':
                    self.listbox_modele_reczne.itemconfig(idx, fg='#ea580c' if not self.dark_mode_var.get() else '#fdba74')
                elif status == 'safe':
                    self.listbox_modele_reczne.itemconfig(idx, fg='#15803d' if not self.dark_mode_var.get() else '#86efac')
            except Exception:
                pass

    def odswiez_liste_modeli_auto(self):
        self.listbox_modele_auto.delete(0, tk.END)
        indeksy = []
        for auto_id in self.auto_order:
            m = self.modele_auto[auto_id]
            opis = m['opis']
            self.listbox_modele_auto.insert(tk.END, f"{opis.get('godzina', '--:--')} | {opis.get('data', 'brak daty')} | {m.get('inercja', 0.0):,.2f} MVA·s | {m.get('zapotrzebowanie', 0.0):,.2f} MW | {opis.get('plik', '')}")
            indeksy.append((self.listbox_modele_auto.size() - 1, auto_id))
        self.listbox_modele_auto.insert(tk.END, '')
        self.listbox_modele_auto.insert(tk.END, f'🤖 RAZEM: {len(self.auto_order)} modeli bieżących')
        for idx, auto_id in indeksy:
            status = self.klasyfikuj_model_wg_progu(self.modele_auto[auto_id].get('inercja', 0.0))
            try:
                if status == 'danger':
                    self.listbox_modele_auto.itemconfig(idx, fg='#dc2626' if not self.dark_mode_var.get() else '#fca5a5')
                elif status == 'warning':
                    self.listbox_modele_auto.itemconfig(idx, fg='#ea580c' if not self.dark_mode_var.get() else '#fdba74')
                elif status == 'safe':
                    self.listbox_modele_auto.itemconfig(idx, fg='#15803d' if not self.dark_mode_var.get() else '#86efac')
            except Exception:
                pass

    def odswiez_panel_h_dyd(self):
        if not hasattr(self, 'tree_h_values'):
            return
        self.tree_h_values.delete(*self.tree_h_values.get_children())
        filtr = self.h_filter_var.get().strip().lower() if hasattr(self, 'h_filter_var') else ''
        keys = sorted(self.H_generatory.keys())
        shown = 0
        modified = 0
        for gen_name in keys:
            if filtr and filtr not in gen_name.lower():
                continue
            shown += 1
            val = float(self.H_generatory.get(gen_name, 0.0))
            src = 'DYD'
            tags = ()
            if gen_name in self.H_generatory_original and abs(val - float(self.H_generatory_original.get(gen_name, 0.0))) > 1e-9:
                src = 'Ręcznie'
                tags = ('h_manual',)
                modified += 1
            self.tree_h_values.insert('', tk.END, iid=gen_name, values=(gen_name, f'{val:.4f}', src), tags=tags)
        total = len(self.H_generatory)
        if total == 0:
            self.label_h_summary.config(text='Brak wczytanego pliku DYD', fg=self.get_palette()['muted'])
        else:
            filt_txt = f' | filtr: {shown}/{total}' if filtr else f' | widoczne: {shown}'
            self.label_h_summary.config(text=f'Baza H: {total} generatorów | ręcznie nadpisane w SQLite: {modified}{filt_txt}', fg=self.get_palette()['accent'])

    def on_h_tree_select(self, event=None):
        if not hasattr(self, 'tree_h_values'):
            return
        sel = self.tree_h_values.selection()
        if not sel:
            self.h_selected_generator_var.set('')
            self.h_edit_value_var.set('')
            return
        gen_name = sel[0]
        self.h_selected_generator_var.set(gen_name)
        self.h_edit_value_var.set(f"{float(self.H_generatory.get(gen_name, 0.0)):.4f}")

    def wyczysc_filtr_h(self):
        self.h_filter_var.set('')
        self.odswiez_panel_h_dyd()

    def zastosuj_reczna_zmiane_h(self):
        gen_name = self.h_selected_generator_var.get().strip()
        if not gen_name:
            messagebox.showwarning('Brak generatora', 'Najpierw wybierz generator z panelu H.')
            return
        txt = self.h_edit_value_var.get().strip().replace(',', '.')
        try:
            val = float(txt)
        except Exception:
            messagebox.showerror('Niepoprawne H', 'Podaj poprawną wartość liczbową H.')
            return
        if val < 0 or val >= 100:
            messagebox.showerror('Niepoprawne H', 'Wartość H musi należeć do zakresu 0 ≤ H < 100 s.')
            return
        self.h_repo.set_manual_value(gen_name, val)
        self.reload_h_from_database()
        self.przelicz_wszystkie_modele_po_zmianie_h()
        self.odswiez_panel_h_dyd()
        try:
            self.tree_h_values.selection_set(gen_name)
            self.tree_h_values.see(gen_name)
        except Exception:
            pass
        self.aktualizuj_status(f'✅ Zmieniono H dla generatora {gen_name} na {val:.4f} s')

    def przywroc_h_dla_zaznaczonego(self):
        gen_name = self.h_selected_generator_var.get().strip()
        if not gen_name:
            messagebox.showwarning('Brak generatora', 'Najpierw wybierz generator z panelu H.')
            return
        if gen_name not in self.H_generatory_original:
            messagebox.showwarning('Brak wartości źródłowej', 'Dla tego generatora nie ma wartości źródłowej z DYD.')
            return
        self.h_repo.restore_dyd_value(gen_name)
        self.reload_h_from_database()
        self.h_edit_value_var.set(f"{float(self.H_generatory.get(gen_name, 0.0)):.4f}")
        self.przelicz_wszystkie_modele_po_zmianie_h()
        self.odswiez_panel_h_dyd()
        try:
            self.tree_h_values.selection_set(gen_name)
            self.tree_h_values.see(gen_name)
        except Exception:
            pass
        self.aktualizuj_status(f'✅ Przywrócono H z DYD dla generatora {gen_name}')

    def przelicz_wszystkie_modele_po_zmianie_h(self):
        for godzina in self.wszystkie_generatory.keys():
            self.przelicz_inercje_dla_modelu('manual', godzina)
        for model_id in self.reczne_order:
            self.przelicz_inercje_dla_modelu('manual_batch', model_id)
        for auto_id in self.auto_order:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_reczne()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        model_type, key = self.get_current_model_ref()
        if model_type is not None:
            self.odswiez_aktywny_model()
        else:
            self.przygotuj_pusty_wykres_grup()

    def aktualizuj_panel_maksimum(self, check_alarm=False):
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if not dodatnie:
            self.label_maksimum_info.config(text='Maksymalna inercja z 24 godzin: brak danych', fg=self.get_palette()['accent'])
            self.label_limit_ocena.config(text='Ocena: brak wartości do porównania', fg=self.get_palette()['muted'])
            return
        godz_max, val_max = max(dodatnie, key=lambda x: x[1])
        godz_min, val_min = min(dodatnie, key=lambda x: x[1])
        opis_max = self.opis_modeli.get(godz_max, {})
        data_txt_max = opis_max.get('data', 'brak daty')
        godz_txt_max = opis_max.get('godzina', f'{godz_max:02d}:30')
        self.label_maksimum_info.config(text=f'Maks: {val_max:,.2f} MVA·s ({data_txt_max} {godz_txt_max}) | Min: {val_min:,.2f} MVA·s ({godz_min:02d}:30)', fg=self.get_palette()['accent'])
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            txt = self.entry_wartosc_graniczna.get().strip()
            if txt:
                self.label_limit_ocena.config(text='Ocena: niepoprawna wartość graniczna', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text='Ocena: brak wartości granicznej', fg=self.get_palette()['muted'])
            return
        if self.email_alarm.config.trigger_mode == 'above':
            if val_max > prog:
                self.label_limit_ocena.config(text=f'Ocena: PRZEKROCZONO wartość graniczną o {val_max - prog:,.2f} MVA·s', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text=f'Ocena: NIE przekroczono wartości granicznej (zapas {prog - val_max:,.2f} MVA·s)', fg=self.get_palette()['success'])
        elif val_min < prog:
                self.label_limit_ocena.config(text=f'Ocena: SPADŁO poniżej wartości granicznej o {prog - val_min:,.2f} MVA·s', fg=self.get_palette()['danger'])
        else:
                self.label_limit_ocena.config(text=f'Ocena: NIE spadło poniżej wartości granicznej (margines {val_min - prog:,.2f} MVA·s)', fg=self.get_palette()['success'])

    def aktualizuj_wykres(self):
        pal = self.get_palette()
        self.figura.patch.set_facecolor(pal['bg'])
        if not hasattr(self, 'os1_load') or self.os1_load is None:
            self.os1_load = self.os1.twinx()
        if not hasattr(self, 'os_manual_load') or self.os_manual_load is None:
            self.os_manual_load = self.os_manual.twinx()
        if not hasattr(self, 'os_auto_load') or self.os_auto_load is None:
            self.os_auto_load = self.os_auto.twinx()

        godziny = list(range(24))
        wartosci = [self.inercja_na_godzine.get(h, 0.0) for h in godziny]
        zapotrzebowania = [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h in godziny]

        self._draw_inertia_load_chart(
            self.os1,
            self.os1_load,
            godziny,
            wartosci,
            zapotrzebowania,
            [f'{h:02d}:30' for h in godziny],
            title='Inercja systemowa E_sys(t) = Σ S_i(t) × H_i',
            xlabel='Czas [HH:30]',
            left_label_coords=(-0.09, 0.78),
            window_size=24,
        )

        manual_payload = self.get_manual_chart_payload()
        manual_rows = manual_payload['rows']
        manual_xs = list(range(len(manual_rows)))
        manual_wartosci = [row['inercja'] for row in manual_rows]
        manual_zapotrzebowania = [row['zapotrzebowanie'] for row in manual_rows]
        manual_labels = [row['label'] for row in manual_rows]
        self._draw_inertia_load_chart(
            self.os_manual,
            self.os_manual_load,
            manual_xs,
            manual_wartosci,
            manual_zapotrzebowania,
            manual_labels,
            title=manual_payload['title'],
            xlabel=manual_payload['xlabel'],
            empty_message=manual_payload['empty_message'],
            left_label_coords=(-0.09, 0.78),
            window_size=manual_payload['window_size'],
        )

        auto_payload = self.get_auto_chart_payload()
        auto_rows = auto_payload['rows']
        auto_xs = list(range(len(auto_rows)))
        auto_wartosci = [row['inercja'] for row in auto_rows]
        auto_zapotrzebowania = [row['zapotrzebowanie'] for row in auto_rows]
        auto_labels = [row['label'] for row in auto_rows]
        self._draw_inertia_load_chart(
            self.os_auto,
            self.os_auto_load,
            auto_xs,
            auto_wartosci,
            auto_zapotrzebowania,
            auto_labels,
            title=auto_payload['title'],
            xlabel=auto_payload['xlabel'],
            empty_message=auto_payload['empty_message'],
            left_label_coords=(-0.09, 0.78),
            window_size=auto_payload['window_size'],
        )

        self.canvas.draw_idle()

    def open_h_database_window(self):
        try:
            if self.h_db_window is not None and self.h_db_window.winfo_exists():
                self.h_db_window.lift()
                self.h_db_window.focus_force()
                self.h_db_window.refresh()
                return
        except Exception:
            pass
        self.h_db_window = HDatabaseWindow(self.root, self)

    def close_rocof_window(self):
        try:
            if self.rocof_window is not None and self.rocof_window.winfo_exists():
                self.rocof_window.destroy()
        except Exception:
            pass
        self.rocof_window = None
        self.rocof_widgets = {}

    def on_rocof_mode_changed(self):
        self.refresh_rocof_window()

    def close_historical_window(self):
        win = self.history_window
        self.history_window = None
        try:
            if win is not None and win.winfo_exists():
                win.destroy()
        except Exception:
            pass

    def refresh_history_window(self):
        try:
            if self.history_window is not None and self.history_window.winfo_exists():
                self.configure_widget_tree(self.history_window)
                self.history_window.refresh_years()
                self.history_window.refresh()
                return
        except Exception:
            pass
        self.history_window = None

    def open_rocof_window(self):
        try:
            if self.rocof_window is not None and self.rocof_window.winfo_exists():
                self.rocof_window.lift()
                self.rocof_window.focus_force()
                self.refresh_rocof_window()
                return
        except Exception:
            pass
        pal = self.get_palette()
        win = tk.Toplevel(self.root)
        self.rocof_window = win
        win.title('Analiza ROCOF systemu i scenariuszy N-1')
        init_managed_toplevel(win, 'rocof_window', 1440, 920)
        win.configure(bg=pal['bg'])
        win.protocol('WM_DELETE_WINDOW', self.close_rocof_window)

        outer = tk.Frame(win, bg=pal['bg'], padx=10, pady=10)
        outer.pack(fill=tk.BOTH, expand=True)

        frame_info = tk.LabelFrame(outer, text='Metodyka', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_info.pack(fill=tk.X)
        info_txt = 'Panel ROCOF używa ujęcia systemowego zgodnego z COI (centre of inertia): df/dt = f0 × ΔP / (2 × Σ(S_i × H_i)). Podział na ODM nie jest tu stosowany, bo ocena dotyczy częstotliwości i bezwładności całego układu synchronicznego.'
        lbl_info = tk.Label(frame_info, text=info_txt, justify='left', anchor='w', wraplength=1280, bg=pal['card'], fg=pal['text'], font=self.font_ui_small)
        lbl_info.pack(fill=tk.X)

        frame_ctrl = tk.LabelFrame(outer, text='Scenariusz zakłócenia', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_ctrl.pack(fill=tk.X, pady=(10, 0))
        tk.Label(frame_ctrl, text='Tryb ΔP:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).pack(side=tk.LEFT)
        option_mode = tk.OptionMenu(frame_ctrl, self.rocof_mode_var, 'Największa jednostka online (Pgen)', 'Własna ΔP [MW]', command=lambda *_: self.on_rocof_mode_changed())
        option_mode.pack(side=tk.LEFT, padx=(8, 10))
        self.stylizuj_option_menu(option_mode)
        tk.Label(frame_ctrl, text='Własna ΔP [MW]:', bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small).pack(side=tk.LEFT)
        entry_dp = tk.Entry(frame_ctrl, textvariable=self.rocof_custom_dp_var, width=14, font=self.font_code_small)
        entry_dp.pack(side=tk.LEFT, padx=(8, 8))
        self.stylizuj_entry(entry_dp)
        entry_dp.bind('<KeyRelease>', lambda *_: self.refresh_rocof_window())
        self.utworz_przycisk(frame_ctrl, 'Przelicz ROCOF', self.refresh_rocof_window, role='primary').pack(side=tk.LEFT, padx=4)
        self.utworz_przycisk(frame_ctrl, 'Zamknij', self.close_rocof_window, role='neutral').pack(side=tk.RIGHT, padx=4)

        frame_summary = tk.LabelFrame(outer, text='Podsumowanie bieżącego modelu', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_summary.pack(fill=tk.X, pady=(10, 0))
        lbl_summary = tk.Label(frame_summary, text='Brak wybranego modelu', justify='left', anchor='w', bg=pal['card'], fg=pal['accent'], font=self.font_code_small)
        lbl_summary.pack(fill=tk.X)

        frame_body = tk.Frame(outer, bg=pal['bg'])
        frame_body.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        frame_left = tk.LabelFrame(frame_body, text='Wykresy ROCOF', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        frame_right = tk.Frame(frame_body, bg=pal['bg'], width=560)
        frame_right.pack(side=tk.LEFT, fill=tk.BOTH, expand=False, padx=(10, 0))
        frame_right.pack_propagate(False)

        figura_rocof, (ax_rocof_system, ax_rocof_n1) = plt.subplots(2, 1, figsize=(9.2, 7.4), gridspec_kw={'height_ratios': [0.95, 1.25]})
        figura_rocof.subplots_adjust(bottom=0.08, left=0.10, right=0.98, top=0.93, hspace=0.44)
        canvas_rocof = FigureCanvasTkAgg(figura_rocof, master=frame_left)
        canvas_rocof.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        frame_system = tk.LabelFrame(frame_right, text='Parametry systemowe', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_system.pack(fill=tk.BOTH, expand=True)
        cols_system = ('metric', 'value', 'desc')
        tree_system = ttk.Treeview(frame_system, columns=cols_system, show='headings', height=7)
        for col, header, width, anchor in [('metric', 'Parametr', 170, 'w'), ('value', 'Wartość', 130, 'center'), ('desc', 'Opis', 220, 'w')]:
            tree_system.heading(col, text=header)
            tree_system.column(col, width=width, anchor=anchor)
        tree_system.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        system_scroll = ttk.Scrollbar(frame_system, orient='vertical', command=tree_system.yview)
        tree_system.configure(yscrollcommand=system_scroll.set)
        system_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        frame_n1 = tk.LabelFrame(frame_right, text='Najgorsze scenariusze N-1', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_n1.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        cols_n1 = ('generator', 'grupa', 'dp', 'rocof', 't49')
        tree_n1 = ttk.Treeview(frame_n1, columns=cols_n1, show='headings', height=9)
        for col, header, width in [('generator', 'Generator', 210), ('grupa', 'Typ', 125), ('dp', 'ΔP [MW]', 92), ('rocof', 'RoCoF [Hz/s]', 110), ('t49', 't do 49 Hz [s]', 118)]:
            tree_n1.heading(col, text=header)
            tree_n1.column(col, width=width, anchor='center' if col != 'generator' else 'w')
        tree_n1.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        n1_scroll = ttk.Scrollbar(frame_n1, orient='vertical', command=tree_n1.yview)
        tree_n1.configure(yscrollcommand=n1_scroll.set)
        n1_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        frame_notes = tk.LabelFrame(outer, text='Komentarz analityczny', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_notes.pack(fill=tk.BOTH, expand=False, pady=(10, 0))
        txt_notes = scrolledtext.ScrolledText(frame_notes, height=8, wrap=tk.WORD, font=self.font_code_small)
        txt_notes.pack(fill=tk.BOTH, expand=True)

        self.rocof_widgets = {
            'info': lbl_info,
            'summary': lbl_summary,
            'custom_entry': entry_dp,
            'tree_system': tree_system,
            'tree_n1': tree_n1,
            'notes': txt_notes,
            'figure': figura_rocof,
            'ax_system': ax_rocof_system,
            'ax_n1': ax_rocof_n1,
            'canvas': canvas_rocof,
        }
        self.refresh_rocof_window()

    def open_historical_window(self):
        try:
            if self.history_window is not None and self.history_window.winfo_exists():
                self.history_window.lift()
                self.history_window.focus_force()
                self.refresh_history_window()
                return
        except Exception:
            pass
        self.history_window = HistoricalInertiaWindow(self.root, self)
        self.history_window.protocol('WM_DELETE_WINDOW', self.close_historical_window)

    def refresh_rocof_window(self):
        try:
            if self.rocof_window is None or not self.rocof_window.winfo_exists():
                return
        except Exception:
            self.rocof_window = None
            self.rocof_widgets = {}
            return
        refs = self.rocof_widgets
        if not refs:
            return
        pal = self.get_palette()
        self.rocof_window.configure(bg=pal['bg'])
        entry = refs.get('custom_entry')
        if entry is not None:
            entry.configure(state='normal' if self.rocof_mode_var.get() == 'Własna ΔP [MW]' else 'disabled')
        model_type, key = self.get_current_model_ref()
        tree_system = refs['tree_system']
        tree_n1 = refs['tree_n1']
        notes = refs['notes']
        for tree in (tree_system, tree_n1):
            for item in tree.get_children():
                tree.delete(item)
        notes.delete('1.0', tk.END)
        figura = refs['figure']
        ax_system = refs['ax_system']
        ax_n1 = refs['ax_n1']
        figura.patch.set_facecolor(pal['bg'])
        ax_system.clear()
        ax_n1.clear()
        self.stylizuj_os(ax_system, grid=True, grid_axis='y')
        self.stylizuj_os(ax_n1, grid=True, grid_axis='x')
        if model_type is None:
            refs['summary'].config(text='Brak wybranego modelu do analizy ROCOF.', fg=pal['muted'])
            ax_system.text(0.5, 0.5, 'Wybierz model 24h, ręczny lub bieżący', ha='center', va='center', transform=ax_system.transAxes, color=pal['muted'])
            ax_n1.text(0.5, 0.5, 'Brak danych scenariuszy N-1', ha='center', va='center', transform=ax_n1.transAxes, color=pal['muted'])
            refs['canvas'].draw_idle()
            return
        custom_dp = None
        mode = self.rocof_mode_var.get()
        custom_error = None
        if mode == 'Własna ΔP [MW]':
            raw = self.rocof_custom_dp_var.get().strip().replace(',', '.')
            if raw:
                try:
                    custom_dp = float(raw)
                    if custom_dp <= 0:
                        raise ValueError()
                except Exception:
                    custom_error = 'Własna ΔP musi być dodatnią liczbą.'
            else:
                custom_error = 'Podaj wartość ΔP dla trybu ręcznego.'
        analiza = self.build_rocof_analysis(model_type, key, custom_delta_p=custom_dp)
        opis = self.get_model_meta(model_type, key) or {}
        if analiza is None:
            refs['summary'].config(text='Brak aktywnych generatorów do analizy ROCOF.', fg=pal['warning'])
            ax_system.text(0.5, 0.5, 'Brak danych systemowych', ha='center', va='center', transform=ax_system.transAxes, color=pal['muted'])
            ax_n1.text(0.5, 0.5, 'Brak danych', ha='center', va='center', transform=ax_n1.transAxes, color=pal['muted'])
            refs['canvas'].draw_idle()
            return
        system_rocof = analiza['rocof']
        time_to_49 = analiza['time_to_49hz']
        dp_txt = analiza['default_dp']
        largest_unit = analiza.get('largest_unit') or {}
        risk_txt = 'niski'
        if system_rocof >= 2.0:
            risk_txt = 'bardzo wysoki'
        elif system_rocof >= 1.0:
            risk_txt = 'wysoki'
        elif system_rocof >= 0.5:
            risk_txt = 'umiarkowany'
        summary_lines = [
            f"Model: {opis.get('data', 'brak daty')} {opis.get('godzina', '--:--')} | {self.get_model_type_label(model_type)}",
            f"Jednostki online area=1 = {analiza['eligible_count']} | z H > 0 = {analiza['inertia_count']} | Σ(S×H) = {analiza['energy']:,.2f} MVA·s",
            f"ΔP = {dp_txt:,.2f} MW | RoCoF0 = {system_rocof:.4f} Hz/s | czas do 49 Hz ≈ {time_to_49:.3f} s",
            f"Ocena orientacyjna: poziom stresu częstotliwościowego = {risk_txt}. {analiza['source_note']}",
        ]
        if custom_error:
            summary_lines.append(custom_error)
        refs['summary'].config(text='\n'.join(summary_lines), fg=pal['danger'] if system_rocof >= 1.0 else pal['accent'])

        system_scope_desc = 'COI / cały system bez podziału na ODM'
        delta_p_desc = 'tryb ręczny' if mode == 'Własna ΔP [MW]' and custom_dp is not None else 'największa jednostka online'
        critical_unit_label = largest_unit.get('reg_name', '-') if largest_unit else '-'
        critical_unit_desc = f"P = {largest_unit.get('delta_p', 0.0):,.2f} MW | typ: {largest_unit.get('group_label', '-')}" if largest_unit else '-'
        system_rows = [
            ('Zakres analizy', 'Cały system', system_scope_desc),
            ('Jednostki online', str(int(analiza['eligible_count'])), 'po filtrze area = 1'),
            ('Jednostki z H > 0', str(int(analiza['inertia_count'])), 'wliczone do Σ(S×H)'),
            ('Σ(S×H)', f"{analiza['energy']:,.2f} MVA·s", 'energia kinetyczna modelu'),
            ('ΔP bazowe', f"{dp_txt:,.2f} MW", delta_p_desc),
            ('Jednostka krytyczna', critical_unit_label, critical_unit_desc),
            ('RoCoF0', f"{system_rocof:.4f} Hz/s", f'poziom {risk_txt}'),
            ('t do 49 Hz', f"{time_to_49:.3f} s" if time_to_49 > 0 else '-', 'przybliżenie liniowe'),
        ]
        for metric, value, desc in system_rows:
            tree_system.insert('', tk.END, values=(metric, value, desc))
        for row in analiza['contingencies'][:12]:
            tree_n1.insert('', tk.END, values=(row['reg_name'], row['group_label'], f"{row['delta_p']:.2f}", f"{row['rocof']:.4f}", f"{row['time_to_49hz']:.3f}" if row['time_to_49hz'] > 0 else '-'))

        bar_color = pal['accent']
        if system_rocof >= 1.0:
            bar_color = pal['danger']
        elif system_rocof >= 0.5:
            bar_color = pal['warning']
        bars_system = ax_system.bar([0], [system_rocof], color=bar_color, edgecolor=pal['border'], linewidth=0.9, width=0.18)
        ax_system.set_title('RoCoF początkowy dla całego systemu (COI)')
        ax_system.set_ylabel('RoCoF [Hz/s]')
        ax_system.set_xticks([0])
        ax_system.set_xticklabels(['System'], color=pal['text'])
        ax_system.set_xlim(-0.42, 0.42)
        ymax_system = max(2.2, system_rocof * 1.35 if system_rocof > 0 else 1.0)
        ax_system.set_ylim(0, ymax_system)
        for limit, label, color in [(0.5, '0.5 Hz/s', pal['line_secondary']), (1.0, '1.0 Hz/s', pal['warning']), (2.0, '2.0 Hz/s', pal['danger'])]:
            ax_system.axhline(limit, color=color, linestyle='--', linewidth=0.9, alpha=0.9)
            ax_system.text(0.98, limit / ymax_system, label, transform=ax_system.transAxes, ha='right', va='bottom', fontsize=7.4, color=color)
        for bar in bars_system:
            ax_system.annotate(f'{system_rocof:.3f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 5), textcoords='offset points', ha='center', va='bottom', fontsize=9.0, color=pal['text'], fontweight='bold')
        ax_system.text(0.02, 0.95, f"Σ(S×H) = {analiza['energy']:,.0f} MVA·s\nΔP = {dp_txt:,.1f} MW", transform=ax_system.transAxes, ha='left', va='top', fontsize=8.1, color=pal['text'], bbox={'boxstyle': 'round,pad=0.28', 'facecolor': pal['card'], 'edgecolor': pal['border'], 'alpha': 0.92})

        top_n1 = list(reversed(analiza['contingencies'][:8]))
        n1_labels = [row['reg_name'][:24] for row in top_n1]
        n1_values = [row['rocof'] for row in top_n1]
        bars_n1 = ax_n1.barh(range(len(top_n1)), n1_values, color=pal['warning'], edgecolor=pal['border'], linewidth=0.8)
        ax_n1.set_title('Najgorsze scenariusze N-1 dla bieżącego modelu')
        ax_n1.set_xlabel('RoCoF [Hz/s]')
        ax_n1.set_yticks(range(len(top_n1)))
        ax_n1.set_yticklabels(n1_labels, fontsize=8.2, color=pal['text'])
        ax_n1.set_xlim(0, max(n1_values) * 1.18 if max(n1_values or [0.0]) > 0 else 1.0)
        for bar, row in zip(bars_n1, top_n1):
            if row['rocof'] > 0:
                ax_n1.annotate(f"{row['rocof']:.3f} | ΔP={row['delta_p']:.0f} MW", xy=(bar.get_width(), bar.get_y() + bar.get_height() / 2), xytext=(6, 0), textcoords='offset points', ha='left', va='center', fontsize=7.4, color=pal['text'])

        note_lines = [
            'Założenia analizy ROCOF:',
            '1. Pokazywany jest początkowy RoCoF po zakłóceniu, bez odpowiedzi pierwotnej i bez tłumienia obciążenia.',
            '2. Panel używa ujęcia systemowego (COI / cały układ synchroniczny), więc nie rozbija wyniku na ODM.',
            '3. Do analizy są brane tylko generatory z area = 1; ΔP domyślnie odpowiada największej aktywnej jednostce online według Pgen.',
            '4. Σ(S×H) jest liczone z jednostek mających dodatnie H w modelu i dodatnią moc pozorną S.',
            '5. Czas do 49 Hz jest przybliżeniem liniowym: 1 Hz / RoCoF.',
            '',
            'Interpretacja:',
            '• około 0.5 Hz/s: łagodniejszy przebieg początkowy',
            '• około 1.0 Hz/s: szybki spadek częstotliwości',
            '• około 2.0 Hz/s i więcej: bardzo szybki spadek, mało czasu na reakcję rezerw i zabezpieczeń',
        ]
        notes.insert(tk.END, '\n'.join(note_lines))
        refs['canvas'].draw_idle()

    def reload_h_from_database(self, only_names=None):
        if only_names is None:
            self.H_generatory = dict(self.h_repo.get_effective_map())
            self.H_generatory_original = dict(self.h_repo.get_dyd_map())
        else:
            effective = self.h_repo.get_effective_map(only_names=only_names)
            original = self.h_repo.get_dyd_map(only_names=only_names)
            self.H_generatory.update(effective)
            for name in only_names:
                if name in original:
                    self.H_generatory_original[name] = original[name]
        self.H_group_defaults = dict(self.h_repo.get_group_default_map())
        if hasattr(self, 'h_db_window') and self.h_db_window is not None:
            try:
                if self.h_db_window.winfo_exists():
                    self.h_db_window.refresh()
            except Exception:
                pass

    def open_archive_browser(self):
        ArchiveBrowserWindow(self.root, self)

    def build_archive_payload(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return None
        opis = dict(self.get_model_meta(model_type, key) or {})
        original_statuses = dict(self.get_model_original_statuses(model_type, key) or {})
        generators = []
        for gen in gen_list:
            H = float(self.get_h_value_for_generator(gen))
            s_apparent = self.get_generator_apparent_power(gen)
            Ei = s_apparent * H
            generators.append({
                'reg_name': gen.get('reg_name', ''),
                'status': int(gen.get('status', 0)),
                'status_original': int(original_statuses.get(gen.get('reg_name', ''), gen.get('status', 0))),
                'area': gen.get('area', '1'),
                'pgen': float(gen.get('pgen', 0.0)),
                'pmax': float(gen.get('pmax', 0.0)),
                'qmax': float(gen.get('qmax', 0.0)),
                's_apparent': float(s_apparent),
                'H': H,
                'Ei': Ei,
                'group_code': gen.get('group_code'),
                'zone': gen.get('zone', ''),
                'group_label': self.formatuj_grupe_generatora(gen),
                'odm_code': gen.get('odm_code', ''),
                'odm_name': self.formatuj_odm_generatora(gen),
                'odm_source': gen.get('odm_source', ''),
            })
        return {
            'archived_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'model_type': model_type,
            'model_date': opis.get('data', ''),
            'model_hour': opis.get('godzina', ''),
            'source_file': opis.get('plik', ''),
            'source_path': opis.get('sciezka', ''),
            'total_inertia': float(self.get_model_energy_now(model_type, key)),
            'original_inertia': float(self.get_model_energy_original(model_type, key)),
            'demand': float(self.get_model_demand_now(model_type, key)),
            'rocof': float(self.get_model_rocof_now(model_type, key)),
            'largest_loss': float(self.get_model_largest_loss(model_type, key)),
            'odm_metrics': self.get_model_odm_metrics(model_type, key),
            'connected_count': int(self.get_model_value(model_type, key, 'liczba_zalaczonych')),
            'opis': opis,
            'generators': generators,
        }

    def save_model_to_archive(self, model_type, key):
        payload = self.build_archive_payload(model_type, key)
        if not payload:
            return None
        try:
            return self.archive_repo.insert_snapshot(payload)
        except Exception:
            return None

    def save_all_models_to_archive(self):
        inserted = 0
        for godzina in sorted(self.wszystkie_generatory.keys()):
            if self.save_model_to_archive('manual', godzina):
                inserted += 1
        for model_id in self.reczne_order:
            if self.save_model_to_archive('manual_batch', model_id):
                inserted += 1
        for auto_id in self.auto_order:
            if self.save_model_to_archive('auto', auto_id):
                inserted += 1
        if inserted > 0:
            self.clear_history_analysis_cache()
            self.refresh_history_window()
        return inserted

    def restore_from_archive_payload(self, payload):
        self.archive_restore_counter += 1
        archive_id = f'archive_{self.archive_restore_counter:05d}'
        generators = []
        original_statuses = {}
        h_map = {}
        for g in payload.get('generators', []):
            reg_name = g.get('reg_name', '')
            area_code = self.normalizuj_kod_odm(g.get('area', '1'))
            zone_code = self.normalizuj_kod_odm(g.get('zone', ''))
            odm_code = self.normalizuj_kod_odm(g.get('odm_code', '') or zone_code)
            generators.append({'reg_name': reg_name, 'status': int(g.get('status', 0)), 'area': area_code, 'pgen': float(g.get('pgen', 0.0)), 'pmax': float(g.get('pmax', 0.0)), 'qmax': float(g.get('qmax', 0.0)), 'group_code': g.get('group_code'), 'zone': zone_code, 'odm_code': odm_code, 'odm_name': g.get('odm_name', self.formatuj_odm_z_kodu(odm_code)), 'odm_source': g.get('odm_source', f'ZONE={zone_code}' if zone_code else '')})
            original_statuses[reg_name] = int(g.get('status_original', g.get('status', 0)))
            if float(g.get('H', 0.0)) > 0:
                h_map[reg_name] = float(g.get('H', 0.0))
        self.H_generatory.update(h_map)
        opis_payload = dict(payload.get('opis') or {})
        opis_payload.setdefault('data', payload.get('model_date', ''))
        opis_payload.setdefault('godzina', payload.get('model_hour', ''))
        opis_payload.setdefault('plik', payload.get('source_file', ''))
        opis_payload['archiwum_id'] = payload.get('id')
        opis_payload['zrodlo'] = 'ARCHIWUM'
        self.modele_reczne[archive_id] = {
            'generatory': generators,
            'statusy_oryginalne': original_statuses,
            'opis': opis_payload,
            'inercja': float(payload.get('total_inertia', 0.0)),
            'inercja_oryginalna': float(payload.get('original_inertia', 0.0)),
            'odm_metrics': payload.get('odm_metrics', self.build_empty_odm_metrics()),
            'rocof': float(payload.get('rocof', 0.0)),
            'largest_loss': float(payload.get('largest_loss', 0.0)),
            'liczba_zalaczonych': int(payload.get('connected_count', 0)),
            'zapotrzebowanie': float(payload.get('demand', 0.0)),
        }
        self.reczne_order.append(archive_id)
        self.przelicz_inercje_dla_modelu('manual_batch', archive_id)
        self.obliczenia_wykonane = True
        self.odswiez_liste_modeli_reczne()
        self.aktywny_typ_modelu = 'manual_batch'
        self.aktualny_reczny_id = archive_id
        self.aktualny_auto_id = None
        self.aktualna_godzina = None
        self.clear_other_model_selections(keep='manual_batch')
        self.odswiez_aktywny_model()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.aktualizuj_status('✅ Odtworzono model z archiwum i wczytano do aplikacji')

    def build_hourly_report_rows(self):
        rows = []
        for godzina in sorted(self.wszystkie_generatory.keys()):
            opis = self.opis_modeli.get(godzina, {})
            godzina_label = opis.get('godzina', f'{godzina:02d}:30')
            e_org = float(self.inercja_oryginalna_na_godzine.get(godzina, 0.0) or 0.0)
            e_now = float(self.inercja_na_godzine.get(godzina, 0.0) or 0.0)
            rows.append({
                'point': godzina_label,
                'data': opis.get('data', ''),
                'godzina': godzina_label,
                'godzina_idx': godzina,
                'inercja_oryginalna': e_org,
                'inercja_po_zmianach': e_now,
                'delta_inercji': e_now - e_org,
                'load': float(self.zapotrzebowanie_na_godzine.get(godzina, 0.0) or 0.0),
                'rocof': float(self.rocof_na_godzine.get(godzina, 0.0) or 0.0),
                'largest_loss': float(self.largest_loss_na_godzine.get(godzina, 0.0) or 0.0),
                'connected_count': int(self.liczba_zalaczonych_na_godzine.get(godzina, 0) or 0),
            })
        return rows

    def eksportuj_do_excel(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        dane = self.pobierz_dane_szczegolowe(model_type, key)
        if dane is None:
            messagebox.showwarning('Brak danych', 'Brak danych do eksportu.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz plik Excel', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not sciezka:
            return
        try:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = 'Szczegóły'
            header_fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            manual_fill = PatternFill(fill_type='solid', fgColor='FFF3CD')
            bold_font = Font(bold=True)
            manual_font = Font(bold=True, color='C77700')
            center = Alignment(horizontal='center', vertical='center')
            opis = self.get_model_meta(model_type, key)
            ws1['A1'] = 'Typ modelu'
            ws1['B1'] = self.get_model_type_label(model_type)
            ws1['A2'] = 'Data modelu'
            ws1['B2'] = opis.get('data', '')
            ws1['A3'] = 'Godzina modelu'
            ws1['B3'] = opis.get('godzina', '')
            ws1['A4'] = 'Filtr'
            ws1['B4'] = self.tekst_filtrowania if self.tekst_filtrowania else '(brak)'
            ws1['A5'] = 'Całkowita inercja systemu [MVA·s]'
            ws1['B5'] = dane['podsumowanie']['energia_calkowita']
            ws1['A6'] = 'Zapotrzebowanie systemu [MW]'
            ws1['B6'] = dane['podsumowanie']['zapotrzebowanie']
            ws1['A7'] = 'ROCOF N-1 area=1 [Hz/s]'
            ws1['B7'] = dane['podsumowanie']['rocof']
            ws1['A8'] = 'Krytyczne ΔP [MW]'
            ws1['B8'] = dane['podsumowanie']['largest_loss']
            for cell in ['A1', 'A2', 'A3', 'A4', 'A5', 'A6', 'A7', 'A8']:
                ws1[cell].font = bold_font
            start_row = 10
            headers = ['LP', 'Kod generatora', 'Grupa', 'ODM', 'Status', 'S [MVA]', 'H [s]', 'E_i = S × H [MVA·s]', 'Uwzględniony w obliczeniach', 'Ręczna zmiana']
            for col, header in enumerate(headers, 1):
                c = ws1.cell(row=start_row, column=col, value=header)
                c.font = bold_font
                c.fill = header_fill
                c.alignment = center
            for r, w in enumerate(dane['wiersze'], start=start_row + 1):
                ws1.cell(row=r, column=1, value=w['lp'])
                ws1.cell(row=r, column=2, value=w['kod_generatora'])
                ws1.cell(row=r, column=3, value=w['grupa'])
                ws1.cell(row=r, column=4, value=w['odm'])
                ws1.cell(row=r, column=5, value=w['status'])
                ws1.cell(row=r, column=6, value=w['s_obliczeniowa'])
                ws1.cell(row=r, column=7, value=w['H'])
                ws1.cell(row=r, column=8, value=w['Ei'])
                ws1.cell(row=r, column=9, value=w['uwzgledniony'])
                ws1.cell(row=r, column=10, value='TAK' if w['reczna_zmiana'] else '-')
                if w['reczna_zmiana']:
                    for col in range(1, 11):
                        ws1.cell(row=r, column=col).fill = manual_fill
                        ws1.cell(row=r, column=col).font = manual_font
            for col, width in {'A': 8, 'B': 28, 'C': 16, 'D': 14, 'E': 12, 'F': 14, 'G': 12, 'H': 22, 'I': 30, 'J': 16}.items():
                ws1.column_dimensions[col].width = width
            ws_odm = wb.create_sheet('ODM')
            headers_odm = ['ODM', 'Inercja [MVA·s]', 'Udział [% E_sys]', 'Liczba generatorów', 'Krytyczne ΔP [MW]', 'ROCOF N-1 area=1 [Hz/s]']
            for col, header in enumerate(headers_odm, 1):
                c = ws_odm.cell(row=1, column=col, value=header)
                c.font = bold_font
                c.fill = header_fill
                c.alignment = center
            for idx, odm_code in enumerate(self.ODM_ORDER, start=2):
                meta = dane['podsumowanie']['odm_metrics'].get(odm_code, {})
                vals = [self.ODM_CODE_TO_NAME[odm_code], meta.get('energy', 0.0), meta.get('share_pct', 0.0), int(meta.get('connected_count', 0)), meta.get('largest_loss', 0.0), meta.get('rocof', 0.0)]
                for col, val in enumerate(vals, 1):
                    ws_odm.cell(row=idx, column=col, value=val)
            for col, width in {'A': 16, 'B': 18, 'C': 16, 'D': 18, 'E': 20, 'F': 18}.items():
                ws_odm.column_dimensions[col].width = width
            if model_type == 'manual':
                ws2 = wb.create_sheet('Podsumowanie godzin')
                headers2 = ['Data', 'Godzina', 'Inercja oryginalna [MVA·s]', 'Inercja po zmianach [MVA·s]', 'Różnica [MVA·s]', 'Zapotrzebowanie [MW]', 'ROCOF N-1 area=1 [Hz/s]']
                for col, header in enumerate(headers2, 1):
                    c = ws2.cell(row=1, column=col, value=header)
                    c.font = bold_font
                    c.fill = header_fill
                    c.alignment = center
                posortowane = sorted(self.wszystkie_generatory.keys())
                for idx, godzina in enumerate(posortowane, start=2):
                    opis_m = self.opis_modeli.get(godzina, {})
                    e_org = self.inercja_oryginalna_na_godzine.get(godzina, 0.0)
                    e_now = self.inercja_na_godzine.get(godzina, 0.0)
                    p_load = self.zapotrzebowanie_na_godzine.get(godzina, 0.0)
                    ws2.cell(row=idx, column=1, value=opis_m.get('data', ''))
                    ws2.cell(row=idx, column=2, value=opis_m.get('godzina', f'{godzina:02d}:30'))
                    ws2.cell(row=idx, column=3, value=e_org)
                    ws2.cell(row=idx, column=4, value=e_now)
                    ws2.cell(row=idx, column=5, value=e_now - e_org)
                    ws2.cell(row=idx, column=6, value=p_load)
                    ws2.cell(row=idx, column=7, value=self.rocof_na_godzine.get(godzina, 0.0))
                for col, width in {'A': 16, 'B': 14, 'C': 24, 'D': 24, 'E': 18, 'F': 18, 'G': 18}.items():
                    ws2.column_dimensions[col].width = width
            arch_rows = self.archive_repo.fetch_records({}, sort_by='archived_at', sort_desc=True)[:2000]
            if arch_rows:
                ws_arch = wb.create_sheet('Archiwum inercji')
                headers_arch = ['ID', 'Archiwizowano', 'Typ', 'Data', 'Godzina', 'Plik', 'Inercja [MVA·s]', 'Inercja org. [MVA·s]', 'Load [MW]', 'L. gen']
                for col, header in enumerate(headers_arch, 1):
                    c = ws_arch.cell(row=1, column=col, value=header)
                    c.font = bold_font
                    c.fill = header_fill
                    c.alignment = center
                for r, z in enumerate(arch_rows, start=2):
                    vals = [z.get('id'), z.get('archived_at', ''), z.get('model_type', ''), z.get('model_date', ''), z.get('model_hour', ''), z.get('source_file', ''), z.get('total_inertia', 0.0), z.get('original_inertia', 0.0), z.get('demand', 0.0), z.get('connected_count', 0)]
                    for cidx, val in enumerate(vals, start=1):
                        ws_arch.cell(row=r, column=cidx, value=val)
                for col, width in {'A': 8, 'B': 22, 'C': 10, 'D': 14, 'E': 12, 'F': 32, 'G': 18, 'H': 18, 'I': 14, 'J': 10}.items():
                    ws_arch.column_dimensions[col].width = width
            if self.alarm_history:
                ws3 = wb.create_sheet('Historia alarmów')
                headers3 = ['Kiedy alarm wystąpił', 'Tryb', 'Typ modelu', 'Data modelu', 'Godzina modelu', 'Wartość [MVA·s]', 'Próg [MVA·s]', 'Odchyłka [MVA·s]', 'Plik', 'Mail', 'Status wysyłki', 'Opis']
                for col, header in enumerate(headers3, 1):
                    c = ws3.cell(row=1, column=col, value=header)
                    c.font = bold_font
                    c.fill = header_fill
                    c.alignment = center
                for r, z in enumerate(self.alarm_history, start=2):
                    wartosc = float(z.get('wartosc', 0.0) or 0.0)
                    prog = float(z.get('prog', 0.0) or 0.0)
                    ws3.cell(row=r, column=1, value=z.get('czas_alarmu', ''))
                    ws3.cell(row=r, column=2, value=z.get('trigger_mode', ''))
                    ws3.cell(row=r, column=3, value=z.get('model_type_label') or self.get_model_type_label(z.get('model_type')))
                    ws3.cell(row=r, column=4, value=z.get('data_modelu', ''))
                    ws3.cell(row=r, column=5, value=z.get('godzina_modelu', ''))
                    ws3.cell(row=r, column=6, value=wartosc)
                    ws3.cell(row=r, column=7, value=prog)
                    ws3.cell(row=r, column=8, value=wartosc - prog)
                    ws3.cell(row=r, column=9, value=z.get('source_file', ''))
                    ws3.cell(row=r, column=10, value='TAK' if z.get('mail_attempted') else 'NIE')
                    ws3.cell(row=r, column=11, value=z.get('mail_status', ''))
                    ws3.cell(row=r, column=12, value=z.get('opis', ''))
                for col, width in {'A': 20, 'B': 10, 'C': 12, 'D': 14, 'E': 14, 'F': 16, 'G': 14, 'H': 16, 'I': 28, 'J': 10, 'K': 24, 'L': 24}.items():
                    ws3.column_dimensions[col].width = width
            wb.save(sciezka)
            messagebox.showinfo('Eksport zakończony', f'Zapisano plik:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu', f'Nie udało się zapisać pliku Excel:\n{e}')

    def eksportuj_do_excel(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        dane = self.pobierz_dane_szczegolowe(model_type, key)
        if dane is None:
            messagebox.showwarning('Brak danych', 'Brak danych do eksportu.')
            return
        sciezka = filedialog.asksaveasfilename(
            title='Zapisz plik Excel',
            defaultextension='.xlsx',
            filetypes=[('Excel', '*.xlsx')],
        )
        if not sciezka:
            return

        try:
            hourly_rows = self.build_hourly_report_rows()
            podsumowanie = dane['podsumowanie']
            opis = self.get_model_meta(model_type, key)
            model_label = self.get_model_type_label(model_type)
            threshold = self.pobierz_wartosc_graniczna()
            original_energy = float(self.get_model_energy_original(model_type, key) or 0.0)
            current_energy = float(podsumowanie['energia_calkowita'] or 0.0)
            energy_delta = current_energy - original_energy
            demand_now = float(podsumowanie['zapotrzebowanie'] or 0.0)
            rocof_now = float(podsumowanie['rocof'] or 0.0)
            largest_loss_now = float(podsumowanie['largest_loss'] or 0.0)
            export_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            wiersze_posortowane = sorted(
                dane['wiersze'],
                key=lambda row: float(row.get('Ei', 0.0) or 0.0),
                reverse=True,
            )
            top_row = wiersze_posortowane[0] if wiersze_posortowane else None
            odm_metrics = podsumowanie.get('odm_metrics', {})
            top_odm_code = (
                max(
                    self.ODM_ORDER,
                    key=lambda code: float((odm_metrics.get(code, {}) or {}).get('energy', 0.0) or 0.0),
                )
                if odm_metrics
                else ''
            )
            top_odm = odm_metrics.get(top_odm_code, {}) if top_odm_code else {}
            manual_change_count = sum((1 for row in dane['wiersze'] if row['reczna_zmiana']))
            uwzglednione_count = sum((1 for row in dane['wiersze'] if row['uwzgledniony'] == 'TAK'))

            report_notes = [
                f'Raport dotyczy modelu [{model_label}] z chwili {opis.get("data", "-")} {opis.get("godzina", "--:--")}.',
                f'Inercja po zmianach wynosi {current_energy:,.2f} MVA·s, a zmiana wzgledem stanu oryginalnego to {energy_delta:+,.2f} MVA·s.',
                f'Zapotrzebowanie modelu wynosi {demand_now:,.2f} MW, a systemowy ROCOF N-1 (area=1) wynosi {rocof_now:.4f} Hz/s przy krytycznym ΔP = {largest_loss_now:,.2f} MW.',
                f'W obliczeniach uwzgledniono {uwzglednione_count} generatorow z dodatnim H; liczba recznych zmian statusu to {manual_change_count}.',
            ]
            if top_row is not None:
                report_notes.append(
                    f'Najwiekszy pojedynczy wklad do E_sys ma jednostka {top_row["kod_generatora"]} '
                    f'({top_row["Ei"]:,.2f} MVA·s, {top_row["udzial"]:.2f}% modelu).'
                )
            if top_odm:
                report_notes.append(
                    f'Dominujacy obszar ODM to {self.ODM_CODE_TO_NAME.get(top_odm_code, top_odm_code)}: '
                    f'{float(top_odm.get("energy", 0.0) or 0.0):,.2f} MVA·s '
                    f'({float(top_odm.get("share_pct", 0.0) or 0.0):.2f}% E_sys).'
                )
            if hourly_rows:
                weakest_hour = min(hourly_rows, key=lambda row: row['inercja_po_zmianach'])
                strongest_hour = max(hourly_rows, key=lambda row: row['inercja_po_zmianach'])
                peak_rocof_hour = max(hourly_rows, key=lambda row: row['rocof'])
                report_notes.append(
                    f'W profilu 24h minimum E_sys wystapilo o {weakest_hour["godzina"]} '
                    f'({weakest_hour["inercja_po_zmianach"]:,.2f} MVA·s), a maksimum o '
                    f'{strongest_hour["godzina"]} ({strongest_hour["inercja_po_zmianach"]:,.2f} MVA·s).'
                )
                report_notes.append(
                    f'Najwyzszy ROCOF w zestawie 24h wystapil o {peak_rocof_hour["godzina"]} '
                    f'({peak_rocof_hour["rocof"]:.4f} Hz/s).'
                )
            else:
                report_notes.append('Profil godzinowy 24h nie jest obecnie dostepny, dlatego raport nie zawiera serii dobowej.')
            if threshold is not None:
                below_rows = [row for row in hourly_rows if row['inercja_po_zmianach'] < threshold]
                if current_energy < threshold:
                    report_notes.append(
                        f'Aktualny model jest ponizej progu {threshold:,.2f} MVA·s o {current_energy - threshold:+,.2f} MVA·s.'
                    )
                else:
                    report_notes.append(
                        f'Aktualny model jest powyzej progu {threshold:,.2f} MVA·s o {current_energy - threshold:+,.2f} MVA·s.'
                    )
                if hourly_rows:
                    report_notes.append(f'Liczba godzin z E_sys ponizej progu: {len(below_rows)} z {len(hourly_rows)}.')
            report_notes.extend(
                [
                    'Metodyka inercji: E_sys = Σ(S_i × H_i), gdzie S_i = sqrt(Pmax^2 + Qmax^2).',
                    'Metodyka ROCOF: uproszczony wariant poczatkowy N-1 dla calego systemu, liczony tylko z jednostek area=1.',
                ]
            )

            wb = Workbook()
            ws_report = wb.active
            ws_report.title = 'Raport'
            ws_details = wb.create_sheet('Szczegoly')
            ws_top = wb.create_sheet('TOP jednostki')
            ws_odm = wb.create_sheet('ODM')

            header_fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            manual_fill = PatternFill(fill_type='solid', fgColor='FFF3CD')
            bold_font = Font(bold=True)
            manual_font = Font(bold=True, color='C77700')
            center = Alignment(horizontal='center', vertical='center')
            wrap = Alignment(vertical='top', wrap_text=True)

            def style_header_row(ws, headers, row=1):
                for col_idx, header in enumerate(headers, start=1):
                    cell = ws.cell(row=row, column=col_idx, value=header)
                    cell.font = bold_font
                    cell.fill = header_fill
                    cell.alignment = center

            ws_report['A1'] = 'Raport analizy inercji systemu'
            ws_report['A1'].font = Font(bold=True, size=15)
            ws_report['A2'] = f'Wygenerowano: {export_time}'
            report_pairs = [
                ('Typ modelu', model_label),
                ('Data modelu', opis.get('data', '')),
                ('Godzina modelu', opis.get('godzina', '')),
                ('Plik zrodlowy', opis.get('plik', '')),
                ('Aktywny filtr', self.tekst_filtrowania if self.tekst_filtrowania else '(brak)'),
                ('Inercja oryginalna [MVA·s]', original_energy),
                ('Inercja po zmianach [MVA·s]', current_energy),
                ('Zmiana inercji [MVA·s]', energy_delta),
                ('Zapotrzebowanie [MW]', demand_now),
                ('ROCOF N-1 area=1 [Hz/s]', rocof_now),
                ('Krytyczne ΔP [MW]', largest_loss_now),
                ('Generatory w obliczeniach', uwzglednione_count),
            ]
            for row_idx, (label_txt, value) in enumerate(report_pairs, start=4):
                ws_report.cell(row=row_idx, column=1, value=label_txt).font = bold_font
                ws_report.cell(row=row_idx, column=2, value=value)
            ws_report['A18'] = 'Komentarz operatorski'
            ws_report['A18'].font = bold_font
            note_row = 19
            for idx, note in enumerate(report_notes, start=1):
                ws_report.cell(row=note_row, column=1, value=f'{idx}.')
                ws_report.cell(row=note_row, column=2, value=note).alignment = wrap
                note_row += 1
            ws_report.column_dimensions['A'].width = 28
            ws_report.column_dimensions['B'].width = 66
            ws_report.column_dimensions['C'].width = 4

            ws_details['A1'] = 'Typ modelu'
            ws_details['B1'] = model_label
            ws_details['A2'] = 'Data modelu'
            ws_details['B2'] = opis.get('data', '')
            ws_details['A3'] = 'Godzina modelu'
            ws_details['B3'] = opis.get('godzina', '')
            ws_details['A4'] = 'Filtr'
            ws_details['B4'] = self.tekst_filtrowania if self.tekst_filtrowania else '(brak)'
            ws_details['A5'] = 'Całkowita inercja systemu [MVA·s]'
            ws_details['B5'] = current_energy
            ws_details['A6'] = 'Inercja oryginalna [MVA·s]'
            ws_details['B6'] = original_energy
            ws_details['A7'] = 'Zapotrzebowanie systemu [MW]'
            ws_details['B7'] = demand_now
            ws_details['A8'] = 'ROCOF N-1 area=1 [Hz/s]'
            ws_details['B8'] = rocof_now
            ws_details['A9'] = 'Krytyczne ΔP [MW]'
            ws_details['B9'] = largest_loss_now
            for cell in ['A1', 'A2', 'A3', 'A4', 'A5', 'A6', 'A7', 'A8', 'A9']:
                ws_details[cell].font = bold_font

            detail_headers = [
                'LP',
                'Kod generatora',
                'Grupa',
                'ODM',
                'Status',
                'Pmax [MW]',
                'Qmax [MVAr]',
                'S [MVA]',
                'H [s]',
                'E_i = S × H [MVA·s]',
                'Uwzględniony w obliczeniach',
                'Ręczna zmiana',
            ]
            detail_start_row = 11
            style_header_row(ws_details, detail_headers, row=detail_start_row)
            for row_idx, row in enumerate(dane['wiersze'], start=detail_start_row + 1):
                ws_details.cell(row=row_idx, column=1, value=row['lp'])
                ws_details.cell(row=row_idx, column=2, value=row['kod_generatora'])
                ws_details.cell(row=row_idx, column=3, value=row['grupa'])
                ws_details.cell(row=row_idx, column=4, value=row['odm'])
                ws_details.cell(row=row_idx, column=5, value=row['status'])
                ws_details.cell(row=row_idx, column=6, value=row['pmax_raw'])
                ws_details.cell(row=row_idx, column=7, value=row['qmax_raw'])
                ws_details.cell(row=row_idx, column=8, value=row['s_obliczeniowa'])
                ws_details.cell(row=row_idx, column=9, value=row['H'])
                ws_details.cell(row=row_idx, column=10, value=row['Ei'])
                ws_details.cell(row=row_idx, column=11, value=row['uwzgledniony'])
                ws_details.cell(row=row_idx, column=12, value='TAK' if row['reczna_zmiana'] else '-')
                if row['reczna_zmiana']:
                    for col_idx in range(1, 13):
                        ws_details.cell(row=row_idx, column=col_idx).fill = manual_fill
                        ws_details.cell(row=row_idx, column=col_idx).font = manual_font
            for col, width in {'A': 8, 'B': 28, 'C': 16, 'D': 14, 'E': 12, 'F': 14, 'G': 14, 'H': 14, 'I': 11, 'J': 22, 'K': 28, 'L': 16}.items():
                ws_details.column_dimensions[col].width = width
            ws_details.freeze_panes = 'A11'
            ws_details.auto_filter.ref = f'A{detail_start_row}:L{max(detail_start_row, detail_start_row + len(dane["wiersze"]))}'

            top_headers = ['Kod generatora', 'ODM', 'Grupa', 'Status', 'E_i [MVA·s]', 'Udzial [%]', 'S [MVA]', 'H [s]', 'Reczna zmiana']
            style_header_row(ws_top, top_headers)
            for row_idx, row in enumerate(wiersze_posortowane[:40], start=2):
                ws_top.cell(row=row_idx, column=1, value=row['kod_generatora'])
                ws_top.cell(row=row_idx, column=2, value=row['odm'])
                ws_top.cell(row=row_idx, column=3, value=row['grupa'])
                ws_top.cell(row=row_idx, column=4, value=row['status'])
                ws_top.cell(row=row_idx, column=5, value=row['Ei'])
                ws_top.cell(row=row_idx, column=6, value=row['udzial'])
                ws_top.cell(row=row_idx, column=7, value=row['s_obliczeniowa'])
                ws_top.cell(row=row_idx, column=8, value=row['H'])
                ws_top.cell(row=row_idx, column=9, value='TAK' if row['reczna_zmiana'] else '-')
            for col, width in {'A': 26, 'B': 14, 'C': 16, 'D': 12, 'E': 16, 'F': 12, 'G': 12, 'H': 10, 'I': 16}.items():
                ws_top.column_dimensions[col].width = width
            ws_top.freeze_panes = 'A2'

            headers_odm = ['ODM', 'Inercja [MVA·s]', 'Udzial [% E_sys]', 'Liczba generatorow', 'Krytyczne ΔP [MW]', 'ROCOF N-1 area=1 [Hz/s]']
            style_header_row(ws_odm, headers_odm)
            for row_idx, odm_code in enumerate(self.ODM_ORDER, start=2):
                meta = podsumowanie['odm_metrics'].get(odm_code, {})
                values = [
                    self.ODM_CODE_TO_NAME[odm_code],
                    meta.get('energy', 0.0),
                    meta.get('share_pct', 0.0),
                    int(meta.get('connected_count', 0)),
                    meta.get('largest_loss', 0.0),
                    meta.get('rocof', 0.0),
                ]
                for col_idx, value in enumerate(values, start=1):
                    ws_odm.cell(row=row_idx, column=col_idx, value=value)
            for col, width in {'A': 16, 'B': 18, 'C': 16, 'D': 18, 'E': 20, 'F': 18}.items():
                ws_odm.column_dimensions[col].width = width
            ws_odm.freeze_panes = 'A2'

            ws_hours = None
            if hourly_rows:
                ws_hours = wb.create_sheet('Serie 24h')
                hour_headers = [
                    'Punkt 24h',
                    'Data',
                    'Godzina',
                    'Inercja oryginalna [MVA·s]',
                    'Inercja po zmianach [MVA·s]',
                    'Różnica [MVA·s]',
                    'Zapotrzebowanie [MW]',
                    'ROCOF N-1 area=1 [Hz/s]',
                    'Krytyczne ΔP [MW]',
                    'Liczba generatorów',
                    'Prog [MVA·s]',
                    'Margines do progu [MVA·s]',
                ]
                style_header_row(ws_hours, hour_headers)
                for row_idx, row in enumerate(hourly_rows, start=2):
                    ws_hours.cell(row=row_idx, column=1, value=row['point'])
                    ws_hours.cell(row=row_idx, column=2, value=row['data'])
                    ws_hours.cell(row=row_idx, column=3, value=row['godzina'])
                    ws_hours.cell(row=row_idx, column=4, value=row['inercja_oryginalna'])
                    ws_hours.cell(row=row_idx, column=5, value=row['inercja_po_zmianach'])
                    ws_hours.cell(row=row_idx, column=6, value=row['delta_inercji'])
                    ws_hours.cell(row=row_idx, column=7, value=row['load'])
                    ws_hours.cell(row=row_idx, column=8, value=row['rocof'])
                    ws_hours.cell(row=row_idx, column=9, value=row['largest_loss'])
                    ws_hours.cell(row=row_idx, column=10, value=row['connected_count'])
                    ws_hours.cell(row=row_idx, column=11, value=threshold if threshold is not None else '')
                    ws_hours.cell(row=row_idx, column=12, value=row['inercja_po_zmianach'] - threshold if threshold is not None else '')
                for col, width in {'A': 14, 'B': 14, 'C': 12, 'D': 24, 'E': 24, 'F': 18, 'G': 18, 'H': 18, 'I': 18, 'J': 16, 'K': 16, 'L': 20}.items():
                    ws_hours.column_dimensions[col].width = width
                ws_hours.freeze_panes = 'A2'

            top_rows_count = max(2, min(41, len(wiersze_posortowane) + 1))
            if top_rows_count > 2:
                chart_top = BarChart()
                chart_top.title = 'TOP jednostki wg E_i'
                chart_top.y_axis.title = 'E_i [MVA*s]'
                chart_top.height = 8.5
                chart_top.width = 18.0
                chart_top.add_data(Reference(ws_top, min_col=5, min_row=1, max_row=top_rows_count), titles_from_data=True)
                chart_top.set_categories(Reference(ws_top, min_col=1, min_row=2, max_row=top_rows_count))
                ws_report.add_chart(chart_top, 'D3')

            if ws_hours is not None:
                max_row_24h = len(hourly_rows) + 1
                chart_inertia = LineChart()
                chart_inertia.title = 'Przebieg inercji 24h'
                chart_inertia.y_axis.title = 'Inercja [MVA*s]'
                chart_inertia.x_axis.title = 'Godzina modelu'
                chart_inertia.height = 8.5
                chart_inertia.width = 18.5
                chart_inertia.add_data(Reference(ws_hours, min_col=4, max_col=5, min_row=1, max_row=max_row_24h), titles_from_data=True)
                chart_inertia.set_categories(Reference(ws_hours, min_col=1, min_row=2, max_row=max_row_24h))
                ws_report.add_chart(chart_inertia, 'D20')

                chart_rocof = LineChart()
                chart_rocof.title = 'ROCOF w profilach godzinowych'
                chart_rocof.y_axis.title = 'ROCOF [Hz/s]'
                chart_rocof.x_axis.title = 'Godzina modelu'
                chart_rocof.height = 7.4
                chart_rocof.width = 18.5
                chart_rocof.add_data(Reference(ws_hours, min_col=8, max_col=8, min_row=1, max_row=max_row_24h), titles_from_data=True)
                chart_rocof.set_categories(Reference(ws_hours, min_col=1, min_row=2, max_row=max_row_24h))
                ws_report.add_chart(chart_rocof, 'D37')

                chart_load = BarChart()
                chart_load.title = 'Zapotrzebowanie godzinowe'
                chart_load.y_axis.title = 'Load [MW]'
                chart_load.height = 7.4
                chart_load.width = 18.5
                chart_load.add_data(Reference(ws_hours, min_col=7, max_col=7, min_row=1, max_row=max_row_24h), titles_from_data=True)
                chart_load.set_categories(Reference(ws_hours, min_col=1, min_row=2, max_row=max_row_24h))
                ws_report.add_chart(chart_load, 'D52')

            chart_odm = BarChart()
            chart_odm.title = 'Inercja wg ODM'
            chart_odm.y_axis.title = 'Inercja [MVA*s]'
            chart_odm.height = 7.6
            chart_odm.width = 18.5
            chart_odm.add_data(Reference(ws_odm, min_col=2, max_col=2, min_row=1, max_row=len(self.ODM_ORDER) + 1), titles_from_data=True)
            chart_odm.set_categories(Reference(ws_odm, min_col=1, min_row=2, max_row=len(self.ODM_ORDER) + 1))
            ws_report.add_chart(chart_odm, 'D69')

            arch_rows = self.archive_repo.fetch_records({}, sort_by='archived_at', sort_desc=True)[:2000]
            if arch_rows:
                ws_arch = wb.create_sheet('Archiwum inercji')
                arch_headers = ['ID', 'Archiwizowano', 'Typ', 'Data', 'Godzina', 'Plik', 'Inercja [MVA·s]', 'Inercja org. [MVA·s]', 'Load [MW]', 'L. gen']
                style_header_row(ws_arch, arch_headers)
                for row_idx, row in enumerate(arch_rows, start=2):
                    values = [
                        row.get('id'),
                        row.get('archived_at', ''),
                        row.get('model_type', ''),
                        row.get('model_date', ''),
                        row.get('model_hour', ''),
                        row.get('source_file', ''),
                        row.get('total_inertia', 0.0),
                        row.get('original_inertia', 0.0),
                        row.get('demand', 0.0),
                        row.get('connected_count', 0),
                    ]
                    for col_idx, value in enumerate(values, start=1):
                        ws_arch.cell(row=row_idx, column=col_idx, value=value)
                for col, width in {'A': 8, 'B': 22, 'C': 10, 'D': 14, 'E': 12, 'F': 32, 'G': 18, 'H': 18, 'I': 14, 'J': 10}.items():
                    ws_arch.column_dimensions[col].width = width
                ws_arch.freeze_panes = 'A2'

            if self.alarm_history:
                ws_alarm = wb.create_sheet('Historia alarmow')
                alarm_headers = ['Kiedy alarm wystąpił', 'Tryb', 'Typ modelu', 'Data modelu', 'Godzina modelu', 'Wartość [MVA·s]', 'Próg [MVA·s]', 'Odchyłka [MVA·s]', 'Plik', 'Mail', 'Status wysyłki', 'Opis']
                style_header_row(ws_alarm, alarm_headers)
                for row_idx, row in enumerate(self.alarm_history, start=2):
                    wartosc = float(row.get('wartosc', 0.0) or 0.0)
                    prog = float(row.get('prog', 0.0) or 0.0)
                    ws_alarm.cell(row=row_idx, column=1, value=row.get('czas_alarmu', ''))
                    ws_alarm.cell(row=row_idx, column=2, value=row.get('trigger_mode', ''))
                    ws_alarm.cell(row=row_idx, column=3, value=row.get('model_type_label') or self.get_model_type_label(row.get('model_type')))
                    ws_alarm.cell(row=row_idx, column=4, value=row.get('data_modelu', ''))
                    ws_alarm.cell(row=row_idx, column=5, value=row.get('godzina_modelu', ''))
                    ws_alarm.cell(row=row_idx, column=6, value=wartosc)
                    ws_alarm.cell(row=row_idx, column=7, value=prog)
                    ws_alarm.cell(row=row_idx, column=8, value=wartosc - prog)
                    ws_alarm.cell(row=row_idx, column=9, value=row.get('source_file', ''))
                    ws_alarm.cell(row=row_idx, column=10, value='TAK' if row.get('mail_attempted') else 'NIE')
                    ws_alarm.cell(row=row_idx, column=11, value=row.get('mail_status', ''))
                    ws_alarm.cell(row=row_idx, column=12, value=row.get('opis', ''))
                for col, width in {'A': 20, 'B': 10, 'C': 12, 'D': 14, 'E': 14, 'F': 16, 'G': 14, 'H': 16, 'I': 28, 'J': 10, 'K': 24, 'L': 24}.items():
                    ws_alarm.column_dimensions[col].width = width
                ws_alarm.freeze_panes = 'A2'

            wb.save(sciezka)
            messagebox.showinfo('Eksport zakończony', f'Zapisano plik:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu', f'Nie udało się zapisać pliku Excel:\n{e}')

    def pobierz_liste_zmian(self):
        zmiany = []
        for godzina, gen_list in self.wszystkie_generatory.items():
            opis = self.opis_modeli.get(godzina, {})
            oryginalne = self.statusy_oryginalne.get(godzina, {})
            for gen in gen_list:
                org = oryginalne.get(gen['reg_name'], gen['status'])
                now = gen['status']
                if org != now:
                    H = self.get_h_value_for_generator(gen)
                    s_apparent = self.get_generator_apparent_power(gen)
                    Ei = s_apparent * H
                    zmiany.append({'typ': '24h', 'data': opis.get('data', ''), 'godzina': opis.get('godzina', f'{godzina:02d}:30'), 'generator': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status_oryginalny': org, 'status_po_zmianie': now, 's_apparent': s_apparent, 'H': H, 'Ei': Ei})
        for model_id in self.reczne_order:
            m = self.modele_reczne.get(model_id, {})
            opis = m.get('opis', {})
            oryginalne = m.get('statusy_oryginalne', {})
            for gen in m.get('generatory', []):
                org = oryginalne.get(gen['reg_name'], gen['status'])
                now = gen['status']
                if org != now:
                    H = self.get_h_value_for_generator(gen)
                    s_apparent = self.get_generator_apparent_power(gen)
                    Ei = s_apparent * H
                    zmiany.append({'typ': 'RECZNY', 'data': opis.get('data', ''), 'godzina': opis.get('godzina', '--:--'), 'generator': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status_oryginalny': org, 'status_po_zmianie': now, 's_apparent': s_apparent, 'H': H, 'Ei': Ei})
        for auto_id in self.auto_order:
            m = self.modele_auto[auto_id]
            opis = m['opis']
            oryginalne = m['statusy_oryginalne']
            for gen in m['generatory']:
                org = oryginalne.get(gen['reg_name'], gen['status'])
                now = gen['status']
                if org != now:
                    H = self.get_h_value_for_generator(gen)
                    s_apparent = self.get_generator_apparent_power(gen)
                    Ei = s_apparent * H
                    zmiany.append({'typ': 'BIEŻĄCE', 'data': opis.get('data', ''), 'godzina': opis.get('godzina', '--:--'), 'generator': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status_oryginalny': org, 'status_po_zmianie': now, 's_apparent': s_apparent, 'H': H, 'Ei': Ei})
        return zmiany

    def eksportuj_raport_zmian(self):
        zmiany = self.pobierz_liste_zmian()
        if not zmiany:
            messagebox.showinfo('Raport zmian', 'Brak ręcznych zmian do eksportu.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz raport zmian', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx'), ('Tekst', '*.txt')])
        if not sciezka:
            return
        try:
            if sciezka.lower().endswith('.txt'):
                with open(sciezka, 'w', encoding='utf-8') as f:
                    f.write('RAPORT RĘCZNYCH ZMIAN STATUSÓW GENERATORÓW\n')
                    f.write('=' * 140 + '\n')
                    for z in zmiany:
                        f.write(f"{z['typ']} | {z['data']} {z['godzina']} | {z['generator']} | {z['grupa']} | {z['status_oryginalny']} -> {z['status_po_zmianie']} | S={z['s_apparent']:.3f} MVA | H={z['H']:.3f} s | Ei={z['Ei']:.3f} MVA·s\n")
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = 'Raport zmian'
                headers = ['Typ modelu', 'Data', 'Godzina', 'Generator', 'Grupa', 'Status oryginalny', 'Status po zmianie', 'S [MVA]', 'H [s]', 'Ei [MVA·s]']
                bold = Font(bold=True)
                fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
                for col, h in enumerate(headers, 1):
                    c = ws.cell(row=1, column=col, value=h)
                    c.font = bold
                    c.fill = fill
                for r, z in enumerate(zmiany, start=2):
                    ws.cell(row=r, column=1, value=z['typ'])
                    ws.cell(row=r, column=2, value=z['data'])
                    ws.cell(row=r, column=3, value=z['godzina'])
                    ws.cell(row=r, column=4, value=z['generator'])
                    ws.cell(row=r, column=5, value=z['grupa'])
                    ws.cell(row=r, column=6, value=z['status_oryginalny'])
                    ws.cell(row=r, column=7, value=z['status_po_zmianie'])
                    ws.cell(row=r, column=8, value=z['s_apparent'])
                    ws.cell(row=r, column=9, value=z['H'])
                    ws.cell(row=r, column=10, value=z['Ei'])
                for col, width in {'A': 12, 'B': 14, 'C': 12, 'D': 28, 'E': 16, 'F': 18, 'G': 18, 'H': 12, 'I': 10, 'J': 14}.items():
                    ws.column_dimensions[col].width = width
                wb.save(sciezka)
            messagebox.showinfo('Eksport raportu', f'Zapisano raport zmian:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu raportu', str(e))

    def eksportuj_aktualny_model_epc_po_zmianach(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        if not self.is_model_epc_exportable(model_type, key):
            messagebox.showwarning('Brak pliku źródłowego', 'Nie znaleziono pliku EPC do eksportu dla wybranego modelu.')
            return
        opis = self.get_model_meta(model_type, key)
        sciezka_zrodlowa = opis.get('sciezka', '')
        nazwa_domyslna = os.path.splitext(os.path.basename(sciezka_zrodlowa))[0] + '_po_zmianach.epc'
        sciezka_docelowa = filedialog.asksaveasfilename(title='Zapisz zmodyfikowany model EPC', defaultextension='.epc', initialfile=nazwa_domyslna, filetypes=[('Plik EPC', '*.epc'), ('Wszystkie', '*.*')])
        if not sciezka_docelowa:
            return
        gen_list = self.get_model_generators(model_type, key)
        status_map = {g['reg_name']: int(g['status']) for g in gen_list}
        try:
            with open(sciezka_zrodlowa, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()
            nowe_linie = []
            in_section = False
            liczba_podmienionych = 0
            for line in lines:
                line_stripped = line.strip()
                line_lower = line_stripped.lower()
                if 'generator data' in line_lower:
                    in_section = True
                    nowe_linie.append(line)
                    continue
                if in_section and 'load data' in line_lower:
                    in_section = False
                    nowe_linie.append(line)
                    continue
                if in_section and ':' in line_stripped:
                    nowa_linia, czy_zmieniono = self.zmien_status_w_linii_generatora(line, status_map)
                    nowe_linie.append(nowa_linia)
                    if czy_zmieniono:
                        liczba_podmienionych += 1
                else:
                    nowe_linie.append(line)
            with open(sciezka_docelowa, 'w', encoding='utf-8', errors='ignore') as f:
                f.writelines(nowe_linie)
            messagebox.showinfo('Eksport EPC zakończony', f'Zapisano zmodyfikowany model EPC:\n{sciezka_docelowa}\n\nZaktualizowane wpisy generatorów w EPC: {liczba_podmienionych}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu EPC', f'Nie udało się zapisać pliku EPC:\n{e}')

    def zmien_status_w_linii_generatora(self, line, status_map):
        try:
            if ':' not in line:
                return (line, False)
            line_no_nl = line.rstrip('\r\n')
            line_ending = line[len(line_no_nl):]
            lewa, prawa = line_no_nl.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            reg_name = quoted[-1].strip() if quoted else None
            if not reg_name or reg_name not in status_map:
                return (line, False)
            m = re.match('^(\\s*)([-+]?\\d+(?:[.,]\\d+)?)(.*)$', prawa)
            if not m:
                return (line, False)
            leading_spaces = m.group(1)
            old_status_token = m.group(2)
            tail = m.group(3)
            nowy_status = int(status_map[reg_name])
            if '.' in old_status_token or ',' in old_status_token:
                separator = ',' if ',' in old_status_token else '.'
                status_txt = f'{float(nowy_status):.1f}'
                if separator == ',':
                    status_txt = status_txt.replace('.', ',')
            else:
                status_txt = str(nowy_status)
            nowa_linia = f'{lewa}:{leading_spaces}{status_txt}{tail}{line_ending}'
            czy_zmieniono = old_status_token != status_txt
            return (nowa_linia, czy_zmieniono)
        except Exception:
            return (line, False)

    def zapisz_sesje(self):
        if not self.wszystkie_generatory and (not self.modele_reczne) and (not self.modele_auto):
            messagebox.showwarning('Brak danych', 'Brak danych do zapisania.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz sesję', defaultextension='.json', filetypes=[('JSON', '*.json')])
        if not sciezka:
            return
        data = {'wszystkie_generatory': self.wszystkie_generatory, 'statusy_oryginalne': self.statusy_oryginalne, 'opis_modeli': self.opis_modeli, 'H_generatory': self.H_generatory, 'H_generatory_original': self.H_generatory_original, 'H_group_defaults': self.H_group_defaults, 'entry_filtr': self.entry_filtr.get(), 'tryb_filtrowania': self.tryb_filtrowania_var.get(), 'sortowanie': self.sortowanie_var.get(), 'wartosc_graniczna': self.entry_wartosc_graniczna.get(), 'obliczenia_wykonane': self.obliczenia_wykonane, 'monitor_folder_path': self.monitor_folder_path, 'modele_reczne': self.modele_reczne, 'reczne_order': self.reczne_order, 'reczne_counter': self.reczne_counter, 'modele_auto': self.modele_auto, 'auto_order': self.auto_order, 'auto_seen_files': self.auto_seen_files, 'auto_counter': self.auto_counter, 'email_alarm_last_key': self.email_alarm.last_alarm_key, 'sent_alarm_tokens': [list(x) for x in self.sent_alarm_tokens], 'alarm_history': self.alarm_history, 'zapotrzebowanie_na_godzine': self.zapotrzebowanie_na_godzine, 'odm_metrics_na_godzine': self.odm_metrics_na_godzine, 'rocof_na_godzine': self.rocof_na_godzine, 'largest_loss_na_godzine': self.largest_loss_na_godzine, 'secondary_chart_scope': self.secondary_chart_scope_var.get()}
        try:
            with open(sciezka, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            messagebox.showinfo('Sesja', f'Zapisano sesję:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd zapisu sesji', str(e))

    def wczytaj_sesje(self):
        sciezka = filedialog.askopenfilename(title='Wczytaj sesję', filetypes=[('JSON', '*.json')])
        if not sciezka:
            return
        try:
            with open(sciezka, 'r', encoding='utf-8') as f:
                data = json.load(f)
            self.wszystkie_generatory = {int(k): v for k, v in data.get('wszystkie_generatory', {}).items()}
            self.statusy_oryginalne = {int(k): v for k, v in data.get('statusy_oryginalne', {}).items()}
            self.opis_modeli = {int(k): v for k, v in data.get('opis_modeli', {}).items()}
            raw_zap = {int(k): v for k, v in data.get('zapotrzebowanie_na_godzine', {}).items()}
            self.zapotrzebowanie_na_godzine = {h: raw_zap.get(h, 0.0) for h in range(24)}
            self.H_generatory = data.get('H_generatory', {})
            self.H_generatory_original = data.get('H_generatory_original', dict(self.H_generatory))
            self.H_group_defaults = data.get('H_group_defaults', dict(self.H_group_defaults))
            self.obliczenia_wykonane = data.get('obliczenia_wykonane', False)
            self.monitor_folder_path = data.get('monitor_folder_path', '')
            self.modele_reczne = data.get('modele_reczne', {})
            for model_id in self.modele_reczne:
                self.modele_reczne[model_id].setdefault('zapotrzebowanie', 0.0)
                self.modele_reczne[model_id].setdefault('odm_metrics', self.build_empty_odm_metrics())
                self.modele_reczne[model_id].setdefault('rocof', 0.0)
                self.modele_reczne[model_id].setdefault('largest_loss', 0.0)
                self.modele_reczne[model_id].setdefault('liczba_zalaczonych', 0)
            self.reczne_order = data.get('reczne_order', list(self.modele_reczne.keys()))
            self.reczne_counter = max(int(data.get('reczne_counter', 0) or 0), len(self.modele_reczne))
            self.modele_auto = data.get('modele_auto', {})
            for auto_id in self.modele_auto:
                self.modele_auto[auto_id].setdefault('zapotrzebowanie', 0.0)
                self.modele_auto[auto_id].setdefault('odm_metrics', self.build_empty_odm_metrics())
                self.modele_auto[auto_id].setdefault('rocof', 0.0)
                self.modele_auto[auto_id].setdefault('largest_loss', 0.0)
            self.auto_order = data.get('auto_order', [])
            self.auto_seen_files = data.get('auto_seen_files', {})
            self.auto_counter = data.get('auto_counter', 0)
            self.email_alarm.last_alarm_key = data.get('email_alarm_last_key')
            self.sent_alarm_tokens = {tuple(x) for x in data.get('sent_alarm_tokens', [])}
            self.alarm_history = data.get('alarm_history', [])
            self.inercja_na_godzine = {h: 0.0 for h in range(24)}
            self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
            self.odm_metrics_na_godzine = {h: self.build_empty_odm_metrics() for h in range(24)}
            self.rocof_na_godzine = {h: 0.0 for h in range(24)}
            self.largest_loss_na_godzine = {h: 0.0 for h in range(24)}
            self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
            self.label_monitor_folder.config(text=f'Folder: {self.monitor_folder_path}' if self.monitor_folder_path else 'Folder: nie wybrano')
            self.secondary_chart_scope_var.set(data.get('secondary_chart_scope', self.secondary_chart_scope_var.get()))
            self.entry_filtr.delete(0, tk.END)
            self.entry_filtr.insert(0, data.get('entry_filtr', ''))
            self.tryb_filtrowania_var.set(data.get('tryb_filtrowania', 'Wszystkie'))
            self.sortowanie_var.set(data.get('sortowanie', 'Kod generatora A-Z'))
            self.entry_wartosc_graniczna.delete(0, tk.END)
            self.entry_wartosc_graniczna.insert(0, data.get('wartosc_graniczna', ''))
            if not self.H_generatory and (self.wszystkie_generatory or self.modele_reczne or self.modele_auto):
                messagebox.showwarning('Sesja', 'Sesja nie zawiera współczynników H albo są puste. Wyniki inercji mogą być zerowe.')
            for godzina in self.wszystkie_generatory.keys():
                self.przelicz_inercje_dla_modelu('manual', godzina)
            for model_id in self.reczne_order:
                self.przelicz_inercje_dla_modelu('manual_batch', model_id)
            for auto_id in self.auto_order:
                self.przelicz_inercje_dla_modelu('auto', auto_id)
            self.aktywny_typ_modelu = None
            self.aktualna_godzina = None
            self.aktualny_reczny_id = None
            self.aktualny_auto_id = None
            self.odswiez_liste_modeli()
            self.odswiez_liste_modeli_reczne()
            self.odswiez_liste_modeli_auto()
            self.odswiez_panel_h_dyd()
            self.aktualizuj_wykres()
            self.aktualizuj_panel_maksimum(check_alarm=False)
            self.aktualizuj_wskazniki_jakosci()
            self.tree_generatory.delete(*self.tree_generatory.get_children())
            self.pole_szczegoly.delete('1.0', tk.END)
            self.label_porownanie.config(text='Brak wybranego modelu')
            self.label_ranking.config(text='Brak wybranego modelu')
            self.label_analiza_inercyjna.config(text='Brak wybranego modelu')
            self.label_reczne.config(text='Brak wybranego modelu')
            self.przygotuj_pusty_wykres_grup()
            self.refresh_alarm_history_window()
            self.refresh_history_window()
            messagebox.showinfo('Sesja', f'Wczytano sesję:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd wczytywania sesji', str(e))

    def pokaz_walidacje_epc_dyd(self):
        if not self.wszystkie_generatory:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj EPC.')
            return
        if not self.H_generatory:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj DYD.')
            return
        epc_names = set()
        for gen_list in self.wszystkie_generatory.values():
            for g in gen_list:
                epc_names.add(g['reg_name'])
        dyd_names = set(self.H_generatory.keys())
        brak_H = sorted(epc_names - dyd_names)
        nadmiar_w_dyd = sorted(dyd_names - epc_names)
        okno = tk.Toplevel(self.root)
        okno.title('Walidacja EPC ↔ DYD')
        init_managed_toplevel(okno, 'validation_epc_dyd', 1200, 800)
        okno.configure(bg=self.get_palette()['bg'])
        txt = scrolledtext.ScrolledText(okno, wrap=tk.WORD, font=self.font_code)
        txt.pack(fill=tk.BOTH, expand=True)
        pal = self.get_palette()
        txt.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'])
        txt.insert(tk.END, 'WALIDACJA ZGODNOŚCI EPC ↔ DYD\n')
        txt.insert(tk.END, '=' * 80 + '\n\n')
        txt.insert(tk.END, f'Generatory w EPC bez współczynnika H w DYD: {len(brak_H)}\n')
        txt.insert(tk.END, '-' * 80 + '\n')
        for x in brak_H:
            txt.insert(tk.END, x + '\n')
        txt.insert(tk.END, '\n')
        txt.insert(tk.END, f'Generatory w DYD nieobecne w EPC: {len(nadmiar_w_dyd)}\n')
        txt.insert(tk.END, '-' * 80 + '\n')
        for x in nadmiar_w_dyd:
            txt.insert(tk.END, x + '\n')
        txt.see('1.0')

    def export_chart(self, figure, suggested_name):
        path = filedialog.asksaveasfilename(title='Zapisz wykres', defaultextension=os.path.splitext(suggested_name)[1], initialfile=suggested_name, filetypes=[('PNG', '*.png'), ('PDF', '*.pdf'), ('Wszystkie', '*.*')])
        if not path:
            return
        try:
            figure.savefig(path, dpi=180, bbox_inches='tight')
            messagebox.showinfo('Eksport wykresu', f'Zapisano:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu wykresu', str(e))

    def export_dashboard_png(self):
        if not PIL_AVAILABLE:
            messagebox.showwarning('Brak Pillow', 'Do eksportu dashboardu PNG potrzebny jest pakiet Pillow.\nZainstaluj: pip install pillow')
            return
        path = filedialog.asksaveasfilename(title='Zapisz dashboard', defaultextension='.png', initialfile='dashboard_inercji.png', filetypes=[('PNG', '*.png')])
        if not path:
            return
        try:
            self.root.update()
            x = self.root.winfo_rootx()
            y = self.root.winfo_rooty()
            w = self.root.winfo_width()
            h = self.root.winfo_height()
            img = ImageGrab.grab(bbox=(x, y, x + w, y + h))
            img.save(path)
            messagebox.showinfo('Eksport dashboardu', f'Zapisano:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu dashboardu', str(e))

    def eksportuj_pelny_raport_pdf(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        path = filedialog.asksaveasfilename(title='Zapisz pełny raport PDF', defaultextension='.pdf', initialfile='raport_inercji.pdf', filetypes=[('PDF', '*.pdf')])
        if not path:
            return
        try:
            with PdfPages(path) as pdf:
                fig1 = plt.figure(figsize=(11.69, 8.27))
                ax = fig1.add_subplot(111)
                ax.axis('off')
                opis = self.get_model_meta(model_type, key)
                E_org = self.get_model_energy_original(model_type, key)
                E_now = self.get_model_energy_now(model_type, key)
                P_load = self.get_model_demand_now(model_type, key)
                diff = E_now - E_org
                diff_proc = diff / E_org * 100 if E_org > 0 else 0.0
                text = []
                text.append('RAPORT ANALIZY INERCJI')
                text.append('=' * 60)
                text.append(f"Typ modelu: {('24h' if model_type == 'manual' else 'Bieżące')}")
                text.append(f"Data modelu: {opis.get('data', '')}")
                text.append(f"Godzina modelu: {opis.get('godzina', '')}")
                text.append(f"Plik: {opis.get('plik', '')}")
                text.append('')
                text.append('PODSUMOWANIE PRZED / PO')
                text.append(f'Inercja oryginalna: {E_org:,.2f} MVA·s')
                text.append(f'Inercja po zmianach: {E_now:,.2f} MVA·s')
                text.append(f'Różnica: {diff:,.2f} MVA·s')
                text.append(f'Zmiana procentowa: {diff_proc:,.2f} %')
                text.append(f'Zapotrzebowanie systemu: {P_load:,.2f} MW')
                text.append('')
                prog = self.pobierz_wartosc_graniczna()
                if prog is not None:
                    text.append(f'Próg alarmowy: {prog:,.2f} MVA·s')
                    tryb = 'powyżej progu' if self.email_alarm.config.trigger_mode == 'above' else 'poniżej progu'
                    text.append(f'Tryb alarmu: {tryb}')
                text.append('')
                dane = self.pobierz_dane_szczegolowe(model_type, key)
                if dane:
                    p = dane['podsumowanie']
                    text.append('PODSUMOWANIE FILTRA')
                    text.append(f"Liczba generatorów po filtrze: {p['liczba_gen_odfiltrowanych']}")
                    text.append(f"Liczba generatorów załączonych: {p['liczba_wlaczonych']}")
                    text.append(f"Energia uwzględnionych: {p['energia_wlaczonych']:,.2f} MVA·s")
                    text.append(f"Energia nieuwzględnionych: {p['energia_wylaczonych']:,.2f} MVA·s")
                    text.append(f"Zapotrzebowanie systemu (LOAD DATA): {p['zapotrzebowanie']:,.2f} MW")
                ax.text(0.02, 0.98, '\n'.join(text), va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig1, bbox_inches='tight')
                plt.close(fig1)
                pdf.savefig(self.figura, bbox_inches='tight')
                pdf.savefig(self.figura_grupy, bbox_inches='tight')
                fig_rank = plt.figure(figsize=(11.69, 8.27))
                axr = fig_rank.add_subplot(111)
                axr.axis('off')
                rank_txt = self.label_ranking.cget('text')
                axr.text(0.02, 0.98, 'RANKING GENERATORÓW\n\n' + rank_txt, va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig_rank, bbox_inches='tight')
                plt.close(fig_rank)
                fig_ana = plt.figure(figsize=(11.69, 8.27))
                axa = fig_ana.add_subplot(111)
                axa.axis('off')
                ana_txt = self.label_analiza_inercyjna.cget('text')
                wsk_txt = self.label_wskazniki.cget('text')
                axa.text(0.02, 0.98, 'ANALIZA INERCYJNA\n\n' + ana_txt + '\n\nWSKAŹNIKI JAKOŚCI 24H\n\n' + wsk_txt, va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig_ana, bbox_inches='tight')
                plt.close(fig_ana)
                fig_alarm = plt.figure(figsize=(11.69, 8.27))
                axal = fig_alarm.add_subplot(111)
                axal.axis('off')
                lines = ['ALARMY\n']
                if self.alarm_history:
                    for z in self.alarm_history[-40:]:
                        lines.append(f"{z.get('czas_alarmu', '')} | {z.get('trigger_mode', '')} | {z.get('data_modelu', '')} {z.get('godzina_modelu', '')} | wartość={z.get('wartosc', 0.0):,.2f} | próg={z.get('prog', 0.0):,.2f} | {z.get('mail_status', '')}")
                else:
                    lines.append('Brak alarmów.')
                axal.text(0.02, 0.98, '\n'.join(lines), va='top', ha='left', family='monospace', fontsize=9)
                pdf.savefig(fig_alarm, bbox_inches='tight')
                plt.close(fig_alarm)
                fig_arch = plt.figure(figsize=(11.69, 8.27))
                axarch = fig_arch.add_subplot(111)
                axarch.axis('off')
                arch_lines = ['ARCHIWUM INERCJI\n']
                arch_rows = self.archive_repo.fetch_records({}, sort_by='archived_at', sort_desc=True)[:40]
                if arch_rows:
                    for z in arch_rows:
                        arch_lines.append(f"{z.get('id')} | {z.get('archived_at', '')} | {z.get('model_date', '')} {z.get('model_hour', '')} | E={z.get('total_inertia', 0.0):,.2f} | Load={z.get('demand', 0.0):,.2f} | {z.get('source_file', '')}")
                else:
                    arch_lines.append('Brak rekordów archiwum.')
                axarch.text(0.02, 0.98, '\n'.join(arch_lines), va='top', ha='left', family='monospace', fontsize=8.5)
                pdf.savefig(fig_arch, bbox_inches='tight')
                plt.close(fig_arch)
                fig_changes = plt.figure(figsize=(11.69, 8.27))
                axc = fig_changes.add_subplot(111)
                axc.axis('off')
                zmiany = self.pobierz_liste_zmian()
                lines = ['RĘCZNE ZMIANY\n']
                if zmiany:
                    for z in zmiany[:80]:
                        lines.append(f"{z['typ']} | {z['data']} {z['godzina']} | {z['generator']} | {z['grupa']} | {z['status_oryginalny']} -> {z['status_po_zmianie']} | S={z['s_apparent']:.3f} MVA | H={z['H']:.3f} s | Ei={z['Ei']:.3f} MVA·s")
                else:
                    lines.append('Brak ręcznych zmian.')
                axc.text(0.02, 0.98, '\n'.join(lines), va='top', ha='left', family='monospace', fontsize=8)
                pdf.savefig(fig_changes, bbox_inches='tight')
                plt.close(fig_changes)
            messagebox.showinfo('Eksport PDF', f'Zapisano raport PDF:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu PDF', str(e))
if __name__ == '__main__':
    root = tk.Tk()
    root.geometry('1500x950')
    try:
        root.state('zoomed')
    except Exception:
        pass
    app = AnalizaInercjiExcelApp(root)

    def _on_close():
        try:
            app.shutdown()
        except Exception:
            pass
        root.destroy()
    root.protocol('WM_DELETE_WINDOW', _on_close)
    root.mainloop()
