from __future__ import annotations
import atexit
import json
import os
import re
import smtplib
import ssl
import statistics
import threading
import tkinter as tk
from dataclasses import asdict, dataclass, field
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
try:
    from PIL import ImageGrab
    PIL_AVAILABLE = True
except Exception:
    PIL_AVAILABLE = False
EMAIL_RE = re.compile('^[^@\\s]+@[^@\\s]+\\.[^@\\s]+$')

@dataclass
class EmailAlarmConfig:
    enabled: bool = False
    smtp_host: str = 'smtp.gmail.com'
    smtp_port: int = 465
    smtp_username: str = ''
    smtp_password: str = ''
    sender_name: str = 'Analiza inercji'
    sender_email: str = ''
    recipients: list[str] = field(default_factory=list)
    trigger_mode: str = 'above'

    def validate(self):
        if not self.enabled:
            return (True, 'Powiadomienia e-mail są wyłączone')
        if not self.smtp_host.strip():
            return (False, 'Brak hosta SMTP')
        if not self.smtp_port:
            return (False, 'Brak portu SMTP')
        if not self.smtp_username.strip():
            return (False, 'Brak loginu SMTP')
        if not self.smtp_password.strip():
            return (False, 'Brak hasła SMTP')
        if not self.sender_email.strip() or not EMAIL_RE.match(self.sender_email.strip()):
            return (False, 'Niepoprawny adres nadawcy')
        if not self.recipients:
            return (False, 'Lista odbiorców jest pusta')
        bad = [x for x in self.recipients if not EMAIL_RE.match(x)]
        if bad:
            return (False, f"Niepoprawne adresy odbiorców: {', '.join(bad)}")
        if self.trigger_mode not in {'above', 'below'}:
            return (False, 'Niepoprawny tryb alarmu')
        return (True, 'OK')

class EmailAlarmService:

    def __init__(self, config_path='~/.analiza_inercji_email.json'):
        self.config_path = Path(config_path).expanduser()
        self.config = EmailAlarmConfig()
        self.last_alarm_key = None
        self.load_config()

    @staticmethod
    def parse_recipients(raw):
        unique = {}
        for item in (x.strip() for x in re.split('[;,\\n]', raw)):
            if item:
                unique.setdefault(item.lower(), item)
        return list(unique.values())

    def load_config(self):
        if not self.config_path.exists():
            return self.config
        try:
            with self.config_path.open('r', encoding='utf-8') as f:
                raw = json.load(f)
            self.config = EmailAlarmConfig(**raw)
        except Exception:
            self.config = EmailAlarmConfig()
        return self.config

    def save_config(self):
        self.config_path.parent.mkdir(parents=True, exist_ok=True)
        with self.config_path.open('w', encoding='utf-8') as f:
            json.dump(asdict(self.config), f, ensure_ascii=False, indent=2)

    def should_trigger(self, prog, wartosc):
        return wartosc > prog if self.config.trigger_mode == 'above' else wartosc < prog

    def send_alarm(self, *, prog, wartosc, data_modelu, godzina_modelu, source_file=''):
        ok, msg = self.config.validate()
        if not ok:
            return (False, msg)
        alarm_key = f'{data_modelu}|{godzina_modelu}|{prog}|{wartosc:.3f}|{self.config.trigger_mode}'
        if alarm_key == self.last_alarm_key:
            return (True, 'Alarm dla tego samego modelu już został wysłany')
        delta = wartosc - prog
        direction = 'przekroczenie' if self.config.trigger_mode == 'above' else 'spadek poniżej'
        when_txt = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        email = EmailMessage()
        email['Subject'] = f'[ALARM INERCJI] {direction} progu | {data_modelu} {godzina_modelu}'
        email['From'] = f'{self.config.sender_name} <{self.config.sender_email}>'
        email['To'] = ', '.join(self.config.recipients)
        email.set_content(f"Data modelu: {data_modelu}\nGodzina modelu: {godzina_modelu}\nWartość inercji: {wartosc:,.2f} MW·s\nPróg: {prog:,.2f} MW·s\nRóżnica: {delta:,.2f} MW·s\nPlik: {source_file or '-'}\nWysłano: {when_txt}\n")
        ctx = ssl.create_default_context()
        with smtplib.SMTP_SSL(self.config.smtp_host, int(self.config.smtp_port), context=ctx, timeout=20) as smtp:
            smtp.login(self.config.smtp_username, self.config.smtp_password)
            smtp.send_message(email)
        self.last_alarm_key = alarm_key
        return (True, 'Mail alarmowy wysłany')

def center_toplevel(window: tk.Toplevel, width: int, height: int):
    try:
        window.update_idletasks()
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()
        x = max(40, int((screen_w - width) / 2))
        y = max(40, int((screen_h - height) / 2))
        window.geometry(f'{width}x{height}+{x}+{y}')
    except Exception:
        window.geometry(f'{width}x{height}')

class EmailAlarmConfigWindow(tk.Toplevel):

    def __init__(self, master, service):
        super().__init__(master)
        self.service = service
        self.title('Konfiguracja powiadomień e-mail')
        center_toplevel(self, 820, 470)
        self.resizable(True, True)
        cfg = service.config
        self.enabled_var = tk.BooleanVar(value=cfg.enabled)
        self.smtp_host_var = tk.StringVar(value=cfg.smtp_host)
        self.smtp_port_var = tk.StringVar(value=str(cfg.smtp_port))
        self.smtp_username_var = tk.StringVar(value=cfg.smtp_username)
        self.smtp_password_var = tk.StringVar(value=cfg.smtp_password)
        self.sender_name_var = tk.StringVar(value=cfg.sender_name)
        self.sender_email_var = tk.StringVar(value=cfg.sender_email)
        self.recipients_var = tk.StringVar(value=', '.join(cfg.recipients))
        self.trigger_mode_var = tk.StringVar(value=cfg.trigger_mode)
        self._build_ui()

    def _style_entry(self, entry: tk.Entry):
        entry.configure(relief='flat', bd=0, highlightthickness=1, highlightbackground='#d0d7e2', highlightcolor='#4f7cff', bg='white', fg='#111827', insertbackground='#111827', font=('Segoe UI', 10))

    def _style_button(self, button: tk.Button, role='neutral'):
        schemes = {'primary': ('#3b82f6', 'white', '#2563eb'), 'success': ('#16a34a', 'white', '#15803d'), 'neutral': ('#eef2f7', '#1f2937', '#dbe4f0')}
        bg, fg, active = schemes.get(role, schemes['neutral'])
        button.configure(bg=bg, fg=fg, activebackground=active, activeforeground=fg, relief='flat', bd=0, padx=12, pady=8, cursor='hand2', font=('Segoe UI Semibold', 10), highlightthickness=0)

    def _build_ui(self):
        self.configure(bg='#f5f7fb')
        fr = tk.Frame(self, padx=16, pady=16, bg='#f5f7fb')
        fr.pack(fill=tk.BOTH, expand=True)
        head = tk.Frame(fr, bg='#eaf1ff', padx=14, pady=12)
        head.pack(fill=tk.X, pady=(0, 14))
        tk.Label(head, text='Powiadomienia e-mail', bg='#eaf1ff', fg='#0f172a', font=('Segoe UI Semibold', 15)).pack(anchor='w')
        tk.Label(head, text='Wprowadź dane SMTP i odbiorców alarmów dla progów inercji.', bg='#eaf1ff', fg='#475467', font=('Segoe UI', 10)).pack(anchor='w', pady=(4, 0))
        card = tk.LabelFrame(fr, text='Ustawienia wysyłki', padx=12, pady=12, bg='white', fg='#334155', font=('Segoe UI Semibold', 10), bd=1, relief='solid')
        card.pack(fill=tk.BOTH, expand=True)
        tk.Checkbutton(card, text='Włącz powiadomienia e-mail', variable=self.enabled_var, bg='white', fg='#0f172a', activebackground='white', activeforeground='#0f172a', selectcolor='white', font=('Segoe UI', 10)).grid(row=0, column=0, columnspan=2, sticky='w', pady=(0, 8))
        fields = [('SMTP host', self.smtp_host_var, None), ('SMTP port', self.smtp_port_var, None), ('SMTP login', self.smtp_username_var, None), ('SMTP hasło / App Password', self.smtp_password_var, '*'), ('Nazwa nadawcy', self.sender_name_var, None), ('Adres nadawcy', self.sender_email_var, None), ('Adresaci (przecinek/średnik)', self.recipients_var, None)]
        for i, (lbl, var, mask) in enumerate(fields, start=1):
            tk.Label(card, text=lbl, bg='white', fg='#344054', font=('Segoe UI Semibold', 9)).grid(row=i, column=0, sticky='w', pady=5)
            ent = tk.Entry(card, textvariable=var, width=72, show=mask)
            ent.grid(row=i, column=1, sticky='we', pady=5)
            self._style_entry(ent)
        tk.Label(card, text='Tryb alarmu', bg='white', fg='#344054', font=('Segoe UI Semibold', 9)).grid(row=8, column=0, sticky='w', pady=5)
        om = tk.OptionMenu(card, self.trigger_mode_var, 'above', 'below')
        om.grid(row=8, column=1, sticky='w')
        om.configure(bg='#eef2f7', fg='#1f2937', activebackground='#dbe4f0', activeforeground='#1f2937', relief='flat', bd=0, cursor='hand2', font=('Segoe UI', 10), highlightthickness=0)
        try:
            om['menu'].configure(bg='white', fg='#111827', activebackground='#eaf1ff', activeforeground='#111827')
        except Exception:
            pass
        row = tk.Frame(fr, bg='#f5f7fb')
        row.pack(anchor='w', pady=(14, 0))
        btn_save = tk.Button(row, text='Zapisz', command=self.on_save)
        btn_test = tk.Button(row, text='Wyślij test', command=self.on_test)
        btn_close = tk.Button(row, text='Zamknij', command=self.destroy)
        self._style_button(btn_save, 'success')
        self._style_button(btn_test, 'primary')
        self._style_button(btn_close, 'neutral')
        btn_save.pack(side=tk.LEFT, padx=4)
        btn_test.pack(side=tk.LEFT, padx=4)
        btn_close.pack(side=tk.LEFT, padx=4)
        card.columnconfigure(1, weight=1)

    def _apply(self):
        c = self.service.config
        c.enabled = self.enabled_var.get()
        c.smtp_host = self.smtp_host_var.get().strip()
        c.smtp_port = int(self.smtp_port_var.get().strip() or '0')
        c.smtp_username = self.smtp_username_var.get().strip()
        c.smtp_password = self.smtp_password_var.get().strip()
        c.sender_name = self.sender_name_var.get().strip() or 'Analiza inercji'
        c.sender_email = self.sender_email_var.get().strip()
        c.recipients = self.service.parse_recipients(self.recipients_var.get())
        c.trigger_mode = self.trigger_mode_var.get().strip() or 'above'

    def on_save(self):
        try:
            self._apply()
            ok, msg = self.service.config.validate()
            if not ok:
                messagebox.showerror('Błąd konfiguracji', msg)
                return
            self.service.save_config()
            messagebox.showinfo('E-mail', 'Zapisano konfigurację.')
        except Exception as e:
            messagebox.showerror('Błąd', str(e))

    def on_test(self):
        try:
            self._apply()
            self.service.last_alarm_key = None
            ok, msg = self.service.send_alarm(prog=10000.0, wartosc=12000.0 if self.service.config.trigger_mode == 'above' else 9000.0, data_modelu='TEST', godzina_modelu='00:30', source_file='test.epc')
            if ok:
                messagebox.showinfo('E-mail', msg)
            else:
                messagebox.showwarning('E-mail', msg)
        except Exception as e:
            messagebox.showerror('Błąd wysyłki', str(e))

class AlarmHistoryWindow(tk.Toplevel):

    def __init__(self, master, history):
        super().__init__(master)
        self.title('Historia alarmów')
        center_toplevel(self, 1320, 660)
        self.configure(bg='#f5f7fb')
        frame = tk.Frame(self, padx=12, pady=12, bg='#f5f7fb')
        frame.pack(fill=tk.BOTH, expand=True)
        cols = ('czas', 'typ_alarmu', 'data', 'godzina', 'wartosc', 'prog', 'plik', 'mail', 'status', 'opis')
        self.tree = ttk.Treeview(frame, columns=cols, show='headings', height=22)
        headers = {'czas': 'Kiedy alarm wystąpił', 'typ_alarmu': 'Tryb', 'data': 'Data modelu', 'godzina': 'Godzina modelu', 'wartosc': 'Wartość [MW·s]', 'prog': 'Próg [MW·s]', 'plik': 'Plik', 'mail': 'Mail', 'status': 'Status wysyłki', 'opis': 'Opis'}
        widths = {'czas': 165, 'typ_alarmu': 80, 'data': 105, 'godzina': 100, 'wartosc': 120, 'prog': 120, 'plik': 220, 'mail': 70, 'status': 170, 'opis': 250}
        for col in cols:
            self.tree.heading(col, text=headers[col])
            self.tree.column(col, width=widths[col], anchor='w')
        yscroll = ttk.Scrollbar(frame, orient='vertical', command=self.tree.yview)
        xscroll = ttk.Scrollbar(frame, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        yscroll.pack(side=tk.RIGHT, fill=tk.Y)
        xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree.tag_configure('mail_ok', background='#e8f5e9', foreground='#1b5e20')
        self.tree.tag_configure('mail_err', background='#ffebee', foreground='#b71c1c')
        self.tree.tag_configure('mail_off', background='#eeeeee', foreground='#424242')
        for alarm in reversed(history):
            mail_status = alarm.get('mail_status', '-')
            if alarm.get('mail_attempted') and alarm.get('mail_ok'):
                tag = 'mail_ok'
            elif alarm.get('mail_attempted') and (not alarm.get('mail_ok')):
                tag = 'mail_err'
            else:
                tag = 'mail_off'
            self.tree.insert('', tk.END, values=(alarm.get('czas_alarmu', ''), alarm.get('trigger_mode', ''), alarm.get('data_modelu', ''), alarm.get('godzina_modelu', ''), f"{alarm.get('wartosc', 0.0):,.2f}", f"{alarm.get('prog', 0.0):,.2f}", alarm.get('source_file', ''), 'TAK' if alarm.get('mail_attempted') else 'NIE', mail_status, alarm.get('opis', '')), tags=(tag,))

class TreeviewTooltip:

    def __init__(self, tree: ttk.Treeview):
        self.tree = tree
        self.tipwindow = None
        self.current_item = None
        self.text = ''
        self.tree.bind('<Motion>', self.on_motion)
        self.tree.bind('<Leave>', self.hide)

    def show(self, x, y, text):
        if self.tipwindow and self.text == text:
            try:
                self.tipwindow.geometry(f'+{x + 18}+{y + 12}')
                return
            except Exception:
                pass
        self.hide()
        self.text = text
        self.tipwindow = tw = tk.Toplevel(self.tree)
        tw.wm_overrideredirect(True)
        tw.wm_geometry(f'+{x + 18}+{y + 12}')
        lbl = tk.Label(tw, text=text, justify='left', background='#fff8e1', relief='solid', borderwidth=1, font=('Cascadia Mono', 9), padx=8, pady=6)
        lbl.pack()

    def hide(self, event=None):
        if self.tipwindow:
            try:
                self.tipwindow.destroy()
            except Exception:
                pass
            self.tipwindow = None
            self.current_item = None
            self.text = ''

    def on_motion(self, event):
        item = self.tree.identify_row(event.y)
        if not item:
            self.hide()
            return
        if hasattr(self.tree, '_tooltip_provider'):
            txt = self.tree._tooltip_provider(item)
            if txt:
                self.current_item = item
                self.show(event.x_root, event.y_root, txt)
            else:
                self.hide()

class AnalizaInercjiExcelApp:
    GENERATOR_GROUP_LABELS = {101: 'JWCDc', 102: 'JWCDp', 201: 'JWCKc', 202: 'JWCKw', 203: 'JWCKp', 204: 'JWCKf', 205: 'JWCKpv', 301: 'MC', 302: 'PR', 303: 'MW', 304: 'FW', 305: 'PV'}
    GROUP_COLORS = ['#2563eb', '#f97316', '#22c55e', '#a855f7', '#ec4899', '#14b8a6', '#8b5cf6', '#0ea5e9', '#eab308', '#ef4444', '#6366f1', '#84cc16']
    LIGHT_PALETTE = {'bg': '#f5f7fb', 'card': '#ffffff', 'card_alt': '#eef2f7', 'header': '#eaf1ff', 'border': '#d7e0ec', 'text': '#0f172a', 'muted': '#667085', 'input_bg': '#ffffff', 'input_fg': '#111827', 'accent': '#3b82f6', 'accent_active': '#2563eb', 'success': '#16a34a', 'success_active': '#15803d', 'warning': '#ea580c', 'warning_active': '#c2410c', 'danger': '#dc2626', 'danger_active': '#b91c1c', 'neutral_btn': '#eef2f7', 'neutral_btn_active': '#dbe4f0', 'chart_grid': '#d9e3f0', 'line_main': '#2563eb', 'line_secondary': '#64748b', 'input_border': '#cfd8e3', 'selection': '#dbeafe', 'selection_fg': '#0f172a'}
    DARK_PALETTE = {'bg': '#0f172a', 'card': '#172033', 'card_alt': '#1e293b', 'header': '#12223f', 'border': '#334155', 'text': '#e5e7eb', 'muted': '#9aa8bd', 'input_bg': '#111827', 'input_fg': '#f8fafc', 'accent': '#60a5fa', 'accent_active': '#3b82f6', 'success': '#4ade80', 'success_active': '#22c55e', 'warning': '#fb923c', 'warning_active': '#f97316', 'danger': '#f87171', 'danger_active': '#ef4444', 'neutral_btn': '#243145', 'neutral_btn_active': '#334155', 'chart_grid': '#334155', 'line_main': '#60a5fa', 'line_secondary': '#94a3b8', 'input_border': '#334155', 'selection': '#1d4ed8', 'selection_fg': '#f8fafc'}

    def __init__(self, root):
        self.root = root
        self.root.title('Analiza inercji – nowoczesny interfejs')
        self.root.configure(bg=self.LIGHT_PALETTE['bg'])
        self.font_ui = ('Segoe UI', 10)
        self.font_ui_small = ('Segoe UI', 9)
        self.font_ui_bold = ('Segoe UI Semibold', 10)
        self.font_ui_bold_small = ('Segoe UI Semibold', 9)
        self.font_title = ('Segoe UI Semibold', 20)
        self.font_subtitle = ('Segoe UI', 10)
        self.font_code = ('Cascadia Mono', 10)
        self.font_code_small = ('Cascadia Mono', 9)
        self.wszystkie_generatory = {}
        self.statusy_oryginalne = {}
        self.opis_modeli = {}
        self.inercja_na_godzine = {h: 0.0 for h in range(24)}
        self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
        self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
        self.zapotrzebowanie_na_godzine = {h: 0.0 for h in range(24)}
        self.modele_auto = {}
        self.auto_order = []
        self.auto_seen_files = {}
        self.auto_counter = 0
        self.obliczenia_wykonane = False
        self.tekst_filtrowania = ''
        self.H_generatory = {}
        self.threshold_update_job = None
        self.sent_alarm_tokens = set()
        self.alarm_history = []
        self.aktywny_typ_modelu = None
        self.aktualna_godzina = None
        self.aktualny_auto_id = None
        self.tryb_filtrowania_var = tk.StringVar(value='Wszystkie')
        self.sortowanie_var = tk.StringVar(value='Kod generatora A-Z')
        self.pokaz_etykiety_inercji_var = tk.BooleanVar(value=True)
        self.pokaz_etykiety_zapotrzebowania_var = tk.BooleanVar(value=True)
        self.pokaz_linie_progu_var = tk.BooleanVar(value=True)
        self.progress_var = tk.DoubleVar(value=0)
        self.monitor_folder_var = tk.BooleanVar(value=False)
        self.monitor_folder_path = ''
        self.monitor_interval_ms = 5000
        self.email_alarm = EmailAlarmService()
        self.ui_config_path = Path('~/.analiza_inercji_ui.json').expanduser()
        self.ui_config = {}
        self.dark_mode_var = tk.BooleanVar(value=False)
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        self.column_filters: dict[str, tk.StringVar] = {}
        self.filter_entries: dict[str, tk.Entry] = {}
        self.bliski_prog_proc = 0.05
        self.ttk_style = ttk.Style(self.root)
        self.button_roles: dict[tk.Button, str] = {}
        self.label_neutral = set()
        self.skonfiguruj_przewijane_okno()
        self.utworz_interfejs()
        self.load_ui_config()
        self.apply_ui_config()
        self.tooltip = TreeviewTooltip(self.tree_generatory)
        self.tree_generatory._tooltip_provider = self.get_generator_tooltip
        atexit.register(self.safe_save_ui_config)

    def get_palette(self):
        return self.DARK_PALETTE if self.dark_mode_var.get() else self.LIGHT_PALETTE

    def get_button_scheme(self, role='neutral'):
        pal = self.get_palette()
        if role == 'primary':
            return (pal['accent'], 'white', pal['accent_active'])
        if role == 'success':
            return (pal['success'], 'white', pal['success_active'])
        if role == 'warning':
            return (pal['warning'], 'white', pal['warning_active'])
        if role == 'danger':
            return (pal['danger'], 'white', pal['danger_active'])
        return (pal['neutral_btn'], pal['text'], pal['neutral_btn_active'])

    def stylizuj_przycisk(self, button: tk.Button, role='neutral'):
        bg, fg, active = self.get_button_scheme(role)
        button.configure(bg=bg, fg=fg, activebackground=active, activeforeground=fg, relief='flat', bd=0, padx=12, pady=8, cursor='hand2', font=self.font_ui_bold, highlightthickness=0)

    def utworz_przycisk(self, parent, text, command, role='neutral', width=None):
        btn = tk.Button(parent, text=text, command=command, width=width)
        self.button_roles[btn] = role
        self.stylizuj_przycisk(btn, role)
        return btn

    def stylizuj_etykiete_neutralna(self, label: tk.Label):
        self.label_neutral.add(label)
        parent_bg = label.master.cget('bg') if label.master else self.get_palette()['bg']
        label.configure(bg=parent_bg, fg=self.get_palette()['text'], font=self.font_ui)

    def stylizuj_entry(self, entry: tk.Entry):
        pal = self.get_palette()
        entry.configure(relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'], font=self.font_code)

    def stylizuj_option_menu(self, om: tk.Widget):
        pal = self.get_palette()
        om.configure(bg=pal['neutral_btn'], fg=pal['text'], activebackground=pal['neutral_btn_active'], activeforeground=pal['text'], relief='flat', bd=0, cursor='hand2', font=self.font_ui, highlightthickness=0)
        try:
            om['menu'].configure(bg=pal['card'], fg=pal['text'], activebackground=pal['selection'], activeforeground=pal['selection_fg'], font=self.font_ui)
        except Exception:
            pass

    def configure_ttk_styles(self):
        pal = self.get_palette()
        try:
            self.ttk_style.theme_use('clam')
        except Exception:
            pass
        self.ttk_style.configure('Treeview', background=pal['input_bg'], fieldbackground=pal['input_bg'], foreground=pal['text'], bordercolor=pal['border'], lightcolor=pal['border'], darkcolor=pal['border'], rowheight=28, relief='flat', font=self.font_ui_small)
        self.ttk_style.map('Treeview', background=[('selected', pal['accent'])], foreground=[('selected', 'white')])
        self.ttk_style.configure('Treeview.Heading', background=pal['card_alt'], foreground=pal['text'], bordercolor=pal['border'], relief='flat', font=self.font_ui_bold, padding=(8, 6))
        self.ttk_style.map('Treeview.Heading', background=[('active', pal['header'])], foreground=[('active', pal['text'])])
        self.ttk_style.configure('Modern.Horizontal.TProgressbar', troughcolor=pal['card_alt'], background=pal['accent'], bordercolor=pal['border'], lightcolor=pal['accent'], darkcolor=pal['accent'], thickness=10)
        self.ttk_style.configure('Vertical.TScrollbar', background=pal['card_alt'], troughcolor=pal['bg'], bordercolor=pal['border'], arrowcolor=pal['muted'], relief='flat')
        self.ttk_style.configure('Horizontal.TScrollbar', background=pal['card_alt'], troughcolor=pal['bg'], bordercolor=pal['border'], arrowcolor=pal['muted'], relief='flat')

    def configure_widget_tree(self, widget):
        pal = self.get_palette()
        cls = widget.winfo_class()
        if isinstance(widget, tk.Button):
            self.stylizuj_przycisk(widget, self.button_roles.get(widget, 'neutral'))
            return
        if widget is self.root:
            widget.configure(bg=pal['bg'])
        elif isinstance(widget, tk.LabelFrame):
            widget.configure(bg=pal['card'], fg=pal['text'], bd=1, relief='solid', padx=8, pady=8, font=self.font_ui_bold, highlightbackground=pal['border'], highlightcolor=pal['border'])
        elif isinstance(widget, tk.Frame):
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            widget.configure(bg=parent_bg)
        elif isinstance(widget, tk.Label):
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            fg = widget.cget('fg')
            if widget in self.label_neutral:
                fg = pal['text']
            widget.configure(bg=parent_bg, fg=fg)
            try:
                if widget.cget('font') in ('TkDefaultFont',):
                    widget.configure(font=self.font_ui)
            except Exception:
                pass
        elif isinstance(widget, tk.Entry):
            self.stylizuj_entry(widget)
        elif isinstance(widget, tk.Listbox):
            widget.configure(bg=pal['input_bg'], fg=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], selectbackground=pal['selection'], selectforeground=pal['selection_fg'], font=self.font_code)
        elif isinstance(widget, (tk.Text, scrolledtext.ScrolledText)):
            widget.configure(bg=pal['input_bg'], fg=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'], highlightcolor=pal['accent'], insertbackground=pal['input_fg'], selectbackground=pal['selection'], selectforeground=pal['selection_fg'], font=self.font_code_small)
        elif isinstance(widget, tk.Canvas):
            if widget is self.main_canvas:
                widget.configure(bg=pal['bg'], highlightthickness=0)
            else:
                parent_bg = pal['bg']
                try:
                    parent_bg = widget.master.cget('bg')
                except Exception:
                    pass
                widget.configure(bg=parent_bg, highlightthickness=0)
        elif cls == 'Checkbutton':
            parent_bg = pal['bg']
            try:
                parent_bg = widget.master.cget('bg')
            except Exception:
                pass
            widget.configure(bg=parent_bg, fg=pal['text'], activebackground=parent_bg, activeforeground=pal['text'], selectcolor=parent_bg, font=self.font_ui)
        elif cls == 'Menubutton':
            self.stylizuj_option_menu(widget)
        for child in widget.winfo_children():
            self.configure_widget_tree(child)

    def stylizuj_text_tags(self):
        pal = self.get_palette()
        self.pole_szczegoly.tag_config('header', foreground=pal['accent'], font=(self.font_code_small[0], self.font_code_small[1], 'bold'))
        self.pole_szczegoly.tag_config('green_line', foreground='#15803d' if not self.dark_mode_var.get() else '#86efac')
        self.pole_szczegoly.tag_config('red_line', foreground='#b91c1c' if not self.dark_mode_var.get() else '#fca5a5')
        self.pole_szczegoly.tag_config('manual_line', background='#fff3cd' if not self.dark_mode_var.get() else '#3a2f12', foreground='#c77700' if not self.dark_mode_var.get() else '#fbbf24')
        self.pole_szczegoly.tag_config('summary', foreground=pal['text'], font=(self.font_code_small[0], self.font_code_small[1], 'bold'))

    def stylizuj_tree_tags(self):
        if self.dark_mode_var.get():
            self.tree_generatory.tag_configure('zal', background='#113126', foreground='#86efac')
            self.tree_generatory.tag_configure('wyl', background='#3a1d22', foreground='#fca5a5')
            self.tree_generatory.tag_configure('manual', background='#3a3212', foreground='#fcd34d')
        else:
            self.tree_generatory.tag_configure('zal', background='#e8f5e9', foreground='#1b5e20')
            self.tree_generatory.tag_configure('wyl', background='#ffebee', foreground='#b71c1c')
            self.tree_generatory.tag_configure('manual', background='#fff3cd', foreground='#e65100')

    def stylizuj_os(self, ax, grid=True, grid_axis='both'):
        pal = self.get_palette()
        ax.set_facecolor(pal['card'])
        for spine in ax.spines.values():
            spine.set_color(pal['border'])
        ax.tick_params(colors=pal['muted'])
        ax.title.set_color(pal['text'])
        ax.xaxis.label.set_color(pal['text'])
        ax.yaxis.label.set_color(pal['text'])
        if grid:
            ax.grid(True, axis=grid_axis, linestyle='--', alpha=0.6, color=pal['chart_grid'])

    def stylizuj_os_prawa(self, ax, color):
        pal = self.get_palette()
        ax.set_facecolor('none')
        ax.grid(False)
        for spine in ax.spines.values():
            spine.set_color(pal['border'])
        try:
            ax.spines['right'].set_color(color)
        except Exception:
            pass
        ax.tick_params(axis='y', colors=color)
        ax.yaxis.label.set_color(color)
        ax.tick_params(axis='x', which='both', bottom=False, labelbottom=False)
        ax.yaxis.set_label_position('left')
        ax.yaxis.label.set_rotation(90)
        ax.yaxis.label.set_horizontalalignment('right')
        ax.yaxis.set_label_coords(-0.13, 0.18)

    def apply_theme(self):
        pal = self.get_palette()
        self.root.configure(bg=pal['bg'])
        self.main_canvas.configure(bg=pal['bg'])
        self.main_frame.configure(bg=pal['bg'])
        self.configure_ttk_styles()
        self.configure_widget_tree(self.root)
        self.stylizuj_tree_tags()
        self.stylizuj_text_tags()
        self.progress_bar.configure(style='Modern.Horizontal.TProgressbar')
        try:
            self.figura.patch.set_facecolor(pal['bg'])
            self.figura_grupy.patch.set_facecolor(pal['bg'])
        except Exception:
            pass
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            self.przygotuj_pusty_wykres_grup()
        else:
            self.aktualizuj_wykres_grup_generatorow(model_type, key)

    def toggle_dark_mode(self):
        self.apply_theme()
        self.save_ui_config()

    def bind_scroll_area(self, widget):
        return

    def is_widget_or_descendant(self, widget, ancestor):
        current = widget
        while current is not None:
            if current == ancestor:
                return True
            current = getattr(current, 'master', None)
        return False

    def get_scroll_target_under_pointer(self, x_root, y_root):
        try:
            widget = self.root.winfo_containing(x_root, y_root)
        except Exception:
            widget = None
        if widget is None:
            return self.main_canvas
        if self.is_widget_or_descendant(widget, self.tree_generatory):
            return self.tree_generatory
        if self.is_widget_or_descendant(widget, self.listbox_modele):
            return self.listbox_modele
        if self.is_widget_or_descendant(widget, self.listbox_modele_auto):
            return self.listbox_modele_auto
        if self.is_widget_or_descendant(widget, self.pole_szczegoly):
            return self.pole_szczegoly
        return self.main_canvas

    def scroll_widget(self, widget, units, horizontal=False):
        try:
            if horizontal:
                widget.xview_scroll(units, 'units')
            else:
                widget.yview_scroll(units, 'units')
            return True
        except Exception:
            return False

    def on_global_mousewheel(self, event):
        units = int(-1 * (event.delta / 120)) if event.delta else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if not self.scroll_widget(target, units):
            self.main_canvas.yview_scroll(units, 'units')

    def on_global_mousewheel_linux(self, event):
        step = -1 if event.num == 4 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if not self.scroll_widget(target, step):
            self.main_canvas.yview_scroll(step, 'units')

    def on_shift_mousewheel_main(self, event):
        units = int(-1 * (event.delta / 120)) if event.delta else 0
        if units == 0:
            units = -1 if getattr(event, 'delta', 0) > 0 else 1
        target = self.get_scroll_target_under_pointer(event.x_root, event.y_root)
        if target is self.tree_generatory:
            if not self.scroll_widget(self.tree_generatory, units, horizontal=True):
                self.main_canvas.xview_scroll(units, 'units')
        else:
            self.main_canvas.xview_scroll(units, 'units')

    def reset_alarm_cache(self):
        self.sent_alarm_tokens.clear()

    def konwertuj_float(self, value):
        return float(str(value).replace(',', '.'))

    def konwertuj_float_soft(self, value, default=0.0):
        try:
            return self.konwertuj_float(value)
        except Exception:
            return default

    def aktualizuj_status(self, tekst):
        try:
            self.label_status.config(text=f'Status: {tekst}')
            self.root.update_idletasks()
        except Exception:
            pass

    def pobierz_wartosc_graniczna(self):
        txt = self.entry_wartosc_graniczna.get().strip()
        if not txt:
            return None
        try:
            return self.konwertuj_float(txt)
        except Exception:
            return None

    def schedule_prog_changed(self, event=None):
        if self.threshold_update_job is not None:
            try:
                self.root.after_cancel(self.threshold_update_job)
            except Exception:
                pass
        self.threshold_update_job = self.root.after(350, self.on_prog_changed)

    def on_prog_changed(self, event=None):
        self.threshold_update_job = None
        self.odswiez_widok_dla_progu(check_alarm=False)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.save_ui_config()

    def sprawdz_prog_teraz(self):
        if self.threshold_update_job is not None:
            try:
                self.root.after_cancel(self.threshold_update_job)
            except Exception:
                pass
            self.threshold_update_job = None
        self.odswiez_widok_dla_progu(check_alarm=False)

    def odswiez_widok_dla_progu(self, check_alarm=False):
        prog = self.pobierz_wartosc_graniczna()
        trigger_mode = self.email_alarm.config.trigger_mode
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if prog is None:
            txt = self.entry_wartosc_graniczna.get().strip()
            if txt:
                self.label_limit_ocena.config(text='Ocena: niepoprawna wartość graniczna', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text='Ocena: brak wartości granicznej', fg=self.get_palette()['muted'])
        elif dodatnie:
            if trigger_mode == 'above':
                _, val_ref = max(dodatnie, key=lambda x: x[1])
                if val_ref > prog:
                    self.label_limit_ocena.config(text=f'Ocena: PRZEKROCZONO wartość graniczną o {val_ref - prog:,.2f} MW·s', fg=self.get_palette()['danger'])
                else:
                    self.label_limit_ocena.config(text=f'Ocena: NIE przekroczono wartości granicznej (zapas {prog - val_ref:,.2f} MW·s)', fg=self.get_palette()['success'])
            else:
                _, val_ref = min(dodatnie, key=lambda x: x[1])
                if val_ref < prog:
                    self.label_limit_ocena.config(text=f'Ocena: SPADŁO poniżej wartości granicznej o {prog - val_ref:,.2f} MW·s', fg=self.get_palette()['danger'])
                else:
                    self.label_limit_ocena.config(text=f'Ocena: NIE spadło poniżej wartości granicznej (margines {val_ref - prog:,.2f} MW·s)', fg=self.get_palette()['success'])
        else:
            self.label_limit_ocena.config(text='Ocena: brak wartości do porównania', fg=self.get_palette()['muted'])
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wykres()
        self.aktualizuj_wskazniki_jakosci()
        if self.aktywny_typ_modelu is not None:
            self.aktualizuj_analize_inercyjna(*self.get_current_model_ref())

    def przygotuj_pusty_wykres(self):
        pal = self.get_palette()
        self.os1.clear()
        self.os2.clear()
        if not hasattr(self, 'os1_load') or self.os1_load is None:
            self.os1_load = self.os1.twinx()
        self.os1_load.clear()
        self.figura.patch.set_facecolor(pal['bg'])
        self.os1.set_title('Inercja systemowa E_sys(t) = Σ Pmax_i(t) × H_i')
        self.os1.set_xlabel('Czas [HH:30]')
        self.os1.set_ylabel('Inercja [MW·s]')
        self.stylizuj_os(self.os1, grid=True, grid_axis='both')
        self.os1.yaxis.set_label_coords(-0.09, 0.78)
        self.os1.set_xticks(range(24))
        self.os1.set_xticklabels([f'{h:02d}:30' for h in range(24)], rotation=45)
        self.os1.set_xlim(-0.5, 23.5)
        self.os1.set_ylim(0, 1)
        self.os1_load.set_ylabel('Zapotrzebowanie [MW]')
        self.stylizuj_os_prawa(self.os1_load, pal['warning'])
        self.os1_load.set_xlim(-0.5, 23.5)
        self.os1_load.set_ylim(0, 1)
        self.os2.set_title('Liczba załączonych generatorów z H')
        self.os2.set_xlabel('Czas [HH:30]')
        self.os2.set_ylabel('Liczba')
        self.stylizuj_os(self.os2, grid=True, grid_axis='both')
        self.os2.set_xticks(range(24))
        self.os2.set_xticklabels([f'{h:02d}:30' for h in range(24)], rotation=45)
        self.os2.set_xlim(-0.5, 23.5)
        self.os2.set_ylim(0, 1)
        self.canvas.draw_idle()

    def przygotuj_pusty_wykres_grup(self, komunikat='Brak wybranego modelu'):
        pal = self.get_palette()
        self.figura_grupy.patch.set_facecolor(pal['bg'])
        self.os_grupy.clear()
        self.os_grupy.set_title('Udział grup generatorów w modelu')
        self.os_grupy.set_xlabel('Grupa generatora')
        self.os_grupy.set_ylabel('Udział [% E_sys]')
        self.stylizuj_os(self.os_grupy, grid=True, grid_axis='y')
        self.os_grupy.set_ylim(0, 100)
        self.os_grupy.text(0.5, 0.5, komunikat, ha='center', va='center', transform=self.os_grupy.transAxes, fontsize=10, color=pal['muted'])
        self.canvas_grupy.draw_idle()

    def skonfiguruj_przewijane_okno(self):
        self.main_canvas = tk.Canvas(self.root, highlightthickness=0, bg=self.get_palette()['bg'])
        self.v_scroll_main = ttk.Scrollbar(self.root, orient='vertical', command=self.main_canvas.yview)
        self.h_scroll_main = ttk.Scrollbar(self.root, orient='horizontal', command=self.main_canvas.xview)
        self.main_canvas.configure(yscrollcommand=self.v_scroll_main.set, xscrollcommand=self.h_scroll_main.set)
        self.v_scroll_main.pack(side=tk.RIGHT, fill=tk.Y)
        self.h_scroll_main.pack(side=tk.BOTTOM, fill=tk.X)
        self.main_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.main_frame = tk.Frame(self.main_canvas, bg=self.get_palette()['bg'])
        self.main_canvas.create_window((0, 0), window=self.main_frame, anchor='nw')
        self.main_frame.bind('<Configure>', lambda e: self.main_canvas.configure(scrollregion=self.main_canvas.bbox('all')))
        self.main_canvas.bind('<Configure>', lambda e: self.main_canvas.configure(scrollregion=self.main_canvas.bbox('all')))
        self.main_canvas.bind_all('<MouseWheel>', self.on_global_mousewheel)
        self.main_canvas.bind_all('<Button-4>', self.on_global_mousewheel_linux)
        self.main_canvas.bind_all('<Button-5>', self.on_global_mousewheel_linux)
        self.main_canvas.bind_all('<Shift-MouseWheel>', self.on_shift_mousewheel_main)

    def utworz_interfejs(self):
        pal = self.get_palette()
        header = tk.Frame(self.main_frame, bg=pal['header'], padx=18, pady=16)
        header.pack(side=tk.TOP, fill=tk.X, padx=10, pady=(10, 12))
        self.header_frame = header
        self.label_title = tk.Label(header, text='Analiza inercji', bg=pal['header'], fg=pal['text'], font=self.font_title)
        self.label_title.pack(anchor='w')
        self.label_subtitle = tk.Label(header, bg=pal['header'], fg=pal['muted'], font=self.font_subtitle, text='Import EPC + DYD, analiza E_sys oraz zapotrzebowania mocy z sekcji LOAD DATA')
        self.label_subtitle.pack(anchor='w', pady=(4, 0))
        panel_gorny = tk.Frame(self.main_frame, padx=10, pady=0, bg=pal['bg'])
        panel_gorny.pack(side=tk.TOP, fill=tk.X)
        przyciski = [('1. Wczytaj wiele plików EPC', self.wczytaj_wiele_plikow_epc, 'primary'), ('2. Wczytaj plik DYD PSLF (H)', self.wczytaj_plik_genrou_H, 'primary'), ('3. Oblicz inercję (Pmax×H)', self.oblicz_inercje_wszystkie, 'success'), ('4. Przelicz wybrany model', self.przelicz_aktywny_model, 'neutral'), ('5. Przywróć oryginalne statusy modelu', self.przywroc_oryginalne_statusy_modelu, 'neutral'), ('6. Przywróć wszystkie modele 24h', self.przywroc_wszystkie_modele_do_oryginalu, 'neutral'), ('7. Eksport tabeli do Excel', self.eksportuj_do_excel, 'neutral'), ('8. Eksportuj model EPC po zmianach', self.eksportuj_aktualny_model_epc_po_zmianach, 'neutral'), ('9. Eksport raportu zmian', self.eksportuj_raport_zmian, 'neutral'), ('10. Zapisz sesję', self.zapisz_sesje, 'neutral'), ('11. Wczytaj sesję', self.wczytaj_sesje, 'neutral'), ('12. Walidacja EPC ↔ DYD', self.pokaz_walidacje_epc_dyd, 'warning'), ('13. Konfiguruj e-mail alarmowy', self.konfiguruj_email_alarm, 'warning'), ('14. Historia alarmów', self.pokaz_historie_alarmow, 'neutral'), ('15. Eksport pełnego raportu PDF', self.eksportuj_pelny_raport_pdf, 'success')]
        frame_przyciski = tk.LabelFrame(panel_gorny, text='Akcje główne', padx=10, pady=10, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_przyciski.pack(side=tk.TOP, fill=tk.X)
        kolumn = 4
        for c in range(kolumn):
            frame_przyciski.grid_columnconfigure(c, weight=1, uniform='btn')
        for i, (txt, cmd, role) in enumerate(przyciski):
            r = i // kolumn
            c = i % kolumn
            btn = self.utworz_przycisk(frame_przyciski, txt, cmd, role=role)
            btn.grid(row=r, column=c, padx=5, pady=5, sticky='ew')
        frame_postep = tk.LabelFrame(panel_gorny, text='Postęp importu', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_postep.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        tk.Label(frame_postep, text='Postęp importu EPC:', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT, padx=(0, 8))
        self.progress_bar = ttk.Progressbar(frame_postep, orient='horizontal', mode='determinate', variable=self.progress_var, length=340, style='Modern.Horizontal.TProgressbar')
        self.progress_bar.pack(side=tk.LEFT, padx=4)
        self.label_progress = tk.Label(frame_postep, text='0%', fg=pal['accent'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_progress.pack(side=tk.LEFT, padx=8)
        frame_monitor = tk.LabelFrame(panel_gorny, text='Monitoring folderu z modelami automatycznymi', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_monitor.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.utworz_przycisk(frame_monitor, 'Wybierz folder', self.wybierz_folder_monitoringu, role='neutral').pack(side=tk.LEFT, padx=4)
        chk_monitor = tk.Checkbutton(frame_monitor, text='Monitoruj folder', variable=self.monitor_folder_var, command=self.toggle_monitor_folder)
        chk_monitor.pack(side=tk.LEFT, padx=8)
        self.label_monitor_folder = tk.Label(frame_monitor, text='Folder: nie wybrano', fg=pal['accent'], bg=pal['card'], font=self.font_code_small)
        self.label_monitor_folder.pack(side=tk.LEFT, padx=10)
        frame_export_plot = tk.LabelFrame(panel_gorny, text='Szybki eksport wykresów', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_export_plot.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.utworz_przycisk(frame_export_plot, 'Eksport wykresu 24h PNG', lambda: self.export_chart(self.figura, 'wykres_24h.png'), role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_export_plot, 'Eksport wykresu grup PNG', lambda: self.export_chart(self.figura_grupy, 'wykres_grup.png'), role='neutral').pack(side=tk.LEFT, padx=3)
        self.utworz_przycisk(frame_export_plot, 'Eksport dashboardu PNG', self.export_dashboard_png, role='neutral').pack(side=tk.LEFT, padx=3)
        frame_opcje = tk.LabelFrame(panel_gorny, text='Opcje widoku', padx=10, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold, bd=1, relief='solid')
        frame_opcje.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        tk.Checkbutton(frame_opcje, text='Etykiety inercji', variable=self.pokaz_etykiety_inercji_var, command=self.on_chart_label_toggle).pack(side=tk.LEFT, padx=10)
        tk.Checkbutton(frame_opcje, text='Etykiety zapotrzebowania', variable=self.pokaz_etykiety_zapotrzebowania_var, command=self.on_chart_label_toggle).pack(side=tk.LEFT, padx=5)
        tk.Checkbutton(frame_opcje, text='Pokaż linię progu', variable=self.pokaz_linie_progu_var, command=self.aktualizuj_wykres).pack(side=tk.LEFT, padx=5)
        tk.Checkbutton(frame_opcje, text='Tryb ciemny', variable=self.dark_mode_var, command=self.toggle_dark_mode).pack(side=tk.LEFT, padx=8)
        self.label_status = tk.Label(frame_opcje, text='Status: Krok 1 → Wczytaj EPC, Krok 2 → DYD, Krok 3 → Oblicz', fg=pal['accent'], bg=pal['card'], font=self.font_ui_bold)
        self.label_status.pack(side=tk.RIGHT, padx=10)
        panel_srodkowy = tk.Frame(self.main_frame, bg=pal['bg'])
        panel_srodkowy.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=10, pady=(12, 10))
        lewa_kolumna = tk.Frame(panel_srodkowy, bg=pal['bg'])
        lewa_kolumna.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        frame_modele = tk.LabelFrame(lewa_kolumna, text='Modele godzinowe 24h', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_modele.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        listbox_frame = tk.Frame(frame_modele, bg=pal['card'])
        listbox_frame.pack(fill=tk.BOTH, expand=True)
        self.listbox_modele = tk.Listbox(listbox_frame, font=self.font_code, width=45, height=10)
        scrollbar_modele = ttk.Scrollbar(listbox_frame, orient='vertical', command=self.listbox_modele.yview)
        self.listbox_modele.configure(yscrollcommand=scrollbar_modele.set)
        self.listbox_modele.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_modele.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox_modele.bind('<<ListboxSelect>>', self.wybierz_model_manual)
        self.bind_scroll_area(self.listbox_modele)
        frame_modele_auto = tk.LabelFrame(lewa_kolumna, text='Modele automatycznie wczytane z folderu', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_modele_auto.pack(side=tk.TOP, fill=tk.BOTH, expand=True, pady=(10, 0))
        listbox_auto_frame = tk.Frame(frame_modele_auto, bg=pal['card'])
        listbox_auto_frame.pack(fill=tk.BOTH, expand=True)
        self.listbox_modele_auto = tk.Listbox(listbox_auto_frame, font=self.font_code, width=45, height=10)
        scrollbar_modele_auto = ttk.Scrollbar(listbox_auto_frame, orient='vertical', command=self.listbox_modele_auto.yview)
        self.listbox_modele_auto.configure(yscrollcommand=scrollbar_modele_auto.set)
        self.listbox_modele_auto.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_modele_auto.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox_modele_auto.bind('<<ListboxSelect>>', self.wybierz_model_auto)
        self.bind_scroll_area(self.listbox_modele_auto)
        frame_filtr = tk.LabelFrame(lewa_kolumna, text='Filtr generatorów', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_filtr.pack(side=tk.TOP, anchor='w', fill=tk.X, pady=(10, 0))
        tk.Label(frame_filtr, text='Kod generatora zawiera (można kilka po przecinku):', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(anchor='w')
        self.entry_filtr = tk.Entry(frame_filtr, font=self.font_code, width=18)
        self.entry_filtr.pack(fill=tk.X, padx=5, pady=4)
        self.entry_filtr.bind('<KeyRelease>', self.zastosuj_filtr)
        row1 = tk.Frame(frame_filtr, bg=pal['card'])
        row1.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(row1, text='Widok:', bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT)
        opcje_trybu = ['Wszystkie', 'Tylko zmienione ręcznie', 'Tylko załączone', 'Tylko wyłączone']
        self.option_tryb = tk.OptionMenu(row1, self.tryb_filtrowania_var, *opcje_trybu, command=lambda _: self.odswiez_aktywny_model())
        self.option_tryb.pack(side=tk.LEFT, padx=(8, 0))
        row2 = tk.Frame(frame_filtr, bg=pal['card'])
        row2.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(row2, text='Sortuj:', bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT)
        opcje_sort = ['Kod generatora A-Z', 'Kod generatora Z-A', 'Pmax malejąco', 'Pmax rosnąco', 'H malejąco', 'H rosnąco', 'Ei malejąco', 'Ei rosnąco', 'Status']
        self.option_sort = tk.OptionMenu(row2, self.sortowanie_var, *opcje_sort, command=lambda _: self.on_dropdown_sort_changed())
        self.option_sort.pack(side=tk.LEFT, padx=(8, 0))
        row3 = tk.Frame(frame_filtr, bg=pal['card'])
        row3.pack(fill=tk.X, padx=5, pady=4)
        self.utworz_przycisk(row3, 'Wyczyść', self.wyczysc_filtr, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Wszystkie', self.pokaz_wszystkie, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Załącz widoczne', self.zalacz_widoczne_z_filtra, role='success').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Wyłącz widoczne', self.wylacz_widoczne_z_filtra, role='warning').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(row3, 'Przywróć widoczne', self.przywroc_widoczne_z_filtra, role='neutral').pack(side=tk.LEFT, padx=2)
        self.label_wynikow = tk.Label(frame_filtr, text='0/0 generatorów po filtrze', fg=pal['success'], bg=pal['card'], font=self.font_code)
        self.label_wynikow.pack(pady=2)
        prawa_kolumna = tk.Frame(panel_srodkowy, bg=pal['bg'])
        prawa_kolumna.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(10, 0))
        frame_limit = tk.LabelFrame(prawa_kolumna, text='Wartość graniczna i maksimum z 24 godzin', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_limit.pack(side=tk.TOP, fill=tk.X, pady=(0, 10))
        tk.Label(frame_limit, text='Wartość graniczna [MW·s]:', font=self.font_ui_bold, bg=pal['card'], fg=pal['text']).pack(side=tk.LEFT, padx=(5, 5))
        self.entry_wartosc_graniczna = tk.Entry(frame_limit, width=14, font=self.font_code)
        self.entry_wartosc_graniczna.pack(side=tk.LEFT, padx=(0, 8))
        self.entry_wartosc_graniczna.bind('<KeyRelease>', self.schedule_prog_changed)
        self.utworz_przycisk(frame_limit, 'Sprawdź', self.sprawdz_prog_teraz, role='primary').pack(side=tk.LEFT, padx=4)
        self.label_maksimum_info = tk.Label(frame_limit, text='Maksymalna inercja z 24 godzin: brak danych', fg=pal['accent'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_maksimum_info.pack(side=tk.LEFT, padx=12)
        self.label_limit_ocena = tk.Label(frame_limit, text='Ocena: brak wartości granicznej', fg=pal['muted'], bg=pal['card'], font=('Cascadia Mono', 10, 'bold'))
        self.label_limit_ocena.pack(side=tk.LEFT, padx=12)
        frame_compare = tk.LabelFrame(prawa_kolumna, text='Porównanie przed / po zmianach dla wybranego modelu', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_compare.pack(side=tk.TOP, fill=tk.X, pady=(0, 10))
        self.label_porownanie = tk.Label(frame_compare, text='Brak wybranego modelu', justify='left', anchor='w', font=('Cascadia Mono', 10, 'bold'), fg=pal['accent'], bg=pal['card'])
        self.label_porownanie.pack(fill=tk.X, padx=5, pady=3)
        frame_metrics = tk.LabelFrame(prawa_kolumna, text='Wskaźniki jakości systemu (24h)', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_metrics.pack(side=tk.TOP, fill=tk.X, pady=(0, 10))
        self.label_wskazniki = tk.Label(frame_metrics, text='Brak danych 24h', justify='left', anchor='w', font=self.font_code_small, fg=pal['accent'], bg=pal['card'])
        self.label_wskazniki.pack(fill=tk.X, padx=5, pady=3)
        frame_tree = tk.LabelFrame(prawa_kolumna, text='Panel generatorów – dwuklik zmienia status', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_tree.pack(side=tk.TOP, fill=tk.BOTH, expand=False)
        tk.Label(frame_tree, text='Dane z generatorów mają H z DYD, a zapotrzebowanie pochodzi z aktywnych rekordów LOAD DATA (suma pola mw).', fg=pal['success'], bg=pal['card'], font=self.font_ui_bold_small).pack(anchor='w', padx=5, pady=(0, 6))
        frame_ops = tk.Frame(frame_tree, bg=pal['card'])
        frame_ops.pack(fill=tk.X, padx=5, pady=(0, 6))
        self.utworz_przycisk(frame_ops, 'Przełącz zaznaczone', self.przelacz_zaznaczone_w_tree, role='neutral').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Załącz zaznaczone', lambda: self.ustaw_status_zaznaczonych_w_tree(1), role='success').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Wyłącz zaznaczone', lambda: self.ustaw_status_zaznaczonych_w_tree(0), role='warning').pack(side=tk.LEFT, padx=2)
        self.utworz_przycisk(frame_ops, 'Przywróć zaznaczone', self.przywroc_zaznaczone_w_tree, role='neutral').pack(side=tk.LEFT, padx=2)
        frame_col_filters = tk.LabelFrame(frame_tree, text='Filtry kolumnowe', padx=4, pady=6, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold_small)
        frame_col_filters.pack(fill=tk.X, padx=5, pady=(0, 6))
        filter_cols = [('kod', 'Kod'), ('grupa', 'Grupa / kod'), ('stan', 'Status'), ('h', 'H min-max'), ('pmax', 'Pmax min-max'), ('ei', 'Ei min-max')]
        for i, (col_key, label) in enumerate(filter_cols):
            cell = tk.Frame(frame_col_filters, bg=pal['card'])
            cell.grid(row=0, column=i, padx=3, pady=2, sticky='ew')
            tk.Label(cell, text=label, font=self.font_ui_bold_small, bg=pal['card'], fg=pal['text']).pack(anchor='w')
            var = tk.StringVar()
            ent = tk.Entry(cell, textvariable=var, width=16, font=self.font_code_small)
            ent.pack(fill=tk.X)
            ent.bind('<KeyRelease>', self.on_column_filter_changed)
            self.column_filters[col_key] = var
            self.filter_entries[col_key] = ent
        for i in range(len(filter_cols)):
            frame_col_filters.grid_columnconfigure(i, weight=1)
        self.label_reczne = tk.Label(frame_tree, text='Brak wybranego modelu', fg=pal['accent'], bg=pal['card'], font=self.font_code)
        self.label_reczne.pack(anchor='w', padx=5, pady=5)
        tree_frame = tk.Frame(frame_tree, bg=pal['card'])
        tree_frame.pack(fill=tk.BOTH, expand=True)
        cols = ('stan', 'kod', 'grupa', 'pmax', 'h', 'ei', 'udzial', 'zmiana')
        self.tree_generatory = ttk.Treeview(tree_frame, columns=cols, show='headings', height=11, selectmode='extended')
        self.tree_generatory.heading('stan', text='Stan', command=lambda: self.sort_tree_by_column('stan'))
        self.tree_generatory.heading('kod', text='Kod generatora', command=lambda: self.sort_tree_by_column('kod'))
        self.tree_generatory.heading('grupa', text='Grupa', command=lambda: self.sort_tree_by_column('grupa'))
        self.tree_generatory.heading('pmax', text='Pmax [MW]', command=lambda: self.sort_tree_by_column('pmax'))
        self.tree_generatory.heading('h', text='H [s]', command=lambda: self.sort_tree_by_column('h'))
        self.tree_generatory.heading('ei', text='Ei [MW·s]', command=lambda: self.sort_tree_by_column('ei'))
        self.tree_generatory.heading('udzial', text='Udział [%]', command=lambda: self.sort_tree_by_column('udzial'))
        self.tree_generatory.heading('zmiana', text='Ręczna zmiana', command=lambda: self.sort_tree_by_column('zmiana'))
        self.tree_generatory.column('stan', width=70, anchor='center')
        self.tree_generatory.column('kod', width=220, anchor='w')
        self.tree_generatory.column('grupa', width=95, anchor='center')
        self.tree_generatory.column('pmax', width=90, anchor='e')
        self.tree_generatory.column('h', width=80, anchor='e')
        self.tree_generatory.column('ei', width=110, anchor='e')
        self.tree_generatory.column('udzial', width=85, anchor='e')
        self.tree_generatory.column('zmiana', width=115, anchor='center')
        scrollbar_tree_y = ttk.Scrollbar(tree_frame, orient='vertical', command=self.tree_generatory.yview)
        scrollbar_tree_x = ttk.Scrollbar(tree_frame, orient='horizontal', command=self.tree_generatory.xview)
        self.tree_generatory.configure(yscrollcommand=scrollbar_tree_y.set, xscrollcommand=scrollbar_tree_x.set)
        self.tree_generatory.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_tree_y.pack(side=tk.RIGHT, fill=tk.Y)
        scrollbar_tree_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_generatory.bind('<Double-1>', self.on_tree_double_click)
        self.bind_scroll_area(self.tree_generatory)
        frame_rank_and_group = tk.Frame(prawa_kolumna, bg=pal['bg'])
        frame_rank_and_group.pack(side=tk.TOP, fill=tk.BOTH, expand=False, pady=(10, 0))
        frame_ranking = tk.LabelFrame(frame_rank_and_group, text='Ranking TOP 10 wkładu generatorów (Pmax×H)', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_ranking.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.label_ranking = tk.Label(frame_ranking, text='Brak wybranego modelu', justify='left', anchor='w', font=self.font_code_small, bg=pal['card'], fg=pal['text'])
        self.label_ranking.pack(fill=tk.BOTH, expand=True, padx=5, pady=3)
        frame_group_chart = tk.LabelFrame(frame_rank_and_group, text='Udział grup generatorów w modelu', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_group_chart.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 0))
        self.figura_grupy, self.os_grupy = plt.subplots(figsize=(6.5, 3.2))
        self.figura_grupy.subplots_adjust(bottom=0.3, left=0.1, right=0.98, top=0.88)
        self.canvas_grupy = FigureCanvasTkAgg(self.figura_grupy, master=frame_group_chart)
        self.canvas_grupy.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.przygotuj_pusty_wykres_grup()
        frame_analiza = tk.LabelFrame(prawa_kolumna, text='Rozszerzona analiza inercyjna', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        frame_analiza.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        self.label_analiza_inercyjna = tk.Label(frame_analiza, text='Brak wybranego modelu', justify='left', anchor='w', font=self.font_code_small, fg=pal['accent'], bg=pal['card'])
        self.label_analiza_inercyjna.pack(fill=tk.X, padx=5, pady=3)
        self.pole_szczegoly = scrolledtext.ScrolledText(prawa_kolumna, width=120, height=22, wrap=tk.WORD, font=self.font_code_small)
        self.pole_szczegoly.pack(side=tk.TOP, fill=tk.BOTH, expand=True, pady=(10, 0))
        self.bind_scroll_area(self.pole_szczegoly)
        panel_wykres = tk.LabelFrame(self.main_frame, text='Wykresy dobowe', padx=8, pady=8, bg=pal['card'], fg=pal['text'], font=self.font_ui_bold)
        panel_wykres.pack(side=tk.BOTTOM, fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))
        self.figura, (self.os1, self.os2) = plt.subplots(2, 1, figsize=(12, 8), gridspec_kw={'height_ratios': [3, 1]})
        self.figura.subplots_adjust(bottom=0.14, hspace=0.35, left=0.14, right=0.9)
        self.os1_load = self.os1.twinx()
        self.canvas = FigureCanvasTkAgg(self.figura, master=panel_wykres)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.przygotuj_pusty_wykres()
        for entry in [self.entry_filtr, self.entry_wartosc_graniczna, *self.filter_entries.values()]:
            self.stylizuj_entry(entry)
        self.stylizuj_option_menu(self.option_tryb)
        self.stylizuj_option_menu(self.option_sort)
        for label in [self.label_status, self.label_maksimum_info, self.label_limit_ocena, self.label_porownanie, self.label_wskazniki, self.label_reczne, self.label_ranking, self.label_analiza_inercyjna, self.label_wynikow, self.label_title, self.label_subtitle, self.label_monitor_folder]:
            self.label_neutral.add(label)
        self.apply_theme()

    def safe_save_ui_config(self):
        try:
            self.save_ui_config()
        except Exception:
            pass

    def load_ui_config(self):
        if not self.ui_config_path.exists():
            self.ui_config = {}
            return
        try:
            with self.ui_config_path.open('r', encoding='utf-8') as f:
                self.ui_config = json.load(f)
        except Exception:
            self.ui_config = {}

    def save_ui_config(self):
        try:
            data = {'window_geometry': self.root.geometry(), 'monitor_folder_path': self.monitor_folder_path, 'threshold': self.entry_wartosc_graniczna.get().strip() if hasattr(self, 'entry_wartosc_graniczna') else '', 'main_filter': self.entry_filtr.get().strip() if hasattr(self, 'entry_filtr') else '', 'tryb_filtrowania': self.tryb_filtrowania_var.get() if hasattr(self, 'tryb_filtrowania_var') else 'Wszystkie', 'sortowanie': self.sortowanie_var.get() if hasattr(self, 'sortowanie_var') else 'Kod generatora A-Z', 'pokaz_etykiety_inercji': self.pokaz_etykiety_inercji_var.get() if hasattr(self, 'pokaz_etykiety_inercji_var') else True, 'pokaz_etykiety_zapotrzebowania': self.pokaz_etykiety_zapotrzebowania_var.get() if hasattr(self, 'pokaz_etykiety_zapotrzebowania_var') else True, 'dark_mode': self.dark_mode_var.get(), 'tree_sort_column': self.tree_sort_column, 'tree_sort_reverse': self.tree_sort_reverse, 'tree_columns': {col: self.tree_generatory.column(col, option='width') for col in self.tree_generatory['columns']} if hasattr(self, 'tree_generatory') else {}, 'column_filters': {col: var.get() for col, var in self.column_filters.items()}}
            self.ui_config_path.parent.mkdir(parents=True, exist_ok=True)
            with self.ui_config_path.open('w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def apply_ui_config(self):
        cfg = self.ui_config or {}
        try:
            geom = cfg.get('window_geometry')
            if geom:
                self.root.geometry(geom)
        except Exception:
            pass
        self.monitor_folder_path = cfg.get('monitor_folder_path', self.monitor_folder_path)
        if self.monitor_folder_path:
            self.label_monitor_folder.config(text=f'Folder: {self.monitor_folder_path}')
        threshold = cfg.get('threshold', '')
        if threshold:
            self.entry_wartosc_graniczna.delete(0, tk.END)
            self.entry_wartosc_graniczna.insert(0, threshold)
        main_filter = cfg.get('main_filter', '')
        if main_filter:
            self.entry_filtr.delete(0, tk.END)
            self.entry_filtr.insert(0, main_filter)
        self.tryb_filtrowania_var.set(cfg.get('tryb_filtrowania', 'Wszystkie'))
        self.sortowanie_var.set(cfg.get('sortowanie', 'Kod generatora A-Z'))
        self.pokaz_etykiety_inercji_var.set(cfg.get('pokaz_etykiety_inercji', True))
        self.pokaz_etykiety_zapotrzebowania_var.set(cfg.get('pokaz_etykiety_zapotrzebowania', True))
        self.dark_mode_var.set(cfg.get('dark_mode', False))
        self.tree_sort_column = cfg.get('tree_sort_column')
        self.tree_sort_reverse = cfg.get('tree_sort_reverse', False)
        widths = cfg.get('tree_columns', {})
        for col, width in widths.items():
            try:
                self.tree_generatory.column(col, width=width)
            except Exception:
                pass
        saved_filters = cfg.get('column_filters', {})
        for col, value in saved_filters.items():
            if col in self.column_filters:
                self.column_filters[col].set(value)
        self.apply_theme()

    def on_chart_label_toggle(self):
        self.aktualizuj_wykres()
        self.save_ui_config()

    def przeskaluj_wartosc_miedzy_osiami(self, wartosc, src_min, src_max, dst_min, dst_max):
        if abs(src_max - src_min) < 1e-09:
            return (dst_min + dst_max) / 2
        return dst_min + (wartosc - src_min) / (src_max - src_min) * (dst_max - dst_min)

    def wyznacz_offset_etykiety(self, indeks, wartosc, seria, *, prefer='up', other_scaled=None):
        base_x = -16 if indeks % 2 == 0 else 16
        step_variant = indeks % 3
        if prefer == 'up':
            base_y = 24 + step_variant * 8
        else:
            base_y = -(26 + step_variant * 8)
        prev_val = seria[indeks - 1] if indeks > 0 else None
        next_val = seria[indeks + 1] if indeks < len(seria) - 1 else None
        near_span = max(abs(wartosc) * 0.03, 600.0)
        if prev_val is not None and abs(wartosc - prev_val) < near_span:
            base_x += -8 if prefer == 'up' else 8
        if next_val is not None and abs(wartosc - next_val) < near_span:
            base_x += 8 if prefer == 'up' else -8
        if other_scaled is not None and abs(wartosc - other_scaled) < near_span:
            base_y += 14 if prefer == 'up' else -14
            base_x += -6 if indeks % 2 == 0 else 6
        if prev_val is not None and next_val is not None:
            if wartosc >= prev_val and wartosc >= next_val:
                base_y += 8 if prefer == 'up' else -4
            elif wartosc <= prev_val and wartosc <= next_val:
                base_y += 4 if prefer == 'up' else -8
        return (base_x, base_y)

    def dodaj_etykiete_punktu(self, ax, x, y, text, x_off, y_off, *, fg, facecolor, edgecolor, valign, fontsize=8):
        ax.annotate(text, xy=(x, y), xytext=(x_off, y_off), textcoords='offset points', ha='center', va=valign, fontsize=fontsize, fontweight='bold', color=fg, bbox=dict(boxstyle='round,pad=0.25', facecolor=facecolor, edgecolor=edgecolor, alpha=0.92), arrowprops=dict(arrowstyle='-', color=edgecolor, lw=0.9, shrinkA=4, shrinkB=4, alpha=0.85), annotation_clip=True, zorder=6)

    def konfiguruj_email_alarm(self):
        EmailAlarmConfigWindow(self.root, self.email_alarm)

    def pokaz_historie_alarmow(self):
        AlarmHistoryWindow(self.root, self.alarm_history)

    def dodaj_do_historii_alarmow(self, **kwargs):
        self.alarm_history.append(kwargs)
        if len(self.alarm_history) > 5000:
            self.alarm_history = self.alarm_history[-5000:]

    def sprawdz_i_wyslij_alarm_progu(self, model_type, key, wartosc):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            return
        if not self.email_alarm.should_trigger(prog, wartosc):
            return
        opis = self.get_model_meta(model_type, key)
        trigger_mode = self.email_alarm.config.trigger_mode
        token = (model_type, str(key), opis.get('plik', ''), opis.get('data', ''), opis.get('godzina', ''), trigger_mode, f'{prog:.6f}', f'{wartosc:.6f}')
        if token in self.sent_alarm_tokens:
            return
        self.sent_alarm_tokens.add(token)
        opis_txt = 'Przekroczenie progu' if trigger_mode == 'above' else 'Spadek poniżej progu'
        base_entry = {'czas_alarmu': datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 'trigger_mode': trigger_mode, 'data_modelu': opis.get('data', 'brak daty'), 'godzina_modelu': opis.get('godzina', '--:--'), 'wartosc': wartosc, 'prog': prog, 'source_file': opis.get('plik', ''), 'mail_attempted': False, 'mail_ok': False, 'mail_status': 'Powiadomienia wyłączone', 'opis': opis_txt}
        if not self.email_alarm.config.enabled:
            self.dodaj_do_historii_alarmow(**base_entry)
            self.aktualizuj_status(f"⚠️ Alarm spełniony dla modelu {opis.get('data', '')} {opis.get('godzina', '')}, ale powiadomienia e-mail są wyłączone")
            return
        base_entry['mail_attempted'] = True

        def worker():
            try:
                ok, msg = self.email_alarm.send_alarm(prog=prog, wartosc=wartosc, data_modelu=opis.get('data', 'brak daty'), godzina_modelu=opis.get('godzina', '--:--'), source_file=opis.get('plik', ''))
                entry = dict(base_entry)
                entry['mail_ok'] = ok
                entry['mail_status'] = msg
                self.root.after(0, lambda e=entry, m=msg, o=ok: self._finalize_alarm_send(e, m, o))
            except Exception as e:
                entry = dict(base_entry)
                entry['mail_ok'] = False
                entry['mail_status'] = str(e)
                self.root.after(0, lambda en=entry, err=e: self._finalize_alarm_send(en, f'Błąd wysyłki e-mail: {err}', False))
        threading.Thread(target=worker, daemon=True).start()

    def sprawdz_alarm_dla_modelu(self, model_type, key):
        wartosc = self.get_model_energy_now(model_type, key)
        if wartosc <= 0:
            return
        self.sprawdz_i_wyslij_alarm_progu(model_type, key, wartosc)

    def sprawdz_alarmy_dla_wszystkich_modeli(self):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            return
        for godzina in sorted(self.wszystkie_generatory.keys()):
            self.sprawdz_alarm_dla_modelu('manual', godzina)
        for auto_id in self.auto_order:
            self.sprawdz_alarm_dla_modelu('auto', auto_id)

    def _finalize_alarm_send(self, entry, status_msg, ok):
        self.dodaj_do_historii_alarmow(**entry)
        self.aktualizuj_status(f'📧 {status_msg}' if ok else f'⚠️ {status_msg}')

    def get_current_model_ref(self):
        if self.aktywny_typ_modelu == 'manual' and self.aktualna_godzina is not None:
            return ('manual', self.aktualna_godzina)
        if self.aktywny_typ_modelu == 'auto' and self.aktualny_auto_id is not None:
            return ('auto', self.aktualny_auto_id)
        return (None, None)

    def _model_mapping(self, model_type, field):
        manual_maps = {'generatory': self.wszystkie_generatory, 'statusy_oryginalne': self.statusy_oryginalne, 'opis': self.opis_modeli, 'inercja': self.inercja_na_godzine, 'inercja_oryginalna': self.inercja_oryginalna_na_godzine, 'liczba_zalaczonych': self.liczba_zalaczonych_na_godzine, 'zapotrzebowanie': self.zapotrzebowanie_na_godzine}
        defaults = {'generatory': [], 'statusy_oryginalne': {}, 'opis': {}, 'inercja': 0.0, 'inercja_oryginalna': 0.0, 'liczba_zalaczonych': 0, 'zapotrzebowanie': 0.0}
        return (manual_maps.get(field) if model_type == 'manual' else None, defaults[field])

    def get_model_value(self, model_type, key, field):
        mapping, default = self._model_mapping(model_type, field)
        if mapping is not None:
            return mapping.get(key, default)
        if model_type == 'auto':
            return self.modele_auto.get(key, {}).get(field, default)
        return default

    def set_model_value(self, model_type, key, field, value):
        mapping, _ = self._model_mapping(model_type, field)
        if mapping is not None:
            mapping[key] = value
        elif model_type == 'auto' and key in self.modele_auto:
            self.modele_auto[key][field] = value

    def get_model_generators(self, model_type, key):
        return self.get_model_value(model_type, key, 'generatory')

    def get_model_original_statuses(self, model_type, key):
        return self.get_model_value(model_type, key, 'statusy_oryginalne')

    def get_model_meta(self, model_type, key):
        return self.get_model_value(model_type, key, 'opis')

    def get_model_energy_now(self, model_type, key):
        return self.get_model_value(model_type, key, 'inercja')

    def get_model_energy_original(self, model_type, key):
        return self.get_model_value(model_type, key, 'inercja_oryginalna')

    def get_model_demand_now(self, model_type, key):
        return self.get_model_value(model_type, key, 'zapotrzebowanie')

    def set_model_energy_now(self, model_type, key, value):
        self.set_model_value(model_type, key, 'inercja', value)

    def set_model_energy_original(self, model_type, key, value):
        self.set_model_value(model_type, key, 'inercja_oryginalna', value)

    def set_model_connected_count(self, model_type, key, value):
        self.set_model_value(model_type, key, 'liczba_zalaczonych', value)

    def is_model_epc_exportable(self, model_type, key):
        meta = self.get_model_meta(model_type, key)
        sciezka = meta.get('sciezka', '')
        return bool(sciezka and os.path.exists(sciezka))

    def wybierz_model_manual(self, event=None):
        selection = self.listbox_modele.curselection()
        if not selection:
            return
        try:
            linia = self.listbox_modele.get(selection[0])
            match = re.match('(?:\\[MIN\\]\\s*)?(?:\\[MAX\\]\\s*)?(\\d{2}):30', linia)
            if not match:
                return
            godzina = int(match.group(1))
            self.aktywny_typ_modelu = 'manual'
            self.aktualna_godzina = godzina
            self.aktualny_auto_id = None
            self.listbox_modele_auto.selection_clear(0, tk.END)
            if not self.obliczenia_wykonane:
                self.pole_szczegoly.delete('1.0', tk.END)
                self.pole_szczegoly.insert(tk.END, '⚠️ Najpierw: 1. EPC → 2. DYD → 3. OBLICZ')
                self.przygotuj_pusty_wykres_grup('Brak obliczeń dla modelu')
                return
            self.odswiez_aktywny_model()
        except Exception:
            pass

    def wybierz_model_auto(self, event=None):
        selection = self.listbox_modele_auto.curselection()
        if not selection:
            return
        try:
            idx = selection[0]
            if idx < 0 or idx >= len(self.auto_order):
                return
            auto_id = self.auto_order[idx]
            self.aktywny_typ_modelu = 'auto'
            self.aktualny_auto_id = auto_id
            self.aktualna_godzina = None
            self.listbox_modele.selection_clear(0, tk.END)
            self.odswiez_aktywny_model()
        except Exception:
            pass

    def wyciagnij_date_i_godzine_z_nazwy(self, sciezka):
        nazwa = os.path.basename(sciezka)
        stem = Path(nazwa).stem
        nazwa_lower = stem.lower()
        data_txt = ''
        matched_span = None
        date_patterns = ['(?<!\\d)(?P<y>(?:20\\d{2}|19\\d{2}))[\\s,._-](?P<m>0?[1-9]|1[0-2])[\\s,._-](?P<d>0?[1-9]|[12]\\d|3[01])(?!\\d)', '(?<!\\d)(?P<d>0?[1-9]|[12]\\d|3[01])[\\s,._-](?P<m>0?[1-9]|1[0-2])[\\s,._-](?P<y>(?:20\\d{2}|19\\d{2}))(?!\\d)', '(?<!\\d)(?P<y>(?:20\\d{2}|19\\d{2}))(?P<m>0[1-9]|1[0-2])(?P<d>0[1-9]|[12]\\d|3[01])(?!\\d)', '(?<!\\d)(?P<d>0[1-9]|[12]\\d|3[01])(?P<m>0[1-9]|1[0-2])(?P<y>(?:20\\d{2}|19\\d{2}))(?!\\d)']
        candidates = []
        for wz in date_patterns:
            for m in re.finditer(wz, nazwa_lower):
                try:
                    d = int(m.group('d'))
                    mm = int(m.group('m'))
                    yyyy = int(m.group('y'))
                    candidates.append((m.start(), m.end(), f'{d:02d}.{mm:02d}.{yyyy:04d}'))
                except Exception:
                    continue
        if candidates:
            candidates.sort(key=lambda x: (x[0], -(x[1] - x[0])))
            start_idx, end_idx, data_txt = candidates[0]
            matched_span = (start_idx, end_idx)
        scope_for_hour = nazwa_lower
        if matched_span:
            a, b = matched_span
            scope_for_hour = nazwa_lower[:a] + ' ' + nazwa_lower[b:]
        godzina = None
        time_patterns = ['(?<!\\d)([01]?\\d|2[0-3])[\\s,._:-]?30(?!\\d)', '(?<!\\d)([01]?\\d|2[0-3])30(?!\\d)']
        for wz in time_patterns:
            matches = list(re.finditer(wz, scope_for_hour))
            if not matches:
                continue
            for match in reversed(matches):
                try:
                    hh = int(match.group(1))
                except Exception:
                    continue
                if 0 <= hh <= 23:
                    godzina = hh
                    break
            if godzina is not None:
                break
        godz_txt = f'{godzina:02d}:30' if godzina is not None else ''
        return (data_txt, godz_txt, godzina)

    def parsuj_linie_generatora(self, line):
        gen = self.parsuj_rekord_generatora(line)
        return [gen] if gen else []

    def znajdz_kod_grupy_generatora(self, full_record):
        try:
            if '/' not in full_record:
                return None
            after_slash = full_record.split('/', 1)[1]
            tokens = re.findall('[-+]?\\d+(?:[.,]\\d+)?', after_slash)
            int_tokens = []
            for t in tokens:
                if '.' in t or ',' in t:
                    continue
                try:
                    int_tokens.append(int(t))
                except Exception:
                    pass
            for value in reversed(int_tokens):
                if value in self.GENERATOR_GROUP_LABELS:
                    return value
        except Exception:
            pass
        return None

    def parsuj_rekord_generatora(self, full_record):
        try:
            if ':' not in full_record:
                return None
            lewa, prawa = full_record.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            if not quoted:
                return None
            reg_name = quoted[-1].strip()
            if not reg_name or reg_name.isdigit():
                return None
            czesci_prawe = [p.strip().strip('"') for p in prawa.split() if p.strip()]
            status = int(float(czesci_prawe[0].replace(',', '.'))) if czesci_prawe else 0
            pmax = 0.0
            if len(czesci_prawe) > 10:
                pmax = self.konwertuj_float_soft(czesci_prawe[10], 0.0)
            else:
                for w in reversed(czesci_prawe):
                    test = self.konwertuj_float_soft(w, 0.0)
                    if test > 0:
                        pmax = test
                        break
            group_code = self.znajdz_kod_grupy_generatora(full_record)
            if pmax > 0:
                return {'reg_name': reg_name, 'status': status, 'pmax': pmax, 'group_code': group_code}
        except Exception:
            pass
        return None

    def parsuj_rekord_loada(self, full_record):
        try:
            if ':' not in full_record:
                return None
            lewa, prawa = full_record.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            load_name = quoted[-1].strip() if quoted else ''
            czesci = [p.strip().strip('"') for p in prawa.split() if p.strip()]
            if len(czesci) < 8:
                return None
            status = int(float(czesci[0].replace(',', '.')))
            mw = float(czesci[1].replace(',', '.'))
            ar = int(float(czesci[7]))
            if status == 1 and ar == 1:
                mw_val = mw
            else:
                mw_val = 0.0
            return {'load_name': load_name, 'status': status, 'mw': mw_val, 'mw_raw': mw, 'area': ar}
        except Exception:
            return None

    def wczytaj_model_epc(self, sciezka):
        gen_data = []
        zapotrzebowanie_mw = 0.0
        current_section = None
        pending_record = None

        def flush_pending(section_name, record_text):
            nonlocal zapotrzebowanie_mw
            if not record_text:
                return
            if section_name == 'generator':
                gen = self.parsuj_rekord_generatora(record_text)
                if gen and gen['reg_name'].strip() and (gen['pmax'] > 0):
                    gen_data.append(gen)
            elif section_name == 'load':
                load = self.parsuj_rekord_loada(record_text)
                if load:
                    zapotrzebowanie_mw += load['mw']
        try:
            with open(sciezka, 'r', encoding='utf-8', errors='ignore') as f:
                for raw_line in f:
                    line = raw_line.strip()
                    if not line:
                        continue
                    line_lower = line.lower()
                    if 'generator data' in line_lower:
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = 'generator'
                        continue
                    if 'load data' in line_lower:
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = 'load'
                        continue
                    if current_section in {'generator', 'load'} and ':' not in line and ('data' in line_lower):
                        flush_pending(current_section, pending_record)
                        pending_record = None
                        current_section = None
                        continue
                    if current_section not in {'generator', 'load'}:
                        continue
                    if ':' in line:
                        flush_pending(current_section, pending_record)
                        pending_record = line
                    elif pending_record:
                        pending_record += ' ' + line
                flush_pending(current_section, pending_record)
        except Exception as e:
            print(f'Błąd EPC {sciezka}: {e}')
        return {'generatory': gen_data, 'zapotrzebowanie_mw': zapotrzebowanie_mw}

    def wczytaj_generatory_z_pliku(self, sciezka):
        return self.wczytaj_model_epc(sciezka)['generatory']

    def wczytaj_wiele_plikow_epc(self):
        sciezki = filedialog.askopenfilenames(title='Wybierz pliki EPC', filetypes=[('*.epc', '*.epc'), ('*.*', '*.*')])
        if not sciezki:
            return
        self.wszystkie_generatory = {}
        self.statusy_oryginalne = {}
        self.opis_modeli = {}
        self.inercja_na_godzine = {h: 0.0 for h in range(24)}
        self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
        self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
        self.zapotrzebowanie_na_godzine = {h: 0.0 for h in range(24)}
        self.obliczenia_wykonane = False
        self.aktywny_typ_modelu = None
        self.aktualna_godzina = None
        self.aktualny_auto_id = None
        self.reset_alarm_cache()
        self.listbox_modele.delete(0, tk.END)
        self.pole_szczegoly.delete('1.0', tk.END)
        self.tree_generatory.delete(*self.tree_generatory.get_children())
        self.progress_var.set(0)
        self.label_progress.config(text='0%')
        self.aktualizuj_status('Wczytuję EPC...')
        przetworzone = 0
        suma = len(sciezki)
        for i, sciezka in enumerate(sciezki, start=1):
            data_txt, godz_txt, godzina = self.wyciagnij_date_i_godzine_z_nazwy(sciezka)
            if godzina is not None:
                model_epc = self.wczytaj_model_epc(sciezka)
                gen_data = model_epc['generatory']
                zapotrzebowanie_mw = model_epc['zapotrzebowanie_mw']
                if gen_data:
                    self.wszystkie_generatory[godzina] = gen_data
                    self.statusy_oryginalne[godzina] = {g['reg_name']: g['status'] for g in gen_data}
                    self.opis_modeli[godzina] = {'data': data_txt if data_txt else 'brak daty', 'godzina': godz_txt if godz_txt else f'{godzina:02d}:30', 'plik': os.path.basename(sciezka), 'sciezka': sciezka}
                    self.zapotrzebowanie_na_godzine[godzina] = zapotrzebowanie_mw
                    przetworzone += 1
            proc = i / suma * 100
            self.progress_var.set(proc)
            self.label_progress.config(text=f'{proc:.0f}% ({i}/{suma})')
            self.aktualizuj_status(f'Wczytywanie EPC... {i}/{suma}')
            self.root.update_idletasks()
        self.odswiez_liste_modeli()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.przygotuj_pusty_wykres_grup()
        self.aktualizuj_status(f'✅ EPC: {przetworzone} modeli. Teraz wczytaj DYD PSLF!')

    def wczytaj_plik_genrou_H(self):
        sciezka = filedialog.askopenfilename(title='Wybierz plik DYD PSLF (.dyd)', filetypes=[('DYD PSLF', '*.dyd'), ('Tekst', '*.txt *.raw *.dyr *.dat'), ('Wszystkie', '*.*')])
        if not sciezka:
            return
        self.aktualizuj_status('Wczytuję współczynniki H...')
        H_map = {}
        current_block = []
        current_name = None
        try:
            with open(sciezka, 'r', encoding='utf-8', errors='ignore') as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    if re.match('GENROU|GENCLS|GENSAL', line.upper()):
                        if current_block and current_name:
                            H = self.parsuj_H_z_bloku_genrou(current_block)
                            if H:
                                H_map[current_name] = H
                        current_name = self.wyciagnij_nazwe_z_genrou(line)
                        current_block = [line]
                        continue
                    if current_name:
                        current_block.append(line)
            if current_block and current_name:
                H = self.parsuj_H_z_bloku_genrou(current_block)
                if H:
                    H_map[current_name] = H
        except Exception as e:
            messagebox.showerror('Błąd DYD', f'Błąd: {e}')
            return
        self.H_generatory = H_map
        self.aktualizuj_status('✅ Wczytano współczynniki H')
        for godzina in self.wszystkie_generatory.keys():
            self.przelicz_inercje_dla_modelu('manual', godzina)
        for auto_id in self.auto_order:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()

    def parsuj_H_z_bloku_genrou(self, block_lines):
        try:
            full_text = ' '.join(block_lines)
            match = re.search('/\\s*([0-9.-]+)', full_text)
            if match:
                H = self.konwertuj_float_soft(match.group(1), 0.0)
                if 0 < H < 20:
                    return H
            for line in block_lines:
                if ':' in line:
                    tokens = re.findall('\\d+\\.?\\d*', line)
                    for t in tokens:
                        H = self.konwertuj_float_soft(t, 0.0)
                        if 0 < H < 20:
                            return H
            return None
        except Exception:
            return None

    def wyciagnij_nazwe_z_genrou(self, line):
        m = re.search('"([^"]+)"', line)
        return m.group(1).strip() if m else None

    def policz_inercje_dla_listy(self, model_type, key, gen_list, use_original=False):
        energia = 0.0
        liczba = 0
        statusy_org = self.get_model_original_statuses(model_type, key)
        for gen in gen_list:
            if gen['reg_name'] not in self.H_generatory:
                continue
            status = statusy_org.get(gen['reg_name'], gen['status']) if use_original else gen['status']
            if status == 1:
                energia += gen['pmax'] * self.H_generatory[gen['reg_name']]
                liczba += 1
        return (energia, liczba)

    def przelicz_inercje_dla_modelu(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return
        energia, liczba = self.policz_inercje_dla_listy(model_type, key, gen_list, use_original=False)
        energia_org, _ = self.policz_inercje_dla_listy(model_type, key, gen_list, use_original=True)
        self.set_model_energy_now(model_type, key, energia)
        self.set_model_energy_original(model_type, key, energia_org)
        self.set_model_connected_count(model_type, key, liczba)

    def oblicz_inercje_wszystkie(self):
        if not self.wszystkie_generatory:
            messagebox.showwarning('!', 'Wczytaj EPC')
            return
        if not self.H_generatory:
            messagebox.showwarning('!', 'Wczytaj DYD PSLF')
            return
        self.aktualizuj_status('Obliczanie E_sys = Σ Pmax_i × H_i ...')
        for godzina in self.wszystkie_generatory.keys():
            self.przelicz_inercje_dla_modelu('manual', godzina)
        for auto_id in self.auto_order:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
        self.obliczenia_wykonane = True
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        if self.aktywny_typ_modelu is not None:
            self.odswiez_aktywny_model()
        else:
            self.przygotuj_pusty_wykres_grup()
        self.sprawdz_alarmy_dla_wszystkich_modeli()
        self.aktualizuj_status('✅ GOTOWE! Kliknij modele po szczegóły.')
        messagebox.showinfo('✅', f'Obliczono dla {len(self.wszystkie_generatory)} godzin oraz {len(self.auto_order)} modeli automatycznych.')

    def przelicz_aktywny_model(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        if not self.H_generatory:
            messagebox.showwarning('Brak DYD', 'Najpierw wczytaj plik DYD z H.')
            return
        self.przelicz_inercje_dla_modelu(model_type, key)
        if model_type == 'manual':
            self.obliczenia_wykonane = True
        self.odswiez_aktywny_model()
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.aktualizuj_status('✅ Przeliczono wybrany model')

    def przelicz_i_odswiez_po_zmianie(self, model_type, key):
        self.przelicz_inercje_dla_modelu(model_type, key)
        if model_type == 'manual':
            self.obliczenia_wykonane = True
        self.odswiez_aktywny_model()
        self.odswiez_liste_modeli()
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.aktualizuj_status('✅ Przeliczono model po zmianie statusów')

    def wybierz_folder_monitoringu(self):
        folder = filedialog.askdirectory(title='Wybierz folder z modelami EPC')
        if not folder:
            return
        self.monitor_folder_path = folder
        self.label_monitor_folder.config(text=f'Folder: {folder}')
        self.aktualizuj_status('Wybrano folder monitoringu EPC')
        self.save_ui_config()

    def toggle_monitor_folder(self):
        if self.monitor_folder_var.get():
            if not self.monitor_folder_path:
                messagebox.showwarning('Brak folderu', 'Najpierw wybierz folder do monitorowania.')
                self.monitor_folder_var.set(False)
                return
            self.aktualizuj_status('Monitoring folderu włączony')
            self.sprawdz_folder_monitoringu()
        else:
            self.aktualizuj_status('Monitoring folderu wyłączony')

    def pobierz_wszystkie_pliki_epc_w_folderze(self):
        if not self.monitor_folder_path or not os.path.isdir(self.monitor_folder_path):
            return []
        pliki = []
        for nazwa in os.listdir(self.monitor_folder_path):
            if nazwa.lower().endswith('.epc'):
                pelna = os.path.join(self.monitor_folder_path, nazwa)
                try:
                    pliki.append((os.path.getmtime(pelna), pelna))
                except Exception:
                    pass
        pliki.sort()
        return pliki

    def sprawdz_folder_monitoringu(self):
        if not self.monitor_folder_var.get():
            return
        pliki = self.pobierz_wszystkie_pliki_epc_w_folderze()
        for mtime, sciezka in pliki:
            last_seen_mtime = self.auto_seen_files.get(sciezka)
            if last_seen_mtime is None or mtime > last_seen_mtime:
                self.auto_seen_files[sciezka] = mtime
                self.wczytaj_model_auto_z_folderu(sciezka)
        self.root.after(self.monitor_interval_ms, self.sprawdz_folder_monitoringu)

    def wczytaj_model_auto_z_folderu(self, sciezka):
        model_epc = self.wczytaj_model_epc(sciezka)
        gen_data = model_epc['generatory']
        zapotrzebowanie_mw = model_epc['zapotrzebowanie_mw']
        if not gen_data:
            return
        data_txt, godz_txt, godzina_idx = self.wyciagnij_date_i_godzine_z_nazwy(sciezka)
        self.auto_counter += 1
        auto_id = f'auto_{self.auto_counter:05d}'
        statusy_org = {g['reg_name']: g['status'] for g in gen_data}
        self.modele_auto[auto_id] = {'generatory': gen_data, 'statusy_oryginalne': statusy_org, 'opis': {'data': data_txt if data_txt else 'brak daty', 'godzina': godz_txt if godz_txt else '--:--', 'plik': os.path.basename(sciezka), 'sciezka': sciezka, 'czas_wykrycia': datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 'godzina_idx': godzina_idx}, 'inercja': 0.0, 'inercja_oryginalna': 0.0, 'liczba_zalaczonych': 0, 'zapotrzebowanie': zapotrzebowanie_mw}
        self.auto_order.append(auto_id)
        if self.H_generatory:
            self.przelicz_inercje_dla_modelu('auto', auto_id)
            self.sprawdz_alarm_dla_modelu('auto', auto_id)
        self.odswiez_liste_modeli_auto()
        self.aktualizuj_status(f'✅ Wykryto nowy lub zaktualizowany model w folderze i dodano do listy automatycznej: {os.path.basename(sciezka)}')

    def czy_status_zmieniony_recznie(self, model_type, key, reg_name, status_biezacy):
        statusy_org = self.get_model_original_statuses(model_type, key)
        status_org = statusy_org.get(reg_name)
        return status_org is not None and status_org != status_biezacy

    def pobierz_widoczne_generatory(self, model_type, key):
        wszystkie_gen = self.get_model_generators(model_type, key)
        gen_z_H = [g for g in wszystkie_gen if g['reg_name'] in self.H_generatory]
        return self.filtrowanie_generatory(model_type, key, gen_z_H)

    def parse_range_filter(self, text):
        txt = text.strip().replace(',', '.')
        if not txt:
            return None
        m = re.match('^\\s*(-?\\d+(?:\\.\\d+)?)\\s*-\\s*(-?\\d+(?:\\.\\d+)?)\\s*$', txt)
        if m:
            a, b = (float(m.group(1)), float(m.group(2)))
            return ('range', min(a, b), max(a, b))
        m = re.match('^\\s*(>=|<=|>|<|=)\\s*(-?\\d+(?:\\.\\d+)?)\\s*$', txt)
        if m:
            return ('cmp', m.group(1), float(m.group(2)))
        try:
            val = float(txt)
            return ('cmp', '=', val)
        except Exception:
            return ('text', txt.lower())

    def match_numeric_filter(self, value, flt):
        if not flt:
            return True
        if flt[0] == 'range':
            return flt[1] <= value <= flt[2]
        if flt[0] == 'cmp':
            op, ref = (flt[1], flt[2])
            if op == '>':
                return value > ref
            if op == '<':
                return value < ref
            if op == '>=':
                return value >= ref
            if op == '<=':
                return value <= ref
            return abs(value - ref) < 1e-09
        return True

    def passes_column_filters(self, gen, model_type, key):
        H = self.H_generatory.get(gen['reg_name'], 0.0)
        Ei = gen['pmax'] * H
        grupa_txt = self.formatuj_grupe_generatora(gen).lower()
        stan_txt = 'zał' if gen['status'] == 1 else 'wył'
        f_kod = self.column_filters['kod'].get().strip().lower()
        if f_kod and f_kod not in gen['reg_name'].lower():
            return False
        f_grupa = self.column_filters['grupa'].get().strip().lower()
        if f_grupa and f_grupa not in grupa_txt:
            return False
        f_stan = self.column_filters['stan'].get().strip().lower()
        if f_stan:
            mapy = {'1': 'zał', 'zal': 'zał', 'zał': 'zał', '0': 'wył', 'wyl': 'wył', 'wył': 'wył'}
            wanted = mapy.get(f_stan, f_stan)
            if wanted not in stan_txt:
                return False
        for field_key, value in [('h', H), ('pmax', gen['pmax']), ('ei', Ei)]:
            raw = self.column_filters[field_key].get().strip()
            if raw:
                flt = self.parse_range_filter(raw)
                if flt and flt[0] in {'range', 'cmp'}:
                    if not self.match_numeric_filter(value, flt):
                        return False
                elif flt and flt[0] == 'text':
                    if flt[1] not in f'{value}'.lower():
                        return False
        return True

    def filtrowanie_generatory(self, model_type, key, gen_data):
        tekst = self.entry_filtr.get().strip().lower()
        self.tekst_filtrowania = tekst
        frazy = [x.strip() for x in tekst.split(',') if x.strip()]
        wynik = list(gen_data)
        if frazy:
            wynik = [g for g in wynik if any((f in g['reg_name'].lower() for f in frazy))]
        tryb = self.tryb_filtrowania_var.get()
        if tryb == 'Tylko zmienione ręcznie':
            wynik = [g for g in wynik if self.czy_status_zmieniony_recznie(model_type, key, g['reg_name'], g['status'])]
        elif tryb == 'Tylko załączone':
            wynik = [g for g in wynik if g['status'] == 1]
        elif tryb == 'Tylko wyłączone':
            wynik = [g for g in wynik if g['status'] == 0]
        wynik = [g for g in wynik if self.passes_column_filters(g, model_type, key)]

        def energia(g):
            return g['pmax'] * self.H_generatory.get(g['reg_name'], 0.0)
        if self.tree_sort_column:
            col = self.tree_sort_column
            rev = self.tree_sort_reverse
            if col == 'stan':
                wynik.sort(key=lambda g: g['status'], reverse=rev)
            elif col == 'kod':
                wynik.sort(key=lambda g: g['reg_name'], reverse=rev)
            elif col == 'grupa':
                wynik.sort(key=lambda g: self.formatuj_grupe_generatora(g), reverse=rev)
            elif col == 'pmax':
                wynik.sort(key=lambda g: g['pmax'], reverse=rev)
            elif col == 'h':
                wynik.sort(key=lambda g: self.H_generatory.get(g['reg_name'], 0.0), reverse=rev)
            elif col == 'ei':
                wynik.sort(key=energia, reverse=rev)
            elif col == 'udzial':
                e_total = max(self.get_model_energy_now(model_type, key), 0.0)
                wynik.sort(key=lambda g: energia(g) / e_total * 100 if e_total > 0 else 0.0, reverse=rev)
            elif col == 'zmiana':
                wynik.sort(key=lambda g: self.czy_status_zmieniony_recznie(model_type, key, g['reg_name'], g['status']), reverse=rev)
            return wynik
        sort_mode = self.sortowanie_var.get()
        if sort_mode == 'Kod generatora A-Z':
            wynik.sort(key=lambda g: g['reg_name'])
        elif sort_mode == 'Kod generatora Z-A':
            wynik.sort(key=lambda g: g['reg_name'], reverse=True)
        elif sort_mode == 'Pmax malejąco':
            wynik.sort(key=lambda g: g['pmax'], reverse=True)
        elif sort_mode == 'Pmax rosnąco':
            wynik.sort(key=lambda g: g['pmax'])
        elif sort_mode == 'H malejąco':
            wynik.sort(key=lambda g: self.H_generatory.get(g['reg_name'], 0.0), reverse=True)
        elif sort_mode == 'H rosnąco':
            wynik.sort(key=lambda g: self.H_generatory.get(g['reg_name'], 0.0))
        elif sort_mode == 'Ei malejąco':
            wynik.sort(key=energia, reverse=True)
        elif sort_mode == 'Ei rosnąco':
            wynik.sort(key=energia)
        elif sort_mode == 'Status':
            wynik.sort(key=lambda g: (g['status'], g['reg_name']), reverse=True)
        return wynik

    def zastosuj_filtr(self, event=None):
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def on_column_filter_changed(self, event=None):
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def sort_tree_by_column(self, col):
        if self.tree_sort_column == col:
            self.tree_sort_reverse = not self.tree_sort_reverse
        else:
            self.tree_sort_column = col
            self.tree_sort_reverse = False
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def on_dropdown_sort_changed(self):
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def wyczysc_filtr(self):
        self.entry_filtr.delete(0, tk.END)
        self.tryb_filtrowania_var.set('Wszystkie')
        self.sortowanie_var.set('Kod generatora A-Z')
        self.tree_sort_column = None
        self.tree_sort_reverse = False
        for var in self.column_filters.values():
            var.set('')
        self.save_ui_config()
        self.odswiez_aktywny_model()

    def pokaz_wszystkie(self):
        self.wyczysc_filtr()

    def odswiez_aktywny_model(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            return
        self.odswiez_panel_generatorow(model_type, key)
        self.pokaz_szczegoly_modelu(model_type, key)
        self.aktualizuj_porownanie_modelu(model_type, key)
        self.aktualizuj_ranking(model_type, key)
        self.aktualizuj_wykres_grup_generatorow(model_type, key)
        self.aktualizuj_analize_inercyjna(model_type, key)

    def formatuj_grupe_generatora(self, gen):
        code = gen.get('group_code')
        if code is None:
            return '-'
        label = self.GENERATOR_GROUP_LABELS.get(code, f'KOD {code}')
        return f'{code} ({label})'

    def odswiez_panel_generatorow(self, model_type, key):
        self.tree_generatory.delete(*self.tree_generatory.get_children())
        gen_data = self.pobierz_widoczne_generatory(model_type, key)
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        typ_txt = '24h' if model_type == 'manual' else 'AUTO'
        demand_txt = self.get_model_demand_now(model_type, key)
        self.label_reczne.config(text=f'Model [{typ_txt}] {data_txt} {godz_txt} | Widoczne: {len(gen_data)} | Load: {demand_txt:,.2f} MW')
        e_total = max(self.get_model_energy_now(model_type, key), 0.0)
        for gen in gen_data:
            H = self.H_generatory.get(gen['reg_name'], 0.0)
            Ei = gen['pmax'] * H
            udzial = Ei / e_total * 100 if e_total > 0 else 0.0
            zmieniony = self.czy_status_zmieniony_recznie(model_type, key, gen['reg_name'], gen['status'])
            stan = 'ZAŁ' if gen['status'] == 1 else 'WYŁ'
            tag = 'manual' if zmieniony else 'zal' if gen['status'] == 1 else 'wyl'
            self.tree_generatory.insert('', tk.END, iid=gen['reg_name'], values=(stan, gen['reg_name'], self.formatuj_grupe_generatora(gen), f"{gen['pmax']:.3f}", f'{H:.3f}', f'{Ei:.2f}', f'{udzial:.2f}', 'TAK' if zmieniony else '-'), tags=(tag,))

    def get_generator_tooltip(self, item_id):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            return ''
        gen_list = self.get_model_generators(model_type, key)
        e_total = max(self.get_model_energy_now(model_type, key), 0.0)
        for gen in gen_list:
            if gen['reg_name'] == item_id:
                H = self.H_generatory.get(gen['reg_name'], 0.0)
                Ei = gen['pmax'] * H
                udzial = Ei / e_total * 100 if e_total > 0 else 0.0
                grupa = self.formatuj_grupe_generatora(gen)
                org = self.get_model_original_statuses(model_type, key).get(gen['reg_name'], gen['status'])
                org_txt = 'ZAŁ' if org == 1 else 'WYŁ'
                now_txt = 'ZAŁ' if gen['status'] == 1 else 'WYŁ'
                return f"Generator: {gen['reg_name']}\nPełna nazwa: {gen['reg_name']}\nStatus oryginalny: {org_txt}\nStatus bieżący: {now_txt}\nTyp grupy: {grupa}\nH: {H:.3f} s\nPmax: {gen['pmax']:.3f} MW\nEi: {Ei:.2f} MW·s\nUdział: {udzial:.2f}%"
        return ''

    def on_tree_double_click(self, event):
        item = self.tree_generatory.identify_row(event.y)
        model_type, key = self.get_current_model_ref()
        if not item or model_type is None:
            return
        gen_list = self.get_model_generators(model_type, key)
        for gen in gen_list:
            if gen['reg_name'] == item:
                gen['status'] = 0 if gen['status'] == 1 else 1
                break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przelacz_zaznaczone_w_tree(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa:
                    gen['status'] = 0 if gen['status'] == 1 else 1
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def ustaw_status_zaznaczonych_w_tree(self, nowy_status):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa:
                    gen['status'] = nowy_status
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przywroc_zaznaczone_w_tree(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        zaznaczone = self.tree_generatory.selection()
        if not zaznaczone:
            messagebox.showwarning('Brak zaznaczenia', 'Zaznacz co najmniej jeden generator.')
            return
        oryg = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        for nazwa in zaznaczone:
            for gen in gen_list:
                if gen['reg_name'] == nazwa and nazwa in oryg:
                    gen['status'] = oryg[nazwa]
                    break
        self.przelicz_i_odswiez_po_zmianie(model_type, key)

    def przywroc_oryginalne_statusy_modelu(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        oryginalne = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        liczba_zmian = 0
        for gen in gen_list:
            nazwa = gen['reg_name']
            if nazwa in oryginalne and gen['status'] != oryginalne[nazwa]:
                gen['status'] = oryginalne[nazwa]
                liczba_zmian += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        self.aktualizuj_status(f'✅ Przywrócono oryginalne statusy modelu ({liczba_zmian} zmian)')

    def przywroc_wszystkie_modele_do_oryginalu(self):
        if not self.wszystkie_generatory or not self.statusy_oryginalne:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj modele EPC.')
            return
        lacznie_zmian = 0
        for godzina, gen_list in self.wszystkie_generatory.items():
            oryginalne = self.statusy_oryginalne.get(godzina, {})
            for gen in gen_list:
                nazwa = gen['reg_name']
                if nazwa in oryginalne and gen['status'] != oryginalne[nazwa]:
                    gen['status'] = oryginalne[nazwa]
                    lacznie_zmian += 1
            self.przelicz_inercje_dla_modelu('manual', godzina)
        self.aktualizuj_wykres()
        self.aktualizuj_panel_maksimum(check_alarm=False)
        self.aktualizuj_wskazniki_jakosci()
        self.odswiez_liste_modeli()
        if self.aktywny_typ_modelu is not None:
            self.odswiez_aktywny_model()
        else:
            self.przygotuj_pusty_wykres_grup()
        self.aktualizuj_status(f'✅ Przywrócono wszystkie modele 24h do statusów domyślnych ({lacznie_zmian} zmian)')
        messagebox.showinfo('Przywracanie zakończone', f'Przywrócono statusy domyślne we wszystkich modelach 24h.\nLiczba zmian: {lacznie_zmian}')

    def zalacz_widoczne_z_filtra(self):
        self.ustaw_status_widocznych(1)

    def wylacz_widoczne_z_filtra(self):
        self.ustaw_status_widocznych(0)

    def przywroc_widoczne_z_filtra(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        widoczne = {g['reg_name'] for g in self.pobierz_widoczne_generatory(model_type, key)}
        oryginalne = self.get_model_original_statuses(model_type, key)
        gen_list = self.get_model_generators(model_type, key)
        liczba = 0
        for gen in gen_list:
            if gen['reg_name'] in widoczne and gen['reg_name'] in oryginalne:
                if gen['status'] != oryginalne[gen['reg_name']]:
                    gen['status'] = oryginalne[gen['reg_name']]
                    liczba += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        self.aktualizuj_status(f'✅ Przywrócono widoczne generatory do oryginału ({liczba} zmian)')

    def ustaw_status_widocznych(self, nowy_status):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        widoczne = {g['reg_name'] for g in self.pobierz_widoczne_generatory(model_type, key)}
        gen_list = self.get_model_generators(model_type, key)
        liczba = 0
        for gen in gen_list:
            if gen['reg_name'] in widoczne and gen['status'] != nowy_status:
                gen['status'] = nowy_status
                liczba += 1
        self.przelicz_i_odswiez_po_zmianie(model_type, key)
        txt = 'załączono' if nowy_status == 1 else 'wyłączono'
        self.aktualizuj_status(f'✅ {txt.capitalize()} widoczne generatory ({liczba} zmian)')

    def pobierz_dane_szczegolowe(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return None
        gen_z_H = [g for g in gen_list if g['reg_name'] in self.H_generatory]
        gen_data = self.filtrowanie_generatory(model_type, key, gen_z_H)
        energia_calkowita = self.get_model_energy_now(model_type, key)
        zapotrzebowanie = self.get_model_demand_now(model_type, key)
        wiersze = []
        wlaczonych = 0
        energia_wlaczonych = 0.0
        energia_filtrowanych = 0.0
        energia_wylaczonych = 0.0
        for i, gen in enumerate(gen_data, 1):
            pmax = gen['pmax']
            H = self.H_generatory[gen['reg_name']]
            Ei = pmax * H
            energia_filtrowanych += Ei
            uwzgledniony = 'TAK' if gen['status'] == 1 else 'NIE'
            status_tekst = 'ZAŁ.' if gen['status'] == 1 else 'WYŁ.'
            status_org = self.get_model_original_statuses(model_type, key).get(gen['reg_name'], gen['status'])
            reczna_zmiana = status_org != gen['status']
            if gen['status'] == 1:
                wlaczonych += 1
                energia_wlaczonych += Ei
            else:
                energia_wylaczonych += Ei
            wiersze.append({'lp': i, 'kod_generatora': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status': status_tekst, 'pmax': pmax, 'H': H, 'Ei': Ei, 'uwzgledniony': uwzgledniony, 'reczna_zmiana': reczna_zmiana, 'status_oryginalny': status_org})
        podsumowanie = {'liczba_gen_odfiltrowanych': len(gen_data), 'liczba_wlaczonych': wlaczonych, 'energia_odfiltrowanych': energia_filtrowanych, 'energia_wlaczonych': energia_wlaczonych, 'energia_wylaczonych': energia_wylaczonych, 'energia_calkowita': energia_calkowita, 'udzial_filtra_proc': energia_wlaczonych / energia_calkowita * 100 if energia_calkowita > 0 else 0.0, 'liczba_gen_przed_filtrem': len(gen_z_H), 'zapotrzebowanie': zapotrzebowanie}
        return {'wiersze': wiersze, 'podsumowanie': podsumowanie}

    def pokaz_szczegoly_modelu(self, model_type, key):
        dane = self.pobierz_dane_szczegolowe(model_type, key)
        if dane is None:
            self.pole_szczegoly.delete('1.0', tk.END)
            self.pole_szczegoly.insert(tk.END, '❌ Brak danych dla wybranego modelu')
            return
        wiersze = dane['wiersze']
        p = dane['podsumowanie']
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        plik_txt = opis.get('plik', '')
        typ_txt = '24h' if model_type == 'manual' else 'AUTO'
        self.pole_szczegoly.delete('1.0', tk.END)
        filtr_info = f"(filtr: '{self.tekst_filtrowania}' – {p['liczba_gen_odfiltrowanych']}/{p['liczba_gen_przed_filtrem']})" if self.tekst_filtrowania else ''
        self.label_wynikow.config(text=f"{p['liczba_gen_odfiltrowanych']}/{p['liczba_gen_przed_filtrem']} generatorów po filtrze")
        self.pole_szczegoly.insert(tk.END, f'🔍 Model [{typ_txt}]: {data_txt} {godz_txt} {filtr_info}\n', 'header')
        if plik_txt:
            self.pole_szczegoly.insert(tk.END, f'📄 Plik: {plik_txt}\n', 'header')
        self.pole_szczegoly.insert(tk.END, f"{'=' * 190}\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"📈 Całkowita inercja E_sys = {p['energia_calkowita']:,.2f} MW·s\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"⚡ Zapotrzebowanie (LOAD DATA, suma mw aktywnych loadów) = {p['zapotrzebowanie']:,.2f} MW\n", 'header')
        self.pole_szczegoly.insert(tk.END, f"{'=' * 190}\n\n", 'header')
        naglowek = f"{'LP':<4} {'Kod generatora':<28} {'Grupa':<14} {'Status':<10} {'Pmax [MW]':<12} {'H [s]':<10} {'E_i=P×H [MW·s]':<18} {'Uwzględniony':<18} {'Ręczna zmiana':<16}\n"
        self.pole_szczegoly.insert(tk.END, naglowek, 'summary')
        self.pole_szczegoly.insert(tk.END, f"{'-' * 190}\n", 'summary')
        for w in wiersze:
            znacznik = 'TAK' if w['reczna_zmiana'] else '-'
            line = f"{w['lp']:<4} {w['kod_generatora'][:26]:<28} {w['grupa'][:12]:<14} {w['status']:<10} {w['pmax']:>10.3f}   {w['H']:>8.3f}   {w['Ei']:>16.3f}   {w['uwzgledniony']:<18} {znacznik:<16}\n"
            if w['reczna_zmiana']:
                tag = 'manual_line'
            elif w['status'] == 'ZAŁ.':
                tag = 'green_line'
            else:
                tag = 'red_line'
            self.pole_szczegoly.insert(tk.END, line, tag)
        self.pole_szczegoly.insert(tk.END, f"{'-' * 190}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, '📊 PODSUMOWANIE:\n', 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Odfiltrowane generatory:                           {p['liczba_gen_odfiltrowanych']:>6}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Załączone (po filtrze):                            {p['liczba_wlaczonych']:>6}\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji odfiltrowanych generatorów:          {p['energia_odfiltrowanych']:>12.2f} MW·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji uwzględnionych generatorów w obliczeniach: {p['energia_wlaczonych']:>12.2f} MW·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Suma inercji nieuwzględnionych generatorów:       {p['energia_wylaczonych']:>12.2f} MW·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Całkowita inercja systemu:                        {p['energia_calkowita']:>12.2f} MW·s\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Zapotrzebowanie systemu (LOAD DATA):              {p['zapotrzebowanie']:>12.2f} MW\n", 'summary')
        self.pole_szczegoly.insert(tk.END, f"• Udział odfiltrowanych w E_sys:                    {p['udzial_filtra_proc']:>11.2f} %\n", 'summary')
        self.pole_szczegoly.see('1.0')

    def aktualizuj_porownanie_modelu(self, model_type, key):
        E_org = self.get_model_energy_original(model_type, key)
        E_now = self.get_model_energy_now(model_type, key)
        P_load = self.get_model_demand_now(model_type, key)
        diff = E_now - E_org
        diff_proc = diff / E_org * 100 if E_org > 0 else 0.0
        opis = self.get_model_meta(model_type, key)
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', '--:--')
        typ_txt = '24h' if model_type == 'manual' else 'AUTO'
        txt = f'Model [{typ_txt}]: {data_txt} {godz_txt}\nInercja oryginalna: {E_org:,.2f} MW·s\nInercja po zmianach: {E_now:,.2f} MW·s\nRóżnica: {diff:,.2f} MW·s\nZmiana procentowa: {diff_proc:,.2f} %\nZapotrzebowanie mocy: {P_load:,.2f} MW'
        self.label_porownanie.config(text=txt, fg=self.get_palette()['success'] if diff >= 0 else self.get_palette()['danger'])

    def aktualizuj_wskazniki_jakosci(self):
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if not dodatnie:
            self.label_wskazniki.config(text='Brak danych 24h', fg=self.get_palette()['accent'])
            return
        wartosci = [v for _, v in dodatnie]
        srednia = statistics.mean(wartosci)
        mediana = statistics.median(wartosci)
        minimum = min(wartosci)
        maksimum = max(wartosci)
        odchylenie = statistics.pstdev(wartosci) if len(wartosci) > 1 else 0.0
        zap = [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h, _ in dodatnie]
        srednie_zap = statistics.mean(zap) if zap else 0.0
        min_zap = min(zap) if zap else 0.0
        max_zap = max(zap) if zap else 0.0
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            godziny_naruszenia = 'brak progu'
            liczba_ponizej = 'brak progu'
        else:
            godziny = [h for h, v in dodatnie if v < prog]
            liczba_ponizej = len(godziny)
            godziny_naruszenia = ', '.join((f'{h:02d}:30' for h in godziny)) if godziny else 'brak'
        udzialy = self.policz_udzialy_top_globalnie()
        txt = f"Średnia inercja dobowa: {srednia:,.2f} MW·s\nMinimum / maksimum: {minimum:,.2f} / {maksimum:,.2f} MW·s\nMediana: {mediana:,.2f} MW·s\nOdchylenie standardowe: {odchylenie:,.2f} MW·s\nŚrednie zapotrzebowanie: {srednie_zap:,.2f} MW\nZapotrzebowanie min / max: {min_zap:,.2f} / {max_zap:,.2f} MW\nLiczba godzin poniżej progu: {liczba_ponizej}\nGodziny poniżej progu: {godziny_naruszenia}\nUdział TOP 1 / TOP 3 / TOP 5 generatorów: {udzialy['top1']:.2f}% / {udzialy['top3']:.2f}% / {udzialy['top5']:.2f}%"
        self.label_wskazniki.config(text=txt, fg=self.get_palette()['accent'])

    def policz_udzialy_top_globalnie(self):
        wkłady = []
        suma = 0.0
        for godzina, gen_list in self.wszystkie_generatory.items():
            for gen in gen_list:
                if gen['status'] == 1 and gen['reg_name'] in self.H_generatory:
                    ei = gen['pmax'] * self.H_generatory[gen['reg_name']]
                    wkłady.append(ei)
                    suma += ei
        if suma <= 0 or not wkłady:
            return {'top1': 0.0, 'top3': 0.0, 'top5': 0.0}
        wkłady.sort(reverse=True)
        return {'top1': sum(wkłady[:1]) / suma * 100, 'top3': sum(wkłady[:3]) / suma * 100, 'top5': sum(wkłady[:5]) / suma * 100}

    def aktualizuj_ranking(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            self.label_ranking.config(text='Brak wybranego modelu')
            return
        ranking = []
        for gen in gen_list:
            if gen['reg_name'] in self.H_generatory:
                Ei = gen['pmax'] * self.H_generatory[gen['reg_name']]
                ranking.append((gen['reg_name'], gen['status'], gen['pmax'], self.H_generatory[gen['reg_name']], Ei, self.formatuj_grupe_generatora(gen)))
        ranking.sort(key=lambda x: x[4], reverse=True)
        ranking = ranking[:10]
        if not ranking:
            self.label_ranking.config(text='Brak generatorów z H dla wybranego modelu')
            return
        linie = [f"{'LP':<3} {'Kod':<20} {'Gr':<10} {'St':<4} {'Pmax':>8} {'H':>8} {'Ei':>12}"]
        for i, (nazwa, status, pmax, H, Ei, grupa_txt) in enumerate(ranking, 1):
            st = 'ZAŁ' if status == 1 else 'WYŁ'
            linie.append(f'{i:<3} {nazwa[:20]:<20} {grupa_txt[:10]:<10} {st:<4} {pmax:>8.2f} {H:>8.2f} {Ei:>12.2f}')
        self.label_ranking.config(text='\n'.join(linie))

    def oblicz_udzialy_grup_generatorow(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            return []
        suma_e = self.get_model_energy_now(model_type, key)
        if suma_e <= 0:
            return []
        grupy = {}
        for gen in gen_list:
            if gen['status'] != 1:
                continue
            if gen['reg_name'] not in self.H_generatory:
                continue
            code = gen.get('group_code')
            if code is None:
                continue
            Ei = gen['pmax'] * self.H_generatory[gen['reg_name']]
            grupy[code] = grupy.get(code, 0.0) + Ei
        if not grupy:
            return []
        wynik = []
        for code, ei_sum in sorted(grupy.items(), key=lambda x: x[1], reverse=True):
            label = self.GENERATOR_GROUP_LABELS.get(code, f'KOD {code}')
            udzial = ei_sum / suma_e * 100 if suma_e > 0 else 0.0
            wynik.append({'code': code, 'label': label, 'ei_sum': ei_sum, 'share_pct': udzial})
        return wynik

    def aktualizuj_wykres_grup_generatorow(self, model_type, key):
        dane = self.oblicz_udzialy_grup_generatorow(model_type, key)
        pal = self.get_palette()
        self.figura_grupy.patch.set_facecolor(pal['bg'])
        self.os_grupy.clear()
        self.os_grupy.set_title('Udział grup generatorów w modelu')
        self.os_grupy.set_xlabel('Grupa generatora')
        self.os_grupy.set_ylabel('Udział [% E_sys]')
        self.stylizuj_os(self.os_grupy, grid=True, grid_axis='y')
        if not dane:
            self.os_grupy.set_ylim(0, 100)
            self.os_grupy.text(0.5, 0.5, 'Brak danych grup generatorów\nlub brak aktywnych generatorów z H', ha='center', va='center', transform=self.os_grupy.transAxes, fontsize=10, color=pal['muted'])
            self.canvas_grupy.draw_idle()
            return
        x_labels = [f"{d['code']}\n{d['label']}" for d in dane]
        y_vals = [d['share_pct'] for d in dane]
        colors = [self.GROUP_COLORS[i % len(self.GROUP_COLORS)] for i in range(len(dane))]
        bars = self.os_grupy.bar(range(len(dane)), y_vals, color=colors, edgecolor=pal['border'], linewidth=0.8)
        self.os_grupy.set_xticks(range(len(dane)))
        self.os_grupy.set_xticklabels(x_labels, rotation=0, ha='center', fontsize=9, color=pal['text'])
        ymax = max(y_vals) if y_vals else 1.0
        self.os_grupy.set_ylim(0, max(100 if ymax < 80 else ymax * 1.2, 5))
        for bar, d in zip(bars, dane):
            h = bar.get_height()
            self.os_grupy.annotate(f"{d['share_pct']:.1f}%", xy=(bar.get_x() + bar.get_width() / 2, h), xytext=(0, 4), textcoords='offset points', ha='center', va='bottom', fontsize=8, fontweight='bold', color=pal['text'])
        self.canvas_grupy.draw_idle()

    def aktualizuj_analize_inercyjna(self, model_type, key):
        gen_list = self.get_model_generators(model_type, key)
        if not gen_list:
            self.label_analiza_inercyjna.config(text='Brak wybranego modelu')
            return
        aktywne = []
        nieaktywne = []
        for gen in gen_list:
            H = self.H_generatory.get(gen['reg_name'])
            if not H:
                continue
            Ei = gen['pmax'] * H
            rekord = (gen['reg_name'], Ei)
            if gen['status'] == 1:
                aktywne.append(rekord)
            else:
                nieaktywne.append(rekord)
        aktywne.sort(key=lambda x: x[1], reverse=True)
        nieaktywne.sort(key=lambda x: x[1], reverse=True)
        E = self.get_model_energy_now(model_type, key)
        P = self.get_model_demand_now(model_type, key)
        prog = self.pobierz_wartosc_graniczna()
        trigger_mode = self.email_alarm.config.trigger_mode
        top3 = sum((v for _, v in aktywne[:3]))
        koncentracja = top3 / E * 100 if E > 0 else 0.0
        linie = [f'Bieżąca inercja: {E:,.2f} MW·s', f'Zapotrzebowanie systemu: {P:,.2f} MW']
        if prog is not None:
            if trigger_mode == 'above':
                if E >= prog:
                    linie.append(f'Margines względem progu: +{E - prog:,.2f} MW·s')
                else:
                    linie.append(f'Niedobór względem progu: {E - prog:,.2f} MW·s')
            elif E < prog:
                linie.append(f'Model jest poniżej progu o {prog - E:,.2f} MW·s')
            else:
                linie.append(f'Model jest powyżej progu o {E - prog:,.2f} MW·s')
        linie.append(f'Koncentracja TOP3 aktywnych: {koncentracja:.2f} % E_sys')
        if aktywne:
            linie.append(f'Najbardziej krytyczny aktywny generator: {aktywne[0][0]} ({aktywne[0][1]:.2f} MW·s)')
        if nieaktywne:
            linie.append(f'Najlepszy kandydat do załączenia: {nieaktywne[0][0]} (+{nieaktywne[0][1]:.2f} MW·s)')
        grupy = self.oblicz_udzialy_grup_generatorow(model_type, key)
        if grupy:
            top_grupa = grupy[0]
            linie.append(f"Dominująca grupa generatorów: {top_grupa['code']} ({top_grupa['label']}) = {top_grupa['share_pct']:.2f}% E_sys")
        if prog is not None and trigger_mode == 'above' and (E < prog):
            brak = prog - E
            suma = 0.0
            rekom = []
            for nazwa, ei in nieaktywne:
                rekom.append(f'{nazwa} (+{ei:.1f})')
                suma += ei
                if suma >= brak:
                    break
            if rekom:
                linie.append('Rekomendowane załączenia do przekroczenia progu: ' + ', '.join(rekom))
        ok_color = True
        if prog is not None:
            ok_color = E >= prog
        self.label_analiza_inercyjna.config(text='\n'.join(linie), fg=self.get_palette()['success'] if ok_color else self.get_palette()['danger'])

    def formatuj_opis_modelu(self, godzina):
        opis = self.opis_modeli.get(godzina, {})
        data_txt = opis.get('data', 'brak daty')
        godz_txt = opis.get('godzina', f'{godzina:02d}:30')
        prefix = ''
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if dodatnie:
            godz_min, _ = min(dodatnie, key=lambda x: x[1])
            godz_max, _ = max(dodatnie, key=lambda x: x[1])
            if godzina == godz_min:
                prefix += '[MIN] '
            if godzina == godz_max:
                prefix += '[MAX] '
        return f'{prefix}{godz_txt} | {data_txt}'

    def klasyfikuj_model_wg_progu(self, wartosc):
        prog = self.pobierz_wartosc_graniczna()
        if prog is None or prog <= 0:
            return None
        if self.email_alarm.config.trigger_mode == 'above':
            if wartosc < prog:
                return 'danger'
            if wartosc <= prog * (1 + self.bliski_prog_proc):
                return 'warning'
            return 'safe'
        else:
            if wartosc > prog:
                return 'safe'
            if wartosc >= prog * (1 - self.bliski_prog_proc):
                return 'warning'
            return 'danger'

    def odswiez_liste_modeli(self):
        self.listbox_modele.delete(0, tk.END)
        indeksy = []
        for godzina in sorted(self.wszystkie_generatory.keys()):
            txt = self.formatuj_opis_modelu(godzina)
            self.listbox_modele.insert(tk.END, txt)
            indeksy.append((self.listbox_modele.size() - 1, godzina))
        self.listbox_modele.insert(tk.END, '')
        self.listbox_modele.insert(tk.END, f'📋 RAZEM: {len(self.wszystkie_generatory)} modeli EPC')
        for idx, godzina in indeksy:
            status = self.klasyfikuj_model_wg_progu(self.inercja_na_godzine.get(godzina, 0.0))
            try:
                if status == 'danger':
                    self.listbox_modele.itemconfig(idx, fg='#dc2626' if not self.dark_mode_var.get() else '#fca5a5')
                elif status == 'warning':
                    self.listbox_modele.itemconfig(idx, fg='#ea580c' if not self.dark_mode_var.get() else '#fdba74')
                elif status == 'safe':
                    self.listbox_modele.itemconfig(idx, fg='#15803d' if not self.dark_mode_var.get() else '#86efac')
            except Exception:
                pass

    def odswiez_liste_modeli_auto(self):
        self.listbox_modele_auto.delete(0, tk.END)
        indeksy = []
        for auto_id in self.auto_order:
            m = self.modele_auto[auto_id]
            opis = m['opis']
            self.listbox_modele_auto.insert(tk.END, f"{opis.get('godzina', '--:--')} | {opis.get('data', 'brak daty')} | {m.get('inercja', 0.0):,.2f} MW·s | {m.get('zapotrzebowanie', 0.0):,.2f} MW | {opis.get('plik', '')}")
            indeksy.append((self.listbox_modele_auto.size() - 1, auto_id))
        self.listbox_modele_auto.insert(tk.END, '')
        self.listbox_modele_auto.insert(tk.END, f'🤖 RAZEM: {len(self.auto_order)} modeli automatycznych')
        for idx, auto_id in indeksy:
            status = self.klasyfikuj_model_wg_progu(self.modele_auto[auto_id].get('inercja', 0.0))
            try:
                if status == 'danger':
                    self.listbox_modele_auto.itemconfig(idx, fg='#dc2626' if not self.dark_mode_var.get() else '#fca5a5')
                elif status == 'warning':
                    self.listbox_modele_auto.itemconfig(idx, fg='#ea580c' if not self.dark_mode_var.get() else '#fdba74')
                elif status == 'safe':
                    self.listbox_modele_auto.itemconfig(idx, fg='#15803d' if not self.dark_mode_var.get() else '#86efac')
            except Exception:
                pass

    def aktualizuj_panel_maksimum(self, check_alarm=False):
        dodatnie = [(h, v) for h, v in self.inercja_na_godzine.items() if v > 0]
        if not dodatnie:
            self.label_maksimum_info.config(text='Maksymalna inercja z 24 godzin: brak danych', fg=self.get_palette()['accent'])
            self.label_limit_ocena.config(text='Ocena: brak wartości do porównania', fg=self.get_palette()['muted'])
            return
        godz_max, val_max = max(dodatnie, key=lambda x: x[1])
        godz_min, val_min = min(dodatnie, key=lambda x: x[1])
        opis_max = self.opis_modeli.get(godz_max, {})
        data_txt_max = opis_max.get('data', 'brak daty')
        godz_txt_max = opis_max.get('godzina', f'{godz_max:02d}:30')
        self.label_maksimum_info.config(text=f'Maks: {val_max:,.2f} MW·s ({data_txt_max} {godz_txt_max}) | Min: {val_min:,.2f} MW·s ({godz_min:02d}:30)', fg=self.get_palette()['accent'])
        prog = self.pobierz_wartosc_graniczna()
        if prog is None:
            txt = self.entry_wartosc_graniczna.get().strip()
            if txt:
                self.label_limit_ocena.config(text='Ocena: niepoprawna wartość graniczna', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text='Ocena: brak wartości granicznej', fg=self.get_palette()['muted'])
            return
        if self.email_alarm.config.trigger_mode == 'above':
            if val_max > prog:
                self.label_limit_ocena.config(text=f'Ocena: PRZEKROCZONO wartość graniczną o {val_max - prog:,.2f} MW·s', fg=self.get_palette()['danger'])
            else:
                self.label_limit_ocena.config(text=f'Ocena: NIE przekroczono wartości granicznej (zapas {prog - val_max:,.2f} MW·s)', fg=self.get_palette()['success'])
        elif val_min < prog:
            self.label_limit_ocena.config(text=f'Ocena: SPADŁO poniżej wartości granicznej o {prog - val_min:,.2f} MW·s', fg=self.get_palette()['danger'])
        else:
            self.label_limit_ocena.config(text=f'Ocena: NIE spadło poniżej wartości granicznej (margines {val_min - prog:,.2f} MW·s)', fg=self.get_palette()['success'])

    def aktualizuj_wykres(self):
        pal = self.get_palette()
        self.os1.clear()
        self.os2.clear()
        if not hasattr(self, 'os1_load') or self.os1_load is None:
            self.os1_load = self.os1.twinx()
        self.os1_load.clear()
        self.figura.patch.set_facecolor(pal['bg'])
        self.os1.set_title('Inercja systemowa E_sys(t) = Σ Pmax_i(t) × H_i')
        self.os1.set_xlabel('Czas [HH:30]')
        self.os1.set_ylabel('Inercja [MW·s]')
        self.stylizuj_os(self.os1, grid=True, grid_axis='both')
        self.os1.yaxis.set_label_coords(-0.09, 0.78)
        godziny = list(range(24))
        wartosci = [self.inercja_na_godzine.get(h, 0.0) for h in godziny]
        zapotrzebowania = [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h in godziny]
        dostepne_godziny = set(self.wszystkie_generatory.keys())
        self.os1.plot(godziny, wartosci, 'o-', linewidth=3, markersize=8, color=pal['line_main'], markerfacecolor=pal['selection'], markeredgecolor=pal['accent_active'], label='Inercja [MW·s]', zorder=3)
        self.os1_load.set_ylabel('Zapotrzebowanie [MW]')
        self.stylizuj_os_prawa(self.os1_load, pal['warning'])
        self.os1_load.plot(godziny, zapotrzebowania, 's--', linewidth=2.2, markersize=6, color=pal['warning'], markerfacecolor=pal['card'], markeredgecolor=pal['warning_active'], label='Zapotrzebowanie [MW]', zorder=2)
        prog = self.pobierz_wartosc_graniczna()
        if self.pokaz_linie_progu_var.get() and prog is not None:
            self.os1.axhline(prog, linestyle='-', linewidth=1.2, color=pal['danger'], label=f'Próg alarmowy: {prog:.0f}', zorder=1)
        self.os1.set_xticks(range(24))
        self.os1.set_xticklabels([f'{h:02d}:30' for h in range(24)], rotation=45)
        self.os1.set_xlim(-0.5, 23.5)
        max_source = wartosci[:]
        if prog is not None:
            max_source.append(prog)
        max_val = max(max_source) * 1.2 if max(max_source) > 0 else 1.0
        self.os1.set_ylim(0, max_val)
        if dostepne_godziny:
            zap_dost = [self.zapotrzebowanie_na_godzine.get(h, 0.0) for h in dostepne_godziny]
            min_load = min(zap_dost)
            max_load = max(zap_dost)
            if abs(max_load - min_load) < 1e-09:
                pad = max(1.0, abs(max_load) * 0.2)
            else:
                pad = max(1.0, (max_load - min_load) * 0.15)
            dol = min(0.0, min_load - pad)
            gora = max(1.0, max_load + pad)
        else:
            dol, gora = (0.0, 1.0)
        self.os1_load.set_xlim(-0.5, 23.5)
        self.os1_load.set_ylim(dol, gora)
        wartosci_dodatnie = [(h, v) for h, v in zip(godziny, wartosci) if v > 0]
        if wartosci_dodatnie:
            godz_min, val_min = min(wartosci_dodatnie, key=lambda x: x[1])
            godz_max, val_max = max(wartosci_dodatnie, key=lambda x: x[1])
            self.os1.plot(godz_min, val_min, 'o', markersize=14, color=pal['danger'], label=f'Minimum: {godz_min:02d}:30', zorder=4)
            self.os1.plot(godz_max, val_max, 'o', markersize=14, color=pal['success'], label=f'Maksimum: {godz_max:02d}:30', zorder=4)
            self.dodaj_etykiete_punktu(self.os1, godz_min, val_min, f'MIN {godz_min:02d}:30\n{val_min:.0f}', -42 if godz_min % 2 == 0 else -30, -34, fg=pal['danger'], facecolor=pal['card'], edgecolor=pal['danger'], valign='top', fontsize=9)
            self.dodaj_etykiete_punktu(self.os1, godz_max, val_max, f'MAX {godz_max:02d}:30\n{val_max:.0f}', 42 if godz_max % 2 == 0 else 30, 30, fg=pal['success'], facecolor=pal['card'], edgecolor=pal['success'], valign='bottom', fontsize=9)
        y1_min, y1_max = self.os1.get_ylim()
        y2_min, y2_max = self.os1_load.get_ylim()
        load_scaled_values = [self.przeskaluj_wartosc_miedzy_osiami(load_val, y2_min, y2_max, y1_min, y1_max) for load_val in zapotrzebowania]
        if self.pokaz_etykiety_inercji_var.get():
            for h, val in enumerate(wartosci):
                if val <= 0:
                    continue
                x_off, y_off = self.wyznacz_offset_etykiety(h, val, wartosci, prefer='up', other_scaled=load_scaled_values[h])
                self.dodaj_etykiete_punktu(self.os1, h, val, f'{val:.0f}', x_off, y_off, fg=pal['text'], facecolor=pal['header'], edgecolor=pal['border'], valign='bottom')
        if self.pokaz_etykiety_zapotrzebowania_var.get():
            for h, load_val in enumerate(zapotrzebowania):
                if load_val <= 0:
                    continue
                x_off, y_off = self.wyznacz_offset_etykiety(h, load_scaled_values[h], load_scaled_values, prefer='down', other_scaled=wartosci[h])
                self.dodaj_etykiete_punktu(self.os1_load, h, load_val, f'{load_val:.0f}', x_off, y_off, fg=pal['warning_active'], facecolor=pal['card_alt'], edgecolor=pal['warning'], valign='top')
        h1, l1 = self.os1.get_legend_handles_labels()
        h2, l2 = self.os1_load.get_legend_handles_labels()
        legend = self.os1.legend(h1 + h2, l1 + l2, loc='lower left')
        if legend:
            legend.get_frame().set_facecolor(pal['card'])
            legend.get_frame().set_edgecolor(pal['border'])
            for txt in legend.get_texts():
                txt.set_color(pal['text'])
        licznosci = [self.liczba_zalaczonych_na_godzine.get(h, 0) for h in godziny]
        self.os2.set_title('Liczba załączonych generatorów z H')
        self.os2.set_xlabel('Czas [HH:30]')
        self.os2.set_ylabel('Liczba')
        self.stylizuj_os(self.os2, grid=True, grid_axis='both')
        self.os2.plot(godziny, licznosci, 'o-', linewidth=2.4, markersize=6, color=pal['line_secondary'])
        self.os2.set_xticks(range(24))
        self.os2.set_xticklabels([f'{h:02d}:30' for h in range(24)], rotation=45)
        self.os2.set_xlim(-0.5, 23.5)
        max_count = max(licznosci) * 1.2 if max(licznosci) > 0 else 1
        self.os2.set_ylim(0, max_count)
        self.canvas.draw_idle()

    def eksportuj_do_excel(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        dane = self.pobierz_dane_szczegolowe(model_type, key)
        if dane is None:
            messagebox.showwarning('Brak danych', 'Brak danych do eksportu.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz plik Excel', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not sciezka:
            return
        try:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = 'Szczegóły'
            header_fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            manual_fill = PatternFill(fill_type='solid', fgColor='FFF3CD')
            bold_font = Font(bold=True)
            manual_font = Font(bold=True, color='C77700')
            center = Alignment(horizontal='center', vertical='center')
            opis = self.get_model_meta(model_type, key)
            ws1['A1'] = 'Typ modelu'
            ws1['B1'] = '24h' if model_type == 'manual' else 'AUTO'
            ws1['A2'] = 'Data modelu'
            ws1['B2'] = opis.get('data', '')
            ws1['A3'] = 'Godzina modelu'
            ws1['B3'] = opis.get('godzina', '')
            ws1['A4'] = 'Filtr'
            ws1['B4'] = self.tekst_filtrowania if self.tekst_filtrowania else '(brak)'
            ws1['A5'] = 'Całkowita inercja systemu [MW·s]'
            ws1['B5'] = dane['podsumowanie']['energia_calkowita']
            ws1['A6'] = 'Zapotrzebowanie systemu [MW]'
            ws1['B6'] = dane['podsumowanie']['zapotrzebowanie']
            for cell in ['A1', 'A2', 'A3', 'A4', 'A5', 'A6']:
                ws1[cell].font = bold_font
            start_row = 8
            headers = ['LP', 'Kod generatora', 'Grupa', 'Status', 'Pmax [MW]', 'H [s]', 'E_i = Pmax × H [MW·s]', 'Uwzględniony w obliczeniach', 'Ręczna zmiana']
            for col, header in enumerate(headers, 1):
                c = ws1.cell(row=start_row, column=col, value=header)
                c.font = bold_font
                c.fill = header_fill
                c.alignment = center
            for r, w in enumerate(dane['wiersze'], start=start_row + 1):
                ws1.cell(row=r, column=1, value=w['lp'])
                ws1.cell(row=r, column=2, value=w['kod_generatora'])
                ws1.cell(row=r, column=3, value=w['grupa'])
                ws1.cell(row=r, column=4, value=w['status'])
                ws1.cell(row=r, column=5, value=w['pmax'])
                ws1.cell(row=r, column=6, value=w['H'])
                ws1.cell(row=r, column=7, value=w['Ei'])
                ws1.cell(row=r, column=8, value=w['uwzgledniony'])
                ws1.cell(row=r, column=9, value='TAK' if w['reczna_zmiana'] else '-')
                if w['reczna_zmiana']:
                    for col in range(1, 10):
                        ws1.cell(row=r, column=col).fill = manual_fill
                        ws1.cell(row=r, column=col).font = manual_font
            for col, width in {'A': 8, 'B': 28, 'C': 16, 'D': 12, 'E': 14, 'F': 12, 'G': 22, 'H': 30, 'I': 16}.items():
                ws1.column_dimensions[col].width = width
            if model_type == 'manual':
                ws2 = wb.create_sheet('Podsumowanie godzin')
                headers2 = ['Data', 'Godzina', 'Inercja oryginalna [MW·s]', 'Inercja po zmianach [MW·s]', 'Różnica [MW·s]', 'Zapotrzebowanie [MW]']
                for col, header in enumerate(headers2, 1):
                    c = ws2.cell(row=1, column=col, value=header)
                    c.font = bold_font
                    c.fill = header_fill
                    c.alignment = center
                posortowane = sorted(self.wszystkie_generatory.keys())
                for idx, godzina in enumerate(posortowane, start=2):
                    opis_m = self.opis_modeli.get(godzina, {})
                    e_org = self.inercja_oryginalna_na_godzine.get(godzina, 0.0)
                    e_now = self.inercja_na_godzine.get(godzina, 0.0)
                    p_load = self.zapotrzebowanie_na_godzine.get(godzina, 0.0)
                    ws2.cell(row=idx, column=1, value=opis_m.get('data', ''))
                    ws2.cell(row=idx, column=2, value=opis_m.get('godzina', f'{godzina:02d}:30'))
                    ws2.cell(row=idx, column=3, value=e_org)
                    ws2.cell(row=idx, column=4, value=e_now)
                    ws2.cell(row=idx, column=5, value=e_now - e_org)
                    ws2.cell(row=idx, column=6, value=p_load)
                for col, width in {'A': 16, 'B': 14, 'C': 24, 'D': 24, 'E': 18, 'F': 18}.items():
                    ws2.column_dimensions[col].width = width
            if self.alarm_history:
                ws3 = wb.create_sheet('Historia alarmów')
                headers3 = ['Kiedy alarm wystąpił', 'Tryb', 'Data modelu', 'Godzina modelu', 'Wartość [MW·s]', 'Próg [MW·s]', 'Plik', 'Mail', 'Status wysyłki', 'Opis']
                for col, header in enumerate(headers3, 1):
                    c = ws3.cell(row=1, column=col, value=header)
                    c.font = bold_font
                    c.fill = header_fill
                    c.alignment = center
                for r, z in enumerate(self.alarm_history, start=2):
                    ws3.cell(row=r, column=1, value=z.get('czas_alarmu', ''))
                    ws3.cell(row=r, column=2, value=z.get('trigger_mode', ''))
                    ws3.cell(row=r, column=3, value=z.get('data_modelu', ''))
                    ws3.cell(row=r, column=4, value=z.get('godzina_modelu', ''))
                    ws3.cell(row=r, column=5, value=z.get('wartosc', 0.0))
                    ws3.cell(row=r, column=6, value=z.get('prog', 0.0))
                    ws3.cell(row=r, column=7, value=z.get('source_file', ''))
                    ws3.cell(row=r, column=8, value='TAK' if z.get('mail_attempted') else 'NIE')
                    ws3.cell(row=r, column=9, value=z.get('mail_status', ''))
                    ws3.cell(row=r, column=10, value=z.get('opis', ''))
                for col, width in {'A': 20, 'B': 10, 'C': 14, 'D': 14, 'E': 16, 'F': 14, 'G': 28, 'H': 10, 'I': 24, 'J': 24}.items():
                    ws3.column_dimensions[col].width = width
            wb.save(sciezka)
            messagebox.showinfo('Eksport zakończony', f'Zapisano plik:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu', f'Nie udało się zapisać pliku Excel:\n{e}')

    def pobierz_liste_zmian(self):
        zmiany = []
        for godzina, gen_list in self.wszystkie_generatory.items():
            opis = self.opis_modeli.get(godzina, {})
            oryginalne = self.statusy_oryginalne.get(godzina, {})
            for gen in gen_list:
                org = oryginalne.get(gen['reg_name'], gen['status'])
                now = gen['status']
                if org != now:
                    H = self.H_generatory.get(gen['reg_name'], 0.0)
                    Ei = gen['pmax'] * H
                    zmiany.append({'typ': '24h', 'data': opis.get('data', ''), 'godzina': opis.get('godzina', f'{godzina:02d}:30'), 'generator': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status_oryginalny': org, 'status_po_zmianie': now, 'pmax': gen['pmax'], 'H': H, 'Ei': Ei})
        for auto_id in self.auto_order:
            m = self.modele_auto[auto_id]
            opis = m['opis']
            oryginalne = m['statusy_oryginalne']
            for gen in m['generatory']:
                org = oryginalne.get(gen['reg_name'], gen['status'])
                now = gen['status']
                if org != now:
                    H = self.H_generatory.get(gen['reg_name'], 0.0)
                    Ei = gen['pmax'] * H
                    zmiany.append({'typ': 'AUTO', 'data': opis.get('data', ''), 'godzina': opis.get('godzina', '--:--'), 'generator': gen['reg_name'], 'grupa': self.formatuj_grupe_generatora(gen), 'status_oryginalny': org, 'status_po_zmianie': now, 'pmax': gen['pmax'], 'H': H, 'Ei': Ei})
        return zmiany

    def eksportuj_raport_zmian(self):
        zmiany = self.pobierz_liste_zmian()
        if not zmiany:
            messagebox.showinfo('Raport zmian', 'Brak ręcznych zmian do eksportu.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz raport zmian', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx'), ('Tekst', '*.txt')])
        if not sciezka:
            return
        try:
            if sciezka.lower().endswith('.txt'):
                with open(sciezka, 'w', encoding='utf-8') as f:
                    f.write('RAPORT RĘCZNYCH ZMIAN STATUSÓW GENERATORÓW\n')
                    f.write('=' * 140 + '\n')
                    for z in zmiany:
                        f.write(f"{z['typ']} | {z['data']} {z['godzina']} | {z['generator']} | {z['grupa']} | {z['status_oryginalny']} -> {z['status_po_zmianie']} | Pmax={z['pmax']:.3f} | H={z['H']:.3f} | Ei={z['Ei']:.3f}\n")
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = 'Raport zmian'
                headers = ['Typ modelu', 'Data', 'Godzina', 'Generator', 'Grupa', 'Status oryginalny', 'Status po zmianie', 'Pmax [MW]', 'H [s]', 'Ei [MW·s]']
                bold = Font(bold=True)
                fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
                for col, h in enumerate(headers, 1):
                    c = ws.cell(row=1, column=col, value=h)
                    c.font = bold
                    c.fill = fill
                for r, z in enumerate(zmiany, start=2):
                    ws.cell(row=r, column=1, value=z['typ'])
                    ws.cell(row=r, column=2, value=z['data'])
                    ws.cell(row=r, column=3, value=z['godzina'])
                    ws.cell(row=r, column=4, value=z['generator'])
                    ws.cell(row=r, column=5, value=z['grupa'])
                    ws.cell(row=r, column=6, value=z['status_oryginalny'])
                    ws.cell(row=r, column=7, value=z['status_po_zmianie'])
                    ws.cell(row=r, column=8, value=z['pmax'])
                    ws.cell(row=r, column=9, value=z['H'])
                    ws.cell(row=r, column=10, value=z['Ei'])
                for col, width in {'A': 12, 'B': 14, 'C': 12, 'D': 28, 'E': 16, 'F': 18, 'G': 18, 'H': 12, 'I': 10, 'J': 14}.items():
                    ws.column_dimensions[col].width = width
                wb.save(sciezka)
            messagebox.showinfo('Eksport raportu', f'Zapisano raport zmian:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu raportu', str(e))

    def eksportuj_aktualny_model_epc_po_zmianach(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model.')
            return
        if not self.is_model_epc_exportable(model_type, key):
            messagebox.showwarning('Brak pliku źródłowego', 'Nie znaleziono pliku EPC do eksportu dla wybranego modelu.')
            return
        opis = self.get_model_meta(model_type, key)
        sciezka_zrodlowa = opis.get('sciezka', '')
        nazwa_domyslna = os.path.splitext(os.path.basename(sciezka_zrodlowa))[0] + '_po_zmianach.epc'
        sciezka_docelowa = filedialog.asksaveasfilename(title='Zapisz zmodyfikowany model EPC', defaultextension='.epc', initialfile=nazwa_domyslna, filetypes=[('Plik EPC', '*.epc'), ('Wszystkie', '*.*')])
        if not sciezka_docelowa:
            return
        gen_list = self.get_model_generators(model_type, key)
        status_map = {g['reg_name']: int(g['status']) for g in gen_list}
        try:
            with open(sciezka_zrodlowa, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()
            nowe_linie = []
            in_section = False
            liczba_podmienionych = 0
            for line in lines:
                line_stripped = line.strip()
                line_lower = line_stripped.lower()
                if 'generator data' in line_lower:
                    in_section = True
                    nowe_linie.append(line)
                    continue
                if in_section and 'load data' in line_lower:
                    in_section = False
                    nowe_linie.append(line)
                    continue
                if in_section and ':' in line_stripped:
                    nowa_linia, czy_zmieniono = self.zmien_status_w_linii_generatora(line, status_map)
                    nowe_linie.append(nowa_linia)
                    if czy_zmieniono:
                        liczba_podmienionych += 1
                else:
                    nowe_linie.append(line)
            with open(sciezka_docelowa, 'w', encoding='utf-8', errors='ignore') as f:
                f.writelines(nowe_linie)
            messagebox.showinfo('Eksport EPC zakończony', f'Zapisano zmodyfikowany model EPC:\n{sciezka_docelowa}\n\nZaktualizowane wpisy generatorów w EPC: {liczba_podmienionych}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu EPC', f'Nie udało się zapisać pliku EPC:\n{e}')

    def zmien_status_w_linii_generatora(self, line, status_map):
        try:
            if ':' not in line:
                return (line, False)
            line_no_nl = line.rstrip('\r\n')
            line_ending = line[len(line_no_nl):]
            lewa, prawa = line_no_nl.split(':', 1)
            quoted = re.findall('"([^"]*)"', lewa)
            reg_name = quoted[-1].strip() if quoted else None
            if not reg_name or reg_name not in status_map:
                return (line, False)
            m = re.match('^(\\s*)([-+]?\\d+(?:[.,]\\d+)?)(.*)$', prawa)
            if not m:
                return (line, False)
            leading_spaces = m.group(1)
            old_status_token = m.group(2)
            tail = m.group(3)
            nowy_status = int(status_map[reg_name])
            if '.' in old_status_token or ',' in old_status_token:
                separator = ',' if ',' in old_status_token else '.'
                status_txt = f'{float(nowy_status):.1f}'
                if separator == ',':
                    status_txt = status_txt.replace('.', ',')
            else:
                status_txt = str(nowy_status)
            nowa_linia = f'{lewa}:{leading_spaces}{status_txt}{tail}{line_ending}'
            czy_zmieniono = old_status_token != status_txt
            return (nowa_linia, czy_zmieniono)
        except Exception:
            return (line, False)

    def zapisz_sesje(self):
        if not self.wszystkie_generatory and (not self.modele_auto):
            messagebox.showwarning('Brak danych', 'Brak danych do zapisania.')
            return
        sciezka = filedialog.asksaveasfilename(title='Zapisz sesję', defaultextension='.json', filetypes=[('JSON', '*.json')])
        if not sciezka:
            return
        data = {'wszystkie_generatory': self.wszystkie_generatory, 'statusy_oryginalne': self.statusy_oryginalne, 'opis_modeli': self.opis_modeli, 'H_generatory': self.H_generatory, 'entry_filtr': self.entry_filtr.get(), 'tryb_filtrowania': self.tryb_filtrowania_var.get(), 'sortowanie': self.sortowanie_var.get(), 'wartosc_graniczna': self.entry_wartosc_graniczna.get(), 'obliczenia_wykonane': self.obliczenia_wykonane, 'monitor_folder_path': self.monitor_folder_path, 'modele_auto': self.modele_auto, 'auto_order': self.auto_order, 'auto_seen_files': self.auto_seen_files, 'auto_counter': self.auto_counter, 'email_alarm_last_key': self.email_alarm.last_alarm_key, 'sent_alarm_tokens': [list(x) for x in self.sent_alarm_tokens], 'alarm_history': self.alarm_history, 'zapotrzebowanie_na_godzine': self.zapotrzebowanie_na_godzine}
        try:
            with open(sciezka, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            messagebox.showinfo('Sesja', f'Zapisano sesję:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd zapisu sesji', str(e))

    def wczytaj_sesje(self):
        sciezka = filedialog.askopenfilename(title='Wczytaj sesję', filetypes=[('JSON', '*.json')])
        if not sciezka:
            return
        try:
            with open(sciezka, 'r', encoding='utf-8') as f:
                data = json.load(f)
            self.wszystkie_generatory = {int(k): v for k, v in data.get('wszystkie_generatory', {}).items()}
            self.statusy_oryginalne = {int(k): v for k, v in data.get('statusy_oryginalne', {}).items()}
            self.opis_modeli = {int(k): v for k, v in data.get('opis_modeli', {}).items()}
            raw_zap = {int(k): v for k, v in data.get('zapotrzebowanie_na_godzine', {}).items()}
            self.zapotrzebowanie_na_godzine = {h: raw_zap.get(h, 0.0) for h in range(24)}
            self.H_generatory = data.get('H_generatory', {})
            self.obliczenia_wykonane = data.get('obliczenia_wykonane', False)
            self.monitor_folder_path = data.get('monitor_folder_path', '')
            self.modele_auto = data.get('modele_auto', {})
            for auto_id in self.modele_auto:
                self.modele_auto[auto_id].setdefault('zapotrzebowanie', 0.0)
            self.auto_order = data.get('auto_order', [])
            self.auto_seen_files = data.get('auto_seen_files', {})
            self.auto_counter = data.get('auto_counter', 0)
            self.email_alarm.last_alarm_key = data.get('email_alarm_last_key')
            self.sent_alarm_tokens = {tuple(x) for x in data.get('sent_alarm_tokens', [])}
            self.alarm_history = data.get('alarm_history', [])
            self.inercja_na_godzine = {h: 0.0 for h in range(24)}
            self.inercja_oryginalna_na_godzine = {h: 0.0 for h in range(24)}
            self.liczba_zalaczonych_na_godzine = {h: 0 for h in range(24)}
            self.label_monitor_folder.config(text=f'Folder: {self.monitor_folder_path}' if self.monitor_folder_path else 'Folder: nie wybrano')
            self.entry_filtr.delete(0, tk.END)
            self.entry_filtr.insert(0, data.get('entry_filtr', ''))
            self.tryb_filtrowania_var.set(data.get('tryb_filtrowania', 'Wszystkie'))
            self.sortowanie_var.set(data.get('sortowanie', 'Kod generatora A-Z'))
            self.entry_wartosc_graniczna.delete(0, tk.END)
            self.entry_wartosc_graniczna.insert(0, data.get('wartosc_graniczna', ''))
            if not self.H_generatory and (self.wszystkie_generatory or self.modele_auto):
                messagebox.showwarning('Sesja', 'Sesja nie zawiera współczynników H albo są puste. Wyniki inercji mogą być zerowe.')
            for godzina in self.wszystkie_generatory.keys():
                self.przelicz_inercje_dla_modelu('manual', godzina)
            for auto_id in self.auto_order:
                self.przelicz_inercje_dla_modelu('auto', auto_id)
            self.aktywny_typ_modelu = None
            self.aktualna_godzina = None
            self.aktualny_auto_id = None
            self.odswiez_liste_modeli()
            self.odswiez_liste_modeli_auto()
            self.aktualizuj_wykres()
            self.aktualizuj_panel_maksimum(check_alarm=False)
            self.aktualizuj_wskazniki_jakosci()
            self.tree_generatory.delete(*self.tree_generatory.get_children())
            self.pole_szczegoly.delete('1.0', tk.END)
            self.label_porownanie.config(text='Brak wybranego modelu')
            self.label_ranking.config(text='Brak wybranego modelu')
            self.label_analiza_inercyjna.config(text='Brak wybranego modelu')
            self.label_reczne.config(text='Brak wybranego modelu')
            self.przygotuj_pusty_wykres_grup()
            messagebox.showinfo('Sesja', f'Wczytano sesję:\n{sciezka}')
        except Exception as e:
            messagebox.showerror('Błąd wczytywania sesji', str(e))

    def pokaz_walidacje_epc_dyd(self):
        if not self.wszystkie_generatory:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj EPC.')
            return
        if not self.H_generatory:
            messagebox.showwarning('Brak danych', 'Najpierw wczytaj DYD.')
            return
        epc_names = set()
        for gen_list in self.wszystkie_generatory.values():
            for g in gen_list:
                epc_names.add(g['reg_name'])
        dyd_names = set(self.H_generatory.keys())
        brak_H = sorted(epc_names - dyd_names)
        nadmiar_w_dyd = sorted(dyd_names - epc_names)
        okno = tk.Toplevel(self.root)
        okno.title('Walidacja EPC ↔ DYD')
        center_toplevel(okno, 1200, 800)
        okno.configure(bg=self.get_palette()['bg'])
        txt = scrolledtext.ScrolledText(okno, wrap=tk.WORD, font=self.font_code)
        txt.pack(fill=tk.BOTH, expand=True)
        pal = self.get_palette()
        txt.configure(bg=pal['input_bg'], fg=pal['input_fg'], insertbackground=pal['input_fg'], relief='flat', bd=0, highlightthickness=1, highlightbackground=pal['input_border'])
        txt.insert(tk.END, 'WALIDACJA ZGODNOŚCI EPC ↔ DYD\n')
        txt.insert(tk.END, '=' * 80 + '\n\n')
        txt.insert(tk.END, f'Generatory w EPC bez współczynnika H w DYD: {len(brak_H)}\n')
        txt.insert(tk.END, '-' * 80 + '\n')
        for x in brak_H:
            txt.insert(tk.END, x + '\n')
        txt.insert(tk.END, '\n')
        txt.insert(tk.END, f'Generatory w DYD nieobecne w EPC: {len(nadmiar_w_dyd)}\n')
        txt.insert(tk.END, '-' * 80 + '\n')
        for x in nadmiar_w_dyd:
            txt.insert(tk.END, x + '\n')
        txt.see('1.0')

    def export_chart(self, figure, suggested_name):
        path = filedialog.asksaveasfilename(title='Zapisz wykres', defaultextension=os.path.splitext(suggested_name)[1], initialfile=suggested_name, filetypes=[('PNG', '*.png'), ('PDF', '*.pdf'), ('Wszystkie', '*.*')])
        if not path:
            return
        try:
            figure.savefig(path, dpi=180, bbox_inches='tight')
            messagebox.showinfo('Eksport wykresu', f'Zapisano:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu wykresu', str(e))

    def export_dashboard_png(self):
        if not PIL_AVAILABLE:
            messagebox.showwarning('Brak Pillow', 'Do eksportu dashboardu PNG potrzebny jest pakiet Pillow.\nZainstaluj: pip install pillow')
            return
        path = filedialog.asksaveasfilename(title='Zapisz dashboard', defaultextension='.png', initialfile='dashboard_inercji.png', filetypes=[('PNG', '*.png')])
        if not path:
            return
        try:
            self.root.update()
            x = self.root.winfo_rootx()
            y = self.root.winfo_rooty()
            w = self.root.winfo_width()
            h = self.root.winfo_height()
            img = ImageGrab.grab(bbox=(x, y, x + w, y + h))
            img.save(path)
            messagebox.showinfo('Eksport dashboardu', f'Zapisano:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu dashboardu', str(e))

    def eksportuj_pelny_raport_pdf(self):
        model_type, key = self.get_current_model_ref()
        if model_type is None:
            messagebox.showwarning('Brak wyboru', 'Najpierw wybierz model z jednej z list.')
            return
        path = filedialog.asksaveasfilename(title='Zapisz pełny raport PDF', defaultextension='.pdf', initialfile='raport_inercji.pdf', filetypes=[('PDF', '*.pdf')])
        if not path:
            return
        try:
            with PdfPages(path) as pdf:
                fig1 = plt.figure(figsize=(11.69, 8.27))
                ax = fig1.add_subplot(111)
                ax.axis('off')
                opis = self.get_model_meta(model_type, key)
                E_org = self.get_model_energy_original(model_type, key)
                E_now = self.get_model_energy_now(model_type, key)
                P_load = self.get_model_demand_now(model_type, key)
                diff = E_now - E_org
                diff_proc = diff / E_org * 100 if E_org > 0 else 0.0
                text = []
                text.append('RAPORT ANALIZY INERCJI')
                text.append('=' * 60)
                text.append(f"Typ modelu: {('24h' if model_type == 'manual' else 'AUTO')}")
                text.append(f"Data modelu: {opis.get('data', '')}")
                text.append(f"Godzina modelu: {opis.get('godzina', '')}")
                text.append(f"Plik: {opis.get('plik', '')}")
                text.append('')
                text.append('PODSUMOWANIE PRZED / PO')
                text.append(f'Inercja oryginalna: {E_org:,.2f} MW·s')
                text.append(f'Inercja po zmianach: {E_now:,.2f} MW·s')
                text.append(f'Różnica: {diff:,.2f} MW·s')
                text.append(f'Zmiana procentowa: {diff_proc:,.2f} %')
                text.append(f'Zapotrzebowanie systemu: {P_load:,.2f} MW')
                text.append('')
                prog = self.pobierz_wartosc_graniczna()
                if prog is not None:
                    text.append(f'Próg alarmowy: {prog:,.2f} MW·s')
                    tryb = 'powyżej progu' if self.email_alarm.config.trigger_mode == 'above' else 'poniżej progu'
                    text.append(f'Tryb alarmu: {tryb}')
                text.append('')
                dane = self.pobierz_dane_szczegolowe(model_type, key)
                if dane:
                    p = dane['podsumowanie']
                    text.append('PODSUMOWANIE FILTRA')
                    text.append(f"Liczba generatorów po filtrze: {p['liczba_gen_odfiltrowanych']}")
                    text.append(f"Liczba generatorów załączonych: {p['liczba_wlaczonych']}")
                    text.append(f"Energia uwzględnionych: {p['energia_wlaczonych']:,.2f} MW·s")
                    text.append(f"Energia nieuwzględnionych: {p['energia_wylaczonych']:,.2f} MW·s")
                    text.append(f"Zapotrzebowanie systemu (LOAD DATA): {p['zapotrzebowanie']:,.2f} MW")
                ax.text(0.02, 0.98, '\n'.join(text), va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig1, bbox_inches='tight')
                plt.close(fig1)
                pdf.savefig(self.figura, bbox_inches='tight')
                pdf.savefig(self.figura_grupy, bbox_inches='tight')
                fig_rank = plt.figure(figsize=(11.69, 8.27))
                axr = fig_rank.add_subplot(111)
                axr.axis('off')
                rank_txt = self.label_ranking.cget('text')
                axr.text(0.02, 0.98, 'RANKING GENERATORÓW\n\n' + rank_txt, va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig_rank, bbox_inches='tight')
                plt.close(fig_rank)
                fig_ana = plt.figure(figsize=(11.69, 8.27))
                axa = fig_ana.add_subplot(111)
                axa.axis('off')
                ana_txt = self.label_analiza_inercyjna.cget('text')
                wsk_txt = self.label_wskazniki.cget('text')
                axa.text(0.02, 0.98, 'ANALIZA INERCYJNA\n\n' + ana_txt + '\n\nWSKAŹNIKI JAKOŚCI 24H\n\n' + wsk_txt, va='top', ha='left', family='monospace', fontsize=10)
                pdf.savefig(fig_ana, bbox_inches='tight')
                plt.close(fig_ana)
                fig_alarm = plt.figure(figsize=(11.69, 8.27))
                axal = fig_alarm.add_subplot(111)
                axal.axis('off')
                lines = ['ALARMY\n']
                if self.alarm_history:
                    for z in self.alarm_history[-40:]:
                        lines.append(f"{z.get('czas_alarmu', '')} | {z.get('trigger_mode', '')} | {z.get('data_modelu', '')} {z.get('godzina_modelu', '')} | wartość={z.get('wartosc', 0.0):,.2f} | próg={z.get('prog', 0.0):,.2f} | {z.get('mail_status', '')}")
                else:
                    lines.append('Brak alarmów.')
                axal.text(0.02, 0.98, '\n'.join(lines), va='top', ha='left', family='monospace', fontsize=9)
                pdf.savefig(fig_alarm, bbox_inches='tight')
                plt.close(fig_alarm)
                fig_changes = plt.figure(figsize=(11.69, 8.27))
                axc = fig_changes.add_subplot(111)
                axc.axis('off')
                zmiany = self.pobierz_liste_zmian()
                lines = ['RĘCZNE ZMIANY\n']
                if zmiany:
                    for z in zmiany[:80]:
                        lines.append(f"{z['typ']} | {z['data']} {z['godzina']} | {z['generator']} | {z['grupa']} | {z['status_oryginalny']} -> {z['status_po_zmianie']} | Pmax={z['pmax']:.3f} | H={z['H']:.3f} | Ei={z['Ei']:.3f}")
                else:
                    lines.append('Brak ręcznych zmian.')
                axc.text(0.02, 0.98, '\n'.join(lines), va='top', ha='left', family='monospace', fontsize=8)
                pdf.savefig(fig_changes, bbox_inches='tight')
                plt.close(fig_changes)
            messagebox.showinfo('Eksport PDF', f'Zapisano raport PDF:\n{path}')
        except Exception as e:
            messagebox.showerror('Błąd eksportu PDF', str(e))
if __name__ == '__main__':
    root = tk.Tk()
    root.geometry('1500x950')
    try:
        root.state('zoomed')
    except Exception:
        pass
    app = AnalizaInercjiExcelApp(root)

    def _on_close():
        try:
            app.save_ui_config()
        except Exception:
            pass
        root.destroy()
    root.protocol('WM_DELETE_WINDOW', _on_close)
    root.mainloop()
